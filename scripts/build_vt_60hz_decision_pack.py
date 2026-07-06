from __future__ import annotations

import json
import sys
from pathlib import Path

import pandas as pd
from analysis.paths import results_dir
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.util import Inches, Pt


REPO_ROOT = Path(__file__).resolve().parents[1]
DOCS_DIR = REPO_ROOT / "docs"
_PROMPT_DIR = REPO_ROOT / "analysis_outputs" / "vt_60hz_prompt_analysis_2026_06_22"
FIGURES_DIR = _PROMPT_DIR / "figures"
TABLES_DIR = _PROMPT_DIR / "tables"
_SLUG = "vt_60hz_decision_pack"
BASE_OUTPUT_DIR = results_dir(_SLUG)
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")


def ensure_dirs() -> None:
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    BASE_OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


def read_csv(name: str) -> pd.DataFrame:
    return pd.read_csv(TABLES_DIR / name)


def load_inputs() -> dict[str, pd.DataFrame]:
    return {
        "focus_summary": read_csv("true60_focus_summary.csv"),
        "axis_summary": read_csv("true60_axis_sensitivity_summary.csv"),
        "axis_hr": read_csv("true60_axis_sensitivity_hr.csv"),
        "threshold_summary": read_csv("failure_stage_threshold_summary.csv"),
        "threshold_tests": read_csv("failure_stage_threshold_tests.csv"),
        "failure_mix": read_csv("vt_true60_failure_mix.csv"),
        "environment": read_csv("vt_true60_environment_vs_rest.csv"),
        "cox_terms": read_csv("lifelines_cox_terms.csv"),
    }


def fmt_days(value: float) -> str:
    return f"{value:.0f} d"


def fmt_pct(value: float) -> str:
    return f"{100.0 * value:.0f}%"


def build_metrics(data: dict[str, pd.DataFrame]) -> dict[str, object]:
    focus = data["focus_summary"]
    axis_summary = data["axis_summary"]
    axis_hr = data["axis_hr"]
    threshold = data["threshold_summary"]
    threshold_tests = data["threshold_tests"]
    failure_mix = data["failure_mix"]
    environment = data["environment"]
    cox_terms = data["cox_terms"]

    def row(df: pd.DataFrame, **filters: object) -> pd.Series:
        mask = pd.Series(True, index=df.index)
        for key, value in filters.items():
            mask &= df[key] == value
        return df.loc[mask].iloc[0]

    vt_base_raw = row(focus, population="Vt", mode="50-55 Hz baseline")
    vt_true_raw = row(focus, population="Vt", mode=">58 Hz true60 proxy")
    vt_true_ttf = row(axis_summary, axis_name="ttf_true_days", population="Vt", mode=">58 Hz true60 proxy")
    vt_base_ttf = row(axis_summary, axis_name="ttf_true_days", population="Vt", mode="50-55 Hz baseline")
    vt_hr_best = row(axis_hr, axis_name="best_available_days", population="Vt")
    vt_hr_true = row(axis_hr, axis_name="ttf_true_days", population="Vt")
    global_hr_true = row(axis_hr, axis_name="ttf_true_days", population="Global")
    vt_trf_hr = row(axis_hr, axis_name="trf_50hz_equiv_days", population="Vt")
    global_trf_hr = row(axis_hr, axis_name="trf_50hz_equiv_days", population="Global")

    thr30_vt = row(threshold, threshold_days=30, population="Vt", mode=">58 Hz true60 proxy")
    thr30_vt_base = row(threshold, threshold_days=30, population="Vt", mode="50-55 Hz baseline")
    thr60_vt = row(threshold, threshold_days=60, population="Vt", mode=">58 Hz true60 proxy")
    thr60_vt_base = row(threshold, threshold_days=60, population="Vt", mode="50-55 Hz baseline")
    thr90_vt = row(threshold, threshold_days=90, population="Vt", mode=">58 Hz true60 proxy")
    thr90_vt_base = row(threshold, threshold_days=90, population="Vt", mode="50-55 Hz baseline")
    thr90_global_test = row(threshold_tests, threshold_days=90, population="Global")
    thr90_vt_test = row(threshold_tests, threshold_days=90, population="Vt")

    h2s_vt = row(environment, metric="H2S effective, mg/L")
    glf_vt = row(environment, metric="GLF")

    vt_high_h2s_term = row(cox_terms, model="vt_true60_vs_baseline_adjusted", term="high_h2s")

    mix_true60 = failure_mix.loc[failure_mix["group"] == ">58 Hz"].copy()
    mix_rest = failure_mix.loc[failure_mix["group"] == "Rest of Vt"].copy()
    mix_merged = mix_true60.merge(mix_rest[["category", "share"]], on="category", how="left", suffixes=("_true60", "_rest"))
    mix_merged["share_diff"] = mix_merged["share_true60"] - mix_merged["share_rest"].fillna(0.0)
    mix_merged = mix_merged.sort_values("share_diff", ascending=False).reset_index(drop=True)
    top_risk_category = mix_merged.iloc[0]
    second_risk_category = mix_merged.iloc[1]

    return {
        "vt_base_raw": vt_base_raw,
        "vt_true_raw": vt_true_raw,
        "vt_base_ttf": vt_base_ttf,
        "vt_true_ttf": vt_true_ttf,
        "vt_hr_best": vt_hr_best,
        "vt_hr_true": vt_hr_true,
        "global_hr_true": global_hr_true,
        "vt_trf_hr": vt_trf_hr,
        "global_trf_hr": global_trf_hr,
        "thr30_vt": thr30_vt,
        "thr30_vt_base": thr30_vt_base,
        "thr60_vt": thr60_vt,
        "thr60_vt_base": thr60_vt_base,
        "thr90_vt": thr90_vt,
        "thr90_vt_base": thr90_vt_base,
        "thr90_global_test": thr90_global_test,
        "thr90_vt_test": thr90_vt_test,
        "h2s_vt": h2s_vt,
        "glf_vt": glf_vt,
        "vt_high_h2s_term": vt_high_h2s_term,
        "top_risk_category": top_risk_category,
        "second_risk_category": second_risk_category,
    }


def build_markdown(metrics: dict[str, object]) -> Path:
    vt_base_raw = metrics["vt_base_raw"]
    vt_true_raw = metrics["vt_true_raw"]
    vt_base_ttf = metrics["vt_base_ttf"]
    vt_true_ttf = metrics["vt_true_ttf"]
    vt_hr_best = metrics["vt_hr_best"]
    vt_hr_true = metrics["vt_hr_true"]
    global_hr_true = metrics["global_hr_true"]
    vt_trf_hr = metrics["vt_trf_hr"]
    global_trf_hr = metrics["global_trf_hr"]
    thr30_vt = metrics["thr30_vt"]
    thr30_vt_base = metrics["thr30_vt_base"]
    thr60_vt = metrics["thr60_vt"]
    thr60_vt_base = metrics["thr60_vt_base"]
    thr90_vt = metrics["thr90_vt"]
    thr90_vt_base = metrics["thr90_vt_base"]
    thr90_global_test = metrics["thr90_global_test"]
    thr90_vt_test = metrics["thr90_vt_test"]
    h2s_vt = metrics["h2s_vt"]
    glf_vt = metrics["glf_vt"]
    vt_high_h2s_term = metrics["vt_high_h2s_term"]
    top_risk_category = metrics["top_risk_category"]
    second_risk_category = metrics["second_risk_category"]

    raw_drop_days = float(vt_base_raw["median_duration_days"] - vt_true_raw["median_duration_days"])
    raw_drop_pct = raw_drop_days / float(vt_base_raw["median_duration_days"]) if float(vt_base_raw["median_duration_days"]) else float("nan")
    ttf_drop_days = float(vt_base_ttf["median_duration"] - vt_true_ttf["median_duration"])
    ttf_drop_pct = ttf_drop_days / float(vt_base_ttf["median_duration"]) if float(vt_base_ttf["median_duration"]) else float("nan")

    lines = [
        "# Vt 60 Hz Decision Summary",
        "",
        "## Executive Answer",
        "",
        "**Short answer:** `60 Hz` should **not** be treated as a default safe operating mode for ESPs in `Vt` if the goal is to maximize pump life.",
        "",
        "The evidence does **not** support a blanket statement that `60 Hz` is universally bad across the whole portfolio. But for `Vt`, the balance of evidence points in one direction: pumps run at `>58 Hz` tend to fail **earlier in calendar time**, and the early-failure burden is higher. The subgroup is still small, so the exact size of the penalty is uncertain, but the direction is consistent enough that `60 Hz` in `Vt` should be treated as a **caution / exception mode**, not a normal default mode.",
        "",
        "## Direct Answer To The Operating Question",
        "",
        "### Is it safe to run ESP at 60 Hz in Vt?",
        "",
        "Not as a default life-maximizing setting.",
        "",
        "Operationally, the best current recommendation is:",
        "",
        "- `Do not ban 60 Hz outright portfolio-wide.`",
        "- `Do not assume 60 Hz is safe-by-default in Vt.`",
        "- `Use 60 Hz in Vt only when there is a clear production reason, and pair it with tighter monitoring.`",
        "",
        "## Will 60 Hz shorten TTF in Vt, and by how much?",
        "",
        "The evidence says **probably yes in calendar time**, but the exact size depends on how TTF is defined and how censoring is handled.",
        "",
        "Best practical range from this study:",
        "",
        f"- Using the raw best-available duration summary, the median in `Vt` moves from `{float(vt_base_raw['median_duration_days']):.1f}` d at `50–55 Hz` to `{float(vt_true_raw['median_duration_days']):.1f}` d at `>58 Hz`: about **{raw_drop_days:.0f} days shorter** ({raw_drop_pct:.0%}).",
        f"- Using the `TTF_true` sensitivity pass, the median moves from `{fmt_days(float(vt_base_ttf['median_duration']))}` to `{fmt_days(float(vt_true_ttf['median_duration']))}`: about **{ttf_drop_days:.0f} days shorter** ({ttf_drop_pct:.0%}).",
        f"- The adjusted hazard ratio inside `Vt` is in the **{float(vt_hr_true['hr_true60']):.2f}–{float(vt_hr_best['hr_true60']):.2f}x** range depending on the time axis, but the Vt-only interval is still wide because the sample is small.",
        "",
        "So the responsible answer is:",
        "",
        f"- `Expected effect in Vt:` likely shorter calendar TTF",
        f"- `Best direct-median estimate:` roughly **30 to 35 days shorter**",
        f"- `Model-based reading:` higher failure hazard in the **1.34–1.54x** range, so the effective life penalty may be larger in some censoring-adjusted views",
        f"- `Confidence in exact magnitude:` moderate-to-low",
        f"- `Confidence in direction:` moderate",
        "",
        "## Does 60 Hz increase early-failure risk in Vt?",
        "",
        "Yes, directionally it does.",
        "",
        f"- At `30` days, early-failure share among Vt failures is `{thr30_vt_base['early_share_failures']:.2f}` at `50–55 Hz` vs `{thr30_vt['early_share_failures']:.2f}` at `>58 Hz`.",
        f"- At `60` days, it is `{thr60_vt_base['early_share_failures']:.2f}` vs `{thr60_vt['early_share_failures']:.2f}`.",
        f"- At `90` days, it is `{thr90_vt_base['early_share_failures']:.2f}` vs `{thr90_vt['early_share_failures']:.2f}`.",
        "",
        f"The `90`-day gap in `Vt` is the most decision-relevant one: **{thr90_vt_test['share_difference_true60_minus_baseline']:.2f} higher early-failure share** at `>58 Hz`. The same pattern also appears globally, where the `90`-day gap is `{thr90_global_test['share_difference_true60_minus_baseline']:.2f}` and statistically clearer.",
        "",
        "## Is the 60 Hz signal in Vt just because those wells were harsher?",
        "",
        "The available evidence says **no, not mainly**.",
        "",
        f"- Effective `H2S` in `Vt >58 Hz` wells is actually **lower**, not higher: median `{float(h2s_vt['true60_median']):.3f}` vs `{float(h2s_vt['rest_median']):.3f}` in the rest of `Vt`.",
        f"- `GLF` is also not higher in the `>58 Hz` Vt subgroup: median `{float(glf_vt['true60_median']):.1f}` vs `{float(glf_vt['rest_median']):.1f}`.",
        "",
        "So the worse `Vt >58 Hz` outcome cannot be dismissed as a simple 'those were just the harshest wells' story based on the measured harshness proxies we have.",
        "",
        "## What equipment is most likely to become problematic at 60 Hz in Vt?",
        "",
        "The first risk signal is **working-organ / clogging-related failure**, and the second is the **motor / electrical package**.",
        "",
        f"- In the `Vt >58 Hz` failure mix, `{top_risk_category['category']}` is the largest single category: `{top_risk_category['share_true60']:.1%}` of failures at `>58 Hz` versus `{top_risk_category['share_rest']:.1%}` in the rest of `Vt`.",
        f"- The next important risk cluster is `{second_risk_category['category']}`: `{second_risk_category['share_true60']:.1%}` at `>58 Hz` versus `{second_risk_category['share_rest']:.1%}` in the rest of `Vt`.",
        "",
        "When we separate early failures from mature failures, the pattern becomes clearer:",
        "",
        "- Earlier `Vt` failures are relatively more concentrated in `Засорение РО` and some fast mechanical/electrical events.",
        "- More mature `Vt` failures lean relatively more toward `ПЭД (R-0)`, `НКТ`, and seal/protection-related categories.",
        "",
        "**Practical interpretation:** if `Vt` wells are pushed to `60 Hz`, the first equipment family to watch is the **working organs / clogging path**. The second family to watch is the **motor / electrical package (ПЭД)**.",
        "",
        "## Is there evidence of a fixed cumulative-work capacity (TRF budget)?",
        "",
        "Only partially.",
        "",
        f"- On the `TTF_true` axis, the global `>58 Hz` penalty is stronger: adjusted HR `{float(global_hr_true['hr_true60']):.2f}`.",
        f"- On the `TRF` axis, the same penalty shrinks materially: adjusted HR `{float(global_trf_hr['hr_true60']):.2f}` globally and `{float(vt_trf_hr['hr_true60']):.2f}` in `Vt`.",
        "",
        "This suggests the pump may **not** simply be wearing out after less total accumulated frequency work. Instead, at higher frequency it appears to get through that usable work capacity faster, so it fails sooner in calendar time.",
        "",
        "But the data do **not** support a single clean universal `TRF` failure threshold, either globally or in `Vt`. If a cumulative-duty budget exists, it is more likely **failure-mechanism-specific** than universal.",
        "",
        "## Stronger Non-Frequency Driver",
        "",
        f"`H2S` remains the stronger and more stable risk signal than `60 Hz` itself. In the adjusted Vt model, high `H2S` has an estimated hazard ratio of `{float(vt_high_h2s_term['exp(coef)']):.2f}`.",
        "",
        "That means the most conservative operational rule is not 'frequency only'. It is:",
        "",
        "- `avoid 60 Hz especially when H2S is high`",
        "- `avoid 60 Hz especially when failure history already points to clogging / working-organ problems`",
        "",
        "## Final Recommendation",
        "",
        "### Recommended operating policy for Vt",
        "",
        "- `Default mode:` stay with the normal `50–55 Hz` regime when life maximization matters.",
        "- `60 Hz mode:` use only when justified by production need, and document it as an exception mode.",
        "- `Monitoring priority at 60 Hz:` watch for working-organ clogging / solids / scale symptoms first, then motor / electrical package stress.",
        "- `Highest-risk combination:` `60 Hz` plus elevated `H2S`.",
        "",
        "### One-sentence decision answer",
        "",
        "**Running ESPs at `60 Hz` in `Vt` is not the safest default choice for TTF. The most likely effect is shorter calendar life, with direct-median estimates around `30–35 days` shorter, and the main additional risk appearing first in working-organ / clogging failures and then in the motor (`ПЭД`) package.**",
        "",
        "## Key Figures",
        "",
        "### Vt TTF_true survival",
        "",
        f"![Vt TTF_true survival](../analysis_outputs/vt_60hz_prompt_analysis_2026_06_22/figures/ttf_true_days_vt_km.png)",
        "",
        "### Early-failure share at 30 / 60 / 90 days",
        "",
        f"![Threshold sensitivity](../analysis_outputs/vt_60hz_prompt_analysis_2026_06_22/figures/true60_threshold_early_share_ttf_true.png)",
        "",
        "### Are >58 Hz Vt wells simply harsher?",
        "",
        f"![Vt environment check](../analysis_outputs/vt_60hz_prompt_analysis_2026_06_22/figures/vt_true60_environment_check.png)",
        "",
        "### Which Vt failures become more common at >58 Hz?",
        "",
        f"![Vt failure mix](../analysis_outputs/vt_60hz_prompt_analysis_2026_06_22/figures/vt_true60_failure_mix.png)",
        "",
    ]
    path = DOCS_DIR / "vt_60hz_vt_decision_summary.md"
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def add_title_slide(prs: Presentation, title: str, subtitle: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(246, 247, 249)
    bg.line.fill.background()
    box = slide.shapes.add_textbox(Inches(0.7), Inches(0.7), Inches(11.6), Inches(1.2))
    p = box.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RGBColor(26, 36, 52)
    sub = slide.shapes.add_textbox(Inches(0.75), Inches(1.8), Inches(11.0), Inches(1.0))
    p2 = sub.text_frame.paragraphs[0]
    p2.text = subtitle
    p2.font.size = Pt(16)
    p2.font.color.rgb = RGBColor(74, 86, 104)


def add_bullets_slide(prs: Presentation, title: str, bullets: list[str], image_path: Path | None = None) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title_box = slide.shapes.add_textbox(Inches(0.45), Inches(0.3), Inches(12.0), Inches(0.6))
    p = title_box.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = RGBColor(26, 36, 52)

    body_width = Inches(5.3) if image_path else Inches(11.9)
    body = slide.shapes.add_textbox(Inches(0.55), Inches(1.0), body_width, Inches(5.9))
    tf = body.text_frame
    tf.word_wrap = True
    for idx, bullet in enumerate(bullets):
        para = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        para.text = bullet
        para.font.size = Pt(17)
        para.font.color.rgb = RGBColor(46, 56, 68)
        para.space_after = Pt(10)
    if image_path is not None and image_path.exists():
        slide.shapes.add_picture(str(image_path), Inches(6.0), Inches(1.0), width=Inches(6.4))


def build_slides(metrics: dict[str, object]) -> Path:
    vt_base_raw = metrics["vt_base_raw"]
    vt_true_raw = metrics["vt_true_raw"]
    vt_base_ttf = metrics["vt_base_ttf"]
    vt_true_ttf = metrics["vt_true_ttf"]
    vt_hr_true = metrics["vt_hr_true"]
    global_hr_true = metrics["global_hr_true"]
    vt_trf_hr = metrics["vt_trf_hr"]
    thr90_vt_test = metrics["thr90_vt_test"]
    vt_high_h2s_term = metrics["vt_high_h2s_term"]
    top_risk_category = metrics["top_risk_category"]
    second_risk_category = metrics["second_risk_category"]

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    add_title_slide(
        prs,
        "Vt 60 Hz Decision Pack",
        "Question: is it safe to run ESPs at 60 Hz in Vt, how much TTF may be lost, and what is most likely to fail first?",
    )

    add_bullets_slide(
        prs,
        "Decision",
        [
            "60 Hz should not be treated as the default life-maximizing mode in Vt.",
            "The direction of the effect is unfavorable: higher early-failure burden and shorter calendar TTF at >58 Hz.",
            "The exact penalty is uncertain because the Vt true60 subgroup is small, but the direct median summaries point to roughly 30 to 35 days shorter life.",
            "Operational policy: use 60 Hz in Vt only as an exception mode with tighter monitoring.",
        ],
        image_path=FIGURES_DIR / "ttf_true_days_vt_km.png",
    )

    add_bullets_slide(
        prs,
        "How Much TTF Is Likely Lost?",
        [
            f"Raw Vt median duration: {float(vt_base_raw['median_duration_days']):.1f} d at 50-55 Hz versus {float(vt_true_raw['median_duration_days']):.1f} d at >58 Hz.",
            f"TTF_true Vt median duration: {float(vt_base_ttf['median_duration']):.0f} d versus {float(vt_true_ttf['median_duration']):.0f} d.",
            f"Adjusted Vt hazard ratio on TTF_true: {float(vt_hr_true['hr_true60']):.2f}x. Global TTF_true signal is clearer at {float(global_hr_true['hr_true60']):.2f}x.",
            "Interpretation: the best reading is not a precise single number, but a directional penalty consistent with shorter calendar life.",
        ],
        image_path=FIGURES_DIR / "true60_axis_sensitivity_hr.png",
    )

    add_bullets_slide(
        prs,
        "Early-Failure Risk",
        [
            "The >58 Hz group carries more early failures, and the gap gets larger as the threshold moves from 30 to 60 to 90 days.",
            f"In Vt, the 90-day early-failure share is higher by {float(thr90_vt_test['share_difference_true60_minus_baseline']):.2f}.",
            "That makes the main operational risk not just lower long-run life, but a greater chance of losing the pump early.",
        ],
        image_path=FIGURES_DIR / "true60_threshold_early_share_ttf_true.png",
    )

    add_bullets_slide(
        prs,
        "Is This Just A Harsher-Well Effect?",
        [
            "No clear evidence of that in Vt.",
            "The >58 Hz Vt subgroup does not show higher H2S or higher GLF than the rest of Vt on the measured proxies we have.",
            "So the worse true60 outcome cannot be explained away as 'those were simply the harshest wells.'",
        ],
        image_path=FIGURES_DIR / "vt_true60_environment_check.png",
    )

    add_bullets_slide(
        prs,
        "What Is Most Likely To Fail?",
        [
            f"The first risk signal is {top_risk_category['category']}: {top_risk_category['share_true60']:.0%} of Vt >58 Hz failures versus {top_risk_category['share_rest']:.0%} in the rest of Vt.",
            f"The second key risk cluster is {second_risk_category['category']}: {second_risk_category['share_true60']:.0%} versus {second_risk_category['share_rest']:.0%}.",
            "Earlier Vt failures lean more toward clogging / working-organ problems, while more mature failures lean relatively more toward the motor package.",
            "Practical watchlist at 60 Hz: working organs first, motor/electrical package second.",
        ],
        image_path=FIGURES_DIR / "vt_true60_failure_mix.png",
    )

    add_bullets_slide(
        prs,
        "TRF Interpretation",
        [
            f"On the TRF axis, the Vt hazard ratio shrinks to {float(vt_trf_hr['hr_true60']):.2f}x.",
            "That means the data do not support a clean universal TRF failure threshold.",
            "The simpler interpretation is: at higher frequency the pump appears to use up its workable life faster in calendar time, not clearly fail after less total accumulated frequency work.",
        ],
        image_path=FIGURES_DIR / "trf_50hz_equiv_days_vt_km.png",
    )

    add_bullets_slide(
        prs,
        "Recommended Operating Rule",
        [
            "Default for Vt: stay in the 50-55 Hz regime when TTF matters.",
            "Use 60 Hz only when there is a strong production reason and the team accepts a likely life penalty.",
            f"Be extra conservative when H2S is high: the adjusted Vt hazard ratio for high-H2S wells is {float(vt_high_h2s_term['exp(coef)']):.2f}x.",
            "If 60 Hz is used, tighten surveillance on clogging / solids / scale indicators and on the motor / electrical package.",
        ],
    )

    path = BASE_OUTPUT_DIR / "vt_60hz_vt_decision_slides.pptx"
    prs.save(path)
    return path


def main() -> None:
    ensure_dirs()
    data = load_inputs()
    metrics = build_metrics(data)
    md_path = build_markdown(metrics)
    slides_path = build_slides(metrics)
    summary = {
        "markdown_path": str(md_path),
        "slides_path": str(slides_path),
    }
    (BASE_OUTPUT_DIR / "vt_60hz_vt_decision_pack.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(f"[decision-pack] Wrote markdown to {md_path}", flush=True)
    print(f"[decision-pack] Wrote slides to {slides_path}", flush=True)


if __name__ == "__main__":
    main()
