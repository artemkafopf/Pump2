from __future__ import annotations

import math
from datetime import date, datetime
from pathlib import Path

import numpy as np
import openpyxl

from analysis.workflows.production_risk import config as C
from analysis.workflows.production_risk.crosswalk import (
    DowntimeStats,
    PlanData,
    load_plan,
    load_current_techregime_status,
)
from analysis.workflows.production_risk.layers import _downtime_days, _impute_age_pmf
from analysis.workflows.production_risk import repair_compat
from analysis.workflows.production_risk.survival import (
    HazardLayer,
    StrataModel,
    WellState,
    current_pump_p_fail,
    current_pump_p_fail_monthly,
    kpod_hazard_theta,
    project_well,
    ql_hazard_theta,
    uncertainty_hazard_theta,
)


def _bisect_b50(params: dict[str, float]) -> float:
    lo, hi = 0.0, 15.0 * max(params["eta1"], params["eta2"])
    for _ in range(200):
        mid = 0.5 * (lo + hi)
        s = (
            params["w1"] * math.exp(-((mid / params["eta1"]) ** params["beta1"]))
            + (1.0 - params["w1"]) * math.exp(-((mid / params["eta2"]) ** params["beta2"]))
        )
        if s > 0.5:
            lo = mid
        else:
            hi = mid
    return 0.5 * (lo + hi)


def _make_plan_workbook(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws_registry = wb.active
    ws_registry.title = "Реестр"
    ws_summary = wb.create_sheet("Сводные данные")

    headers = [
        "Месторождение",
        "УН",
        "Объект сбора",
        "Обьект подготовки",
        "Слой",
        "Дата ввода",
        "Резерв",
        "Тип нефти",
        "Показатель_2",
        "ЕИ",
        "Показатель",
        "Well_ID*",
        "№ куста*",
        datetime(2026, 1, 1),
    ]
    for ws in (ws_registry, ws_summary):
        for idx, value in enumerate(headers, start=1):
            ws.cell(row=3, column=idx, value=value)

    summary_rows = [
        ["Ярактинское НГКМ", "УН-Яракта", "", "", "", "", "", "", "Добыча нефти", "т", "", "Ya_001", "1", 310.0],
        ["Ярактинское НГКМ", "УН-Яракта", "", "", "", "", "", "", "Добыча нефти м3", "м3", "", "Ya_001", "1", 310.0],
        ["Ярактинское НГКМ", "УН-Яракта", "", "", "", "", "", "", "Добыча жидкости", "м3", "", "Ya_001", "1", 620.0],
        ["Ярактинское НГКМ", "УН-Яракта", "", "", "", "", "", "", "Добыча воды (ППД)", "м3", "", "Ya_001", "1", 999.0],
        ["Ярактинское НГКМ", "УН-Яракта", "", "", "", "", "", "", "Отработанное время", "дн", "", "Ya_001", "1", 40.0],
    ]
    registry_rows = [
        ["Ярактинское НГКМ", "УН-Яракта", "", "", "", "", "", "", "Добыча нефти", "т/сут", "", "Ya_001", "1", 10.0],
        ["Ярактинское НГКМ", "УН-Яракта", "", "", "", "", "", "", "Добыча нефти м3", "м3/сут", "", "Ya_001", "1", 10.0],
        ["Ярактинское НГКМ", "УН-Яракта", "", "", "", "", "", "", "Добыча жидкости", "м3/сут", "", "Ya_001", "1", 20.0],
        ["Ярактинское НГКМ", "УН-Яракта", "", "", "", "", "", "", "Добыча воды (ППД)", "м3/сут", "", "Ya_001", "1", 999.0],
    ]
    for row_idx, row in enumerate(summary_rows, start=4):
        for col_idx, value in enumerate(row, start=1):
            ws_summary.cell(row=row_idx, column=col_idx, value=value)
    for row_idx, row in enumerate(registry_rows, start=4):
        for col_idx, value in enumerate(row, start=1):
            ws_registry.cell(row=row_idx, column=col_idx, value=value)
    wb.save(path)


def test_strata_resolution_matches_vba_fallback():
    model = StrataModel()
    _, key_exact = model.resolve("Ya", "nonsour", "brt")
    _, key_pooled = model.resolve("Vt", "sour", "missing")
    _, key_global = model.resolve("Unknown", "nonsour", "brt")
    assert key_exact == "Ya_nonsour_brt"
    assert key_pooled == "Vt_sour_Pooled"
    assert key_global == "Global_Pooled"


def test_b50_and_conditional_probability_from_registry():
    model = StrataModel()
    params, _ = model.resolve("Ya", "nonsour", "brt")
    b50 = _bisect_b50(params)
    assert abs(b50 - params["b50"]) / params["b50"] <= 0.01
    p30 = current_pump_p_fail({100: 1.0}, params, 30.0, model)
    p90 = current_pump_p_fail({100: 1.0}, params, 90.0, model)
    assert 0.0 <= p30 <= p90 <= 1.0


def test_theta_reference_and_missing_inputs_are_neutral():
    hazard = HazardLayer()
    coeffs = hazard.coeffs
    ref = {str(row["covariate"]): float(row["reference_value"]) for _, row in coeffs.iterrows()}
    assert abs(hazard.theta(ref) - 1.0) < 1e-12
    assert abs(hazard.theta({}) - 1.0) < 1e-12


def test_project_well_zero_uptime_zero_losses():
    model = StrataModel()
    params, _ = model.resolve("Ya", "nonsour", "brt")
    well = WellState(
        code="Ya_001",
        raw_id="Ya_001",
        plan_field="Ярактинское НГКМ",
        model_field="Ya",
        stratum="Ya_nonsour_brt",
        params=params,
        age_pmf={100: 1.0},
        age_source="history_running",
        age_mean=100.0,
        state_label="exact_active",
        scope_label="included",
        confidence="high",
        contractor_group="brt",
        contractor_source="exact_history",
        sour_class="nonsour",
        sour_source="exact_history",
        covariates={},
        cov_source="none",
        current_status=None,
        current_in_operation=None,
        current_status_date=None,
        include_primary=True,
        oil_volume=np.array([0.0, 0.0]),
        liquid_volume=np.array([0.0, 0.0]),
        op_days=np.array([0.0, 0.0]),
        cal_days=np.array([31.0, 28.0]),
        first_active_month=None,
        first_gtm_date=None,
        event_90d_flag=False,
    )
    failures, lost = project_well(well, params, 16, model)
    assert np.allclose(failures, 0.0)
    assert np.allclose(lost, 0.0)


def test_project_well_deterministic_outputs_are_bounded():
    model = StrataModel()
    params, _ = model.resolve("Ya", "nonsour", "brt")
    well = WellState(
        code="Ya_002",
        raw_id="Ya_002",
        plan_field="Ярактинское НГКМ",
        model_field="Ya",
        stratum="Ya_nonsour_brt",
        params=params,
        age_pmf={100: 1.0},
        age_source="history_running",
        age_mean=100.0,
        state_label="exact_active",
        scope_label="included",
        confidence="high",
        contractor_group="brt",
        contractor_source="exact_history",
        sour_class="nonsour",
        sour_source="exact_history",
        covariates={},
        cov_source="none",
        current_status=None,
        current_in_operation=None,
        current_status_date=None,
        include_primary=True,
        oil_volume=np.array([310.0]),
        liquid_volume=np.array([620.0]),
        op_days=np.array([31.0]),
        cal_days=np.array([31.0]),
        first_active_month="2026-07",
        first_gtm_date=None,
        event_90d_flag=False,
    )
    failures, lost = project_well(well, params, 16, model)
    assert failures.shape == (1,)
    assert lost.shape == (1,)
    assert 0.0 <= failures[0] <= 31.0
    assert 0.0 <= lost[0] <= 31.0


def test_project_well_ql_hazard_changes_failure_path(monkeypatch):
    monkeypatch.setattr(C, "QL_HAZARD_ENABLED", True)  # dial off by default; test the mechanism
    model = StrataModel()
    params, _ = model.resolve("Ya", "nonsour", "brt")
    well = WellState(
        code="Ya_003",
        raw_id="Ya_003",
        plan_field="Ярактинское НГКМ",
        model_field="Ya",
        stratum="Ya_nonsour_brt",
        params=params,
        age_pmf={100: 1.0},
        age_source="history_running",
        age_mean=100.0,
        state_label="exact_active",
        scope_label="included",
        confidence="high",
        contractor_group="brt",
        contractor_source="exact_history",
        sour_class="nonsour",
        sour_source="exact_history",
        covariates={},
        cov_source="none",
        current_status=None,
        current_in_operation=None,
        current_status_date=None,
        include_primary=True,
        oil_volume=np.array([310.0]),
        liquid_volume=np.array([620.0]),
        op_days=np.array([31.0]),
        cal_days=np.array([31.0]),
        first_active_month="2026-07",
        first_gtm_date=None,
        event_90d_flag=False,
    )
    base, _ = project_well(well, params, 16, model)
    high_ql, _ = project_well(
        well,
        params,
        16,
        model,
        log_ql_monthly=np.array([math.log1p(400.0)]),
        model_field="Ya",
    )
    assert not np.allclose(base, high_ql)


def test_ql_hazard_theta_identities_and_direction(monkeypatch):
    monkeypatch.setattr(C, "QL_HAZARD_ENABLED", True)  # dial off by default; test the mechanism
    field = "Ya"
    ref = C.QL_HAZARD_FIELD_REF_LOG[field]
    ages = np.array([1.0, 100.0, 533.0, 1000.0])

    # at the field reference Ql (z=0) the dial is neutral at every age; NaN is neutral too
    assert np.allclose(ql_hazard_theta(ref, field, ages), 1.0)
    assert ql_hazard_theta(ref, field, 100.0) == 1.0
    assert ql_hazard_theta(float("nan"), field, 100.0) == 1.0
    assert np.allclose(ql_hazard_theta(float("nan"), field, ages), np.ones_like(ages))

    # capped at +/- log(5)
    clipped = ql_hazard_theta(ref + 10.0, field, ages)
    at_cap = ql_hazard_theta(ref + C.QL_HAZARD_CAP_LOG_RATIO, field, ages)
    assert np.allclose(clipped, at_cap)

    # sign-corrected plain PH (gamma=0): higher liquid rate raises hazard at EVERY age,
    # lower liquid rate lowers it, and there is no age crossover (theta constant in age).
    high_ql = ref + math.log(5.0)
    low_ql = ref - math.log(5.0)
    high = ql_hazard_theta(high_ql, field, ages)
    low = ql_hazard_theta(low_ql, field, ages)
    assert np.all(high > 1.0)
    assert np.all(low < 1.0)
    assert np.allclose(high, high[0])  # no age dependence when gamma == 0
    assert math.isclose(ql_hazard_theta(high_ql, field, 100.0),
                        math.exp(C.QL_HAZARD_CAP_LOG_RATIO * C.QL_HAZARD_BETA), rel_tol=1e-9)


def test_kpod_hazard_theta_identities_caps_and_age_interaction():
    ages = np.array([1.0, 100.0, 1000.0])
    old = {
        "enabled": C.KPOD_HAZARD_ENABLED,
        "beta_under": C.KPOD_HAZARD_BETA_UNDER,
        "gamma_under": C.KPOD_HAZARD_GAMMA_UNDER,
        "gamma_over": C.KPOD_HAZARD_GAMMA_OVER,
        "field_params": dict(C.KPOD_HAZARD_FIELD_PARAMS),
    }
    try:
        C.KPOD_HAZARD_ENABLED = True
        C.KPOD_HAZARD_BETA_UNDER = 0.5
        C.KPOD_HAZARD_GAMMA_UNDER = 0.0
        C.KPOD_HAZARD_GAMMA_OVER = 0.02
        C.KPOD_HAZARD_FIELD_PARAMS = {}

        assert np.allclose(kpod_hazard_theta(0.9, ages), 1.0)
        assert kpod_hazard_theta(float("nan"), 100.0) == 1.0
        assert np.all(kpod_hazard_theta(C.KPOD_HAZARD_K_HI + 0.1, ages) > 1.0)
        assert np.all(kpod_hazard_theta(C.KPOD_HAZARD_K_LO - 0.1, ages) > 1.0)

        clipped = kpod_hazard_theta(C.KPOD_HAZARD_K_HI + 100.0, ages)
        at_cap = kpod_hazard_theta(C.KPOD_HAZARD_K_HI + C.KPOD_HAZARD_CAP_OVER, ages)
        assert np.allclose(clipped, at_cap)

        over = kpod_hazard_theta(C.KPOD_HAZARD_K_HI + 0.2, ages)
        assert over[-1] > over[0]

        C.KPOD_HAZARD_ENABLED = False
        assert np.allclose(kpod_hazard_theta(C.KPOD_HAZARD_K_HI + 0.2, ages), 1.0)
    finally:
        C.KPOD_HAZARD_ENABLED = old["enabled"]
        C.KPOD_HAZARD_BETA_UNDER = old["beta_under"]
        C.KPOD_HAZARD_GAMMA_UNDER = old["gamma_under"]
        C.KPOD_HAZARD_GAMMA_OVER = old["gamma_over"]
        C.KPOD_HAZARD_FIELD_PARAMS = old["field_params"]


def test_kpod_hazard_theta_uses_field_specific_overrides():
    old_enabled = C.KPOD_HAZARD_ENABLED
    old = dict(C.KPOD_HAZARD_FIELD_PARAMS)
    try:
        C.KPOD_HAZARD_ENABLED = True
        C.KPOD_HAZARD_FIELD_PARAMS = {
            "Bt": {
                "k_hi": 1.0,
                "beta_over": 2.0,
                "cap_over": 1.0,
            }
        }
        global_theta = kpod_hazard_theta(1.1, 100.0, "Vt")
        bt_theta = kpod_hazard_theta(1.1, 100.0, "Bt")
        assert math.isclose(global_theta, 1.0, rel_tol=1e-12)
        assert bt_theta > 1.0
    finally:
        C.KPOD_HAZARD_ENABLED = old_enabled
        C.KPOD_HAZARD_FIELD_PARAMS = old


def test_uncertainty_hazard_theta_decays_and_is_state_gated():
    old_enabled = C.UNCERTAINTY_HAZARD_ENABLED
    old_apply = set(C.UNCERTAINTY_HAZARD_APPLY_TO)
    old_field = dict(C.UNCERTAINTY_HAZARD_FIELD_PARAMS)
    try:
        C.UNCERTAINTY_HAZARD_ENABLED = True
        C.UNCERTAINTY_HAZARD_APPLY_TO = {"new_plan_only"}
        C.UNCERTAINTY_HAZARD_FIELD_PARAMS = {"Bt": {"u0": 0.5, "tau_days": 10.0, "cap": 1.4}}
        ages = np.array([0.0, 10.0, 100.0])
        theta = uncertainty_hazard_theta(ages, "new_plan_only", "Bt")
        assert math.isclose(theta[0], 1.4, rel_tol=1e-12)  # capped below 1 + u0
        assert theta[1] > theta[2] > 1.0
        assert np.allclose(uncertainty_hazard_theta(ages, "exact_active", "Bt"), 1.0)
        C.UNCERTAINTY_HAZARD_ENABLED = False
        assert np.allclose(uncertainty_hazard_theta(ages, "new_plan_only", "Bt"), 1.0)
    finally:
        C.UNCERTAINTY_HAZARD_ENABLED = old_enabled
        C.UNCERTAINTY_HAZARD_APPLY_TO = old_apply
        C.UNCERTAINTY_HAZARD_FIELD_PARAMS = old_field


def test_ql_hazard_reference_is_wiring_neutral_for_projection_and_90d():
    model = StrataModel()
    params, _ = model.resolve("Ya", "nonsour", "brt")
    well = WellState(
        code="Ya_004",
        raw_id="Ya_004",
        plan_field="Ярактинское НГКМ",
        model_field="Ya",
        stratum="Ya_nonsour_brt",
        params=params,
        age_pmf={100: 0.25, 500: 0.75},
        age_source="imputed_stratum",
        age_mean=400.0,
        state_label="exact_stale_imputed",
        scope_label="included",
        confidence="medium",
        contractor_group="brt",
        contractor_source="exact_history",
        sour_class="nonsour",
        sour_source="exact_history",
        covariates={},
        cov_source="none",
        current_status=None,
        current_in_operation=None,
        current_status_date=None,
        include_primary=True,
        oil_volume=np.array([310.0, 300.0, 320.0]),
        liquid_volume=np.array([620.0, 600.0, 640.0]),
        op_days=np.array([31.0, 28.0, 31.0]),
        cal_days=np.array([31.0, 28.0, 31.0]),
        first_active_month="2026-07",
        first_gtm_date=None,
        event_90d_flag=False,
    )
    ref = C.QL_HAZARD_FIELD_REF_LOG["Ya"]
    ref_vector = np.array([ref, ref, ref], dtype=float)

    base_fail, base_lost = project_well(well, params, 16, model)
    ql_fail, ql_lost = project_well(well, params, 16, model, log_ql_monthly=ref_vector, model_field="Ya")
    assert np.allclose(base_fail, ql_fail, atol=1e-12)
    assert np.allclose(base_lost, ql_lost, atol=1e-12)

    base_90 = current_pump_p_fail(well.age_pmf, params, 90.0, model)
    ql_90 = current_pump_p_fail_monthly(
        well.age_pmf,
        params,
        well.op_days,
        well.cal_days,
        model,
        horizon_calendar_days=90,
        log_ql_monthly=ref_vector,
        model_field="Ya",
    )
    assert abs(base_90 - ql_90) < 1e-6


def test_current_pump_monthly_probability_is_bounded():
    model = StrataModel()
    params, _ = model.resolve("Ya", "nonsour", "brt")
    p_fail = current_pump_p_fail_monthly(
        {100: 0.4, 900: 0.6},
        params,
        np.array([31.0, 28.0, 31.0]),
        np.array([31.0, 28.0, 31.0]),
        model,
        horizon_calendar_days=90,
        log_ql_monthly=np.array([C.QL_HAZARD_FIELD_REF_LOG["Ya"] + math.log(5.0)] * 3),
        model_field="Ya",
    )
    assert 0.0 <= p_fail <= 1.0


def test_project_well_kpod_hazard_changes_failure_path():
    model = StrataModel()
    params, _ = model.resolve("Ya", "nonsour", "brt")
    well = WellState(
        code="Ya_005",
        raw_id="Ya_005",
        plan_field="Ярактинское НГКМ",
        model_field="Ya",
        stratum="Ya_nonsour_brt",
        params=params,
        age_pmf={100: 1.0},
        age_source="history_running",
        age_mean=100.0,
        state_label="exact_active",
        scope_label="included",
        confidence="high",
        contractor_group="brt",
        contractor_source="exact_history",
        sour_class="nonsour",
        sour_source="exact_history",
        covariates={},
        cov_source="none",
        current_status=None,
        current_in_operation=None,
        current_status_date=None,
        include_primary=True,
        oil_volume=np.array([310.0]),
        liquid_volume=np.array([620.0]),
        op_days=np.array([31.0]),
        cal_days=np.array([31.0]),
        first_active_month="2026-07",
        first_gtm_date=None,
        event_90d_flag=False,
    )
    base, _ = project_well(well, params, 16, model)
    overloaded, _ = project_well(well, params, 16, model, kpod_monthly=np.array([2.0]))
    assert overloaded[0] > base[0]


def test_project_well_uncertainty_hazard_changes_only_new_launch_path(monkeypatch):
    monkeypatch.setattr(C, "UNCERTAINTY_HAZARD_ENABLED", True)  # dial off by default; test the mechanism
    model = StrataModel()
    params, _ = model.resolve("Ya", "nonsour", "brt")
    base_kwargs = dict(
        code="Ya_006",
        raw_id="Ya_006",
        plan_field="Ярактинское НГКМ",
        model_field="Ya",
        stratum="Ya_nonsour_brt",
        params=params,
        age_pmf={0: 1.0},
        age_source="new",
        age_mean=0.0,
        scope_label="included",
        confidence="medium",
        contractor_group="brt",
        contractor_source="planned",
        sour_class="nonsour",
        sour_source="planned",
        covariates={},
        cov_source="none",
        current_status=None,
        current_in_operation=None,
        current_status_date=None,
        include_primary=True,
        oil_volume=np.array([310.0]),
        liquid_volume=np.array([620.0]),
        op_days=np.array([31.0]),
        cal_days=np.array([31.0]),
        first_active_month="2026-07",
        first_gtm_date=None,
        event_90d_flag=False,
    )
    new_well = WellState(state_label="new_plan_only", **base_kwargs)
    exact_well = WellState(state_label="exact_active", **base_kwargs)

    base, _ = project_well(new_well, params, 16, model)
    uncertain, _ = project_well(
        new_well,
        params,
        16,
        model,
        uncertainty_state_label="new_plan_only",
        uncertainty_field="Ya",
    )
    exact_base, _ = project_well(exact_well, params, 16, model)
    exact_uncertain, _ = project_well(
        exact_well,
        params,
        16,
        model,
        uncertainty_state_label="exact_active",
        uncertainty_field="Ya",
    )
    assert uncertain[0] > base[0]
    assert np.allclose(exact_uncertain, exact_base)


def test_historical_ql_replay_uses_daily_age_varying_theta():
    import pandas as pd
    from analysis.workflows.production_risk import failure_rate as fr

    model = StrataModel()
    params, _ = model.resolve("Ya", "nonsour", "brt")
    params = dict(params)
    params["uptime_factor"] = 1.0
    month = "2026-01"

    class _Plan:
        months = [month]
        liquid_volume = pd.DataFrame([[400.0 * 31.0]], index=["YA_001"], columns=[month])
        op_days = pd.DataFrame([[31.0]], index=["YA_001"], columns=[month])
        op_days_raw = op_days
        cal_days = pd.Series({month: 31.0})

    rows: list[dict] = []
    fr._append_interval_predictions(
        rows,
        plan=_Plan(),
        active=pd.DataFrame(True, index=["YA_001"], columns=[month]),
        model=model,
        field="УН-Яракта",
        code="YA_001",
        months=[month],
        start=datetime(2026, 1, 1),
        end=datetime(2026, 2, 1),
        params=params,
        total_op=None,
        use_ql_hazard=True,
    )

    q_arr = model.daily_fail_prob_array(params, 40)
    survived = 1.0
    for age in range(31):
        theta = ql_hazard_theta(math.log1p(400.0), "Ya", float(age))
        q_eff = 1.0 - (1.0 - float(q_arr[age])) ** theta
        survived *= 1.0 - q_eff
    expected = 1.0 - survived

    assert len(rows) == 1
    assert abs(rows[0]["predicted_failures"] - expected) < 1e-12


def test_historical_kpod_replay_uses_raw_ql_over_qnom_daily():
    import pandas as pd
    from analysis.workflows.production_risk import failure_rate as fr

    model = StrataModel()
    params, _ = model.resolve("Ya", "nonsour", "brt")
    params = dict(params)
    params["uptime_factor"] = 1.0
    month = "2026-01"

    class _Plan:
        months = [month]
        liquid_volume = pd.DataFrame([[200.0 * 31.0]], index=["YA_001"], columns=[month])
        op_days = pd.DataFrame([[31.0]], index=["YA_001"], columns=[month])
        op_days_raw = op_days
        cal_days = pd.Series({month: 31.0})

    rows: list[dict] = []
    fr._append_interval_predictions(
        rows,
        plan=_Plan(),
        active=pd.DataFrame(True, index=["YA_001"], columns=[month]),
        model=model,
        field="УН-Яракта",
        code="YA_001",
        months=[month],
        start=datetime(2026, 1, 1),
        end=datetime(2026, 2, 1),
        params=params,
        total_op=None,
        use_kpod_hazard=True,
        qnominal=100.0,
    )

    q_arr = model.daily_fail_prob_array(params, 40)
    survived = 1.0
    for age in range(31):
        theta = kpod_hazard_theta(2.0, float(age))
        q_eff = 1.0 - (1.0 - float(q_arr[age])) ** theta
        survived *= 1.0 - q_eff
    expected = 1.0 - survived

    assert len(rows) == 1
    assert abs(rows[0]["predicted_failures"] - expected) < 1e-12


def test_age_imputation_prefers_stratum_then_field_then_global():
    pmf, source, conf = _impute_age_pmf("Ya_nonsour_brt", "Ya", {"Ya_nonsour_brt": [10, 20]}, {"Ya": [30]}, [40])
    assert source == "imputed_stratum"
    assert conf == "medium"
    assert abs(sum(pmf.values()) - 1.0) < 1e-12

    _, source, _ = _impute_age_pmf("Za_nonsour_brt", "Ya", {}, {"Ya": [30]}, [40])
    assert source == "imputed_field"

    _, source, conf = _impute_age_pmf("Za_nonsour_brt", "Za", {}, {}, [40])
    assert source == "imputed_global"
    assert conf == "low"


def test_downtime_days_use_field_when_n_is_enough_else_global():
    global_stats = DowntimeStats(n=100, p25=7, p50=16, p75=44, mean=22)
    field_stats = {
        "Ya": DowntimeStats(n=31, p25=8, p50=17, p75=47, mean=24),
        "Mc": DowntimeStats(n=10, p25=4, p50=10, p75=20, mean=12),
    }
    assert _downtime_days("Ya", global_stats, field_stats, "p50") == 17
    assert _downtime_days("Mc", global_stats, field_stats, "p50") == 16
    assert _downtime_days("None", global_stats, field_stats, "p75") == 44


def test_load_plan_uses_summary_sheet_and_clips_runtime(tmp_path: Path):
    path = tmp_path / "plan.xlsx"
    _make_plan_workbook(path)
    plan = load_plan(date(2026, 1, 1), date(2026, 1, 1), master_path=path)
    assert plan.oil_volume.at["YA_001", "2026-01"] == 310.0
    assert plan.oil_volume_m3.at["YA_001", "2026-01"] == 310.0
    assert abs(plan.produced_water_rate.at["YA_001", "2026-01"] - 10.0) < 1e-12
    assert plan.registry_oil_rate.at["YA_001", "2026-01"] == 10.0
    assert plan.registry_water_rate.at["YA_001", "2026-01"] == 0.0
    assert plan.op_days.at["YA_001", "2026-01"] == 31.0
    assert plan.producer_meta.at["YA_001", "license_area"] == "УН-Яракта"
    assert not plan.anomalies.empty
    assert "runtime_gt_calendar" in set(plan.anomalies["reason"])


def test_load_current_techregime_status_marks_off_today(tmp_path: Path):
    path = tmp_path / "tr.xlsm"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ТР НЕФТЬ"
    ws.cell(row=3, column=4, value=datetime(2026, 7, 1))
    headers = {
        4: "ID скважины",
        46: "Способ эксплуатации",
        48: "Время наработки (ННО)",
        100: "Состояние по фонду",
    }
    for col, name in headers.items():
        ws.cell(row=8, column=col, value=name)
    ws.cell(row=9, column=4, value="Ya_001")
    ws.cell(row=9, column=46, value="ЭЦН")
    ws.cell(row=9, column=48, value=123)
    ws.cell(row=9, column=100, value="В бездействии прошлых лет")
    ws.cell(row=10, column=4, value="Ya_002")
    ws.cell(row=10, column=46, value="ЭЦН")
    ws.cell(row=10, column=48, value=45)
    ws.cell(row=10, column=100, value="В работе")
    wb.save(path)

    statuses = load_current_techregime_status(path)
    assert statuses["YA_001"].in_operation is False
    assert statuses["YA_001"].age_op == 123.0
    assert statuses["YA_002"].in_operation is True


def test_repair_compat_preserves_legacy_status_shape(tmp_path: Path):
    model = StrataModel()
    params, _ = model.resolve("Ya", "nonsour", "brt")
    state = WellState(
        code="YA_001",
        raw_id="Ya_001",
        plan_field="Ярактинское НГКМ",
        model_field="Ya",
        stratum="Ya_nonsour_brt",
        params=params,
        age_pmf={100: 1.0},
        age_source="history_running",
        age_mean=100.0,
        state_label="exact_active",
        scope_label="included",
        confidence="high",
        contractor_group="brt",
        contractor_source="exact_history",
        sour_class="nonsour",
        sour_source="exact_history",
        covariates={},
        cov_source="none",
        current_status="В работе",
        current_in_operation=True,
        current_status_date="2026-07-01",
        include_primary=True,
        oil_volume=np.array([310.0]),
        liquid_volume=np.array([620.0]),
        op_days=np.array([31.0]),
        cal_days=np.array([31.0]),
        first_active_month="2026-07",
        first_gtm_date=None,
        event_90d_flag=False,
    )
    import pandas as pd

    df = pd.DataFrame(
        [
            {
                "scenario": C.PRIMARY_SCENARIO_ID,
                "wid": "YA_001",
                "raw_id": "Ya_001",
                "plan_field": "Ярактинское НГКМ",
                "month": "2026-07",
                "planned_oil_t": 310.0,
                "planned_op_days": 31.0,
                "expected_failures": 1.2,
                "downtime_days": 3,
            }
        ]
    )
    cfg = C.RunConfig(
        forecast_start=date(2026, 7, 1),
        horizon_end=date(2026, 7, 1),
        fact_through_month="2026-06",
        write_excel=False,
    )
    payload = repair_compat.build(df, [state], cfg)
    assert payload["rows"][0]["used_prediction_source"] == "survival"
    assert payload["rows"][0]["catboost_nno"] is None
    assert len(payload["rows"][0]["statuses"]) == len(payload["dates"])
    # Failures-only view: one status==0 per allocated failure event (E=1.2 -> 1 event),
    # NOT a painted downtime block.
    row0 = payload["rows"][0]
    assert row0["statuses"].count(0) == len(row0["event_dates"]) == 1
    # each 0 sits exactly on a failure event date
    for iso in row0["event_dates"]:
        idx = payload["dates"].index(iso)
        assert row0["statuses"][idx] == 0
    assert row0["downtime_days"] == 3

    # «Вероятностный прогноз» = наработка сегодня + расстояние до подъёма, т.е. та же
    # величина, что «Факт ННО» (полная наработка), а не остаток.  Легаси-потребитель
    # (app/services/repair_forecast.py) сравнивает `actual_nno >= catboost_nno` —
    # это осмысленно только для полной наработки.  Прогноз обязан быть >= факта.
    assert row0["probabilistic_nno"] == row0["actual_nno"] + row0["days_to_workover"]
    assert row0["probabilistic_nno"] >= row0["actual_nno"]

    path = repair_compat.write_excel(tmp_path / "compat.xlsx", payload)
    assert path.exists()
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        assert "Repair forecast" in wb.sheetnames
        assert "Summary" in wb.sheetnames
    finally:
        wb.close()


def test_event_allocation_preserves_fleet_totals_and_month_profile():
    """Fleet event count must equal round(sum expected); months must not collapse
    into waves (the old floor-walk produced 0-event and 146-event months)."""
    import pandas as pd

    months = [f"2026-{m:02d}" for m in range(7, 13)] + [f"2027-{m:02d}" for m in range(1, 13)]
    rows = []
    n_wells = 200
    per_month = 0.06  # E_w = 1.08 per well; fleet expected = 216
    for i in range(n_wells):
        for month in months:
            rows.append(
                {
                    "scenario": C.PRIMARY_SCENARIO_ID,
                    "wid": f"YA_{i:03d}",
                    "raw_id": f"Ya_{i:03d}",
                    "plan_field": "Ярактинское НГКМ",
                    "month": month,
                    "planned_oil_t": 300.0,
                    "planned_op_days": 30.0,
                    "expected_failures": per_month,
                    "downtime_days": 1,
                }
            )
    df = pd.DataFrame(rows)
    cfg = C.RunConfig(forecast_start=date(2026, 7, 1), horizon_end=date(2027, 12, 1), write_excel=False)
    payload = repair_compat.build(df, [], cfg)

    total_expected = per_month * len(months) * n_wells
    total_allocated = sum(len(r["event_dates"]) for r in payload["rows"])
    assert total_allocated == round(total_expected)

    by_month: dict[str, int] = {}
    for row in payload["rows"]:
        for iso in row["event_dates"]:
            key = iso[:7]
            by_month[key] = by_month.get(key, 0) + 1
    expected_monthly = total_expected / len(months)  # 12 per month
    assert set(by_month) <= set(months)
    for month in months:
        got = by_month.get(month, 0)
        # uniform profile -> every month populated, no waves.  The old floor-walk
        # produced months with 0 and months with 146 events; hash-level fluctuation
        # around the mean (12) is expected and fine.
        assert got > 0, f"month {month} got zero events"
        assert got <= 3.0 * expected_monthly, (month, got)


def test_off_today_wells_are_not_painted_down_in_failures_only_view():
    """Failures-only view: an off-today well with no predicted failures shows all 1s;
    the old repair-in-progress downtime block is intentionally gone."""
    import pandas as pd

    model = StrataModel()
    params, _ = model.resolve("Ya", "nonsour", "brt")
    state = WellState(
        code="YA_001",
        raw_id="Ya_001",
        plan_field="Ярактинское НГКМ",
        model_field="Ya",
        stratum="Ya_nonsour_brt",
        params=params,
        age_pmf={100: 1.0},
        age_source="history_running",
        age_mean=100.0,
        state_label="exact_active_off_today",
        scope_label="included",
        confidence="high",
        contractor_group="brt",
        contractor_source="exact_history",
        sour_class="nonsour",
        sour_source="exact_history",
        covariates={},
        cov_source="none",
        current_status="Остановлена",
        current_in_operation=False,
        current_status_date="2026-07-01",
        include_primary=True,
        oil_volume=np.array([310.0, 300.0]),
        liquid_volume=np.array([620.0, 600.0]),
        op_days=np.array([31.0, 31.0]),
        cal_days=np.array([31.0, 31.0]),
        first_active_month="2026-07",
        first_gtm_date=None,
        event_90d_flag=False,
    )
    df = pd.DataFrame(
        [
            {
                "scenario": C.PRIMARY_SCENARIO_ID,
                "wid": "YA_001",
                "raw_id": "Ya_001",
                "plan_field": "Ярактинское НГКМ",
                "month": "2026-07",
                "planned_oil_t": 310.0,
                "planned_op_days": 31.0,
                "expected_failures": 0.0,
                "downtime_days": 16,
            },
            {
                "scenario": C.PRIMARY_SCENARIO_ID,
                "wid": "YA_001",
                "raw_id": "Ya_001",
                "plan_field": "Ярактинское НГКМ",
                "month": "2026-08",
                "planned_oil_t": 300.0,
                "planned_op_days": 31.0,
                "expected_failures": 0.0,
                "downtime_days": 16,
            },
        ]
    )
    cfg = C.RunConfig(forecast_start=date(2026, 7, 1), horizon_end=date(2026, 8, 1), write_excel=False)
    payload = repair_compat.build(df, [state], cfg)
    statuses = payload["rows"][0]["statuses"]
    assert set(statuses) == {1}
    assert payload["rows"][0]["event_dates"] == []


def test_techregime_age_used_for_stale_and_new_wells():
    from analysis.workflows.production_risk.crosswalk import CurrentTechregimeStatus
    from analysis.workflows.production_risk.layers import _techregime_age

    class _PlanStub:
        op_days = __import__("pandas").DataFrame()
        cal_days = __import__("pandas").Series(dtype=float)

    tr = CurrentTechregimeStatus(
        well_code="YA_001",
        status="В работе",
        in_operation=True,
        lift="ЭЦН",
        age_op=234.0,
        report_date=datetime(2026, 7, 1),
    )
    assert _techregime_age(tr, _PlanStub(), "YA_001", date(2026, 7, 1)) == 234
    assert _techregime_age(None, _PlanStub(), "YA_001", date(2026, 7, 1)) is None
    tr_no_age = CurrentTechregimeStatus(
        well_code="YA_002", status="В работе", in_operation=True, lift="ЭЦН", age_op=None, report_date=None
    )
    assert _techregime_age(tr_no_age, _PlanStub(), "YA_002", date(2026, 7, 1)) is None


def test_model_prefix_mapping_and_explicit_global_fallback():
    from analysis.workflows.production_risk import crosswalk

    # MR (Мирнинский Мр pad) is the one non-legacy prefix we map — it is the same
    # field/stratum as Mc and the driver of the refit.
    assert crosswalk.map_model_field_from_well("MR_1001") == "Mc"
    # Established legacy mappings still resolve.
    assert crosswalk.map_model_field_from_well("VT_1001") == "Vt"
    assert crosswalk.map_model_field_from_well("AU_1001") == "Za"
    # Minor / unclear prefixes are intentionally pooled, not force-mapped.
    for code in ("NE_1001", "AM_1001", "YAY_1", "ZYI_1", "KI_1", "BT_1", "MSH_1"):
        assert crosswalk.map_model_field_from_well(code) is None
        assert crosswalk.is_explicit_global_fallback(code)


def test_resolve_qnominal_prefers_installed_then_field_typical(monkeypatch):
    from analysis.workflows.production_risk import crosswalk

    def _fake_tables(path_key: str = ""):
        return {"YA_001": 160.0}, {"Ya": 100.0, "Vt": 220.0}

    monkeypatch.setattr(crosswalk, "_qnominal_tables", _fake_tables)
    assert crosswalk.resolve_qnominal("YA_001", "Ya") == 160.0
    assert crosswalk.resolve_qnominal("YA_999", "Ya") == 100.0
    assert crosswalk.resolve_qnominal("UNKNOWN_1", None) is None


def test_strata_model_exposes_uptime_factor(tmp_path: Path):
    from analysis.workflows.production_risk.survival import StrataModel

    registry = tmp_path / "esp_models.csv"
    registry.write_text(
        "stratum,field,h2s_class,contractor_group,model_kind,w1,beta1,eta1,beta2,eta2,b20,b50,b80,b50_lo,b50_hi,uptime_factor,n_runs,n_failures,pct_mixed_clock,clock,fit_date\n"
        "Mc_nonsour_Pooled,Mc,nonsour,Pooled,k2,0.1,0.8,30,1.4,360,50,220,500,190,280,0.74,10,8,0,ttf_mix,2026-07-13\n"
        "Global_Pooled,,,,k1,0,1,100,1,100,20,70,150,,,1,10,5,0,ttf_mix,2026-07-13\n",
        encoding="utf-8-sig",
    )
    params, stratum = StrataModel(registry_path=registry).resolve("Mc", "nonsour", "Pooled")
    assert stratum == "Mc_nonsour_Pooled"
    assert params["uptime_factor"] == 0.74


def _fr_fixtures():
    """Minimal PlanData/EspSource/projection for failure-rate tests."""
    import pandas as pd
    from analysis.workflows.production_risk import crosswalk

    months = ["2026-05", "2026-06", "2026-07"]
    wids = ["YA_001", "YA_002", "YA_003"]
    op_raw = pd.DataFrame(30.0, index=pd.Index(wids, name="wid"), columns=months)
    oil = pd.DataFrame(
        [[100.0, 100.0, 100.0], [90.0, 90.0, 90.0], [0.0, 0.0, 0.0]],
        index=pd.Index(wids, name="wid"),
        columns=months,
    )
    liq = pd.DataFrame(
        [[200.0, 200.0, 200.0], [180.0, 180.0, 180.0], [0.0, 0.0, 0.0]],
        index=pd.Index(wids, name="wid"),
        columns=months,
    )
    meta = pd.DataFrame(
        {
            "plan_field": ["Ярактинское НГКМ", "Ярактинское НГКМ", "Ярактинское НГКМ"],
            "license_area": ["УН-Яракта", "УН-Яракта", "УН-Яракта"],
        },
        index=pd.Index(wids, name="wid"),
    )

    class _Plan:
        pass

    plan = _Plan()
    plan.months = months
    plan.op_days_raw = op_raw
    plan.oil_volume = oil
    plan.liquid_volume = liq
    plan.producer_meta = meta

    def _run(code, mount, stop, ff, age):
        return crosswalk.EspRun(
            well_code=code, field_raw="Ya", ctr_raw="Борец", sour_raw="Некислый",
            run_seq=1, mount=mount, stop=stop, demo=None, age_op=age,
            failure_flag=ff, glf=None, load_mean=None, curvature=None,
        )

    class _Src:
        pass

    src = _Src()
    src.source_cutoff = datetime(2026, 6, 4)
    src.runs_by_well = {
        "YA_001": [_run("YA_001", datetime(2024, 6, 1), datetime(2026, 5, 20), 1, 500.0)],
        "YA_002": [_run("YA_002", datetime(2025, 1, 1), None, 0, 400.0)],
    }

    projection = pd.DataFrame(
        [
            {"scenario": C.PRIMARY_SCENARIO_ID, "wid": w, "plan_field": "Ярактинское НГКМ",
             "month": "2026-07", "expected_failures": ef}
            for w, ef in [("YA_001", 0.10), ("YA_002", 0.05)]
        ]
    )
    return plan, src, projection


def test_failure_rate_fleet_observed_and_forecast(monkeypatch):
    from analysis.workflows.production_risk import failure_rate as fr

    monkeypatch.setattr(fr, "_bundle_observed_failures", lambda *args, **kwargs: None)
    plan, src, projection = _fr_fixtures()
    cfg = C.RunConfig(
        forecast_start=date(2026, 7, 1),
        horizon_end=date(2026, 7, 1),
        fact_through_month="2026-06",
        write_excel=False,
    )
    res = fr.compute(plan, src, projection, cfg)

    g = res.monthly[res.monthly["field"] == fr.GLOBAL_LABEL].set_index("month")
    # fleet size = active producing wells: non-zero op time and oil/liquid > 0.
    # YA_003 has runtime but zero production, so it must not inflate the denominator.
    assert g.at["2026-05", "fleet_size"] == 2
    # one actual failure (Failure Flag=1, stop in 2026-05)
    assert g.at["2026-05", "observed_failures"] == 1.0
    assert abs(g.at["2026-05", "observed_rate"] - 0.5) < 1e-9
    # forecast month uses the sanctioned projection expected_failures; observed is blanked.
    # YA wells carry the model-field calibration.  The tiny fixture uses a synthetic
    # reporting УН name, so no production reporting-field factor applies.
    expected_ya = (
        0.15
        * fr._CALIBRATION_FACTORS.get("Ya", 1.0)
        * fr._REPORTING_FIELD_CALIBRATION_FACTORS.get("УН-Яракта", 1.0)
    )
    assert abs(g.at["2026-07", "predicted_failures"] - expected_ya) < 1e-9
    assert math.isnan(g.at["2026-07", "observed_rate"])
    # history months carry a Weibull prediction driven by observed install intervals
    assert g.loc[["2026-05", "2026-06"], "predicted_failures"].sum() >= 0.0
    # The still-running YA_002 interval contributes before forecast start; stopped
    # YA_001 is not synthetically renewed unless an observed Big/Svod interval exists.
    assert g.at["2026-06", "predicted_failures"] > 0.0
    assert fr.GLOBAL_LABEL in res.fields
    assert "УН-Яракта" in set(res.monthly["field"])
    assert "global_pooled_share_by_field" in res.coverage
    assert res.coverage["fleet_denominator"] == "interval_before_plan_then_plan_active_producing"


def test_failure_rate_sheets_written(tmp_path: Path):
    from openpyxl import Workbook
    from analysis.workflows.production_risk import failure_rate as fr

    plan, src, projection = _fr_fixtures()
    cfg = C.RunConfig(
        forecast_start=date(2026, 7, 1),
        horizon_end=date(2026, 7, 1),
        fact_through_month="2026-06",
        write_excel=False,
    )
    res = fr.compute(plan, src, projection, cfg)

    wb = Workbook()
    fr.write_sheets(wb, res)
    assert fr.RATE_SHEET in wb.sheetnames
    assert fr.DATA_SHEET in wb.sheetnames
    path = tmp_path / "fr.xlsx"
    wb.save(path)
    assert path.exists()
    # charts embedded only for material fields; the audit sheet still includes all fields.
    ws = wb[fr.RATE_SHEET]
    assert len(ws._charts) == len(res.chart_fields)
    assert fr.GLOBAL_LABEL in res.chart_fields


def test_lazy_analysis_import_keeps_scipy_out():
    """Guards the frozen-EXE size: importing the production-risk chain must not drag
    in the heavy scientific stack via analysis/__init__."""
    import subprocess
    import sys as _sys

    code = (
        "import sys; sys.path.insert(0, r'" + str(Path(__file__).resolve().parents[1]) + "'); "
        "import analysis; import analysis.workflows.production_risk.survival; "
        "assert 'scipy' not in sys.modules, 'scipy leaked into the lazy import chain'; "
        "from analysis import weibull_survival; import scipy; print('ok')"
    )
    proc = subprocess.run([_sys.executable, "-c", code], capture_output=True, text=True)
    assert proc.returncode == 0, proc.stderr
    assert "ok" in proc.stdout


def test_repair_compat_spreads_same_month_events_by_well():
    import pandas as pd

    df = pd.DataFrame(
        [
            {
                "scenario": C.PRIMARY_SCENARIO_ID,
                "wid": "YA_001",
                "raw_id": "Ya_001",
                "plan_field": "Ярактинское НГКМ",
                "month": "2026-07",
                "planned_oil_t": 310.0,
                "planned_op_days": 31.0,
                "expected_failures": 1.0,
                "downtime_days": 1,
            },
            {
                "scenario": C.PRIMARY_SCENARIO_ID,
                "wid": "YA_002",
                "raw_id": "Ya_002",
                "plan_field": "Ярактинское НГКМ",
                "month": "2026-07",
                "planned_oil_t": 310.0,
                "planned_op_days": 31.0,
                "expected_failures": 1.0,
                "downtime_days": 1,
            },
        ]
    )
    cfg = C.RunConfig(forecast_start=date(2026, 7, 1), horizon_end=date(2026, 7, 1), write_excel=False)
    payload = repair_compat.build(df, [], cfg)
    dates = [row["event_dates"][0] for row in payload["rows"]]
    assert len(set(dates)) == len(dates)


def test_drop_stale_open_runs_kills_contradicted_open_rows():
    """An open run is stale when the same well has a newer install (VT_2704 pattern:
    brine-bore runs folded onto the oil code) or a closed run ending after the open
    run began (YA_601 pattern: Big open row never closed).  A fresh open run after a
    same-week pump swap must survive."""
    import pandas as pd

    from analysis.workflows.production_risk.esp_population import drop_stale_open_runs

    pop = pd.DataFrame(
        [
            # YA_601 pattern: open row shadowed by a later closed run.
            {"code": "A_1", "install": pd.Timestamp("2016-06-24"), "end": pd.NaT},
            {"code": "A_1", "install": pd.Timestamp("2016-01-01"), "end": pd.Timestamp("2017-09-02")},
            # VT_2704 pattern: open row shadowed by a later INSTALL (still open or not).
            {"code": "B_2", "install": pd.Timestamp("2023-02-13"), "end": pd.NaT},
            {"code": "B_2", "install": pd.Timestamp("2023-11-06"), "end": pd.Timestamp("2024-05-06")},
            # Healthy: closed run, then a swap within tolerance, current run open.
            {"code": "C_3", "install": pd.Timestamp("2024-01-01"), "end": pd.Timestamp("2025-05-01")},
            {"code": "C_3", "install": pd.Timestamp("2025-05-03"), "end": pd.NaT},
        ]
    )
    out = drop_stale_open_runs(pop)
    open_codes = set(out.loc[out["end"].isna(), "code"])
    assert open_codes == {"C_3"}
    # Closed history is never touched.
    assert (out["end"].notna().sum()) == 3
