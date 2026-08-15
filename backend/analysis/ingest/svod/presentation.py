"""Оформление готового регистра: пометка справочных колонок и лист-справка.

Заказчик просил не выбрасывать телеметрические средние из Свода — пусть
смотрящий видит режим по скважине, — но и не давать принять их за вход расчёта.

⚠⚠ Пометка сделана **оформлением и легендой, а не переименованием**. Шесть
модулей Pump2 читают эти заголовки из регистра дословно
(``production_risk.freq_bins.SVOD_COVARIATES``, ``svod_nno_decomposition``,
``svod_ttf_by_design``, ``svod_ttf_vs_ql``, ``esp_survival.data``,
``production_risk.crosswalk``), поэтому префикс вроде ``СПР.`` не защитил бы их,
а молча обнулил бы им ковариаты.

Отсюда два механизма:

* в листе «Свод» заголовки справочного блока получают заливку и примечание;
* лист «Справка о колонках» перечисляет КАЖДУЮ колонку с её ролью и источником —
  это машиночитаемо, в отличие от заливки.
"""

from __future__ import annotations

from typing import Optional

from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook

from ..config import (
    BIG_COLUMN_RULES,
    LAB_CHEMISTRY_COLUMNS,
    OPZ_COLUMNS,
    SVOD_CAUSE_COLUMNS,
    SVOD_COLUMN_ROLES,
    SVOD_COMMENT_COLUMN,
    SVOD_FRAC_COLUMNS,
    SVOD_LEGEND_SHEET_NAME,
    SVOD_REFERENCE_COLUMNS,
    SVOD_REFERENCE_NOTE,
    TECHREGIME_COLUMN_RULES,
)

#: Amber fill on the reference block's header cells.
REFERENCE_HEADER_FILL = PatternFill(fill_type="solid", start_color="FFFFE9B0", end_color="FFFFE9B0")
#: Green fill on the cause-taxonomy header cells (new, derived here).
DERIVED_HEADER_FILL = PatternFill(fill_type="solid", start_color="FFD9EAD3", end_color="FFD9EAD3")

#: Where each column's values come from, for the legend sheet.
COLUMN_SOURCES = {
    **{name: "ПДК / паспорт оборудования" for name in SVOD_COLUMN_ROLES["идентификация"]},
    **{name: "выведено сборщиком из ПДК" for name in SVOD_CAUSE_COLUMNS},
    **{name: "WellsArtificialLiftBig (паспорт)" for name in BIG_COLUMN_RULES},
    **{name: "телеметрия / ТехРежим" for name in TECHREGIME_COLUMN_RULES},
    **{name: "ОПЗ" for name in OPZ_COLUMNS.values()},
    **{name: "Свод ГРП" for name in SVOD_FRAC_COLUMNS},
    **{name: "лабораторная химия" for name in LAB_CHEMISTRY_COLUMNS},
    SVOD_COMMENT_COLUMN: "провенанс сборщика",
    "ГЖФ": "телеметрия / ТехРежим (расчёт)",
    "Дельта Дебита Ж и номинала, м3/сут": "расчёт: Дебит жидк. − Ном. Произв.",
}

_ROLE_BY_COLUMN = {
    column: role
    for role, columns in SVOD_COLUMN_ROLES.items()
    for column in columns
}

ROLE_NOTES = {
    "справочная": SVOD_REFERENCE_NOTE,
    "причина (разметка)": (
        "Разметка причины подъёма. Известна только ПОСЛЕ подъёма ⇒ законна как "
        "фильтр популяции, не как ковариата прогноза."
    ),
    "провенанс": "Что сборщик восстановил или ограничил в этой строке.",
}


def column_role(column: str) -> str:
    """Role label shown on the legend sheet."""
    return _ROLE_BY_COLUMN.get(column, "прочее")


def mark_reference_headers(worksheet, *, header_row: int = 1) -> int:
    """Fill and annotate the header cells of the reference / derived blocks.

    Returns the number of header cells marked.
    """
    marked = 0
    reference = set(SVOD_REFERENCE_COLUMNS)
    derived = set(SVOD_CAUSE_COLUMNS)
    for cell in worksheet[header_row]:
        name = str(cell.value).strip() if cell.value is not None else ""
        if name in reference:
            cell.fill = REFERENCE_HEADER_FILL
            cell.comment = Comment(SVOD_REFERENCE_NOTE, "analysis.ingest")
            marked += 1
        elif name in derived:
            cell.fill = DERIVED_HEADER_FILL
            cell.comment = Comment(ROLE_NOTES["причина (разметка)"], "analysis.ingest")
            marked += 1
    return marked


def write_legend_sheet(
    workbook: Workbook,
    columns,
    *,
    sheet_name: str = SVOD_LEGEND_SHEET_NAME,
    build_note: Optional[str] = None,
) -> None:
    """Write the machine-readable column legend as a second sheet."""
    if sheet_name in workbook.sheetnames:
        del workbook[sheet_name]
    worksheet = workbook.create_sheet(sheet_name)

    headers = ["№", "Колонка", "Роль", "Источник", "Примечание"]
    for index, title in enumerate(headers, start=1):
        cell = worksheet.cell(row=1, column=index, value=title)
        cell.font = Font(bold=True)

    for row_index, column in enumerate(columns, start=2):
        role = column_role(column)
        worksheet.cell(row=row_index, column=1, value=row_index - 1)
        worksheet.cell(row=row_index, column=2, value=column)
        worksheet.cell(row=row_index, column=3, value=role)
        worksheet.cell(row=row_index, column=4, value=COLUMN_SOURCES.get(column, ""))
        note_cell = worksheet.cell(row=row_index, column=5, value=ROLE_NOTES.get(role, ""))
        note_cell.alignment = Alignment(wrap_text=True, vertical="top")
        if role == "справочная":
            worksheet.cell(row=row_index, column=3).fill = REFERENCE_HEADER_FILL
        elif role == "причина (разметка)":
            worksheet.cell(row=row_index, column=3).fill = DERIVED_HEADER_FILL

    if build_note:
        footer_row = len(list(columns)) + 3
        cell = worksheet.cell(row=footer_row, column=2, value=build_note)
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    for letter, width in zip("ABCDE", (6, 46, 22, 34, 70)):
        worksheet.column_dimensions[letter].width = width
    worksheet.freeze_panes = "A2"


__all__ = [
    "COLUMN_SOURCES",
    "DERIVED_HEADER_FILL",
    "REFERENCE_HEADER_FILL",
    "ROLE_NOTES",
    "column_role",
    "mark_reference_headers",
    "write_legend_sheet",
]
