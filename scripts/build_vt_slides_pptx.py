"""Convert docs/notes/vt_model_slides.md → a 16:9 PPTX deck.

Usage: python build_pptx.py [input.md] [output.pptx]  (defaults to the RU deck).
"""
import os, re, sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

MD = sys.argv[1] if len(sys.argv) > 1 else r"D:\GitHub\Pump2\docs\notes\vt_model_slides.md"
OUT = sys.argv[2] if len(sys.argv) > 2 else r"D:\GitHub\Pump2\docs\notes\vt_model_slides.pptx"
BASE = os.path.dirname(MD)
_EN = os.path.basename(MD).endswith("_en.md")   # language of this deck

NAVY = RGBColor(0x1F, 0x3A, 0x5F); DARK = RGBColor(0x22, 0x22, 0x22)
GREY = RGBColor(0x66, 0x66, 0x66); ACC = RGBColor(0x2E, 0x5E, 0x8C)
HDR = RGBColor(0xE8, 0xEE, 0xF4)

EMU_IN = 914400
SW, SH = 13.333, 7.5
prs = Presentation(); prs.slide_width = Inches(SW); prs.slide_height = Inches(SH)
BLANK = prs.slide_layouts[6]


def img_size(path):
    from struct import unpack
    with open(path, "rb") as f:
        head = f.read(26)
    if head[:8] == b"\x89PNG\r\n\x1a\n":
        w, h = unpack(">II", head[16:24]); return w, h
    return 1200, 600


def runs(parush):
    """Split inline markdown into (text, bold, italic, mono) runs."""
    out = []
    for chunk in re.split(r"(\*\*.*?\*\*|\*.*?\*|`.*?`)", parush):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**") and len(chunk) > 4:
            out.append((chunk[2:-2], True, False, False))
        elif chunk.startswith("`") and chunk.endswith("`"):
            out.append((chunk[1:-1], False, False, True))
        elif chunk.startswith("*") and chunk.endswith("*") and len(chunk) > 2:
            out.append((chunk[1:-1], False, True, False))
        else:
            out.append((chunk.replace("**", ""), False, False, False))
    return out


def add_runs(para, text, sz, color=DARK, bold_all=False):
    for t, b, it, m in runs(text):
        r = para.add_run(); r.text = t
        r.font.size = Pt(sz); r.font.bold = b or bold_all; r.font.italic = it
        r.font.color.rgb = color
        r.font.name = "Consolas" if m else "Calibri"


def parse_sections(md):
    lines = md.split("\n")
    secs, cur = [], None
    for l in lines:
        if l.startswith("# ") or l.startswith("## "):
            if cur:
                secs.append(cur)
            cur = {"h": l.lstrip("#").strip(), "lvl": 1 if l.startswith("# ") else 2, "body": []}
        elif cur is not None:
            cur["body"].append(l)
    if cur:
        secs.append(cur)
    return secs


def parse_blocks(body):
    blocks, i, n = [], 0, len(body)
    while i < n:
        l = body[i]
        if l.strip() == "" or re.match(r"^[-*_]{3,}$", l.strip()):
            i += 1; continue
        if l.startswith("```"):
            j = i + 1; code = []
            while j < n and not body[j].startswith("```"):
                code.append(body[j]); j += 1
            blocks.append(("code", code)); i = j + 1; continue
        m = re.match(r"!\[.*?\]\((.*?)\)", l.strip())
        if m:
            blocks.append(("image", m.group(1))); i += 1; continue
        if l.lstrip().startswith("|"):
            tbl = []
            while i < n and body[i].lstrip().startswith("|"):
                tbl.append(body[i].strip()); i += 1
            blocks.append(("table", tbl)); continue
        if l.startswith(">"):
            q = []
            while i < n and body[i].startswith(">"):
                q.append(body[i][1:].strip()); i += 1
            blocks.append(("quote", " ".join(x for x in q if x))); continue
        if re.match(r"^\s*([-*]|\d+\.)\s", l):
            b = []
            while i < n and body[i].strip() and not body[i].startswith(("#", ">", "|", "!", "```")):
                b.append(body[i]); i += 1
            blocks.append(("bullets", b)); continue
        p = []
        while i < n and body[i].strip() and not re.match(r"^\s*([-*]|\d+\.)\s", body[i]) \
                and not body[i].startswith(("#", ">", "|", "!", "```")):
            p.append(body[i]); i += 1
        blocks.append(("para", " ".join(p))); continue
    return blocks


def merge_bullets(raw):
    """Turn raw bullet lines into (level, text) items, joining wrapped continuations."""
    items = []
    for l in raw:
        indent = len(l) - len(l.lstrip())
        s = l.strip()
        mm = re.match(r"^([-*]|\d+\.)\s+(.*)", s)
        if mm:
            lvl = 1 if indent >= 2 else 0
            items.append([lvl, mm.group(2)])
        elif items:
            items[-1][1] += " " + s
    return items


# ---- estimate block heights (inches) for vertical flow ----
def est_h(kind, data):
    if kind == "para":
        return 0.02 + 0.22 * max(1, (len(data) // 120) + 1)
    if kind == "quote":
        return 0.10 + 0.21 * max(1, (len(data) // 125) + 1)
    if kind == "bullets":
        h = 0.04
        for lvl, t in merge_bullets(data):
            h += 0.22 * max(1, (len(t) // (108 if lvl == 0 else 96)) + 1)
        return h
    if kind == "table":
        rows = [r for r in data if not re.match(r"^\|[\s:|-]+\|$", r)]
        return 0.08 + 0.26 * len(rows)
    if kind == "code":
        return 0.14 + 0.19 * len(data)
    return 0.0


def add_title_bar(slide, text, num=None):
    bar = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.92))
    bar.fill.solid(); bar.fill.fore_color.rgb = NAVY; bar.line.fill.background()
    bar.shadow.inherit = False
    tf = bar.text_frame; tf.margin_left = Inches(0.4); tf.margin_top = Inches(0.12)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
    add_runs(p, text, 22, RGBColor(0xFF, 0xFF, 0xFF), bold_all=True)


def place_textbox(slide, y, kind, data, x=0.45, w=12.45):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(0.4))
    tf = tb.text_frame; tf.word_wrap = True
    first = True
    if kind == "para":
        p = tf.paragraphs[0]; add_runs(p, data, 13)
    elif kind == "quote":
        p = tf.paragraphs[0]
        r0 = p.add_run(); r0.text = "▍ "; r0.font.size = Pt(11); r0.font.color.rgb = ACC
        for t, b, it, m in runs(data):
            r = p.add_run(); r.text = t; r.font.size = Pt(10.5); r.font.italic = True
            r.font.color.rgb = GREY; r.font.name = "Consolas" if m else "Calibri"
    elif kind == "bullets":
        for lvl, t in merge_bullets(data):
            p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
            p.level = lvl
            b = p.add_run(); b.text = ("• " if lvl == 0 else "– ")
            b.font.size = Pt(13); b.font.color.rgb = ACC
            add_runs(p, t, 13 if lvl == 0 else 12)
    return tb


def place_code(slide, y, code, fs=9.0, x=0.45):
    lh = fs / 9.0 * 0.19
    h = 0.12 + lh * len(code)
    w = min(12.45, 0.0106 * fs * max((len(c) for c in code), default=20) + 0.5)
    box = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    box.fill.solid(); box.fill.fore_color.rgb = RGBColor(0xF4, 0xF6, 0xF8)
    box.line.color.rgb = RGBColor(0xD0, 0xD7, 0xDE); box.line.width = Pt(0.75)
    box.shadow.inherit = False
    tf = box.text_frame; tf.word_wrap = False; tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = Inches(0.12); tf.margin_top = Inches(0.06); tf.margin_bottom = Inches(0.05)
    for k, c in enumerate(code):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.line_spacing = 1.0; p.alignment = PP_ALIGN.LEFT
        r = p.add_run(); r.text = c if c else " "
        r.font.name = "Consolas"; r.font.size = Pt(fs); r.font.color.rgb = DARK
    return h


def _clen(s):
    return len(re.sub(r"[*`]", "", s))


def place_table(slide, y, tbl, x=0.45):
    rows = [r for r in tbl if not re.match(r"^\|[\s:|-]+\|$", r)]
    grid = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    nr, nc = len(grid), max(len(r) for r in grid)
    for r in grid:
        r += [""] * (nc - len(r))
    w = min(12.6, max(3.0, 1.05 * nc + 1.2))
    # proportional column widths from max content length (floor so labels don't vanish)
    colmax = [max(_clen(grid[ri][ci]) for ri in range(nr)) for ci in range(nc)]
    weights = [max(m, 5) for m in colmax]
    tot = sum(weights)
    colw = [max(0.85, w * wt / tot) for wt in weights]
    w = sum(colw)
    # per-row height from wrapping (chars-per-line depends on that column's width)
    rowh = []
    for ri, row in enumerate(grid):
        cw = 0.072 if ri == 0 else 0.068
        lines = 1
        for ci in range(nc):
            cpl = max(4, int(colw[ci] / cw))
            lines = max(lines, -(-_clen(row[ci]) // cpl))  # ceil
        rowh.append(max(0.30, 0.10 + 0.155 * lines))
    H = sum(rowh)
    gt = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(H)).table
    gt.first_row = False; gt.horz_banding = False
    for ci in range(nc):
        gt.columns[ci].width = Inches(colw[ci])
    for ri, row in enumerate(grid):
        gt.rows[ri].height = Inches(rowh[ri])
        for ci in range(nc):
            cell = gt.cell(ri, ci)
            cell.margin_left = Inches(0.06); cell.margin_right = Inches(0.05)
            cell.margin_top = Inches(0.02); cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            add_runs(p, row[ci], 9.5 if ri == 0 else 9,
                     RGBColor(0xFF, 0xFF, 0xFF) if ri == 0 else DARK, bold_all=(ri == 0))
            cell.fill.solid()
            cell.fill.fore_color.rgb = ACC if ri == 0 else (HDR if ri % 2 else RGBColor(0xFF, 0xFF, 0xFF))
    return H + 0.12


def place_images(slide, imgs, y_top, y_bot):
    band = max(1.6, y_bot - y_top)
    n = len(imgs)
    gap = 0.2
    colw = (12.5 - gap * (n - 1)) / n
    dims = [img_size(p) for p in imgs]
    # width-limited height per image, then cap to band
    heights = [colw / (w / h) for (w, h) in dims]
    scale = min(1.0, band / max(heights))
    xs = 0.42
    for (p, (w, h)) in zip(imgs, dims):
        iw = colw * scale
        ih = iw / (w / h)
        y = y_top + (band - ih) / 2
        slide.shapes.add_picture(p, Inches(xs), Inches(y), width=Inches(iw))
        xs += colw + gap


def content_slide(sec):
    state = {"slide": prs.slides.add_slide(BLANK), "y": 1.08}
    add_title_bar(state["slide"], sec["h"])
    blocks = parse_blocks(sec["body"])
    imgs = [os.path.normpath(os.path.join(BASE, d)) for k, d in blocks if k == "image"]
    textblocks = [(k, d) for k, d in blocks if k != "image"]
    BOT = 7.28

    def newpage():
        state["slide"] = prs.slides.add_slide(BLANK)
        add_title_bar(state["slide"], sec["h"] + (" (continued)" if _EN else " (продолжение)"))
        state["y"] = 1.18

    for idx, (k, d) in enumerate(textblocks):
        if k == "code":
            # fit whole block on this slide: shrink font under the space left, reserving
            # room for any blocks that follow (small font is OK — meant for copy-paste).
            follow = sum(est_h(kk, dd) for kk, dd in textblocks[idx + 1:])
            reserve = (follow + 0.2) if follow > 0 else 0.1
            avail = BOT - state["y"] - 0.15 - reserve
            need_lh = avail / max(1, len(d))
            fs = max(4.5, min(9.0, need_lh / 0.19 * 9.0))
            state["y"] += place_code(state["slide"], state["y"], d, fs) + 0.1
        elif k == "table":
            if state["y"] + est_h(k, d) > BOT and state["y"] > 1.3:
                newpage()
            state["y"] += place_table(state["slide"], state["y"], d) + 0.12
        else:
            if state["y"] + est_h(k, d) > BOT and state["y"] > 1.3:
                newpage()
            place_textbox(state["slide"], state["y"], k, d)
            state["y"] += est_h(k, d) + 0.06
    if imgs:
        min_band = 2.05 if len(imgs) >= 3 else 1.85
        top = min(state["y"] + 0.08, 7.32 - min_band)
        place_images(state["slide"], imgs, top, 7.32)


def divider_slide(text, sub=""):
    slide = prs.slides.add_slide(BLANK)
    bg = slide.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid(); bg.fill.fore_color.rgb = NAVY; bg.line.fill.background(); bg.shadow.inherit = False
    tb = slide.shapes.add_textbox(Inches(1), Inches(3.0), Inches(11.3), Inches(1.6))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    add_runs(p, text, 34, RGBColor(0xFF, 0xFF, 0xFF), bold_all=True)


def title_slide(title, subtitle):
    slide = prs.slides.add_slide(BLANK)
    band = slide.shapes.add_shape(1, 0, Inches(2.3), prs.slide_width, Inches(2.9))
    band.fill.solid(); band.fill.fore_color.rgb = NAVY; band.line.fill.background(); band.shadow.inherit = False
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(2.7), Inches(11.7), Inches(2.1))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    add_runs(p, title, 34, RGBColor(0xFF, 0xFF, 0xFF), bold_all=True)
    if subtitle:
        tb2 = slide.shapes.add_textbox(Inches(0.8), Inches(5.4), Inches(11.7), Inches(1.4))
        tf2 = tb2.text_frame; tf2.word_wrap = True
        p2 = tf2.paragraphs[0]; p2.alignment = PP_ALIGN.CENTER
        add_runs(p2, subtitle, 14, GREY)


# ---------------- build ----------------
md = open(MD, encoding="utf-8").read()
secs = parse_sections(md)

# section 0 = H1 title + preamble (glossary blockquote)
title = secs[0]["h"]
pre = secs[0]["body"]
title_slide(title, "ESP operating-life model · dependence on regime and contractor · v3.1"
            if _EN else "Модель срока службы УЭЦН · зависимость от режима и подрядчика · v3.1")

# glossary slide from preamble quote lines (keep the glossary part)
gloss = [l[1:].strip() for l in pre if l.startswith(">")]
gtext = " ".join(x for x in gloss if x)
_gmark = "Glossary.**" if "Glossary.**" in gtext else ("расшифровка).**" if "расшифровка).**" in gtext else None)
if _gmark:
    sl = prs.slides.add_slide(BLANK)
    add_title_bar(sl, "Glossary" if _EN else "Словарь терминов")
    body = gtext.split(_gmark, 1)[1].strip()
    tb = sl.shapes.add_textbox(Inches(0.5), Inches(1.15), Inches(12.4), Inches(6.0))
    tf = tb.text_frame; tf.word_wrap = True
    for part in re.split(r"(?<=\.)\s+(?=\*\*)", body):
        if not part.strip():
            continue
        p = tf.add_paragraph() if tf.paragraphs[0].runs else tf.paragraphs[0]
        r = p.add_run(); r.text = "• "; r.font.size = Pt(13); r.font.color.rgb = ACC
        add_runs(p, part.strip(), 13)
        p.space_after = Pt(4)

for sec in secs[1:]:
    h = sec["h"]
    if h.startswith(("Раздел", "Part ")):
        divider_slide(h)
    elif re.match(r"^\d+\s+—", h):
        content_slide(sec)
    elif h.startswith("Приложение"):
        content_slide(sec)
    else:
        content_slide(sec)

prs.save(OUT)
print("saved", OUT, "| slides:", len(prs.slides._sldIdLst))
