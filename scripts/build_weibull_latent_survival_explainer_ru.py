from __future__ import annotations

import math
from pathlib import Path
import sys

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Cm, Pt


REPO_ROOT = Path(__file__).resolve().parents[1]
BACKEND_DIR = REPO_ROOT / "backend"
for candidate in (str(REPO_ROOT), str(BACKEND_DIR)):
    if candidate not in sys.path:
        sys.path.insert(0, candidate)

from analysis import (
    TwoComponentLatentWeibullModel,
    WeibullParameters,
    latent_curve_frame,
    latent_life_quantile,
    latent_posterior_given_failure,
    latent_survival,
    weibull_survival,
)
from analysis.paths import results_dir

_SLUG = "weibull_latent_explainer_ru"
_OUT_ROOT = results_dir(_SLUG)
GENERATED_DIR = _OUT_ROOT / "figures"
PPTX_PATH = _OUT_ROOT / "weibull_latent_survival_explainer_ru.pptx"
MD_PATH = _OUT_ROOT / "weibull_latent_survival_explainer_ru.md"


W = Cm(33.87)
H = Cm(19.05)
MARGIN = Cm(1.0)

C_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
C_TITLE = RGBColor(0x1A, 0x27, 0x44)
C_BODY = RGBColor(0x22, 0x22, 0x22)
C_ACCENT = RGBColor(0xB3, 0x6A, 0x00)
C_BLUE = RGBColor(0x1F, 0x5C, 0x99)
C_GREEN = RGBColor(0x2E, 0x7D, 0x32)
C_RED = RGBColor(0xC2, 0x41, 0x0C)


def weibull_quantile(failure_probability: float, beta: float, eta: float) -> float:
    return float(eta * (-math.log(max(1.0 - failure_probability, 1e-12))) ** (1.0 / beta))


def weibull_mean(beta: float, eta: float) -> float:
    return float(eta * math.gamma(1.0 + 1.0 / beta))


def _white_bg(slide) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = C_WHITE


def _add_title(slide, title: str) -> None:
    box = slide.shapes.add_textbox(int(MARGIN), int(Cm(0.35)), int(W - 2 * MARGIN), int(Cm(1.4)))
    frame = box.text_frame
    p = frame.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.name = "Calibri"
    r.font.size = Pt(24)
    r.font.bold = True
    r.font.color.rgb = C_TITLE

    rule = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        0,
        int(Cm(1.95)),
        int(W),
        int(Cm(0.08)),
    )
    rule.fill.solid()
    rule.fill.fore_color.rgb = C_TITLE
    rule.line.fill.background()


def _add_text_block(slide, title: str, lines: list[str], left: Cm, top: Cm, width: Cm, height: Cm) -> None:
    box = slide.shapes.add_textbox(int(left), int(top), int(width), int(height))
    frame = box.text_frame
    frame.word_wrap = True

    header = frame.paragraphs[0]
    run = header.add_run()
    run.text = title
    run.font.name = "Calibri"
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = C_BLUE

    for line in lines:
        p = frame.add_paragraph()
        p.level = 1
        p.space_before = Pt(1)
        p.space_after = Pt(1)
        r = p.add_run()
        r.text = f"• {line}"
        r.font.name = "Calibri"
        r.font.size = Pt(15)
        r.font.color.rgb = C_BODY


def _add_caption(slide, text: str, left: Cm, top: Cm, width: Cm, color: RGBColor = C_ACCENT) -> None:
    box = slide.shapes.add_textbox(int(left), int(top), int(width), int(Cm(1.1)))
    frame = box.text_frame
    frame.word_wrap = True
    p = frame.paragraphs[0]
    r = p.add_run()
    r.text = text
    r.font.name = "Calibri"
    r.font.size = Pt(13)
    r.font.italic = True
    r.font.color.rgb = color


def _add_image(slide, path: Path, left: Cm, top: Cm, width: Cm, height: Cm) -> None:
    slide.shapes.add_picture(str(path), int(left), int(top), int(width), int(height))


def _save_single_weibull_figure(path: Path) -> dict[str, float]:
    beta = 1.6
    eta = 420.0
    times = np.linspace(0.0, 900.0, 500)
    survival = np.asarray(weibull_survival(times, WeibullParameters(beta=beta, eta=eta)), dtype=float)

    b10 = weibull_quantile(0.10, beta, eta)
    b50 = weibull_quantile(0.50, beta, eta)
    b90 = weibull_quantile(0.90, beta, eta)
    mean_life = weibull_mean(beta, eta)
    s180 = float(weibull_survival(180.0, WeibullParameters(beta=beta, eta=eta)))

    plt.figure(figsize=(10, 6))
    plt.plot(times, survival, color="#1F5C99", linewidth=3, label="S(t) = exp(-(t/η)^β)")
    for x_value, label, color in [
        (b10, "B10", "#2E7D32"),
        (b50, "B50 (медиана)", "#C2410C"),
        (eta, "η", "#8E24AA"),
        (b90, "B90", "#455A64"),
    ]:
        y_value = float(weibull_survival(x_value, WeibullParameters(beta=beta, eta=eta)))
        plt.axvline(x_value, color=color, linestyle="--", linewidth=1.4)
        plt.scatter([x_value], [y_value], color=color, s=40, zorder=5)
        plt.text(x_value + 8, y_value + 0.03, f"{label} = {x_value:.0f} д", color=color, fontsize=11)

    mean_survival = float(weibull_survival(mean_life, WeibullParameters(beta=beta, eta=eta)))
    plt.scatter([mean_life], [mean_survival], color="#D4A017", s=45, zorder=5, label="Средняя наработка E[T]")
    plt.text(mean_life + 8, mean_survival - 0.08, f"E[T] = {mean_life:.0f} д", color="#B36A00", fontsize=11)

    plt.scatter([180.0], [s180], color="#00838F", s=45, zorder=5)
    plt.text(188, s180 + 0.03, f"S(180) = {s180:.2f}", color="#00838F", fontsize=11)

    plt.title("Классическая 2-параметрическая Weibull-модель", fontsize=16)
    plt.xlabel("Время, дни")
    plt.ylabel("Вероятность безотказной работы S(t)")
    plt.ylim(0.0, 1.02)
    plt.xlim(0.0, 900.0)
    plt.grid(alpha=0.25)
    plt.legend(loc="upper right")
    plt.tight_layout()
    plt.savefig(path, dpi=180)
    plt.close()

    return {
        "beta": beta,
        "eta": eta,
        "b10": b10,
        "b50": b50,
        "b90": b90,
        "mean": mean_life,
        "s180": s180,
    }


def _save_latent_figure(path: Path) -> dict[str, float]:
    model = TwoComponentLatentWeibullModel(
        weight_1=0.35,
        component_1=WeibullParameters(beta=0.85, eta=120.0, label="Скрытый класс 1: ранние отказы"),
        component_2=WeibullParameters(beta=1.90, eta=420.0, label="Скрытый класс 2: нормальный ресурс"),
    )
    frame = latent_curve_frame(model, max_time=900.0, num_points=500)
    times = frame["time"].to_numpy(dtype=float)
    mixture = frame["survival"].to_numpy(dtype=float)

    b10 = float(latent_life_quantile(0.10, model))
    b50 = float(latent_life_quantile(0.50, model))
    b90 = float(latent_life_quantile(0.90, model))
    mean_1 = weibull_mean(model.component_1.beta, model.component_1.eta)
    mean_2 = weibull_mean(model.component_2.beta, model.component_2.eta)
    mean_mix = model.weight_1 * mean_1 + (1.0 - model.weight_1) * mean_2
    s180 = float(latent_survival(180.0, model))
    post_early_fail = float(latent_posterior_given_failure(90.0, model)[0])

    plt.figure(figsize=(10, 6))
    plt.plot(times, mixture, color="#1F5C99", linewidth=3, label="Смесь: π·S1(t) + (1-π)·S2(t)")
    plt.plot(times, frame["component_1_survival"], color="#C2410C", linestyle="--", linewidth=2, label="Класс 1: ранние отказы")
    plt.plot(times, frame["component_2_survival"], color="#2E7D32", linestyle="--", linewidth=2, label="Класс 2: нормальный ресурс")

    for x_value, label, color in [
        (b10, "B10", "#2E7D32"),
        (b50, "B50 смеси", "#C2410C"),
        (b90, "B90", "#455A64"),
    ]:
        y_value = float(latent_survival(x_value, model))
        plt.axvline(x_value, color=color, linestyle=":", linewidth=1.5)
        plt.scatter([x_value], [y_value], color=color, s=40, zorder=5)
        plt.text(x_value + 8, y_value + 0.03, f"{label} = {x_value:.0f} д", color=color, fontsize=11)

    mean_survival = float(latent_survival(mean_mix, model))
    plt.scatter([mean_mix], [mean_survival], color="#B36A00", s=45, zorder=5)
    plt.text(mean_mix + 8, mean_survival - 0.08, f"E[T] смеси = {mean_mix:.0f} д", color="#B36A00", fontsize=11)

    plt.scatter([180.0], [s180], color="#00838F", s=45, zorder=5)
    plt.text(188, s180 + 0.03, f"S(180) = {s180:.2f}", color="#00838F", fontsize=11)

    plt.title("Латентная Weibull-модель: смесь скрытых классов", fontsize=16)
    plt.xlabel("Время, дни")
    plt.ylabel("Вероятность безотказной работы S(t)")
    plt.ylim(0.0, 1.02)
    plt.xlim(0.0, 900.0)
    plt.grid(alpha=0.25)
    plt.legend(loc="upper right")
    plt.tight_layout()
    plt.savefig(path, dpi=180)
    plt.close()

    return {
        "pi": model.weight_1,
        "beta_1": model.component_1.beta,
        "eta_1": model.component_1.eta,
        "beta_2": model.component_2.beta,
        "eta_2": model.component_2.eta,
        "b10": b10,
        "b50": b50,
        "b90": b90,
        "mean": mean_mix,
        "s180": s180,
        "posterior_early_fail_90": post_early_fail,
    }


def _build_pptx(single_path: Path, latent_path: Path, single_stats: dict[str, float], latent_stats: dict[str, float]) -> None:
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    # Cover
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _white_bg(slide)
    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, int(Cm(1.2)), int(H))
    bar.fill.solid()
    bar.fill.fore_color.rgb = C_TITLE
    bar.line.fill.background()
    title_box = slide.shapes.add_textbox(int(Cm(2.1)), int(Cm(3.8)), int(Cm(28)), int(Cm(4)))
    frame = title_box.text_frame
    p = frame.paragraphs[0]
    r = p.add_run()
    r.text = "Weibull-модель выживаемости\nи латентная модель выживаемости"
    r.font.name = "Calibri"
    r.font.size = Pt(30)
    r.font.bold = True
    r.font.color.rgb = C_TITLE
    p.alignment = PP_ALIGN.LEFT
    subtitle = frame.add_paragraph()
    subtitle.space_before = Pt(10)
    rs = subtitle.add_run()
    rs.text = "Краткое русское объяснение основных параметров и чтения кривых S(t)"
    rs.font.name = "Calibri"
    rs.font.size = Pt(18)
    rs.font.color.rgb = C_BODY
    _add_caption(
        slide,
        "Иллюстративные параметры на графиках можно заменить на ваши реальные оценки из модели.",
        Cm(2.1),
        Cm(11.0),
        Cm(26.5),
    )

    # Slide 1: classic Weibull
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _white_bg(slide)
    _add_title(slide, "Слайд 1 — Классическая Weibull-модель выживаемости")
    _add_text_block(
        slide,
        "Что задаёт модель",
        [
            "Функция выживания: S(t) = exp(-(t/η)^β).",
            "β — параметр формы: β < 1 ранние отказы, β ≈ 1 случайные отказы, β > 1 износ.",
            "η — параметр масштаба (characteristic life): при t = η имеем S(t) = e^-1 ≈ 36.8%.",
            "Средняя наработка E[T] и медиана B50 обычно не совпадают: это разные характеристики ресурса.",
        ],
        Cm(1.1),
        Cm(2.4),
        Cm(12.0),
        Cm(10.5),
    )
    _add_text_block(
        slide,
        "Что отмечено на кривой",
        [
            f"B10 = {single_stats['b10']:.0f} д: к этому сроку отказало ~10% парка.",
            f"B50 = {single_stats['b50']:.0f} д: медианный срок службы.",
            f"η = {single_stats['eta']:.0f} д: characteristic life, S(η) ≈ 36.8%.",
            f"E[T] = {single_stats['mean']:.0f} д: средняя наработка до отказа.",
            f"S(180) = {single_stats['s180']:.2f}: надёжность на выбранном миссионном горизонте.",
        ],
        Cm(1.1),
        Cm(10.7),
        Cm(12.0),
        Cm(6.5),
    )
    _add_image(slide, single_path, Cm(14.0), Cm(2.5), Cm(18.2), Cm(13.4))
    _add_caption(
        slide,
        "Иллюстративный пример: β = 1.60, η = 420 д. При β > 1 риск отказа растёт по мере старения.",
        Cm(14.0),
        Cm(16.2),
        Cm(18.2),
    )

    # Slide 2: latent model
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _white_bg(slide)
    _add_title(slide, "Слайд 2 — Латентная Weibull-модель: смесь скрытых классов")
    _add_text_block(
        slide,
        "Идея модели",
        [
            "Одна наблюдаемая популяция может состоять из нескольких скрытых подгрупп с разной надёжностью.",
            "Смесь выживаемости: S(t) = π·S1(t) + (1-π)·S2(t).",
            "π — доля short-life класса; у каждого класса свои β_k и η_k.",
            "Модель удобна, когда часть фонда умирает рано, а другая часть имеет нормальный ресурс.",
        ],
        Cm(1.1),
        Cm(2.4),
        Cm(12.0),
        Cm(10.0),
    )
    _add_text_block(
        slide,
        "Как читать параметры",
        [
            f"π = {latent_stats['pi']:.2f}: доля скрытого класса ранних отказов.",
            f"Класс 1: β1 = {latent_stats['beta_1']:.2f}, η1 = {latent_stats['eta_1']:.0f} д.",
            f"Класс 2: β2 = {latent_stats['beta_2']:.2f}, η2 = {latent_stats['eta_2']:.0f} д.",
            f"B50 смеси = {latent_stats['b50']:.0f} д, E[T] смеси = {latent_stats['mean']:.0f} д.",
            f"P(ранний класс | отказ на 90 дн.) = {latent_stats['posterior_early_fail_90']:.2f}.",
        ],
        Cm(1.1),
        Cm(10.6),
        Cm(12.0),
        Cm(6.7),
    )
    _add_image(slide, latent_path, Cm(14.0), Cm(2.5), Cm(18.2), Cm(13.4))
    _add_caption(
        slide,
        "Иллюстративный пример: смесь short-life и normal-life классов. Общая кривая может быть изогнутой сильнее, чем у одной Weibull-модели.",
        Cm(14.0),
        Cm(16.2),
        Cm(18.2),
    )

    prs.save(str(PPTX_PATH))


def _build_markdown(single_stats: dict[str, float], latent_stats: dict[str, float]) -> None:
    text = f"""cover
## Weibull-модель выживаемости и латентная модель выживаемости

Краткое русское объяснение основных параметров и чтения кривых `S(t)`.

> Иллюстративные параметры на графиках можно заменить на ваши реальные оценки из модели.

---

## Слайд 1 — Классическая Weibull-модель выживаемости

![Кривая классической Weibull-модели](generated/weibull_latent_survival_explainer_ru/weibull_single_curve_ru.png)

### Что задаёт модель

- Функция выживания: `S(t) = exp(-(t/η)^β)`.
- `β` — параметр формы: `β < 1` ранние отказы, `β ≈ 1` случайные отказы, `β > 1` износ.
- `η` — параметр масштаба: при `t = η` имеем `S(t) = e^-1 ≈ 36.8%`.
- Средняя наработка `E[T]` и медиана `B50` не совпадают: это разные характеристики ресурса.

### Что отмечено на кривой

- `B10 = {single_stats['b10']:.0f} д`: к этому сроку отказало около 10% парка.
- `B50 = {single_stats['b50']:.0f} д`: медианный срок службы.
- `η = {single_stats['eta']:.0f} д`: characteristic life.
- `E[T] = {single_stats['mean']:.0f} д`: средняя наработка до отказа.
- `S(180) = {single_stats['s180']:.2f}`: надёжность на горизонте 180 дней.

---

## Слайд 2 — Латентная Weibull-модель: смесь скрытых классов

![Кривая латентной Weibull-модели](generated/weibull_latent_survival_explainer_ru/weibull_latent_curve_ru.png)

### Идея модели

- Наблюдаемая популяция может состоять из нескольких скрытых подгрупп с разной надёжностью.
- Смесь выживаемости: `S(t) = π·S1(t) + (1-π)·S2(t)`.
- `π` — доля short-life класса; у каждого класса свои `β_k` и `η_k`.
- Модель полезна, когда одна часть фонда умирает рано, а другая часть имеет нормальный ресурс.

### Как читать параметры

- `π = {latent_stats['pi']:.2f}` — доля скрытого класса ранних отказов.
- Класс 1: `β1 = {latent_stats['beta_1']:.2f}`, `η1 = {latent_stats['eta_1']:.0f} д`.
- Класс 2: `β2 = {latent_stats['beta_2']:.2f}`, `η2 = {latent_stats['eta_2']:.0f} д`.
- `B50 смеси = {latent_stats['b50']:.0f} д`, `E[T] смеси = {latent_stats['mean']:.0f} д`.
- `P(ранний класс | отказ на 90 дн.) = {latent_stats['posterior_early_fail_90']:.2f}`.
"""
    MD_PATH.write_text(text, encoding="utf-8")


def main() -> None:
    GENERATED_DIR.mkdir(parents=True, exist_ok=True)
    single_path = GENERATED_DIR / "weibull_single_curve_ru.png"
    latent_path = GENERATED_DIR / "weibull_latent_curve_ru.png"

    single_stats = _save_single_weibull_figure(single_path)
    latent_stats = _save_latent_figure(latent_path)
    _build_markdown(single_stats, latent_stats)
    _build_pptx(single_path, latent_path, single_stats, latent_stats)

    print(f"Wrote markdown: {MD_PATH}")
    print(f"Wrote presentation: {PPTX_PATH}")
    print(f"Wrote figures: {single_path}, {latent_path}")


if __name__ == "__main__":
    main()
