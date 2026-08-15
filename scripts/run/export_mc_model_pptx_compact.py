"""Компактная ТЕХНИЧЕСКАЯ СПРАВКА по модели отказов УЭЦН Мирнинского — ОДИН слайд.

Сжатая версия трёхслайдового `export_mc_model_pptx.py`: оставлены фигура разреза по Qж,
одна общая таблица параметров (Вейбулл + слой Cox) и укороченное описание того, как из
модели собирается график ремонтов. Это справочник, а не презентация, поэтому шрифт
мелкий и плотность высокая.

Терминология выровнена по правкам пользователя в `..._v1.pptx`: «наработка НА отказ»,
дебит жидкости — **Qж** (не Ql), «цензурированы», β «не является триггером отказа».

Фигура перестраивается здесь же (только Mc — быстро), чтобы подписи на самой картинке
тоже были Qж; Vt-фигура из большого набора для этого слайда не нужна.

Запуск:  python scripts/run/export_mc_model_pptx_compact.py
"""
from __future__ import annotations

import sys
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO_ROOT / "backend"))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt

from analysis.paths import results_dir
from analysis.workflows.production_risk import config as C
from analysis.workflows.production_risk import esp_population as P
from analysis.workflows.production_risk import esp_forecast as M

from export_mc_model_pptx import _bullets, _pic, _table, _title, SLIDE_H, SLIDE_W
from export_mc_model_slides import slide2

HORIZON = 18


def main() -> None:
    asof = pd.Timestamp(str(C.RunConfig().forecast_start))
    pop = P.build(str(C.RunConfig().forecast_start))
    model = M.fit_mc_model(pop)
    mc = P.select(pop, "Mc")
    mc = mc[mc["tte"] > 0].copy()

    out = results_dir("production_risk_mc_model_slides")
    figs = out / "figures"
    figs.mkdir(parents=True, exist_ok=True)
    s2 = slide2(mc, model, figs)                       # перерисовываем с подписями Qж
    fig = s2["path"]

    # цифры для блока «как строится график» — из живого прогноза, не из памяти
    codes = set(pop[pop["field"] == "Mc"]["code"])
    live = M.live_fleet(pop, asof, field="Mc")
    ent = M.plan_entrants(set(live["code"]), codes, asof, HORIZON)
    fleet = M.active_fleet_by_month(codes=codes | set(ent["code"]))
    monthly, per_well = M.forecast(model, live, asof=asof, months=HORIZON,
                                   fleet_by_month=fleet, entrants=ent)
    by_code = M.allocate_events(monthly, per_well)
    n_events = sum(len(v) for v in by_code.values())
    sum_e = float(monthly["ожидаемые события"].sum())
    f0, f1 = int(monthly["в работе"].iloc[0]), int(monthly["в работе"].iloc[-1])

    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    s = prs.slides.add_slide(prs.slide_layouts[6])

    _title(s, "Модель наработки на отказ УЭЦН — Мирнинский УН · техническая справка",
           "Монтажи 2024+, некислые. Событие — ОТКАЗ; ГТМ и работающие цензурированы. "
           "Вейбулл-смесь cal c0 k2 + сценарный слой Cox по Qж")

    # ── слева: фигура + примечания по β ──────────────────────────────────────
    _pic(s, fig, Inches(0.30), Inches(0.95), Inches(7.55))
    _bullets(s, Inches(0.32), Inches(5.32), Inches(7.5), Inches(2.0), [
        ("Откуда β.", "Landmark-схема: ковариата из окна [0, 30) суток, риск-набор — "
                      "дожившие до 30 суток. Оценено глобально по 19 стратам, n = 1741."),
        ("Проверка на Vt_nonsour.", "Cox на 476 пробегах / 209 отказах даёт β = 0.296 "
                                    "(95% ДИ 0.116…0.476, p = 0.001) — интервал накрывает 0.226."),
        ("Чего β НЕ значит.", "Годится для ранжирования, но частично несёт неизмеренные "
                              "свойства скважины/насоса и не является триггером отказа."),
        ("Корреляция ковариат.", "ГЖФ, Qн, Qв, Qг, WCUT сильно коррелируют с Qж на Mc, поэтому "
                                 "их отдельный вклад не оценить — поправка на Qж косвенно "
                                 "учитывает влияние остальных факторов."),
    ], size=8.5)

    # ── справа: одна таблица параметров ──────────────────────────────────────
    _table(s, [
        ["Параметр", "Значение"],
        ["Вейбулл cal c0 k2 — S(t) = w₁·exp(−(t/η₁)^β₁) + (1−w₁)·exp(−(t/η₂)^β₂)", ""],
        ["w₁ — доля 1-й компоненты", "%.4f" % model.w1],
        ["β₁ / η₁, сут — форма и масштаб 1-й", "%.3f / %.2f" % (model.beta1, model.eta1)],
        ["β₂ / η₂, сут — форма и масштаб 2-й", "%.3f / %.1f" % (model.beta2, model.eta2)],
        ["τ, сут — горизонт RMST (95-й проц.)", "%.0f" % model.tau],
        ["Пробегов / отказов / цензурировано", "%d / %d / %d"
         % (model.n_runs, model.n_events, model.n_censored)],
        ["Отсечка c, сут — ранние НЕ вырезаны", "%.0f" % model.cut_days],
        ["RMST(0) / MRL(0), сут", "%.0f / %.0f" % (model.rmst(0.0), model.mrl(0.0))],
        ["Слой Cox по Qж (сценарный) — S(t | Qж) = S₀(t)^θ", ""],
        ["β для log-Qж", "ln(1.253) = 0.226"],
        ["θ(Qж)", "exp(β·[ln(1+Qж) − ref])"],
        ["θ при Qж = 3× / 5× медианы", "1.28 / 1.44"],
    ], Inches(7.85), Inches(0.95), Inches(5.18), Inches(3.30),
        col_w=[2.6, 1.15], font=8.0)

    # ── справа снизу: как строится график (сжато) ────────────────────────────
    _bullets(s, Inches(7.90), Inches(4.82), Inches(5.15), Inches(2.5), [
        ("Как строится график ремонтов.",
         "Модель даёт МЕСЯЧНУЮ вероятность, а не дату; отказ ставится на 1-е число месяца."),
        ("Сколько.", "Для насоса за месяц p = 1 − S(a+Δ)/S(a), a — возраст; сумма по парку = "
                     "ожидание E(m), дробное. В график идёт n(m) = ⌊E(m) + остаток⌋: дробный "
                     "хвост копится по парку и добавляет отказ при переходе через 1. "
                     "За %d мес ΣE = %.1f → %d целых событий." % (HORIZON, sum_e, n_events)),
        ("Кто и в каком порядке.", "risk_30d — ожидаемое число отказов насоса за месяц; ДОЛГ — их "
                                   "сумма с даты расчёта (у всех с нуля). Событие достаётся насосам "
                                   "с наибольшим долгом, долг уменьшается на 1. Очередь строго по "
                                   "риску отдавала бы отказ только старым и морила ВНС. Порядок "
                                   "«старые первыми» задаёт форма опасности Mc (risk_30d растёт с "
                                   "возрастом), а не само правило. Отказавший насос сразу меняется "
                                   "на новый с возраста 0, поэтому скважина может отказать дважды."),
        ("Парк.", "Активные добывающие скважины плана ПП: %d → %d за горизонт "
                  "(%d живых + %d вводимых)." % (f0, f1, len(live), len(ent))),
    ], size=8.5)

    path = out / "Модель_отказов_Мирнинский_компакт.pptx"
    prs.save(str(path))
    print("записано:", path)
    print("ΣE=%.1f → %d событий | парк %d→%d | живых %d + ВНС %d"
          % (sum_e, n_events, f0, f1, len(live), len(ent)))


if __name__ == "__main__":
    main()
