"""Tests for Свод well identity, non-oil-bore filtering and censored (running) runs.

Covers the two builder-reliability tasks:

* T1-T6 -- bore identity is preserved and non-oil (brine / injection /
  piezometric / water-intake) bores never enter the oil-production register.
* R1-R7 -- living (censored) runs are added with the survival invariants:
  one live run per well, monotonic age, runtime <= calendar.
"""

# Перенесено из db_builder/tests/test_svod_well_identity.py вместе с кодом сборщика: тесты
# обязаны переезжать с тем, что они держат, иначе перенос молча теряет именно
# те инварианты, ради которых код и писался.

from pathlib import Path

import pandas as pd
from openpyxl import Workbook

from analysis.ingest.svod.well_identity import (
    bore_suffix,
    is_non_oil_purpose,
    is_rassol_bore,
    normalize_well_key,
    well_number_key,
)
from analysis.ingest.svod.artificial_lift_processor import (
    extract_running_wells_from_artificial_lift,
    runtime_within_calendar,
)


BIG_HEADER_ROW1 = [
    "Скважина",
    "Цель спуска",
    "Дата отказа",
    "Дата монтажа",
    "Дата запуска",
    "Дата демонтажа",
    "ННО",
    "Дизайн УЭЦН",
    "Насос (50Гц)",
]
BIG_HEADER_ROW2 = [None, None, None, None, None, None, None, None, "Собственник оборудования"]


def _write_big_workbook(path: Path, data_rows) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Скважинное оборудование"
    for col_idx, value in enumerate(BIG_HEADER_ROW1, start=1):
        ws.cell(row=1, column=col_idx, value=value)
    for col_idx, value in enumerate(BIG_HEADER_ROW2, start=1):
        ws.cell(row=2, column=col_idx, value=value)
    for row_idx, values in enumerate(data_rows, start=3):
        for col_idx, value in enumerate(values, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)
    wb.save(path)


def _row(well, purpose, failure_date, installation_date, dismantling_date, nno):
    """Build one Big data row in header order.

    Header order: well, purpose, failure, installation, launch, dismantling,
    ННО, design, pump. Launch mirrors installation (irrelevant to these tests).
    """
    return [
        well, purpose, failure_date, installation_date, installation_date,
        dismantling_date, nno, "MT5A-DP", "Борец",
    ]


# --------------------------------------------------------------------------- #
# Task 1 -- identity & non-oil filter
# --------------------------------------------------------------------------- #

def test_t1_bore_suffix_is_not_stripped_and_variants_collapse():
    # The brine bore keeps a distinct identity from the oil well of the same number.
    assert normalize_well_key("Vt_2704рс") != normalize_well_key("Vt_2704")
    # Case / space / latin-cyrillic homoglyph variants of the SAME bore collapse.
    variants = {
        normalize_well_key("Vt_2704рс"),
        normalize_well_key("VT_2704РС"),
        normalize_well_key("VT_2704 РС"),
        normalize_well_key("VT_2704pc"),  # latin homoglyphs of рс
    }
    assert len(variants) == 1
    assert bore_suffix("Vt_2704рс") == "рс"
    assert bore_suffix("Vt_2704") == ""
    assert is_rassol_bore("Vt_2704рс") and not is_rassol_bore("Vt_2704")
    # Trailing separators must not hide the brine suffix.
    assert is_rassol_bore("Yaw_400рс_")
    assert is_rassol_bore("Yaw_400 рс")
    # Number key groups bores of one physical well but not two different numbers.
    assert well_number_key("Vt_2704рс") == well_number_key("Vt_2704")
    assert well_number_key("Vt_2704") != well_number_key("Vt_2705")


def test_t4_non_oil_purposes_recognized():
    for purpose in ["Нагнетательная", "Пьезометр", "водозаборная ВД", "водозаборная НД",
                    "консервация", "ликвидирована", "Техн. операции"]:
        assert is_non_oil_purpose(purpose), purpose
    assert not is_non_oil_purpose("Мех. добыча")


def test_t2_t3_t5_vt_2704_brine_bore_never_enters_population(tmp_path):
    big_path = tmp_path / "big_vt2704.xlsx"
    _write_big_workbook(big_path, [
        # Stale suffixless open oil row -- the phantom "living pump" (T2).
        _row("Vt_2704", "Мех. добыча", None, "2023-02-13", None, 30),
        # Real brine bore runs under the same number, closed (T3/T5).
        _row("Vt_2704рс", "Мех. добыча", "2023-04-28", "2023-02-13", "2023-04-28", 74),
        _row("Vt_2704рс", "Мех. добыча", "2024-05-06", "2023-11-06", "2024-05-06", 182),
        # Last brine state is injection, still open -> excluded (rassol + purpose).
        _row("Vt_2704рс", "Нагнетательная", None, "2024-05-08", None, 40),
        # A genuine oil well that must survive.
        _row("Vt_100", "Мех. добыча", None, "2025-09-01", None, 20),
    ])

    running = extract_running_wells_from_artificial_lift(str(big_path), as_of_date="2026-07-21")

    wells = set(running["well"])
    # No brine bore of 2704 in the population (T3), and no phantom oil 2704 (T2/T5).
    assert not any(is_rassol_bore(w) for w in wells)
    number_keys = {well_number_key(w) for w in wells}
    assert well_number_key("Vt_2704") not in number_keys
    assert "Vt_2704" not in wells and "Vt_2704рс" not in wells
    # The genuine oil well survives.
    assert "Vt_100" in wells


def test_t4_t6_non_oil_active_runs_are_dropped(tmp_path, capsys):
    big_path = tmp_path / "big_nonoil.xlsx"
    _write_big_workbook(big_path, [
        _row("Vt_10", "Мех. добыча", None, "2025-01-01", None, 100),
        _row("Vt_11", "Нагнетательная", None, "2025-01-01", None, 100),
        _row("Vt_12", "Пьезометр", None, "2025-01-01", None, 100),
        _row("Vt_13", "водозаборная ВД", None, "2025-01-01", None, 100),
        _row("Vt_14рс", "Мех. добыча", None, "2025-01-01", None, 100),
    ])

    running = extract_running_wells_from_artificial_lift(str(big_path), as_of_date="2026-01-01")

    assert set(running["well"]) == {"Vt_10"}
    # T6: dropped active runs are logged with a per-reason breakdown.
    out = capsys.readouterr().out
    assert "non-oil filter dropped" in out


# --------------------------------------------------------------------------- #
# Task 2 -- censored (running) runs & invariants
# --------------------------------------------------------------------------- #

def test_r1_r6_running_runs_are_open_and_carry_no_failure(tmp_path):
    big_path = tmp_path / "big_open.xlsx"
    _write_big_workbook(big_path, [
        _row("Vt_100", "Мех. добыча", None, "2025-09-01", None, 20),
        _row("Vt_101", "Мех. добыча", None, "2025-08-01", None, 50),
    ])

    running = extract_running_wells_from_artificial_lift(str(big_path), as_of_date="2026-01-01")

    # R1: open runs exist; R6: they are non-failures (censored).
    assert len(running) == 2
    assert running["failure_date"].isna().all()
    assert running["Дата остановки"].isna().all()
    assert (running["failure_flag"] == 0).all()


def test_r3_stale_open_row_dropped_by_monotonicity(tmp_path):
    # YA_601-style: an open run mounted 2016 while a closed run ended 2017.
    big_path = tmp_path / "big_ya601.xlsx"
    _write_big_workbook(big_path, [
        _row("Ya_601", "Мех. добыча", None, "2016-01-01", None, 3600),
        _row("Ya_601", "Мех. добыча", "2017-09-02", "2015-06-01", "2017-09-02", 458),
        _row("Ya_700", "Мех. добыча", None, "2025-01-01", None, 100),
    ])

    running = extract_running_wells_from_artificial_lift(str(big_path), as_of_date="2026-01-01")

    assert "Ya_601" not in set(running["well"])
    assert "Ya_700" in set(running["well"])


def test_r3_live_run_after_last_closed_run_survives(tmp_path):
    big_path = tmp_path / "big_relaunch.xlsx"
    _write_big_workbook(big_path, [
        _row("Vt_800", "Мех. добыча", "2024-05-06", "2023-11-06", "2024-05-06", 182),
        _row("Vt_800", "Мех. добыча", None, "2024-05-08", None, 60),  # new live run
    ])

    running = extract_running_wells_from_artificial_lift(str(big_path), as_of_date="2026-01-01")

    assert list(running["well"]) == ["Vt_800"]
    assert running.iloc[0]["installation_date"] == pd.Timestamp("2024-05-08")


def test_r5_one_live_run_per_well(tmp_path):
    big_path = tmp_path / "big_dup.xlsx"
    _write_big_workbook(big_path, [
        _row("Vt_500", "Мех. добыча", None, "2024-01-01", None, 100),
        _row("Vt_500", "Мех. добыча", None, "2024-06-01", None, 40),
    ])

    running = extract_running_wells_from_artificial_lift(str(big_path), as_of_date="2026-01-01")

    assert len(running) == 1
    assert running.iloc[0]["installation_date"] == pd.Timestamp("2024-06-01")


def test_r4_runtime_within_calendar_helper():
    # Runtime under the calendar age is fine; over it (beyond tolerance) fails.
    assert runtime_within_calendar(100, "2025-01-01", "2025-06-01") is True
    assert runtime_within_calendar(100000, "2025-01-01", "2025-06-01") is False
    assert runtime_within_calendar(None, "2025-01-01", "2025-06-01") is None


def test_r4_runtime_violation_is_logged_not_silent(tmp_path, capsys):
    big_path = tmp_path / "big_runtime.xlsx"
    _write_big_workbook(big_path, [
        _row("Vt_700", "Мех. добыча", None, "2025-01-01", None, 100000),
    ])

    extract_running_wells_from_artificial_lift(str(big_path), as_of_date="2025-06-01")

    out = capsys.readouterr().out
    assert "runtime > calendar age" in out
