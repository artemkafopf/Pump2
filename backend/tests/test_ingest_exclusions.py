"""Отсев с причиной и решение «дописать или перестроить».

Две вещи, которых сборщику не хватало:

* каждый отброшенный ряд должен быть НАЗВАН и объяснён — до этого семь этапов
  отсева уходили в `print` со счётчиком, и проверить решение было нельзя;
* полная пересборка нужна только когда меняется ПОДХОД. Прирост данных — повод
  дописать, а не собрать заново.
"""
from __future__ import annotations

import pandas as pd

from analysis.ingest.svod.exclusions import (
    MOUNTED_AFTER_HORIZON,
    NON_OIL_PURPOSE,
    REASON_COLUMN,
    STAGE_COLUMN,
    USABLE_COLUMN,
    ExclusionLedger,
    build_full_register,
    write_full_register,
)
from analysis.ingest.svod.versioning import (
    BUILD_RULES_VERSION,
    BuildMetadata,
    decide_build_mode,
    describe_source,
    rules_fingerprint,
    write_metadata_sheet,
)


def _rows(wells):
    return pd.DataFrame({"Скв.": wells, "Дата монтажа": ["2026-01-01"] * len(wells)})


# --------------------------------------------------------------------------- #
# Реестр отсева
# --------------------------------------------------------------------------- #

def test_every_dropped_row_carries_its_reason():
    ledger = ExclusionLedger()
    ledger.add(_rows(["Vt_1", "Vt_2"]), MOUNTED_AFTER_HORIZON)
    ledger.add(_rows(["Vt_3"]), NON_OIL_PURPOSE)

    assert ledger.total == 3
    frame = ledger.frame()
    assert set(frame[REASON_COLUMN]) == {MOUNTED_AFTER_HORIZON.code, NON_OIL_PURPOSE.code}
    assert set(frame[USABLE_COLUMN]) == {0}
    assert set(frame[STAGE_COLUMN]) == {MOUNTED_AFTER_HORIZON.stage, NON_OIL_PURPOSE.stage}


def test_the_summary_explains_each_reason_not_just_counts_it():
    """Счётчика мало: причина должна объяснять, почему строка непригодна."""
    ledger = ExclusionLedger()
    ledger.add(_rows(["Vt_1"]), MOUNTED_AFTER_HORIZON)
    summary = ledger.summary().set_index("причина")
    assert summary.loc[MOUNTED_AFTER_HORIZON.code, "строк"] == 1
    assert "ПДК догонит" in summary.loc[MOUNTED_AFTER_HORIZON.code, "объяснение"]


def test_the_full_register_keeps_usable_rows_first_and_marks_the_rest():
    ledger = ExclusionLedger()
    ledger.add(_rows(["Vt_9"]), MOUNTED_AFTER_HORIZON)
    full = build_full_register(_rows(["Vt_1", "Vt_2"]), ledger)

    assert len(full) == 3
    assert full[USABLE_COLUMN].tolist() == [1, 1, 0]
    assert full.loc[full["Скв."] == "Vt_9", REASON_COLUMN].iloc[0] == MOUNTED_AFTER_HORIZON.code
    assert full.loc[full["Скв."] == "Vt_1", REASON_COLUMN].isna().all()


def test_the_full_register_warns_in_the_file_itself(tmp_path):
    """⚠⚠ Файл обязан сам сообщать, что он не вход расчёта."""
    from openpyxl import load_workbook

    ledger = ExclusionLedger()
    ledger.add(_rows(["Vt_9"]), MOUNTED_AFTER_HORIZON)
    path = tmp_path / "full.xlsx"
    write_full_register(path, _rows(["Vt_1"]), ledger, horizon="2026-07-12")

    workbook = load_workbook(path)
    assert "Свод полный" in workbook.sheetnames
    assert "Причины отсева" in workbook.sheetnames
    banner = workbook["Свод полный"].cell(row=1, column=1).value
    assert "НЕ ВХОД РАСЧЁТА" in banner
    assert "2026-07-12" in banner
    # Пометки — первыми колонками, их видно без прокрутки.
    header = [cell.value for cell in workbook["Свод полный"][2]]
    assert header[:3] == [USABLE_COLUMN, REASON_COLUMN, STAGE_COLUMN]


# --------------------------------------------------------------------------- #
# Дописать или перестроить
# --------------------------------------------------------------------------- #

def _write_register(path, metadata: BuildMetadata):
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.active.title = "Свод"
    workbook.active.cell(row=1, column=1, value="Скв.")
    write_metadata_sheet(workbook, metadata)
    workbook.save(path)
    return path


def test_a_missing_register_means_full_rebuild(tmp_path):
    decision = decide_build_mode(tmp_path / "нет.xlsx")
    assert decision.full_rebuild
    assert "нет" in decision.reason


def test_a_register_without_a_passport_means_full_rebuild(tmp_path):
    """Регистр прежнего сборщика: чем размечены его строки — неизвестно."""
    from openpyxl import Workbook

    path = tmp_path / "старый.xlsx"
    workbook = Workbook()
    workbook.active.title = "Свод"
    workbook.save(path)

    decision = decide_build_mode(path)
    assert decision.full_rebuild
    assert "паспорт" in decision.reason


def test_unchanged_rules_and_a_moved_horizon_mean_appending(tmp_path):
    """Прирост данных — повод ДОПИСАТЬ, а не собирать заново."""
    path = _write_register(
        tmp_path / "свод.xlsx",
        BuildMetadata.current("2026-06-30", {"ПДК": "pdk"}),
    )
    decision = decide_build_mode(path, horizon="2026-07-12")
    assert not decision.full_rebuild
    assert "горизонт сдвинулся" in decision.reason


def test_an_unmoved_horizon_says_there_is_nothing_to_add(tmp_path):
    path = _write_register(
        tmp_path / "свод.xlsx",
        BuildMetadata.current("2026-07-12", {"ПДК": "pdk"}),
    )
    decision = decide_build_mode(path, horizon="2026-07-12")
    assert not decision.full_rebuild
    assert "дописывать нечего" in decision.reason


def test_changed_rules_force_a_full_rebuild(tmp_path):
    """⚠⚠ Иначе в одном файле окажутся две разные разметки."""
    stale = BuildMetadata.current("2026-06-30", {"ПДК": "pdk"})
    stale.rules_fingerprint = "0000000000000000"
    path = _write_register(tmp_path / "свод.xlsx", stale)

    decision = decide_build_mode(path, horizon="2026-07-12")
    assert decision.full_rebuild
    assert "изменился подход" in decision.reason


def test_the_fingerprint_follows_the_rule_tables(monkeypatch):
    """Отпечаток берётся с ТАБЛИЦ: правка правила его двигает, правка текста — нет."""
    before = rules_fingerprint()
    assert before == rules_fingerprint()          # устойчив

    import analysis.ingest.svod.causes as causes
    monkeypatch.setattr(causes, "CAUSE_RULES", causes.CAUSE_RULES + (("новое правило", "эксплуатация"),))
    assert rules_fingerprint() != before


def test_the_passport_records_source_versions(tmp_path):
    """«Почему числа поехали» без списка источников не отвечается."""
    metadata = BuildMetadata.current(
        "2026-07-12",
        {"паспорт оборудования": r"d:\big\WellsArtificialLiftBig 20260815.xlsx"},
    )
    path = _write_register(tmp_path / "свод.xlsx", metadata)
    restored = BuildMetadata.read(path)

    assert restored.rules_version == BUILD_RULES_VERSION
    assert restored.horizon == "2026-07-12"
    assert "версия 2026-08-15" in restored.sources["паспорт оборудования"]


def test_describe_source_reads_the_version_from_the_name():
    assert describe_source("big/WellsArtificialLiftBig 20260815.xlsx") == (
        "WellsArtificialLiftBig 20260815.xlsx (версия 2026-08-15)"
    )
    assert describe_source("big/WellsArtificialLiftBig.xlsx") == "WellsArtificialLiftBig.xlsx"
