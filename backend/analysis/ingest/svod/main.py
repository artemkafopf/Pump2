"""Main workflow for ESP failure register update"""

import pandas as pd
import openpyxl
from pathlib import Path
from typing import Dict, Optional, Tuple, List
from datetime import datetime
from copy import copy

from ..config import (
    COLUMN_SYNONYMS,
    DEFAULT_EVENT_SCOPE,
    DEFAULT_INCLUDE_EVENT_FAILURE_FLAG,
    DEFAULT_OPERATING_DATA_PRIMARY_SOURCE,
    DEFAULT_OPERATING_TIME_DERIVED_COLUMNS,
    DEFAULT_TECHREGIME_INTERVAL_DAYS,
    DEFAULT_TECHREGIME_USE_SQL_IF_AVAILABLE,
    FIELD_NAME_TO_CODE,
    LAB_CHEMISTRY_COLUMNS,
    OPERATING_DATA_PRIMARY_SOURCE_OPTIONS,
    SVOD_COMMENT_COLUMN,
    SVOD_RUNNING_COLUMN,
)
from ..io import (
    load_excel_file, load_workbook_sheet, save_workbook,
    append_column_to_worksheet, append_row_to_worksheet, create_audit_sheet, write_dataframe_to_sheet,
    highlight_runtime_cell_if_needed, apply_autofilter_and_autofit,
)
from ..normalize import normalize_text, normalize_well, parse_date, parse_number, get_column_by_synonym, strip_header_prefix
from .matching import DuplicateStatus, MatchResult, detect_duplicates as reconcile_new_records
from .well_identity import normalize_well_key, is_non_oil_purpose, is_rassol_bore
from .enrich import (
    enrich_from_telemetry,
)
from .pdk_processor import (
    derive_failure_flag,
    describe_horizon_completeness,
    extract_new_failures_from_pdk,
    pdk_observation_horizon,
)
from .enrich import derive_acid_type, derive_field_code_from_well, derive_runtime_group, normalize_contractor
from .causes import CAUSE_GROUP_COLUMN, DISPUTED_COLUMN, attach_cause_columns
from .presentation import mark_reference_headers, write_legend_sheet
from .versioning import BuildMetadata, write_metadata_sheet
from .exclusions import (
    BRINE_BORE,
    CLOSED_AFTER_HORIZON,
    MOUNTED_AFTER_HORIZON,
    NON_OIL_PURPOSE,
    RUNTIME_BEYOND_HORIZON,
    STALE_OPEN_RUN,
    SUPERSEDED_OPEN_RUN,
    ExclusionLedger,
    mark_reason,
)
from ..frac import frac_flags_for_runs
from .pump_specs import append_comment, fill_missing_pump_specs
from .artificial_lift_processor import (
    enrich_from_artificial_lift,
    enrich_big_derived_fields,
    extract_running_wells_from_artificial_lift,
)
from .techregime_processor import enrich_from_techregime
from .techregime_processor import get_default_techregime_selected_columns
from .opz_processor import enrich_from_opz
from .lab_processor import (
    enrich_from_lab,
    enrich_lab_chemistry
)
from ..operating_data import build_operating_features
from .audit import AuditReport, generate_field_fill_matrix
from .esp_specs import build_esp_specs_from_target, fill_specs_from_esp_type


def default_updated_workbook_path(target_path: str, update_mode: int = 0) -> str:
    """Build the default output path for the updated workbook copy."""
    path = Path(target_path)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return str(path.with_name(f"{path.stem}_{timestamp}_update={update_mode}.xlsx"))


class FailureUpdateWorkflow:
    """Main workflow for updating ESP failure register"""
    EVENT_FAILURE_FLAG_COLUMN = "Failure Flag"
    ENRICHMENT_QUERY_DATE_COLUMN = "techregime_query_date"
    
    def __init__(
        self,
        target_path: str,
        pdk_path: str,
        artificial_lift_path: Optional[str] = None,
        artificial_lift_manual_path: Optional[str] = None,
        big_path: Optional[str] = None,
        manual_path: Optional[str] = None,
        small_path: Optional[str] = None,
        techregime_path: Optional[str] = None,
        opz_clean_path: Optional[str] = None,
        opz_standard_path: Optional[str] = None,
        telemetry_vt_path: Optional[str] = None,
        telemetry_dir: Optional[str] = None,
        telemetry_db_path: Optional[str] = None,
        lab_path: Optional[str] = None,
        frac_path: Optional[str] = None,
        as_of=None,
        update_mode: int = 1,
        output_workbook: Optional[str] = None,
        audit_xlsx: Optional[str] = None,
        audit_md: Optional[str] = None,
        refresh_techregime_existing: bool = False,
        dry_run: bool = False,
        event_scope: str = DEFAULT_EVENT_SCOPE,
        exclude_ambiguous_events: bool = False,
        techregime_use_sql_if_available: bool = DEFAULT_TECHREGIME_USE_SQL_IF_AVAILABLE,
        techregime_interval=DEFAULT_TECHREGIME_INTERVAL_DAYS,
        techregime_selected_columns: Optional[List[str]] = None,
        include_event_failure_flag: bool = DEFAULT_INCLUDE_EVENT_FAILURE_FLAG,
        operating_data_primary_source: str = DEFAULT_OPERATING_DATA_PRIMARY_SOURCE,
        include_operating_time_derived: bool = False,
        apply_presentation: bool = False,
        normalize_svod_fields: bool = False,
    ):
        """
        Initialize workflow.
        
        Args:
            target_path: Path to target workbook
            pdk_path: Path to PDK source
            artificial_lift_path: Path to well artificial lift enrichment file
            artificial_lift_manual_path: Path to manual well artificial lift enrichment file
            big_path: Deprecated alias for artificial_lift_path
            manual_path: Deprecated alias for artificial_lift_manual_path
            small_path: Path to Small enrichment file
            techregime_path: Path to TechRegime enrichment file
            opz_clean_path: Path to OPZ clean database
            opz_standard_path: Path to OPZ standard database
            telemetry_vt_path: Path to Vt telemetry data
            telemetry_dir: Path to telemetry folder with one file per field
            lab_path: Path to lab chemistry workbook
            update_mode: 0=newer-than-last-date, 1=full update, 2=rebuild from scratch
            output_workbook: Output workbook path for the copied updated workbook
            audit_xlsx: Audit report Excel path
            audit_md: Audit report Markdown path
            refresh_techregime_existing: Refresh TechRegime-managed fields on existing target rows
            dry_run: If True, don't modify workbooks
            event_scope: Either "all" or "failures_only"
            exclude_ambiguous_events: Exclude rows where `failure_flag == -1`
            techregime_use_sql_if_available: Prefer TechRegime SQLite store over Excel when available
            techregime_interval: Lookback interval used to aggregate TechRegime data
            techregime_selected_columns: Target-facing TechRegime columns to populate
            include_event_failure_flag: Add/update `Failure Flag` based on the failure-flag rule
            operating_data_primary_source: Preferred source for overlapping operating metrics
            include_operating_time_derived: Add richer time-window operating features from telemetry/TechRegime
        """
        self.target_path = target_path
        self.pdk_path = pdk_path
        self.artificial_lift_path = artificial_lift_path or big_path
        self.artificial_lift_manual_path = artificial_lift_manual_path or manual_path
        self.small_path = small_path
        self.techregime_path = techregime_path
        self.opz_clean_path = opz_clean_path
        self.opz_standard_path = opz_standard_path
        self.telemetry_vt_path = telemetry_vt_path
        self.telemetry_dir = telemetry_dir
        # ⚠ The telemetry SQLite store does not live inside the export drop (the
        # drop is one dated sub-folder per export, the store sits beside them), so
        # it has to be passed explicitly rather than resolved from telemetry_dir.
        self.telemetry_db_path = telemetry_db_path
        self.lab_path = lab_path
        self.frac_path = frac_path
        # ⚠⚠ Горизонт наблюдения. None => берётся последняя дата ПДК; задать
        # явно имеет смысл, когда последний месяц выгрузки заведомо неполон.
        self.as_of = as_of
        self.observation_horizon = None
        self.update_mode = update_mode
        self.output_workbook = output_workbook or default_updated_workbook_path(target_path, update_mode)
        self.audit_xlsx = audit_xlsx or str(Path(target_path).parent / "audit.xlsx")
        self.audit_md = audit_md or str(Path(target_path).parent / "audit.md")
        self.refresh_techregime_existing = refresh_techregime_existing
        self.dry_run = dry_run
        self.event_scope = event_scope
        self.exclude_ambiguous_events = exclude_ambiguous_events
        self.techregime_use_sql_if_available = techregime_use_sql_if_available
        self.techregime_interval = techregime_interval
        self.techregime_selected_columns = techregime_selected_columns
        self.include_event_failure_flag = include_event_failure_flag
        self.operating_data_primary_source = (
            operating_data_primary_source
            if operating_data_primary_source in OPERATING_DATA_PRIMARY_SOURCE_OPTIONS
            else DEFAULT_OPERATING_DATA_PRIMARY_SOURCE
        )
        self.include_operating_time_derived = include_operating_time_derived
        self.apply_presentation = apply_presentation
        self.normalize_svod_fields = normalize_svod_fields

        self.audit = AuditReport()
        #: Что не попало в выходной Свод и почему — из него строится полный Свод.
        self.exclusions = ExclusionLedger()
        
        # Load target and PDK
        self.target_df = None
        self.target_sheet = None
        self.target_wb = None
        self.target_ws = None
        self.pdk_df = None
        self.new_failures = None
        self.esp_specs = {}
        self.last_target_failure_date = None
        self._append_style_template = None
        self._added_techregime_target_columns: List[str] = []
    
    def execute(self) -> Tuple[int, int, int]:
        """
        Execute full workflow.
        
        Returns:
            Tuple of (new_records_count, skipped_duplicates, likely_duplicates)
        """
        print(f"Starting failure update workflow...")
        print(f"Target: {self.target_path}")
        print(f"PDK source: {self.pdk_path}")
        print(f"Update mode: {self.update_mode}")
        
        # Step 1: Load target workbook
        print("\n[1/8] Loading target workbook...")
        self._load_target()
        
        # Step 2: Extract new failures from PDK
        print("[2/8] Extracting new failures from PDK...")
        self._extract_pdk()
        
        # Step 3: Detect duplicates
        print("[3/8] Detecting duplicates...")
        duplicates = self._detect_duplicates()
        
        # Step 4: Enrich from Big
        if self.artificial_lift_path:
            print("[4/8] Enriching from well artificial lift data...")
            self._enrich_artificial_lift()
        else:
            print("[4/8] Skipping well artificial lift enrichment (file not provided)")
        self._fill_esp_specs()
        
        # Step 5: Enrich from telemetry / TechRegime operating data
        print("[5/8] Enriching from operating data...")
        self._enrich_operating_data()
        self._fill_big_derived_fields()
        
        # Step 6: Enrich from OPZ
        if self.opz_clean_path:
            print("[6/8] Enriching from OPZ...")
            self._enrich_opz()
        else:
            print("[6/8] Skipping OPZ enrichment")
        
        # Step 7: Enrich from lab chemistry
        if self.lab_path:
            print("[7/8] Enriching from lab chemistry...")
            self._enrich_lab()
        else:
            print("[7/8] Skipping lab chemistry enrichment")
        
        # Step 8: Ensure lab chemistry columns exist
        print("[8/8] Finalizing lab chemistry columns...")
        self._mark_lab_chemistry()

        # Sour/non-sour flag from the raw PDK Месторождение ("…кислая") — must run
        # before _normalize_svod_fields rewrites the field to a code.
        self._apply_sour_flag()

        # Canonicalize register identity fields and backfill nominal specs.
        self._normalize_svod_fields()

        # Remove physically-impossible date/runtime values (the known ПДК defects)
        # so they never reach the produced register.
        self._sanitize_run_dates()

        # Keep only oil-production runs — drop injection/piezometric/brine wells
        # (the running-well set is not oil-filtered upstream like the PDK is).
        self._filter_oil_wells()

        # ⚠⚠ Обрезать регистр по горизонту наблюдения ПДК. Всё правее него —
        # экспозиция, в которой отказ записать нечем.
        self._apply_observation_horizon()

        # Fracturing history, joined on (скважина, дата монтажа).
        # ⚠ Runs AFTER _sanitize_run_dates on purpose: the "frac before the mount"
        # test anchors on the mount date, and sanitation is what reconstructs a
        # mount that the source recorded later than its own stop date. Anchoring
        # on the un-sanitized value would judge those runs against a date that
        # never happened.
        self._enrich_frac()

        # Mark the runs that are still turning, and label every stop's cause.
        # Both run last: they read the sanitized, filtered rows.
        self._mark_running_runs()
        self._attach_cause_taxonomy()

        # Append and report
        if not self.dry_run:
            print("\nAppending records to workbook...")
            self._append_records(duplicates)
        else:
            print("\nDry run: skipping append")
        
        print("\nGenerating audit reports...")
        self._generate_audit()
        
        # Summary
        new_count = len(self.new_failures)
        skip_exact = sum(1 for r in duplicates.values() if r.status == DuplicateStatus.EXACT_EXISTING)
        skip_likely = sum(1 for r in duplicates.values() if r.status == DuplicateStatus.LIKELY_EXISTING)
        
        print(f"\n✓ Complete!")
        print(f"  - Records extracted from PDK: {len(self.pdk_df)}")
        print(f"  - Already exist (exact): {skip_exact}")
        print(f"  - Likely exist (review): {skip_likely}")
        print(f"  - New records to append: {new_count - skip_exact - skip_likely}")
        
        return new_count, skip_exact, skip_likely
    
    def _load_target(self):
        """Load target workbook.

        ``load_workbook_sheet`` already parses the sheet with pandas (auto-detecting
        the sheet) *and* opens the openpyxl workbook, so we call it once instead of
        also parsing separately with ``load_excel_file`` first.
        """
        self.target_df, self.target_sheet, self.target_wb, self.target_ws = load_workbook_sheet(
            self.target_path
        )
        
        print(f"  Loaded {len(self.target_df)} existing records from sheet '{self.target_sheet}'")
        
        self.audit.add_summary_stat("Target workbook rows", len(self.target_df))
        self.audit.add_summary_stat("Target sheet", self.target_sheet)
        self.esp_specs = build_esp_specs_from_target(self.target_df)
        self.audit.add_summary_stat("ESP spec mappings", len(self.esp_specs))
        self.last_target_failure_date = self._detect_last_target_failure_date()
        if self.last_target_failure_date is not None:
            self.audit.add_summary_stat("Last target failure date", self.last_target_failure_date.date())
        if self.include_event_failure_flag:
            added = self._ensure_event_failure_flag_column()
            self._populate_existing_event_failure_flag_column()
            if added:
                print(f"  Added target column '{self.EVENT_FAILURE_FLAG_COLUMN}'")
        if self.update_mode == 2:
            self._prepare_rework_target()

    def _extract_pdk(self):
        """Extract new failures from PDK"""
        min_failure_date = self.last_target_failure_date if self.update_mode == 0 else None
        self.pdk_df = extract_new_failures_from_pdk(self.pdk_path, min_failure_date, audit=self.audit)
        self.new_failures = self.pdk_df.copy()
        self._resolve_observation_horizon()
        self._append_running_wells_from_artificial_lift()
        if self.include_event_failure_flag and "failure_flag" in self.new_failures.columns:
            self.new_failures[self.EVENT_FAILURE_FLAG_COLUMN] = self.new_failures["failure_flag"]

        if self.event_scope == "failures_only" and "failure_flag" in self.new_failures.columns:
            before_count = len(self.new_failures)
            self.new_failures = self.new_failures[self.new_failures["failure_flag"] == 1].copy()
            self.pdk_df = self.new_failures.copy()
            print(f"  Failure-only filter kept {len(self.new_failures)} of {before_count} extracted row(s)")
            self.audit.add_summary_stat("PDK rows after failure-only filter", len(self.new_failures))
        elif self.exclude_ambiguous_events and "failure_flag" in self.new_failures.columns:
            before_count = len(self.new_failures)
            self.new_failures = self.new_failures[self.new_failures["failure_flag"] != -1].copy()
            self.pdk_df = self.new_failures.copy()
            print(f"  Excluded ambiguous events (`failure_flag = -1`), kept {len(self.new_failures)} of {before_count} row(s)")
            self.audit.add_summary_stat("Rows after excluding ambiguous events", len(self.new_failures))

        if min_failure_date is not None:
            print(f"  Extracted {len(self.pdk_df)} records from PDK after {min_failure_date.date()}")
            self.audit.add_summary_stat("PDK records after last target date", len(self.pdk_df))
        else:
            print(f"  Extracted {len(self.pdk_df)} records from PDK")
            self.audit.add_summary_stat("PDK records extracted", len(self.pdk_df))

    def source_versions(self) -> dict:
        """Какие именно файлы пошли в сборку — с датой актуализации из имени.

        Регистр обязан помнить не только ЧТО в нём, но и ИЗ ЧЕГО он собран:
        иначе «почему числа поехали» не отвечается вообще.
        """
        return {
            name: value
            for name, value in (
                ("ПДК", self.pdk_path),
                ("паспорт оборудования", self.artificial_lift_path),
                ("ТехРежим", self.techregime_path),
                ("телеметрия", self.telemetry_dir),
                ("телеметрия (БД)", self.telemetry_db_path),
                ("ОПЗ", self.opz_clean_path),
                ("лаборатория", self.lab_path),
                ("ГРП (БД)", self.frac_path),
            )
            if value
        }

    def _resolve_observation_horizon(self) -> None:
        """Set the date the whole register is observed up to.

        ⚠⚠ Отказы приходят ТОЛЬКО из ПДК, а живые пуски — из паспорта
        оборудования, который обновляется отдельно. Значит горизонт наблюдения
        задаёт ПДК, и только он: правее его последней записи можно накапливать
        наработку, но записать отказ нечем. Цензурировать живые пуски сегодняшним
        днём — значит выдать им бесплатную выживаемость.
        """
        if self.as_of is not None:
            self.observation_horizon = parse_date(self.as_of)
            print(f"  Observation horizon: {self.observation_horizon.date()} (задан явно)")
            return

        self.observation_horizon = pdk_observation_horizon(self.pdk_df)
        if self.observation_horizon is None:
            self.observation_horizon = pd.Timestamp.now().normalize()
            print(
                "  ⚠ В ПДК нет ни одной даты остановки — горизонт взят по сегодняшнему дню "
                f"({self.observation_horizon.date()}); живые пуски получат наработку без покрытия отказами"
            )
            return

        print(f"  Observation horizon (последняя дата ПДК): {self.observation_horizon.date()}")
        warning = describe_horizon_completeness(self.pdk_df, self.observation_horizon)
        if warning:
            print(f"  ⚠ {warning}")
            self.audit.add_summary_stat("Horizon completeness warning", warning)
        self.audit.add_summary_stat("Observation horizon", self.observation_horizon.date())

    def _append_running_wells_from_artificial_lift(self) -> None:
        """Append active pump runs from the artificial-lift register for `include all`."""
        if self.event_scope != "all" or not self.artificial_lift_path:
            return

        dropped_active: dict = {}
        running_records = extract_running_wells_from_artificial_lift(
            self.artificial_lift_path,
            manual_path=self.artificial_lift_manual_path,
            # ⚠ Живой пуск цензурируется горизонтом ПДК, а не сегодняшним днём.
            as_of_date=self.observation_horizon,
            excluded=dropped_active,
        )
        for key, reason in (
            ("brine", BRINE_BORE),
            ("non_oil", NON_OIL_PURPOSE),
            ("stale", STALE_OPEN_RUN),
            ("superseded", SUPERSEDED_OPEN_RUN),
        ):
            self.exclusions.add(dropped_active.get(key), reason)
        if running_records.empty:
            return

        if self.new_failures is not None and not self.new_failures.empty:
            existing_pairs = {
                (
                    normalize_well(row.get("well", row.get("Скв."))),
                    parse_date(row.get("installation_date", row.get("Дата монтажа"))),
                )
                for _, row in self.new_failures.iterrows()
            }
            running_records = running_records[
                ~running_records.apply(
                    lambda row: (
                        normalize_well(row.get("well", row.get("Скв."))),
                        parse_date(row.get("installation_date", row.get("Дата монтажа"))),
                    ) in existing_pairs,
                    axis=1,
                )
            ].copy()
        if running_records.empty:
            return

        if self.new_failures is None or self.new_failures.empty:
            self.new_failures = running_records.copy()
        else:
            # Concatenating frames that carry all-NA columns is deprecated: pandas
            # currently ignores such columns when inferring the result dtype and
            # will stop doing so. Drop them before the concat and restore them
            # afterwards by reindexing to the union — same columns, same values,
            # dtypes inferred from the rows that actually hold data.
            columns = list(dict.fromkeys([*self.new_failures.columns, *running_records.columns]))
            self.new_failures = pd.concat(
                [
                    self.new_failures.dropna(axis=1, how="all"),
                    running_records.dropna(axis=1, how="all"),
                ],
                ignore_index=True,
                sort=False,
            ).reindex(columns=columns)
        print(f"  Added {len(running_records)} running well row(s) from artificial lift")
        self.audit.add_summary_stat("Running wells added from artificial lift", len(running_records))

    def _detect_duplicates(self) -> Dict:
        """Reconcile incoming rows against the target through a single well-keyed index.

        Returns ``{new_failures_index: MatchResult}``. Matching is delegated to
        :func:`matching.detect_duplicates`, which builds one index keyed by the
        suffix-preserving :func:`normalize_well_key` and classifies each candidate
        as EXACT_EXISTING (same well + failure/install date), LIKELY_EXISTING
        (within tolerance -> needs review), or NEW. In-batch echoes of a NEW row
        are EXACT_EXISTING with ``matching_row_idx`` None.
        """
        if self.update_mode == 2:
            print("  Rework mode: skipping duplicate detection against target")
            return {}

        well_col = self._find_target_well_column()
        failure_date_col = self._find_target_failure_date_column()

        if not well_col or not failure_date_col:
            print("  Warning: Could not identify well or date columns in target")
            return {}

        install_col = next(
            (c for c in self.target_df.columns if "дата монтажа" in str(c).lower()),
            None,
        )
        nno_col = next(
            (c for c in self.target_df.columns if "наработка" in str(c).lower()),
            None,
        )

        duplicates: Dict[int, MatchResult] = reconcile_new_records(
            self.new_failures,
            self.target_df,
            well_col=well_col,
            failure_date_col=failure_date_col,
            install_date_col=install_col,
            nno_col=nno_col,
        )

        exact_count = sum(1 for r in duplicates.values() if r.status == DuplicateStatus.EXACT_EXISTING)
        likely_count = sum(1 for r in duplicates.values() if r.status == DuplicateStatus.LIKELY_EXISTING)

        print(
            f"  Reconciled against target: {exact_count} exact match(es), "
            f"{likely_count} within-tolerance (needs review)"
        )

        self.audit.add_summary_stat("Exact duplicates found", exact_count)
        self.audit.add_summary_stat("Likely duplicates found", likely_count)

        return duplicates

    def _find_target_well_column(self):
        return next(
            (c for c in self.target_df.columns if "скв" in str(c).lower()),
            None
        )

    def _find_target_failure_date_column(self):
        candidate_columns = [
            c for c in self.target_df.columns
            if "дата" in str(c).lower() and ("останов" in str(c).lower() or "отказ" in str(c).lower())
        ]
        if candidate_columns:
            return candidate_columns[0]
        return next(
            (c for c in self.target_df.columns if "дата" in str(c).lower()),
            None
        )
    
    def _enrich_artificial_lift(self):
        """Enrich from well artificial lift data."""
        if self.artificial_lift_path:
            self.new_failures = enrich_from_artificial_lift(
                self.new_failures,
                self.artificial_lift_path,
                manual_path=self.artificial_lift_manual_path
            )
            print("  Enriched from well artificial lift data")
    
    def _enrich_techregime(self):
        """Enrich from TechRegime"""
        if self.techregime_path:
            self._ensure_enrichment_query_dates()
            self.new_failures = enrich_from_techregime(
                self.new_failures,
                self.techregime_path,
                failure_date_col=self.ENRICHMENT_QUERY_DATE_COLUMN,
                prefer_sqlite=self.techregime_use_sql_if_available,
                interval=self.techregime_interval,
                selected_columns=self._selected_techregime_target_columns(),
            )
            print("  Enriched from TechRegime")

    def _enrich_operating_data(self):
        """Enrich overlapping operating metrics from telemetry and TechRegime with source priority."""
        telemetry_available = bool(self.telemetry_dir or self.telemetry_vt_path or self.telemetry_db_path)
        techregime_available = bool(self.techregime_path)

        if not telemetry_available and not techregime_available:
            print("  Skipping operating-data enrichment (no telemetry or TechRegime source)")
            return

        if techregime_available:
            self._added_techregime_target_columns = self._ensure_selected_techregime_target_columns()
            if self._added_techregime_target_columns:
                print(
                    "  Added target column(s) for selected TechRegime data: "
                    + ", ".join(self._added_techregime_target_columns)
                )

        order: list[str] = []
        if self.operating_data_primary_source == "telemetry":
            if telemetry_available:
                order.append("telemetry")
            if techregime_available:
                order.append("techregime")
        else:
            if techregime_available:
                order.append("techregime")
            if telemetry_available:
                order.append("telemetry")

        if not order:
            print("  Skipping operating-data enrichment (source selection resolved to none)")
            return

        print(
            "  Operating-data source priority: "
            + " -> ".join(order)
            + " (later source only fills missing values when possible)"
        )
        for source_name in order:
            if source_name == "telemetry":
                self._enrich_telemetry()
            elif source_name == "techregime":
                self._enrich_techregime()

        if techregime_available and (self.refresh_techregime_existing or self._added_techregime_target_columns):
            self._refresh_existing_target_techregime_fields()

        if self.include_operating_time_derived:
            self._enrich_operating_time_derived()

    def _selected_techregime_target_columns(self) -> list[str]:
        if self.techregime_selected_columns is None:
            return list(get_default_techregime_selected_columns())
        return [str(column).strip() for column in self.techregime_selected_columns if str(column).strip()]

    def _ensure_selected_techregime_target_columns(self) -> list[str]:
        """Create any selected TechRegime columns that are not yet present in the target."""
        if self.target_df is None:
            return []

        existing_by_key = {
            strip_header_prefix(str(column)): str(column)
            for column in self.target_df.columns
        }
        added_columns: list[str] = []
        for selected_column in self._selected_techregime_target_columns():
            target_key = strip_header_prefix(selected_column)
            if target_key in existing_by_key:
                continue
            self.target_df[selected_column] = None
            if self.target_ws is not None:
                source_column_idx = self.target_ws.max_column if self.target_ws.max_column >= 1 else None
                append_column_to_worksheet(
                    self.target_ws,
                    selected_column,
                    header_row=1,
                    copy_format_from_column=source_column_idx,
                )
            added_columns.append(selected_column)
            existing_by_key[target_key] = selected_column
        return added_columns

    def _refresh_existing_target_techregime_fields(self):
        """Refresh TechRegime-managed fields on existing target rows in memory."""
        if not self.techregime_path or self.target_df is None or self.target_df.empty:
            print("  Skipping existing-target TechRegime refresh (no target rows)")
            return

        def has_value(value) -> bool:
            if value is None:
                return False
            if isinstance(value, str):
                return normalize_text(value) not in {"", "nan", "none", "-", "–"}
            return not pd.isna(value)

        def get_prefixed_target_series(canonical_name: str, fallback_synonyms: list[str] | None = None):
            synonym_values = list(COLUMN_SYNONYMS.get(canonical_name, []))
            if fallback_synonyms:
                synonym_values.extend(fallback_synonyms)
            normalized_synonyms = {strip_header_prefix(value) for value in synonym_values}
            for column in self.target_df.columns:
                if strip_header_prefix(column) in normalized_synonyms:
                    return self.target_df[column]
            return None

        well_series = get_prefixed_target_series("well", ["скв", "скважина", "№ скважины", "well"])
        failure_date_series = get_prefixed_target_series("failure_date")
        runtime_series = get_prefixed_target_series("runtime_nno", ["наработка", "наработка (сут)", "наработка сут"])

        if well_series is None or failure_date_series is None:
            print("  Skipping existing-target TechRegime refresh (well/date columns not found)")
            return

        refresh_seed_data = {
            "well": well_series,
            "failure_date": failure_date_series,
        }
        if runtime_series is not None:
            refresh_seed_data["runtime_nno"] = runtime_series
        refresh_seed = pd.DataFrame(refresh_seed_data, index=self.target_df.index)
        refreshed = enrich_from_techregime(
            refresh_seed,
            self.techregime_path,
            well_col="well",
            failure_date_col="failure_date",
            prefer_sqlite=self.techregime_use_sql_if_available,
            interval=self.techregime_interval,
            selected_columns=self._selected_techregime_target_columns(),
        )

        managed_target_keys = {
            strip_header_prefix(column)
            for column in self._selected_techregime_target_columns()
        }
        header_lookup = {
            str(cell.value).strip(): col_idx
            for col_idx, cell in enumerate(self.target_ws[1], start=1)
            if cell.value is not None
        } if self.target_ws is not None else {}

        updated_cells = 0
        updated_rows = set()
        for row_idx, refreshed_row in refreshed.iterrows():
            mapped_values = self._map_columns(refreshed_row.to_dict())
            worksheet_row = row_idx + 2
            for target_col, mapped_value in zip(self.target_df.columns, mapped_values):
                target_key = strip_header_prefix(str(target_col))
                if target_key not in managed_target_keys:
                    continue
                if not has_value(mapped_value):
                    continue
                self.target_df.at[row_idx, target_col] = mapped_value
                updated_rows.add(row_idx)
                if self.target_ws is not None and target_col in header_lookup:
                    self.target_ws.cell(row=worksheet_row, column=header_lookup[target_col]).value = mapped_value
                updated_cells += 1

        print(
            "  Refreshed TechRegime-managed fields on existing target rows "
            f"({len(updated_rows)} rows, {updated_cells} cells)"
        )
        self.audit.add_summary_stat("Existing target rows refreshed from TechRegime", len(updated_rows))
        self.audit.add_summary_stat("Existing target cells refreshed from TechRegime", updated_cells)
    
    def _enrich_opz(self):
        """Enrich from OPZ"""
        if self.opz_clean_path:
            self.new_failures = enrich_from_opz(
                self.new_failures,
                self.opz_clean_path
            )
            print("  Enriched OPZ flags")
        elif self.opz_standard_path:
            self.new_failures = enrich_from_opz(
                self.new_failures,
                self.opz_standard_path
            )
            print("  Enriched OPZ flags (using standard database)")
    
    def _enrich_telemetry(self):
        """Enrich from telemetry"""
        telemetry_source = self.telemetry_dir or self.telemetry_vt_path
        if telemetry_source or self.telemetry_db_path:
            self._ensure_enrichment_query_dates()
            self.new_failures = enrich_from_telemetry(
                self.new_failures,
                telemetry_source,
                failure_date_col=self.ENRICHMENT_QUERY_DATE_COLUMN,
                sqlite_path=self.telemetry_db_path,
            )
            print("  Enriched from telemetry")

    def _fill_esp_specs(self):
        """Fill gabarit/productivity from ESP type mapping derived from target data."""
        if self.esp_specs:
            self.new_failures = fill_specs_from_esp_type(self.new_failures, self.esp_specs)
            print("  Filled ESP specs from current register mapping")

    def _ensure_enrichment_query_dates(self) -> None:
        """Create a query date used for TechRegime/telemetry extraction.

        Failures keep their actual stop date. Running wells use installation date
        plus current runtime so the averaging window spans the active run.
        """
        if self.new_failures is None or self.new_failures.empty:
            return
        query_dates = []
        for _, row in self.new_failures.iterrows():
            failure_date = parse_date(row.get("failure_date"))
            if failure_date is not None:
                query_dates.append(failure_date)
                continue
            installation_date = parse_date(row.get("installation_date", row.get("Дата монтажа")))
            runtime_nno = row.get("runtime_nno", row.get("Наработка (сут)"))
            runtime_days = pd.to_numeric(pd.Series([runtime_nno]), errors="coerce").iloc[0]
            if installation_date is not None and pd.notna(runtime_days):
                query_dates.append(installation_date + pd.Timedelta(days=float(runtime_days)))
                continue
            launch_date = parse_date(row.get("launch_date", row.get("Дата запуска")))
            query_dates.append(launch_date or installation_date)
        self.new_failures[self.ENRICHMENT_QUERY_DATE_COLUMN] = query_dates

    def _get_target_column_name_by_key(self, normalized_key: str) -> Optional[str]:
        if self.target_df is None:
            return None
        for column in self.target_df.columns:
            if strip_header_prefix(str(column)) == normalized_key:
                return str(column)
        return None

    def _ensure_event_failure_flag_column(self) -> bool:
        """Ensure an event failure-flag column exists in the target.

        Only add the `Failure Flag` column when the incoming target carries no
        event-flag column at all -- an existing English `Failure Flag` or
        Russian `Флаг отказа` is treated as already present and left untouched.
        """
        if self.target_df is None:
            return False
        existing_flag_keys = {
            strip_header_prefix(self.EVENT_FAILURE_FLAG_COLUMN),
            strip_header_prefix("Флаг отказа"),
        }
        if any(self._get_target_column_name_by_key(key) is not None for key in existing_flag_keys):
            return False
        self.target_df[self.EVENT_FAILURE_FLAG_COLUMN] = None
        if self.target_ws is not None:
            source_column_idx = self.target_ws.max_column if self.target_ws.max_column >= 1 else None
            append_column_to_worksheet(
                self.target_ws,
                self.EVENT_FAILURE_FLAG_COLUMN,
                header_row=1,
                copy_format_from_column=source_column_idx,
            )
        return True

    def _populate_existing_event_failure_flag_column(self) -> None:
        """Fill `Failure Flag` on existing target rows."""
        if not self.include_event_failure_flag or self.target_df is None or self.target_df.empty:
            return

        event_column = self._get_target_column_name_by_key(strip_header_prefix(self.EVENT_FAILURE_FLAG_COLUMN))
        if event_column is None:
            return

        failure_flag_column = self._get_target_column_name_by_key("флаг отказа")
        failure_reason_column = self._get_target_column_name_by_key("причина остановки")
        failed_node_column = self._get_target_column_name_by_key("отказавший узел")
        header_lookup = {
            str(cell.value).strip(): col_idx
            for col_idx, cell in enumerate(self.target_ws[1], start=1)
            if cell.value is not None
        } if self.target_ws is not None else {}

        for row_idx, row in self.target_df.iterrows():
            if failure_flag_column is not None:
                event_value = row.get(failure_flag_column)
                if pd.isna(event_value):
                    event_value = derive_failure_flag(
                        row.get(failure_reason_column) if failure_reason_column is not None else None,
                        row.get(failed_node_column) if failed_node_column is not None else None,
                    )
            else:
                event_value = derive_failure_flag(
                    row.get(failure_reason_column) if failure_reason_column is not None else None,
                    row.get(failed_node_column) if failed_node_column is not None else None,
                )
            self.target_df.at[row_idx, event_column] = event_value
            if self.target_ws is not None and event_column in header_lookup:
                self.target_ws.cell(row=row_idx + 2, column=header_lookup[event_column]).value = event_value

    def _fill_big_derived_fields(self):
        """Fill Big-derived computed fields such as liquid-rate delta."""
        self.new_failures = enrich_big_derived_fields(self.new_failures)

    def _apply_sour_flag(self):
        """Copy the PDK-derived sour flag (``acid_type``) into ``Кислый/Некислый``.

        ``extract_new_failures_from_pdk`` computes ``acid_type`` from the raw PDK
        field ("ВТЛУ кислая" → Кислый) *before* Месторождение is reduced to a
        field code, so it is the ground-truth sour flag for every run — it just
        needs to land in the register column, which had been left Некислый.
        """
        if self.new_failures is None or self.new_failures.empty:
            return
        if "acid_type" not in self.new_failures.columns:
            return
        per_row = self.new_failures["acid_type"].fillna("Некислый")
        # Sour is a property of the WELL. Running wells arrive with a field *code*
        # ("Vt", no "кислая") so their per-row acid_type is Некислый — propagate the
        # flag across every run of any well that is sour on any row.
        well = self.new_failures["Скв."].map(normalize_well)
        sour_wells = set(well[per_row == "Кислый"]) - {""}
        flag = well.map(lambda w: "Кислый" if w in sour_wells else "Некислый")
        self.new_failures["Кислый/Некислый"] = flag
        sour_n = int((flag == "Кислый").sum())
        print(f"  Sour flag: {sour_n} run(s) across {len(sour_wells)} sour well(s) "
              f"(well-propagated from PDK acid_type)")

    def _apply_observation_horizon(self):
        """Обрезать регистр по горизонту наблюдения ПДК.

        Три операции, и все три — про одно: не отдавать в расчёт экспозицию, в
        которой событие не могло быть записано.

        * **закрытие за горизонтом ГАСИТСЯ, а не удаляется вместе со строкой.**
          Дата остановки или демонтажа правее горизонта — это событие, которого мы
          не наблюдаем; строка становится цензурированной на горизонте, а её
          прожитая жизнь сохраняется. Удалять пуск было бы хуже всего: он реально
          отработал, и выбросить его — значит потерять экспозицию, которую мы как
          раз наблюдали;
        * **пуск смонтирован после горизонта** — выбрасывается целиком. Тут терять
          нечего: наблюдаемой жизни у него нет вовсе, а в выборке он работает
          чистым разбавителем и тянет выживаемость вверх в ранней полосе;
        * **живой пуск с наработкой за горизонт** — наработка обрезается по
          горизонту. В нынешнем паспорте ННО нет вовсе, и все 100 % живых строк
          получают её как календарный возраст, так что без обрезки каждая несла бы
          разницу «сегодня − горизонт» бесплатной выживаемостью.
        """
        frame = self.new_failures
        if frame is None or frame.empty or self.observation_horizon is None:
            return
        horizon = self.observation_horizon

        # --- 1. Погасить закрытия, которых регистр не наблюдает ----------------
        stop_columns = [c for c in ("Дата остановки", "failure_date") if c in frame.columns]
        dismantle_columns = [c for c in ("Дата демонтажа", "dismantling_date") if c in frame.columns]
        flag_columns = [c for c in ("Флаг отказа", "failure_flag") if c in frame.columns]

        def _beyond(columns) -> pd.Series:
            mask = pd.Series(False, index=frame.index)
            for column in columns:
                values = pd.to_datetime(frame[column], errors="coerce")
                mask = mask | (values.notna() & (values > horizon))
            return mask

        stopped_after_mask = _beyond(stop_columns)

        # ⚠ Демонтаж за горизонтом гасится ТОЛЬКО там, где он единственный признак
        # закрытия. Демонтаж — не событие: событие это остановка. Если остановка
        # наблюдаема, поздний демонтаж — просто физический подъём через несколько
        # дней, и он записан самим ПДК. Безусловное гашение стирало 5 совершенно
        # настоящих дат (Vt_5202, Au_389, Vt_015вз, Vt_3304, Ya_388 — остановки
        # 2026-05-31…2026-07-12, подъёмы 13–14 июля). Гасить надо лишь тогда, когда
        # строка осталась цензурированной: иначе дата демонтажа противоречила бы
        # колонке «Работает».
        observed_stop = pd.Series(False, index=frame.index)
        for column in stop_columns:
            values = pd.to_datetime(frame[column], errors="coerce")
            observed_stop = observed_stop | (values.notna() & (values <= horizon))
        dismantled_after_mask = _beyond(dismantle_columns) & ~observed_stop

        censored = int(stopped_after_mask.sum())
        for column, mask in (
            *((c, stopped_after_mask) for c in stop_columns),
            *((c, dismantled_after_mask) for c in dismantle_columns),
        ):
            if not mask.any():
                continue
            # An all-empty date column arrives as float64 and rejects NaT under
            # pandas' incompatible-dtype deprecation; make it object first.
            if frame[column].dtype != object and not pd.api.types.is_datetime64_any_dtype(frame[column]):
                frame[column] = frame[column].astype(object)
            frame.loc[mask, column] = pd.NaT
        # Событие за горизонтом — не событие: строка цензурирована.
        for column in flag_columns:
            frame.loc[stopped_after_mask, column] = 0
        mark_reason(frame, stopped_after_mask | dismantled_after_mask, CLOSED_AFTER_HORIZON)
        self.exclusions.note(CLOSED_AFTER_HORIZON, censored + int(dismantled_after_mask.sum()))
        if censored or dismantled_after_mask.any():
            print(
                f"  Horizon rule: censored {censored} row(s) whose stop is after "
                f"{horizon.date()} and cleared {int(dismantled_after_mask.sum())} "
                "unobservable dismantling date(s)"
            )
        self.audit.add_summary_stat("Rows censored at horizon", censored)

        # --- 2. Выбросить пуски, начавшиеся за горизонтом ----------------------
        mount = pd.to_datetime(frame.get("Дата монтажа"), errors="coerce") if "Дата монтажа" in frame.columns else None
        mounted_after = 0
        if mount is not None:
            mounted_after_mask = mount.notna() & (mount > horizon)
            mounted_after = int(mounted_after_mask.sum())
            if mounted_after:
                self.exclusions.add(frame.loc[mounted_after_mask], MOUNTED_AFTER_HORIZON)
                frame = frame.loc[~mounted_after_mask]
                self.new_failures = frame
                print(
                    f"  Horizon rule: dropped {mounted_after} run(s) mounted after "
                    f"{horizon.date()} (запуск без наблюдаемого исхода); {len(frame)} remain"
                )
        self.audit.add_summary_stat("Runs dropped: mounted after horizon", mounted_after)

        # Обрезка наработки открытых пусков по горизонту.
        runtime_columns = [c for c in ("Наработка (сут)", "runtime_nno") if c in frame.columns]
        group_columns = [c for c in ("группа наработок", "runtime_group") if c in frame.columns]
        if not runtime_columns:
            return
        launch = pd.to_datetime(frame.get("Дата запуска"), errors="coerce") if "Дата запуска" in frame.columns else None
        mount = pd.to_datetime(frame.get("Дата монтажа"), errors="coerce") if "Дата монтажа" in frame.columns else None
        stop = pd.to_datetime(frame.get("Дата остановки"), errors="coerce") if "Дата остановки" in frame.columns else None
        if mount is None:
            return
        anchor = launch.fillna(mount) if launch is not None else mount
        open_mask = stop.isna() if stop is not None else pd.Series(True, index=frame.index)
        max_runtime = (horizon - anchor).dt.days

        trimmed = 0
        trimmed_days = 0.0
        for idx in frame.index[open_mask & anchor.notna()]:
            limit = max_runtime.at[idx]
            if pd.isna(limit) or limit < 0:
                continue
            current = parse_number(frame.at[idx, runtime_columns[0]])
            if current is None or current <= limit:
                continue
            for column in runtime_columns:
                frame.at[idx, column] = float(limit)
            new_group = derive_runtime_group(limit)
            for column in group_columns:
                frame.at[idx, column] = new_group
            trimmed += 1
            trimmed_days += float(current - limit)

        if trimmed:
            self.exclusions.note(RUNTIME_BEYOND_HORIZON, trimmed)
            print(
                f"  Horizon filter: trimmed runtime on {trimmed} open run(s) to "
                f"{horizon.date()} (−{trimmed_days:.0f} run-day(s) of unobservable exposure)"
            )
        self.audit.add_summary_stat("Open runs trimmed to horizon", trimmed)
        self.audit.add_summary_stat("Unobservable run-days removed", int(trimmed_days))

    def _mark_running_runs(self):
        """Flag the runs that are still turning, by an EMPTY stop date.

        ⚠⚠ ``Флаг отказа == 0`` is NOT "running": non-failure pulls (ГТМ, ППР)
        also carry a zero, and there are ~1 340 of them. Reading the zero as
        "alive" inflates the live fleet from 511 to 1 851 and silently turns
        completed runs into censored ones. The empty stop date is the only
        signal that means the pump has not been pulled.
        """
        frame = self.new_failures
        if frame is None or frame.empty:
            return
        stop_columns = [c for c in ("Дата остановки", "failure_date") if c in frame.columns]
        dismantle_columns = [c for c in ("Дата демонтажа", "dismantling_date") if c in frame.columns]

        def _is_open(row) -> int:
            for column in stop_columns + dismantle_columns:
                if parse_date(row.get(column)) is not None:
                    return 0
            return 1

        frame[SVOD_RUNNING_COLUMN] = [_is_open(row) for _, row in frame.iterrows()]
        running = int(frame[SVOD_RUNNING_COLUMN].sum())
        zero_flag = (
            int((pd.to_numeric(frame.get("failure_flag"), errors="coerce") == 0).sum())
            if "failure_flag" in frame.columns else 0
        )
        print(
            f"  Running runs (empty stop date): {running} "
            f"— for contrast, {zero_flag} row(s) carry failure_flag == 0"
        )
        self.audit.add_summary_stat("Running runs (empty stop date)", running)
        self.audit.add_summary_stat("Rows with failure_flag == 0", zero_flag)

    def _attach_cause_taxonomy(self):
        """Label every row with its cause group, disputed flag and node category."""
        if self.new_failures is None or self.new_failures.empty:
            return
        self.new_failures = attach_cause_columns(self.new_failures)
        counts = self.new_failures[CAUSE_GROUP_COLUMN].value_counts()
        for group, count in counts.items():
            self.audit.add_summary_stat(f"Причина: {group}", int(count))
        self.audit.add_summary_stat(
            "Причина: спорная зона",
            int(pd.to_numeric(self.new_failures[DISPUTED_COLUMN], errors="coerce").fillna(0).sum()),
        )

    def _enrich_frac(self):
        """Join per-run fracturing flags from the ГРП register."""
        if self.new_failures is None or self.new_failures.empty:
            return
        if not self.frac_path:
            print("  Skipping frac enrichment (no ГРП source)")
            return
        try:
            flags = frac_flags_for_runs(self.new_failures, sqlite_path=self.frac_path)
        except Exception as exc:
            print(f"  Skipped frac enrichment: {exc}")
            return
        for column in flags.columns:
            self.new_failures[column] = flags[column]
        with_frac = int(pd.to_numeric(flags["ГРП"], errors="coerce").fillna(0).sum())
        before_mount = int(pd.to_numeric(flags["ГРП до монтажа"], errors="coerce").fillna(0).sum())
        print(f"  ГРП: {with_frac} run(s) on fracced wells, {before_mount} with a frac before the mount")
        self.audit.add_summary_stat("Runs on fracced wells", with_frac)
        self.audit.add_summary_stat("Runs with a frac before the mount", before_mount)

    def _filter_oil_wells(self):
        """Drop non-oil-production runs so the register keeps only oil wells.

        PDK rows are oil-filtered upstream, but the running wells appended from the
        artificial-lift register are not. Remove injection / piezometric / etc.
        bores (by wellbore type or назначение) and brine (рс) bores.
        """
        if self.new_failures is None or self.new_failures.empty:
            return
        df = self.new_failures
        purpose_col = next(
            (c for c in df.columns if "тип ствол" in str(c).lower() or "назначен" in str(c).lower()),
            None,
        )
        non_oil = pd.Series(False, index=df.index)
        if purpose_col is not None:
            non_oil = non_oil | df[purpose_col].map(is_non_oil_purpose)
        well = df["Скв."] if "Скв." in df.columns else df.get("well")
        if well is not None:
            non_oil = non_oil | well.map(is_rassol_bore)
        removed = int(non_oil.sum())
        if removed:
            # Раздельно: рассольный ствол и не-нефтяное назначение — разные причины.
            brine = well.map(is_rassol_bore) if well is not None else pd.Series(False, index=df.index)
            self.exclusions.add(df.loc[non_oil & brine], BRINE_BORE)
            self.exclusions.add(df.loc[non_oil & ~brine], NON_OIL_PURPOSE)
            # ⚠ Индекс НЕ сбрасывается. `_append_records` ищет статус дубля как
            # `duplicates.get(idx)`, где ключи — индексы кадра ДО фильтрации.
            # Пересчёт индекса здесь разъезжал бы эти два кадра и в режиме 1
            # (сверка с существующим регистром) раздавал строкам чужие статусы:
            # часть новых записей уходила бы как дубли, а часть дублей —
            # дописывалась второй раз. В режиме 2 словарь пуст, поэтому дефект
            # не проявлялся.
            self.new_failures = df.loc[~non_oil]
            print(f"  Oil-well filter: removed {removed} non-oil run(s) "
                  f"(injection/piezometric/brine); {len(self.new_failures)} remain")

    def _normalize_svod_fields(self):
        """Canonicalize register identity fields and backfill nominal pump specs.

        - `Месторождение` is set to the well-prefix field code (``Ic``, ``Ya`` ...)
          so long field names (e.g. the TechRegime `М/р` `Ичёдинское нефтяное
          месторождение`) never leak into the register.
        - Contractor / ownership strings are folded to their canonical short name.
        - Missing nominal pump parameters are filled from the pump-type code and
          the fleet nominal database.
        """
        if not self.normalize_svod_fields or self.new_failures is None or self.new_failures.empty:
            return

        frame = self.new_failures
        well_column = next((c for c in ("well", "Скв.") if c in frame.columns), None)
        if well_column is not None:
            field_codes = frame[well_column].apply(derive_field_code_from_well)

            # Learn a field-name -> code map from rows whose well prefix resolves,
            # so rows without a derivable prefix still map their long field name
            # (e.g. TechRegime `Ичёдинское нефтяное месторождение`) to the code.
            # The explicit map seeds fields whose wells never carry a prefix.
            name_to_code: Dict[str, str] = dict(FIELD_NAME_TO_CODE)
            existing_field = frame["Месторождение"] if "Месторождение" in frame.columns else None
            if existing_field is not None:
                learn = pd.DataFrame({
                    "name": existing_field.apply(lambda v: normalize_text(v)),
                    "code": field_codes,
                })
                learn = learn[(learn["name"] != "") & learn["code"].notna()]
                for name, group in learn.groupby("name"):
                    name_to_code[name] = group["code"].value_counts().index[0]

            def _resolve_field(code, current):
                if pd.notna(code) and code:
                    return code
                mapped = name_to_code.get(normalize_text(current))
                return mapped if mapped else current

            for field_column in ("Месторождение", "field"):
                if field_column in frame.columns:
                    frame[field_column] = [
                        _resolve_field(code, current)
                        for code, current in zip(field_codes, frame[field_column])
                    ]
                else:
                    frame[field_column] = field_codes

        for contractor_column in ("Принадлежность", "contractor"):
            if contractor_column in frame.columns:
                frame[contractor_column] = frame[contractor_column].apply(normalize_contractor)

        self.new_failures = fill_missing_pump_specs(frame, comment_column=SVOD_COMMENT_COLUMN)
        print("  Normalized field codes / contractor and backfilled pump nominal specs")

    def _sanitize_run_dates(self):
        """Correct physically-impossible date/runtime values on the outgoing rows.

        These are the known ПДК data defects the register must never carry:
          * a dismantling date earlier than the stop date, or in the future ->
            the dismantling date is a typo; drop it (keep the row).
          * a runtime («Наработка (сут)») exceeding the mount->stop calendar span
            by >30 days (a pump cannot run longer than it was installed) -> cap it
            to the calendar span and rebuild its runtime group.
        Every correction is recorded old->new in the audit ``data_fixes`` sheet.
        """
        frame = self.new_failures
        if frame is None or frame.empty:
            return
        # ⚠ «Демонтаж в будущем» — это проверка на ОПЕЧАТКУ, и мерить её надо
        # сегодняшним днём, а не горизонтом наблюдения. Это разные вещи: горизонт
        # говорит, до какой даты мы видим ОТКАЗЫ, а невозможной дата становится
        # только тогда, когда она ещё не наступила. В нынешнем ПДК между горизонтом
        # (2026-07-12) и сегодня лежат 7 совершенно настоящих демонтажей — привязка
        # этой проверки к горизонту молча стирала бы их.
        as_of = pd.Timestamp.now().normalize()

        # A register field is carried under several aliases (the Russian header
        # and the English working name); the exported value coalesces across them
        # (e.g. a blank PDK «Дата демонтажа» falls back to Big's dismantling_date).
        # Sanitize the *effective* value and clear every alias, so an impossible
        # value can't survive in a column the check didn't look at.
        stop_cols = [c for c in ("Дата остановки", "failure_date") if c in frame.columns]
        mount_cols = [c for c in ("Дата монтажа", "installation_date") if c in frame.columns]
        launch_cols = [c for c in ("Дата запуска", "launch_date") if c in frame.columns]
        dismantle_cols = [c for c in ("Дата демонтажа", "dismantling_date") if c in frame.columns]
        runtime_cols = [c for c in ("Наработка (сут)", "runtime_nno") if c in frame.columns]
        group_cols = [c for c in ("группа наработок", "runtime_group") if c in frame.columns]
        well_cols = [c for c in ("Скв.", "well") if c in frame.columns]

        def effective_date(row, cols):
            for c in cols:
                value = parse_date(row.get(c))
                if value is not None:
                    return value
            return None

        def effective_number(row, cols):
            for c in cols:
                value = parse_number(row.get(c))
                if value is not None:
                    return value
            return None

        if SVOD_COMMENT_COLUMN not in frame.columns:
            frame[SVOD_COMMENT_COLUMN] = None

        def note(idx, text):
            # Record the correction on the entry itself, in its «Комментарий»
            # cell, so every fix is traceable per row (not only in the audit).
            frame.at[idx, SVOD_COMMENT_COLUMN] = append_comment(
                frame.at[idx, SVOD_COMMENT_COLUMN], text
            )

        dismantle_fixes = 0
        runtime_fixes = 0
        mount_fixes = 0
        for idx, row in frame.iterrows():
            well = row.get(well_cols[0]) if well_cols else None
            stop = effective_date(row, stop_cols)
            mount = effective_date(row, mount_cols)

            # Impossible mount date (a stop date before the mount) -> the mount is
            # the typo. Reconstruct it from the launch date when available (mount ~
            # launch), else from stop - runtime, so the row becomes consistent.
            if stop is not None and mount is not None and stop < mount:
                launch = effective_date(row, launch_cols)
                runtime = effective_number(row, runtime_cols)
                new_mount = None
                if launch is not None and launch <= stop:
                    new_mount = launch
                elif runtime is not None and runtime >= 0:
                    new_mount = stop.normalize() - pd.Timedelta(days=float(runtime))
                if new_mount is not None:
                    for c in mount_cols:
                        frame.at[idx, c] = new_mount
                    self.audit.add_data_fix(
                        well, "Дата монтажа", mount.date(), new_mount.date(),
                        "остановки раньше монтажа (монтаж восстановлен)",
                    )
                    note(idx, f"монтаж восстановлен {mount.date()}→{new_mount.date()} (остановка раньше монтажа)")
                    mount = new_mount
                    mount_fixes += 1

            # Impossible dismantling date (before the stop, or in the future) -> drop it.
            dismantle = effective_date(row, dismantle_cols)
            if dismantle is not None and (
                (stop is not None and dismantle < stop) or dismantle > as_of
            ):
                reason = (
                    "демонтаж раньше остановки"
                    if (stop is not None and dismantle < stop)
                    else "демонтаж в будущем"
                )
                for c in dismantle_cols:
                    frame.at[idx, c] = None
                self.audit.add_data_fix(well, "Дата демонтажа", dismantle.date(), None, reason)
                note(idx, f"демонтаж {dismantle.date()} удалён ({reason})")
                dismantle_fixes += 1

            # Runtime exceeding the mount->stop calendar span -> cap to the span.
            if stop is not None and mount is not None:
                runtime = effective_number(row, runtime_cols)
                calendar_days = (stop.normalize() - mount.normalize()).days
                if runtime is not None and calendar_days >= 0 and runtime - calendar_days > 30:
                    for c in runtime_cols:
                        frame.at[idx, c] = float(calendar_days)
                    new_group = derive_runtime_group(calendar_days)
                    for c in group_cols:
                        frame.at[idx, c] = new_group
                    self.audit.add_data_fix(
                        well, "Наработка (сут)", runtime, float(calendar_days),
                        "Наработка > calendar (ПДК defect)",
                    )
                    note(idx, f"Наработка ограничена по календарю {runtime:g}→{calendar_days} сут")
                    runtime_fixes += 1

        if dismantle_fixes or runtime_fixes or mount_fixes:
            print(
                f"  Sanitized impossible values: {dismantle_fixes} dismantling date(s) dropped, "
                f"{mount_fixes} mount date(s) reconstructed, "
                f"{runtime_fixes} runtime(s) capped to calendar span"
            )
        self.audit.add_summary_stat("Impossible dismantling dates dropped", dismantle_fixes)
        self.audit.add_summary_stat("Impossible mount dates reconstructed", mount_fixes)
        self.audit.add_summary_stat("Runtimes capped to calendar", runtime_fixes)

    def _ensure_target_columns_exist(self, columns: List[str]) -> list[str]:
        """Create target columns that are not present yet."""
        if self.target_df is None:
            return []
        existing_by_key = {
            strip_header_prefix(str(column)): str(column)
            for column in self.target_df.columns
        }
        added_columns: list[str] = []
        for column_name in columns:
            target_key = strip_header_prefix(column_name)
            if target_key in existing_by_key:
                continue
            self.target_df[column_name] = None
            if self.target_ws is not None:
                source_column_idx = self.target_ws.max_column if self.target_ws.max_column >= 1 else None
                append_column_to_worksheet(
                    self.target_ws,
                    column_name,
                    header_row=1,
                    copy_format_from_column=source_column_idx,
                )
            added_columns.append(column_name)
            existing_by_key[target_key] = column_name
        return added_columns

    def _enrich_operating_time_derived(self) -> None:
        """Append richer time-window features from telemetry/TechRegime operating histories."""
        if self.new_failures is None or self.new_failures.empty:
            print("  Skipping time-derived operating features (no rows to enrich)")
            return
        if not (self.telemetry_dir or self.telemetry_vt_path or self.techregime_path):
            print("  Skipping time-derived operating features (no operating-data source)")
            return

        telemetry_source = self.telemetry_dir or self.telemetry_vt_path
        techregime_source = self.techregime_path
        try:
            result = build_operating_features(
                self.new_failures,
                telemetry=telemetry_source,
                techregime=techregime_source,
            )
        except Exception as exc:
            print(f"  Skipped time-derived operating features: {exc}")
            return

        if result.features.empty:
            print("  No time-derived operating features were produced from the current sources")
            return

        feature_frame = result.features.copy()
        if "well" in feature_frame.columns:
            feature_frame["well"] = feature_frame["well"].apply(normalize_well)
        join_columns = ["well", "failure_date"]
        for column_name in join_columns:
            if column_name not in self.new_failures.columns:
                self.new_failures[column_name] = None
        if "well" in self.new_failures.columns:
            self.new_failures["well"] = self.new_failures["well"].apply(normalize_well)
        if "failure_date" in self.new_failures.columns:
            self.new_failures["failure_date"] = self.new_failures["failure_date"].apply(parse_date)
        if "failure_date" in feature_frame.columns:
            feature_frame["failure_date"] = feature_frame["failure_date"].apply(parse_date)

        selected_feature_columns = [
            column_name
            for column_name in DEFAULT_OPERATING_TIME_DERIVED_COLUMNS
            if column_name in feature_frame.columns
        ]
        if not selected_feature_columns:
            print("  No configured time-derived feature columns were available to append")
            return

        added_columns = self._ensure_target_columns_exist(selected_feature_columns)
        if added_columns:
            print(
                "  Added target column(s) for time-derived operating features: "
                + ", ".join(added_columns)
            )

        feature_subset = feature_frame[join_columns + selected_feature_columns].copy()
        merged = self.new_failures.merge(
            feature_subset,
            how="left",
            on=join_columns,
            suffixes=("", "__oper"),
        )
        for column_name in selected_feature_columns:
            if column_name in merged.columns:
                self.new_failures[column_name] = merged[column_name]
        print(f"  Added {len(selected_feature_columns)} time-derived operating feature column(s)")
        self.audit.add_summary_stat("Time-derived operating feature columns", len(selected_feature_columns))
    
    def _mark_lab_chemistry(self):
        """Ensure lab chemistry columns exist and track missing-source state."""
        self.new_failures = enrich_lab_chemistry(self.new_failures)
        if self.lab_path:
            print(f"  Ensured {len(LAB_CHEMISTRY_COLUMNS)} lab chemistry columns are present")
            self.audit.lab_chemistry_missing = []
        else:
            print(f"  Marked {len(LAB_CHEMISTRY_COLUMNS)} lab chemistry columns as unavailable")
            self.audit.lab_chemistry_missing = LAB_CHEMISTRY_COLUMNS

    def _enrich_lab(self):
        """Enrich from lab chemistry workbook."""
        if self.lab_path:
            self.new_failures = enrich_from_lab(
                self.new_failures,
                self.lab_path,
            )
            print("  Enriched from lab chemistry")
    
    def _append_records(self, duplicates: Dict):
        """Reconcile incoming rows into the workbook: update / needs-review / append.

        In reconciliation mode (mode 1, the default): an EXACT_EXISTING match to a
        real target row updates that row in place (only changed, non-empty values;
        old->new is audited); a LIKELY_EXISTING (within-tolerance) match is routed
        to the audit's needs-review sheet instead of appended; a NEW row is
        appended. In fast mode (mode 0) and mode 2, EXACT/LIKELY are treated as
        duplicates and skipped, as before, and NEW rows are appended.
        """
        if self.target_wb is None:
            return

        reconcile = self.update_mode == 1
        header_lookup = {
            str(cell.value).strip(): col_idx
            for col_idx, cell in enumerate(self.target_ws[1], start=1)
            if cell.value is not None
        } if self.target_ws is not None else {}

        appended = 0
        updated = 0
        needs_review = 0
        skipped = 0
        copy_format_from_row = self.target_ws.max_row

        for idx, record in self.new_failures.iterrows():
            result = duplicates.get(idx)
            status = result.status if result is not None else DuplicateStatus.NEW
            target_idx = result.matching_row_idx if result is not None else None

            if status in (DuplicateStatus.EXACT_EXISTING, DuplicateStatus.LIKELY_EXISTING):
                if reconcile and status == DuplicateStatus.EXACT_EXISTING and target_idx is not None:
                    if self._update_existing_row(target_idx, record, header_lookup):
                        updated += 1
                    else:
                        skipped += 1
                    continue
                if reconcile and status == DuplicateStatus.LIKELY_EXISTING and target_idx is not None:
                    self.audit.add_needs_review(
                        record.to_dict(),
                        self.target_df.loc[target_idx].to_dict(),
                        result.match_reason if result is not None else None,
                    )
                    needs_review += 1
                    continue
                # Fast/rework mode, or an in-batch echo (no target row): skip.
                self.audit.add_duplicate(record.to_dict(), status)
                skipped += 1
                continue

            # NEW -> append.
            row_dict = record.to_dict()
            mapped_row = self._map_columns(row_dict)
            mapped_row_dict = dict(zip(self.target_df.columns, mapped_row))
            inserted_row = append_row_to_worksheet(
                self.target_ws,
                mapped_row_dict,
                header_row=1,
                copy_format_from_row=copy_format_from_row if self.update_mode != 2 else None,
            )
            if self.update_mode == 2 and self._append_style_template:
                for col_idx, style_parts in self._append_style_template.items():
                    target_cell = self.target_ws.cell(row=inserted_row, column=col_idx)
                    target_cell.font = copy(style_parts["font"])
                    target_cell.border = copy(style_parts["border"])
                    target_cell.fill = copy(style_parts["fill"])
                    target_cell.number_format = style_parts["number_format"]
                    target_cell.protection = copy(style_parts["protection"])
                    target_cell.alignment = copy(style_parts["alignment"])
            highlight_runtime_cell_if_needed(
                self.target_ws,
                inserted_row,
                "I. Наработка (сут)",
                mapped_row_dict.get("I. Наработка (сут)", mapped_row_dict.get("Наработка (сут)")),
                header_row=1,
                threshold_days=90,
            )
            appended += 1
            self.audit.add_new_record(row_dict)

        if self.apply_presentation and self.target_ws is not None:
            apply_autofilter_and_autofit(self.target_ws)
            marked = mark_reference_headers(self.target_ws)
            write_metadata_sheet(
                self.target_wb,
                BuildMetadata.current(self.observation_horizon, self.source_versions()),
            )
            write_legend_sheet(
                self.target_wb,
                [str(cell.value) for cell in self.target_ws[1] if cell.value is not None],
                build_note=(
                    "Регистр собран analysis.ingest.svod из сырых выгрузок "
                    f"({datetime.now():%Y-%m-%d %H:%M}). Справочные колонки помечены "
                    "заливкой в шапке листа «Свод»: они показывают режим по скважине "
                    "и в динамическом прогнозе не участвуют."
                ),
            )
            print(f"  Applied header filter, autofit, and marked {marked} reference/derived header(s)")

        if not self.dry_run:
            save_workbook(self.target_wb, self.output_workbook)
            print(
                f"  Saved to {self.output_workbook}: "
                f"{appended} appended, {updated} updated in place, "
                f"{needs_review} routed to needs-review, {skipped} skipped"
            )

        self.audit.add_summary_stat("Records appended", appended)
        self.audit.add_summary_stat("Existing rows updated in place", updated)
        self.audit.add_summary_stat("Rows routed to needs-review", needs_review)

    @staticmethod
    def _is_empty_value(candidate) -> bool:
        """True for None/NaN/NaT and blank/sentinel strings."""
        if candidate is None:
            return True
        try:
            if pd.isna(candidate):
                return True
        except (TypeError, ValueError):
            pass
        if isinstance(candidate, str):
            return normalize_text(candidate) in {"", "nan", "none", "-", "–"}
        return False

    def _values_differ(self, old, new) -> bool:
        """Compare a target value against an incoming one, tolerant of type drift.

        Numbers compare numerically, dates by calendar day, everything else by
        normalized text — so a re-run that carries identical data (Timestamp vs
        date, ``"45"`` vs ``45.0``) is correctly seen as unchanged. ``new`` is
        assumed non-empty (the caller never overwrites with a blank).
        """
        if self._is_empty_value(old):
            return True
        old_num, new_num = parse_number(old), parse_number(new)
        if old_num is not None and new_num is not None:
            return abs(old_num - new_num) > 1e-9
        old_date, new_date = parse_date(old), parse_date(new)
        if old_date is not None and new_date is not None:
            return old_date.normalize() != new_date.normalize()
        return normalize_text(old) != normalize_text(new)

    def _update_existing_row(self, target_idx, record, header_lookup: Dict) -> bool:
        """Apply an incoming row's enriched values onto an existing target row.

        Only non-empty values that actually differ are written (both to the
        in-memory frame and the worksheet cell); each change is recorded old->new
        in the audit. Returns True when at least one field changed.
        """
        mapped_row = self._map_columns(record.to_dict())
        mapped_dict = dict(zip(self.target_df.columns, mapped_row))
        worksheet_row = target_idx + 2  # header on row 1, data 0-indexed below it

        changes: List[Dict] = []
        for target_col in self.target_df.columns:
            new_value = mapped_dict.get(target_col)
            if self._is_empty_value(new_value):
                continue
            old_value = self.target_df.at[target_idx, target_col]
            if self._values_differ(old_value, new_value):
                self.target_df.at[target_idx, target_col] = new_value
                if self.target_ws is not None and target_col in header_lookup:
                    self.target_ws.cell(row=worksheet_row, column=header_lookup[target_col]).value = new_value
                changes.append({"field": str(target_col), "old": old_value, "new": new_value})

        if changes:
            well = record.get("well", record.get("Скв."))
            self.audit.add_row_update(well, changes)
            return True
        return False

    def _detect_last_target_failure_date(self):
        """Return the latest failure date currently present in the target register."""
        candidate_columns = [
            c for c in self.target_df.columns
            if "дата" in str(c).lower() and ("останов" in str(c).lower() or "отказ" in str(c).lower())
        ]
        if not candidate_columns:
            candidate_columns = [c for c in self.target_df.columns if "дата" in str(c).lower()]
        if not candidate_columns:
            return None
        parsed_dates = self.target_df[candidate_columns[0]].apply(parse_date).dropna()
        if parsed_dates.empty:
            return None
        return parsed_dates.max()

    def _prepare_rework_target(self):
        """Clear existing target data rows while keeping workbook structure and row style template."""
        if self.target_ws is None:
            return
        if self.target_ws.max_row >= 2:
            template_row = self.target_ws.max_row
            self._append_style_template = {}
            for col_idx in range(1, self.target_ws.max_column + 1):
                source_cell = self.target_ws.cell(row=template_row, column=col_idx)
                self._append_style_template[col_idx] = {
                    "font": copy(source_cell.font),
                    "border": copy(source_cell.border),
                    "fill": copy(source_cell.fill),
                    "number_format": source_cell.number_format,
                    "protection": copy(source_cell.protection),
                    "alignment": copy(source_cell.alignment),
                }
            self.target_ws.delete_rows(2, self.target_ws.max_row - 1)
        self.target_df = self.target_df.iloc[0:0].copy()
    
    def _map_columns(self, row_dict: Dict) -> list:
        """Map enriched columns to target columns and return ordered row values."""
        normalized_input = {
            strip_header_prefix(str(key)): value
            for key, value in row_dict.items()
            if key is not None
        }
        
        # Build reverse lookup for synonyms: normalized synonym -> canonical key
        synonym_lookup = {}
        for canonical, synonyms in COLUMN_SYNONYMS.items():
            for synonym in synonyms:
                synonym_lookup[strip_header_prefix(synonym)] = canonical
        
        special_targets = {
            "кислый/некислый": "acid_type",
            "глубина спуска уэцн, по нкт": "descent_depth",
            "габарит уэцн": "габарит",
            "ном. произв. м₃/сут": "производительность",
            "частота": "frequency",
            "загр, двиг,": "motor_load",
            "загр, двиг": "motor_load",
            "failure flag": "failure_flag",
        }
        
        def _is_empty(candidate) -> bool:
            if candidate is None:
                return True
            # Catches NaN and NaT (PDK rows carry NaT in the concatenated
            # `Дата запуска` column, which must not shadow the `launch_date` alias).
            try:
                if pd.isna(candidate):
                    return True
            except (TypeError, ValueError):
                pass
            if isinstance(candidate, str):
                return normalize_text(candidate) in {"", "nan", "none", "-", "–"}
            return False

        mapped_row = []
        for target_col in self.target_df.columns:
            target_key = strip_header_prefix(str(target_col))

            # A literal target column present but empty (e.g. running-well rows
            # carry a blank `Дата запуска`) must not shadow a populated alias such
            # as `launch_date` -- fall through to synonyms whenever the direct
            # match is empty.
            value = normalized_input.get(target_key)

            if _is_empty(value) and target_key in special_targets:
                value = normalized_input.get(special_targets[target_key])

            if _is_empty(value) and target_key in synonym_lookup:
                source_key = normalize_text(synonym_lookup[target_key])
                value = normalized_input.get(source_key)

            mapped_row.append(value)

        return mapped_row
    
    def _generate_audit(self):
        """Generate audit reports"""
        # Generate field fill matrix
        target_cols = list(self.target_df.columns)
        matrix = generate_field_fill_matrix(self.new_failures, target_cols)
        for col, stats in matrix.items():
            self.audit.add_field_fill(col, stats["filled"], stats["empty"])
        
        # Write audit Excel
        self.audit.to_excel(self.audit_xlsx)
        print(f"  Wrote audit Excel: {self.audit_xlsx}")
        
        # Write audit Markdown
        self.audit.to_markdown(self.audit_md)
        print(f"  Wrote audit Markdown: {self.audit_md}")


def run_workflow(
    target_path: str,
    pdk_path: str,
    **kwargs
) -> FailureUpdateWorkflow:
    """
    Run the failure update workflow.
    
    Args:
        target_path: Path to target workbook
        pdk_path: Path to PDK source
        **kwargs: Additional arguments for FailureUpdateWorkflow
        
    Returns:
        FailureUpdateWorkflow instance
    """
    workflow = FailureUpdateWorkflow(target_path, pdk_path, **kwargs)
    workflow.execute()
    return workflow
