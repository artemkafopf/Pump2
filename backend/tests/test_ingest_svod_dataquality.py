"""Regression tests for the Свод data-quality rework.

Guards the fixes for issues reported against the produced register:
  * a non-brine side-track's later closure must NOT drop a genuine oil running run
    (bore-keyed monotonicity), while a brine (рс) bore still blocks a suffixless
    oil phantom;
  * open (running) rows must carry «Наработка (сут)» (calendar age fallback);
  * physically-impossible dates/runtime are corrected on the outgoing rows.
"""

# Перенесено из db_builder/tests/test_svod_dataquality_fixes.py вместе с кодом сборщика: тесты
# обязаны переезжать с тем, что они держат, иначе перенос молча теряет именно
# те инварианты, ради которых код и писался.

from pathlib import Path

import pandas as pd

from openpyxl import Workbook

from analysis.ingest.svod.artificial_lift_processor import (
    extract_running_wells_from_artificial_lift,
)
from analysis.ingest.normalize import normalize_well
from analysis.ingest.svod.main import FailureUpdateWorkflow


BIG_HEADER_ROW1 = [
    "Скважина", "Цель спуска", "Дата отказа", "Дата монтажа", "Дата запуска",
    "Дата демонтажа", "ННО", "Дизайн УЭЦН", "Насос (50Гц)", "Насос (50Гц)",
]
BIG_HEADER_ROW2 = [None, None, None, None, None, None, None, None, "Собственник оборудования", "Тип ГНО"]


def _write_big(path: Path, rows) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Скважинное оборудование"
    for c, v in enumerate(BIG_HEADER_ROW1, start=1):
        ws.cell(row=1, column=c, value=v)
    for c, v in enumerate(BIG_HEADER_ROW2, start=1):
        ws.cell(row=2, column=c, value=v)
    for r, values in enumerate(rows, start=3):
        for c, v in enumerate(values, start=1):
            ws.cell(row=r, column=c, value=v)
    wb.save(path)


def _row(well, purpose, fail, install, dismantle, nno, launch=None, gno="30.2 ЭЦНДИК Э"):
    return [well, purpose, fail, install, launch or install, dismantle, nno, "MT5A-DP", "Борец", gno]


def test_sidetrack_closure_does_not_drop_genuine_oil_running_run(tmp_path):
    # Oil bore Vt_3305 last closed 2026-02-09; a side-track Vt_3305ш closes LATER
    # (2026-03-03); the genuine oil run mounts 2026-02-10 and is still running.
    big = tmp_path / "big.xlsx"
    _write_big(big, [
        _row("Vt_3305", "Мех. добыча", "2026-01-25", "2025-10-07", "2026-02-09", 120),
        _row("Vt_3305", "Мех. добыча", None, "2026-02-10", None, None),          # live oil run
        _row("Vt_3305ш", "Мех. добыча", "2026-02-27", "2025-12-05", "2026-03-03", 90),  # side-track
    ])
    running = extract_running_wells_from_artificial_lift(str(big), as_of_date="2026-07-22")
    assert normalize_well("Vt_3305") in set(running["normalized_well"])


def test_brine_bore_still_blocks_suffixless_oil_phantom(tmp_path):
    # A suffixless oil open row whose physical well is really a brine bore is a
    # phantom and must be dropped (the рс closure blocks it).
    big = tmp_path / "big.xlsx"
    _write_big(big, [
        _row("Vt_2704", "Мех. добыча", None, "2023-02-13", None, 30),            # phantom
        _row("Vt_2704рс", "Мех. добыча", "2024-05-06", "2023-11-06", "2024-05-06", 182),
    ])
    running = extract_running_wells_from_artificial_lift(str(big), as_of_date="2026-07-22")
    keys = {normalize_well(w) for w in running["well"]} if not running.empty else set()
    assert normalize_well("Vt_2704") not in keys


def test_open_row_gets_calendar_runtime_when_no_nno(tmp_path):
    big = tmp_path / "big.xlsx"
    _write_big(big, [
        _row("Ya_785", "Мех. добыча", None, "2026-06-02", None, None, launch="2026-06-02"),
    ])
    running = extract_running_wells_from_artificial_lift(str(big), as_of_date="2026-07-22")
    row = running[running["normalized_well"] == normalize_well("Ya_785")].iloc[0]
    # 2026-06-02 -> 2026-07-22 is 50 days.
    assert row["Наработка (сут)"] == 50


def test_flowing_run_with_esp_is_kept_but_without_esp_is_dropped(tmp_path):
    big = tmp_path / "big.xlsx"
    _write_big(big, [
        # Фонтанная WITH an ESP in Тип ГНО -> a genuine artificial-lift oil run.
        _row("Vt_10", "Фонтанная экспл.", None, "2026-05-01", None, None, gno="30.2 ЭЦНМИК Э"),
        # Фонтанная WITHOUT an ESP (funnel only) -> not artificial lift.
        _row("Vt_20", "Фонтанная экспл.", None, "2026-05-01", None, None, gno="Воронка НКТ-73"),
    ])
    running = extract_running_wells_from_artificial_lift(str(big), as_of_date="2026-07-22")
    keys = set(running["normalized_well"]) if not running.empty else set()
    assert normalize_well("Vt_10") in keys
    assert normalize_well("Vt_20") not in keys


def test_conversion_to_piezometer_closes_oil_at_its_mount(tmp_path):
    big = tmp_path / "big.xlsx"
    _write_big(big, [
        # A stale oil open row that predates a later пьезометр conversion -> dropped.
        _row("Vt_30", "Мех. добыча", None, "2020-01-01", None, None),
        _row("Vt_30", "Пьезометр", "2022-01-01", "2021-01-01", "2022-01-01", None, gno="Воронка"),
        # A genuine oil run mounting AFTER the well's last conversion -> kept (Ya_785 shape).
        _row("Vt_40", "Пьезометр", "2026-06-03", "2025-11-26", "2026-06-03", None, gno="Воронка"),
        _row("Vt_40", "Мех. добыча", None, "2026-06-02", None, None),
    ])
    running = extract_running_wells_from_artificial_lift(str(big), as_of_date="2026-07-22")
    keys = set(running["normalized_well"]) if not running.empty else set()
    assert normalize_well("Vt_30") not in keys   # stale oil predates the пьезометр
    assert normalize_well("Vt_40") in keys        # oil run postdates the conversion


def _sanitize_workflow(new_failures):
    wf = FailureUpdateWorkflow("target.xlsx", "pdk.xlsx", dry_run=True)
    wf.new_failures = new_failures
    wf._sanitize_run_dates()
    return wf


def test_sanitize_drops_dismantle_before_stop():
    wf = _sanitize_workflow(pd.DataFrame([{
        "Скв.": "Vt_1", "Дата монтажа": pd.Timestamp("2020-01-01"),
        "Дата остановки": pd.Timestamp("2021-01-01"), "Дата демонтажа": pd.Timestamp("2020-06-01"),
        "Наработка (сут)": 200,
    }]))
    assert pd.isna(wf.new_failures.iloc[0]["Дата демонтажа"])
    assert any(f["reason"] == "демонтаж раньше остановки" for f in wf.audit.data_fixes)


def test_sanitize_drops_future_dismantle():
    future = (pd.Timestamp.now().normalize() + pd.Timedelta(days=120))
    wf = _sanitize_workflow(pd.DataFrame([{
        "Скв.": "Vt_7805", "Дата монтажа": pd.Timestamp("2025-01-01"),
        "Дата остановки": pd.Timestamp("2026-04-14"), "Дата демонтажа": future,
        "Наработка (сут)": 400,
    }]))
    assert pd.isna(wf.new_failures.iloc[0]["Дата демонтажа"])
    assert any(f["reason"] == "демонтаж в будущем" for f in wf.audit.data_fixes)


def test_sanitize_caps_runtime_exceeding_calendar():
    wf = _sanitize_workflow(pd.DataFrame([{
        "Скв.": "Vt_2", "Дата монтажа": pd.Timestamp("2024-01-01"),
        "Дата остановки": pd.Timestamp("2024-02-01"), "Дата демонтажа": pd.Timestamp("2024-02-02"),
        "Наработка (сут)": 500,  # calendar is 31 days -> way over
    }]))
    assert wf.new_failures.iloc[0]["Наработка (сут)"] == 31.0
    assert any(f["reason"] == "Наработка > calendar (ПДК defect)" for f in wf.audit.data_fixes)


def test_sanitize_reconstructs_mount_when_stop_precedes_it():
    # Stop before mount -> the mount is the typo; rebuild it from the launch date.
    wf = _sanitize_workflow(pd.DataFrame([{
        "Скв.": "Ya_631", "Дата монтажа": pd.Timestamp("2021-10-15"),
        "Дата запуска": pd.Timestamp("2020-10-17"), "Дата остановки": pd.Timestamp("2021-06-09"),
        "Наработка (сут)": 235,
    }]))
    assert wf.new_failures.iloc[0]["Дата монтажа"] == pd.Timestamp("2020-10-17")
    assert any("остановки раньше монтажа" in f["reason"] for f in wf.audit.data_fixes)


def test_sanitize_reconstructs_mount_from_runtime_when_no_launch():
    wf = _sanitize_workflow(pd.DataFrame([{
        "Скв.": "Gor_5", "Дата монтажа": pd.Timestamp("2022-02-25"),
        "Дата остановки": pd.Timestamp("2021-07-01"), "Наработка (сут)": 125,
    }]))
    # 2021-07-01 - 125 days = 2021-02-26
    assert wf.new_failures.iloc[0]["Дата монтажа"] == pd.Timestamp("2021-02-26")


def test_corrections_are_recorded_in_comment_column():
    # Two corrections on one row -> both appear in its «Комментарий», separated by ";".
    wf = _sanitize_workflow(pd.DataFrame([{
        "Скв.": "Vt_9", "Дата монтажа": pd.Timestamp("2024-01-01"),
        "Дата остановки": pd.Timestamp("2024-02-01"), "Дата демонтажа": pd.Timestamp("2020-06-01"),
        "Наработка (сут)": 500,
    }]))
    comment = str(wf.new_failures.iloc[0]["Комментарий"])
    assert "демонтаж" in comment and "удалён" in comment
    assert "Наработка ограничена" in comment


def test_sanitize_leaves_valid_rows_untouched():
    wf = _sanitize_workflow(pd.DataFrame([{
        "Скв.": "Vt_3", "Дата монтажа": pd.Timestamp("2024-01-01"),
        "Дата остановки": pd.Timestamp("2024-06-01"), "Дата демонтажа": pd.Timestamp("2024-06-05"),
        "Наработка (сут)": 150,  # < calendar (152) -> fine
    }]))
    row = wf.new_failures.iloc[0]
    assert row["Дата демонтажа"] == pd.Timestamp("2024-06-05")
    assert row["Наработка (сут)"] == 150
    assert wf.audit.data_fixes == []
