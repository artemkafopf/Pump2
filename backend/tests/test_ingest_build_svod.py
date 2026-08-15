r"""Сборка Свода с нуля — регистр строится из сырых источников, без целевой книги.

Перенесено из ``d:\GitHub\db_builder\tests\test_build_svod.py``. Тест
streamlit-приложения не переносился: GUI остался в db_builder, здесь его роль
играют тонкие обёртки ``scripts/ingest/``.
"""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook

from analysis.ingest.config import (
    SVOD_SHEET_NAME,
    build_svod_target_columns,
)
from analysis.ingest.svod.build_svod import (
    build_svod_from_scratch,
    write_svod_seed_workbook,
)


def _write_pdk_workbook(path: Path) -> None:
    frame = pd.DataFrame([
        {
            "Скв.": "Vt_100",
            "Тип скважины": "НФ",
            "Дата монтажа": "2025-09-01",
            "Дата остановки": "2025-09-20",
            "Наработка (сут)": 19,
            "Причина остановки": "Отказ",
            "Отказавший узел": "Насос",
        },
        {
            # Brine bore -- must be filtered out of the oil register.
            "Скв.": "Vt_100рс",
            "Тип скважины": "НФ",
            "Дата монтажа": "2025-09-01",
            "Дата остановки": "2025-09-15",
            "Наработка (сут)": 14,
            "Причина остановки": "Отказ",
            "Отказавший узел": "Насос",
        },
    ])
    frame.to_excel(path, index=False)


BIG_HEADER_ROW1 = [
    "Скважина", "Цель спуска", "Дата отказа", "Дата монтажа",
    "Дата запуска", "Дата демонтажа", "ННО", "Насос (50Гц)", "Насос (50Гц)",
]
BIG_HEADER_ROW2 = [None, None, None, None, None, None, None, "Модель ГНО", "Собственник оборудования"]


def _write_big_workbook(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Скважинное оборудование"
    for col_idx, value in enumerate(BIG_HEADER_ROW1, start=1):
        ws.cell(row=1, column=col_idx, value=value)
    for col_idx, value in enumerate(BIG_HEADER_ROW2, start=1):
        ws.cell(row=2, column=col_idx, value=value)
    data_rows = [
        # Active oil run -> becomes a censored (running) row.
        ["Vt_200", "Мех. добыча", None, "2025-08-01", "2025-08-02", None, 60, "MT5A-200DP", "Борец"],
        # Closed run for the PDK failure well -> enrichment source for launch date.
        ["Vt_100", "Мех. добыча", "2025-09-20", "2025-09-01", "2025-09-02", "2025-09-21", 19, "MT5A-100DP", "Борец"],
    ]
    for row_idx, values in enumerate(data_rows, start=3):
        for col_idx, value in enumerate(values, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)
    wb.save(path)


def test_write_svod_seed_workbook_has_full_schema(tmp_path):
    seed = tmp_path / "seed.xlsx"
    write_svod_seed_workbook(seed)
    frame = pd.read_excel(seed, sheet_name=SVOD_SHEET_NAME)
    assert list(frame.columns) == build_svod_target_columns()


def test_build_svod_from_scratch_without_techregime_keeps_columns_empty(tmp_path):
    pdk_path = tmp_path / "pdk.xlsx"
    big_path = tmp_path / "big.xlsx"
    output = tmp_path / "svod.xlsx"
    _write_pdk_workbook(pdk_path)
    _write_big_workbook(big_path)

    workflow = build_svod_from_scratch(
        pdk_path=[str(pdk_path)],
        artificial_lift_path=[str(big_path)],
        include_techregime=False,
        include_telemetry=False,
        output_workbook=str(output),
    )

    assert output.exists()
    result = pd.read_excel(output, sheet_name=SVOD_SHEET_NAME)

    # The full schema is preserved even though TechRegime/telemetry were skipped.
    assert list(result.columns) == build_svod_target_columns()

    # TechRegime operating columns exist but are empty (extraction was omitted).
    for techregime_column in ["Дебит жидк.", "Частота", "Обводненность"]:
        assert techregime_column in result.columns
        assert result[techregime_column].isna().all()

    # The failure row is present; the brine bore was filtered out.
    wells = set(result["Скв."].dropna())
    assert "Vt_100" in wells
    assert "Vt_100рс" not in wells

    # The active oil run entered as a censored (open) row: no stop date, flag 0.
    running = result[result["Скв."] == "Vt_200"]
    assert len(running) == 1
    assert pd.isna(running.iloc[0]["Дата остановки"])

    # The seed schema workbook was cleaned up by default.
    assert not (tmp_path / f"_svod_seed_{output.stem}.xlsx").exists()


def test_build_svod_column_order_and_launch_date(tmp_path):
    pdk_path = tmp_path / "pdk.xlsx"
    big_path = tmp_path / "big.xlsx"
    output = tmp_path / "svod.xlsx"
    _write_pdk_workbook(pdk_path)
    _write_big_workbook(big_path)

    build_svod_from_scratch(
        pdk_path=[str(pdk_path)],
        artificial_lift_path=[str(big_path)],
        include_techregime=False,
        include_telemetry=False,
        output_workbook=str(output),
    )

    result = pd.read_excel(output, sheet_name=SVOD_SHEET_NAME)
    # A = well, B = field, C = pad.
    assert list(result.columns[:3]) == ["Скв.", "Месторождение", "Куст"]
    # Дата запуска is filled from the Big `launch_date` alias even though the
    # concatenated column arrives as NaT on PDK rows.
    row = result[result["Скв."] == "Vt_100"].iloc[0]
    assert pd.notna(row["Дата запуска"])
    assert pd.Timestamp(row["Дата запуска"]) == pd.Timestamp("2025-09-02")
    # Месторождение is the well-prefix code, not a long field name.
    assert row["Месторождение"] == "Vt"
    # Тип УЭЦН carries the pump model, not a design filename.
    assert row["Тип УЭЦН"] == "MT5A-100DP"


def test_build_svod_normalizes_contractor_and_fills_pump_nominal(tmp_path):
    pdk_path = tmp_path / "pdk.xlsx"
    big_path = tmp_path / "big.xlsx"
    output = tmp_path / "svod.xlsx"
    # PDK failure well with a parseable type code but no nominal productivity.
    pd.DataFrame([{
        "Скв.": "Vt_100", "Тип скважины": "НФ", "Дата монтажа": "2025-09-01",
        "Дата остановки": "2025-09-20", "Наработка (сут)": 19,
        "Причина остановки": "Отказ", "Отказавший узел": "Насос",
        "Тип УЭЦН": "5а-200-2350", "Принадлежность": "Новые-технологии ООО «ИНК»",
    }]).to_excel(pdk_path, index=False)
    _write_big_workbook(big_path)

    build_svod_from_scratch(
        pdk_path=[str(pdk_path)],
        artificial_lift_path=[str(big_path)],
        include_techregime=False,
        include_telemetry=False,
        output_workbook=str(output),
    )

    row = pd.read_excel(output, sheet_name=SVOD_SHEET_NAME)
    row = row[row["Скв."] == "Vt_100"].iloc[0]
    assert row["Принадлежность"] == "Новые технологии"
    # Nominal productivity backfilled from the Тип УЭЦН code (5а-200-2350).
    assert row["Производительность"] == 200


def test_build_svod_applies_header_filter_and_autofit(tmp_path):
    from openpyxl import load_workbook

    pdk_path = tmp_path / "pdk.xlsx"
    big_path = tmp_path / "big.xlsx"
    output = tmp_path / "svod.xlsx"
    _write_pdk_workbook(pdk_path)
    _write_big_workbook(big_path)

    build_svod_from_scratch(
        pdk_path=[str(pdk_path)],
        artificial_lift_path=[str(big_path)],
        include_techregime=False,
        include_telemetry=False,
        output_workbook=str(output),
    )

    ws = load_workbook(output)[SVOD_SHEET_NAME]
    # Header row carries an AutoFilter over the used range and is frozen.
    assert ws.auto_filter.ref is not None
    assert ws.freeze_panes == "A2"
    # Columns were sized to content (not left at the openpyxl default of None).
    assert any(dim.width for dim in ws.column_dimensions.values())
    # Date columns render date-only and data cells are right-aligned.
    header = [cell.value for cell in ws[1]]
    montage_col = header.index("Дата монтажа") + 1
    montage_cell = ws.cell(row=2, column=montage_col)
    assert montage_cell.number_format == "DD.MM.YYYY"
    assert montage_cell.alignment.horizontal == "right"


def test_build_svod_from_scratch_records_running_wells_as_failure_flag_zero(tmp_path):
    pdk_path = tmp_path / "pdk.xlsx"
    big_path = tmp_path / "big.xlsx"
    output = tmp_path / "svod.xlsx"
    _write_pdk_workbook(pdk_path)
    _write_big_workbook(big_path)

    build_svod_from_scratch(
        pdk_path=[str(pdk_path)],
        artificial_lift_path=[str(big_path)],
        include_techregime=False,
        include_telemetry=False,
        output_workbook=str(output),
    )

    result = pd.read_excel(output, sheet_name=SVOD_SHEET_NAME)
    # The event flag lives in `Флаг отказа`; the English `Failure Flag` is not
    # emitted (it would duplicate this column).
    assert "Failure Flag" not in result.columns
    running = result[result["Скв."] == "Vt_200"].iloc[0]
    assert running["Флаг отказа"] == 0
