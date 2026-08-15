
# Перенесено из db_builder/tests/test_processor.py вместе с кодом сборщика: тесты
# обязаны переезжать с тем, что они держат, иначе перенос молча теряет именно
# те инварианты, ради которых код и писался.
import pandas as pd
from openpyxl import Workbook
from copy import copy
from datetime import date, datetime
from pathlib import Path
import re

# ⚠⚠ В db_builder половина этих имён импортировалась из ``failure_update.enrich``,
# где лежала ВТОРАЯ, мёртвая копия процессоров. То есть тесты проверяли код,
# который в проде не выполнялся, и разойтись он мог сколько угодно. При переносе
# дубликаты удалены, а импорты переставлены на живые реализации — те самые,
# которые действительно собирают регистр.
from analysis.ingest.svod.enrich import (
    build_techregime_columns_from_header_rows,
    derive_acid_type,
    derive_field_code_from_well,
    derive_runtime_group,
    enrich_from_telemetry,
    is_kdmnu_marker,
    is_missing_failed_component,
    is_non_failure_reason_marker,
)
from analysis.ingest.svod.artificial_lift_processor import (
    enrich_big_derived_fields,
    enrich_from_big,
    extract_running_wells_from_artificial_lift,
)
from analysis.ingest.svod.lab_processor import enrich_from_lab, enrich_lab_chemistry
from analysis.ingest.lab.processor import flatten_lab_column_label, select_lab_rows_for_window
from analysis.ingest.svod.opz_processor import enrich_from_opz as current_opz_enrich
from analysis.ingest.svod.techregime_processor import (
    detect_techregime_header_rows,
    discover_available_techregime_columns,
    enrich_from_techregime,
    enrich_from_techregime as techregime_enrich_from_techregime,
)
from analysis.ingest.svod.pdk_processor import (
    _clean_pad_value,
    derive_failure_flag,
    extract_new_failures_from_pdk,
    repair_pdk_well_ids,
)
from analysis.ingest.io import apply_autofilter_and_autofit
from analysis.ingest.svod.esp_specs import (
    build_esp_specs_from_target,
    fill_specs_from_esp_type,
)
from analysis.ingest.svod.lab_processor import collapse_duplicate_lab_columns
from analysis.ingest.config import default_source
from analysis.ingest.svod.build_svod import default_sources
from analysis.ingest.svod.matching import DuplicateStatus
from analysis.ingest.svod.main import (
    FailureUpdateWorkflow,
    default_updated_workbook_path,
)
from analysis.ingest.detect_headers import auto_detect_workbook_header
from analysis.ingest.io import append_row_to_worksheet, highlight_runtime_cell_if_needed, discover_source_files
from analysis.ingest.normalize import strip_header_prefix


def test_strip_header_prefix_removes_excel_style_prefix():
    assert strip_header_prefix("A. Месторождение") == "месторождение"
    assert strip_header_prefix("BW. Кислый/Некислый") == "кислый/некислый"


def test_default_updated_workbook_path_uses_target_name_plus_update_suffix():
    target = Path(r"d:\data\register.xlsx")
    main_path_mode_0 = default_updated_workbook_path(str(target), 0)

    assert Path(main_path_mode_0).parent == target.parent
    assert re.fullmatch(r"register_\d{8}_\d{6}_update=0\.xlsx", Path(main_path_mode_0).name)


def test_every_source_resolves_through_analysis_paths():
    """Каждый вход сборщика приходит из ``analysis.paths``, а не из хардкода.

    В db_builder источники лежали в словаре ``DEFAULT_PATHS`` с абсолютными
    путями внутри самого пакета. Здесь действует правило проекта: путь — только
    через резолвер, чтобы переезд выгрузки правился в одном месте, а не в двух
    репозиториях.
    """
    sources = default_sources()
    resolved = (
        sources.pdk, sources.artificial_lift, sources.opz, sources.lab,
        sources.techregime, sources.telemetry_dir, sources.telemetry_db, sources.frac_db,
    )
    assert all(isinstance(path, Path) for path in resolved)
    assert Path(default_source("pdk")) == sources.pdk
    assert Path(default_source("telemetry")) == sources.telemetry_dir


def test_discover_source_files_ignores_nested_subfolders_for_regular_datasets():
    parent = Path(".pytest-discover-source-files")
    parent.mkdir(exist_ok=True)
    nested = parent / "archive"
    nested.mkdir(exist_ok=True)
    top_file = parent / "current.xlsx"
    nested_file = nested / "old.xlsx"
    top_file.write_text("x", encoding="utf-8")
    nested_file.write_text("x", encoding="utf-8")

    try:
        discovered = discover_source_files(parent, "pdk")
        assert discovered == [top_file.resolve()]
    finally:
        if nested_file.exists():
            try:
                nested_file.unlink()
            except PermissionError:
                pass
        if top_file.exists():
            try:
                top_file.unlink()
            except PermissionError:
                pass
        if nested.exists():
            try:
                nested.rmdir()
            except OSError:
                pass
        if parent.exists():
            try:
                parent.rmdir()
            except OSError:
                pass


def test_auto_detect_workbook_header_prefers_keyword_header_over_banner_row():
    parent = Path(".pytest-auto-detect-header")
    workbook_path = parent / "archive_like.xlsx"
    parent.mkdir(exist_ok=True)
    try:
        with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
            pd.DataFrame(
                [
                    ["", "", "", ""],
                    ["Месторождение", "Скв.", "Дата остановки", "Причина остановки"],
                    ["ЯНГКМ", "Ya_100", "2020-01-01", "Отказ"],
                ]
            ).to_excel(writer, sheet_name="Свод", header=False, index=False)

        sheet_name, header_row, df = auto_detect_workbook_header(
            str(workbook_path),
            required_keywords=["скв", "дата", "отказ", "причина"],
        )

        assert sheet_name == "Свод"
        assert header_row == 1
        assert list(df.columns[:4]) == [
            "Месторождение",
            "Скв.",
            "Дата остановки",
            "Причина остановки",
        ]
    finally:
        if workbook_path.exists():
            try:
                workbook_path.unlink()
            except PermissionError:
                pass
        if parent.exists():
            try:
                parent.rmdir()
            except OSError:
                pass


def test_derive_field_code_from_well_uses_letter_before_underscore():
    assert derive_field_code_from_well("Vt_123") == "Vt"
    assert derive_field_code_from_well("ABC_45") == "ABC"


def test_derive_acid_type_uses_existing_field_text():
    assert derive_acid_type("месторождение кислый фонд") == "Кислый"
    assert derive_acid_type("месторождение некислый фонд") == "Некислый"
    assert derive_acid_type("месторождение без признака") == "Некислый"


def test_is_non_failure_reason_marker_detects_operational_stop_reasons():
    assert is_non_failure_reason_marker("ГТМ")
    assert is_non_failure_reason_marker(" ппр ")
    assert is_non_failure_reason_marker("Прочие")
    assert not is_non_failure_reason_marker("Отказ")


def test_is_missing_failed_component_detects_empty_and_net_values():
    assert is_missing_failed_component(None)
    assert is_missing_failed_component("")
    assert is_missing_failed_component("нет")
    assert is_missing_failed_component(" - ")
    assert not is_missing_failed_component("Кабель")


def test_derive_failure_flag_uses_reason_and_failed_node_rules():
    assert derive_failure_flag("Отказ", "Насос") == 1
    assert derive_failure_flag("ГТМ", "нет") == 0
    assert derive_failure_flag("ППР", "") == 0
    assert derive_failure_flag("Прочие", "нет") == -1
    # A genuine failure reason confirms a failure even when the failed
    # component has not been diagnosed yet (common for recent events).
    assert derive_failure_flag("Отказ", "нет") == 1
    assert derive_failure_flag("Снижение изоляции", "") == 1
    # No usable reason and no failed component stays ambiguous.
    assert derive_failure_flag("", "нет") == -1
    assert derive_failure_flag("ГТМ", "Насос") == -1


def test_derive_failure_flag_trusts_a_diagnosis_recorded_under_prochie():
    """``Прочие`` is a blank reason, not a planned operation.

    ``ГТМ``/``ППР`` name a job that was scheduled; ``Прочие`` is what the operator
    picks when the reason box does not fit. Treating the three alike sent every
    row that already carried a diagnosed component to ``-1``, and the
    ``failures_only`` / ``exclude_ambiguous_events`` filters then dropped the
    diagnosis together with the row.
    """
    assert derive_failure_flag("Прочие", "Кабельная линия") == 1
    assert derive_failure_flag("Прочие", "ПЭД") == 1
    # No component: nothing to go on, still ambiguous.
    assert derive_failure_flag("Прочие", "нет") == -1
    # A component naming the *job* rather than the equipment does not promote it.
    assert derive_failure_flag("Прочие", "непроход УЭЦН") == -1
    # ГТМ/ППР semantics are untouched.
    assert derive_failure_flag("ГТМ", "нет") == 0
    assert derive_failure_flag("ППР", "Кабельная линия") == -1


def test_is_kdmnu_marker_detects_filtered_rows():
    assert is_kdmnu_marker("КДМНУ")
    assert is_kdmnu_marker("Vt_кдмну_оппн_птв-081-2")
    assert not is_kdmnu_marker("Борец")


def test_derive_runtime_group_uses_target_buckets():
    assert derive_runtime_group(0) == "0-3 суток"
    assert derive_runtime_group(3) == "0-3 суток"
    assert derive_runtime_group(4) == "4-30 суток"
    assert derive_runtime_group(30) == "4-30 суток"
    assert derive_runtime_group(31) == "31-180 сут"
    assert derive_runtime_group(180) == "31-180 сут"
    assert derive_runtime_group(181) == "181-365 сут"
    assert derive_runtime_group(365) == "181-365 сут"
    assert derive_runtime_group(366) == "366-765 сут"
    assert derive_runtime_group(765) == "366-765 сут"
    assert derive_runtime_group(766) == "больше 766 сут"


def test_extract_new_failures_keeps_pdk_runtime_group_and_only_falls_back_when_missing():
    source_df = pd.DataFrame([
        {
            "Месторождение": "Vt",
            "Куст": "1",
            "Скв.": "Vt_100",
            "Дата остановки": "2025-09-10",
            "Наработка (сут)": 45,
            "Причина остановки": "Отсутствие подачи",
            "группа наработок": "AJ-группа",
            "Тип скважины": "НФ",
        },
        {
            "Месторождение": "Vt",
            "Куст": "2",
            "Скв.": "Vt_101",
            "Дата остановки": "2025-09-11",
            "Наработка (сут)": 2,
            "Причина остановки": "Отсутствие подачи",
            "группа наработок": None,
            "по назначению": "неф.",
        },
    ])

    original_loader = extract_new_failures_from_pdk.__globals__["load_dataset_frames"]
    extract_new_failures_from_pdk.__globals__["load_dataset_frames"] = lambda path, dataset_type: [
        (Path("pdk_a.xlsx"), source_df.iloc[[0]].copy(), "Sheet1", 0),
        (Path("pdk_b.xlsx"), source_df.iloc[[1]].copy(), "Sheet1", 2),
    ]
    try:
        extracted = extract_new_failures_from_pdk("dummy.xlsx", pd.Timestamp("2025-08-28"))
    finally:
        extract_new_failures_from_pdk.__globals__["load_dataset_frames"] = original_loader

    assert extracted.iloc[0]["группа наработок"] == "AJ-группа"
    assert extracted.iloc[0]["runtime_group"] == "AJ-группа"
    assert extracted.iloc[1]["группа наработок"] == "0-3 суток"
    assert extracted.iloc[1]["runtime_group"] == "0-3 суток"


def test_extract_new_failures_from_pdk_keeps_operational_rows_and_sets_failure_flags():
    source_df = pd.DataFrame([
        {
            "Скв.": "Vt_100",
            "Дата остановки": "2025-09-10",
            "Причина остановки": "ГТМ",
            "Признак отказа": "Плановая",
            "Отказавший узел": "нет",
            "Отказавший элемент": None,
            "Тип скважины": "НФ",
        },
        {
            "Скв.": "Vt_101",
            "Дата остановки": "2025-09-11",
            "Причина остановки": "ППР",
            "Признак отказа": "Диагностика",
            "Отказавший узел": "Насос",
            "Отказавший элемент": None,
            "Тип скважины": "НФ",
        },
        {
            "Скв.": "Vt_102",
            "Дата остановки": "2025-09-12",
            "Причина остановки": "Прочие",
            "Признак отказа": "Прочее",
            "Отказавший узел": "нет",
            "Отказавший элемент": "нет",
            "Тип скважины": "НФ",
        },
    ])

    original_loader = extract_new_failures_from_pdk.__globals__["load_dataset_frames"]
    extract_new_failures_from_pdk.__globals__["load_dataset_frames"] = lambda path, dataset_type: [
        (Path("pdk.xlsx"), source_df, "Sheet1", 0),
    ]
    try:
        extracted = extract_new_failures_from_pdk("dummy.xlsx", None)
    finally:
        extract_new_failures_from_pdk.__globals__["load_dataset_frames"] = original_loader

    assert extracted["Скв."].tolist() == ["Vt_100", "Vt_101", "Vt_102"]
    assert extracted["Флаг отказа"].tolist() == [0, -1, -1]
    assert extracted["failure_flag"].tolist() == [0, -1, -1]
    assert extracted["Признак отказа"].tolist() == ["Плановая", "Диагностика", "Прочее"]
    assert extracted["failure_marker"].tolist() == ["Плановая", "Диагностика", "Прочее"]


def test_extract_new_failures_from_pdk_collapses_same_pump_run_to_failure_row():
    source_df = pd.DataFrame(
        [
            {
                "Скв.": "Vt_100",
                "Дата монтажа": "2025-09-01",
                "Дата остановки": "2025-09-10",
                "Наработка (сут)": 12,
                "Причина остановки": "ГТМ",
                "Отказавший узел": "нет",
                "Тип скважины": "НФ",
            },
            {
                "Скв.": "Vt_100",
                "Дата монтажа": "2025-09-01",
                "Дата остановки": "2025-09-15",
                "Наработка (сут)": 18,
                "Причина остановки": "Отказ",
                "Отказавший узел": "Насос",
                "Тип скважины": "НФ",
            },
        ]
    )

    original_loader = extract_new_failures_from_pdk.__globals__["load_dataset_frames"]
    extract_new_failures_from_pdk.__globals__["load_dataset_frames"] = lambda path, dataset_type: [
        (Path("pdk.xlsx"), source_df, "Sheet1", 0),
    ]
    try:
        extracted = extract_new_failures_from_pdk("dummy.xlsx", None)
    finally:
        extract_new_failures_from_pdk.__globals__["load_dataset_frames"] = original_loader

    assert len(extracted) == 1
    assert extracted.iloc[0]["Скв."] == "Vt_100"
    assert extracted.iloc[0]["failure_flag"] == 1
    assert extracted.iloc[0]["Дата остановки"] == pd.Timestamp("2025-09-15")
    assert extracted.iloc[0]["Наработка (сут)"] == 18.0


def test_extract_new_failures_from_pdk_prefers_the_diagnosed_row_when_dates_tie():
    """One stop recorded twice, component filled in on only one of the rows.

    PDK does this routinely: same well, same installation, same failure date, one
    row carrying ``Отказавший узел`` and its twin carrying "Нет". Date and runtime
    cannot separate them, so the empty row used to win on sort order alone and the
    diagnosis was lost. The tie-break keeps the row that carries evidence.
    """
    source_df = pd.DataFrame(
        [
            {
                "Скв.": "Vt_320",
                "Дата монтажа": "2022-09-14",
                "Дата остановки": "2022-12-08",
                "Наработка (сут)": 85,
                "Причина остановки": "Отсутствие подачи",
                "Отказавший узел": "нет",
                "Тип скважины": "НФ",
            },
            {
                "Скв.": "Vt_320",
                "Дата монтажа": "2022-09-14",
                "Дата остановки": "2022-12-08",
                "Наработка (сут)": 85,
                "Причина остановки": "Отсутствие подачи",
                "Отказавший узел": "ЭЦН",
                "Отказавший элемент": "Рабочие органы",
                "Тип скважины": "НФ",
            },
        ]
    )

    original_loader = extract_new_failures_from_pdk.__globals__["load_dataset_frames"]
    extract_new_failures_from_pdk.__globals__["load_dataset_frames"] = lambda path, dataset_type: [
        (Path("pdk.xlsx"), source_df, "Sheet1", 0),
    ]
    try:
        extracted = extract_new_failures_from_pdk("dummy.xlsx", None)
    finally:
        extract_new_failures_from_pdk.__globals__["load_dataset_frames"] = original_loader

    assert len(extracted) == 1
    assert extracted.iloc[0]["Отказавший узел"] == "ЭЦН"


def test_extract_new_failures_from_pdk_carries_free_text_columns():
    """``Примечание`` and ``Осложнения при ТКРС`` are the only record of the fault
    for stops where ``Отказавший узел`` was never filled in -- 92 % and 78 % of PDK
    rows carry them, and they used to stop at this boundary."""
    source_df = pd.DataFrame(
        [
            {
                "Скв.": "Ya_724",
                "Дата монтажа": "2024-01-10",
                "Дата остановки": "2024-06-01",
                "Наработка (сут)": 143,
                "Причина остановки": "Отсутствие подачи",
                "Отказавший узел": "нет",
                "Примечание": "По дефектации износ рабочих органов ЭЦН.",
                "Осложнения при ТКРС, ДЖ, Рекомендации ТКРС": "Приемная сетка забита мех.примесями",
                "Тип скважины": "НФ",
            },
        ]
    )

    original_loader = extract_new_failures_from_pdk.__globals__["load_dataset_frames"]
    extract_new_failures_from_pdk.__globals__["load_dataset_frames"] = lambda path, dataset_type: [
        (Path("pdk.xlsx"), source_df, "Sheet1", 0),
    ]
    try:
        extracted = extract_new_failures_from_pdk("dummy.xlsx", None)
    finally:
        extract_new_failures_from_pdk.__globals__["load_dataset_frames"] = original_loader

    assert "износ рабочих органов" in extracted.iloc[0]["Примечание"]
    assert "мех.примесями" in extracted.iloc[0]["Осложнения при ТКРС, ДЖ, Рекомендации ТКРС"]


def test_extract_new_failures_from_pdk_collapses_non_failure_pump_run_to_latest_row_and_max_runtime():
    source_df = pd.DataFrame(
        [
            {
                "Скв.": "Vt_101",
                "Дата монтажа": "2025-09-01",
                "Дата остановки": "2025-09-12",
                "Наработка (сут)": 21,
                "Причина остановки": "ГТМ",
                "Отказавший узел": "нет",
                "Тип скважины": "НФ",
            },
            {
                "Скв.": "Vt_101",
                "Дата монтажа": "2025-09-01",
                "Дата остановки": "2025-09-18",
                "Наработка (сут)": 19,
                "Причина остановки": "ППР",
                "Отказавший узел": "нет",
                "Тип скважины": "НФ",
            },
        ]
    )

    original_loader = extract_new_failures_from_pdk.__globals__["load_dataset_frames"]
    extract_new_failures_from_pdk.__globals__["load_dataset_frames"] = lambda path, dataset_type: [
        (Path("pdk.xlsx"), source_df, "Sheet1", 0),
    ]
    try:
        extracted = extract_new_failures_from_pdk("dummy.xlsx", None)
    finally:
        extract_new_failures_from_pdk.__globals__["load_dataset_frames"] = original_loader

    assert len(extracted) == 1
    assert extracted.iloc[0]["Скв."] == "Vt_101"
    assert extracted.iloc[0]["failure_flag"] == 0
    assert extracted.iloc[0]["Дата остановки"] == pd.Timestamp("2025-09-18")
    assert extracted.iloc[0]["Наработка (сут)"] == 21.0
    assert extracted.iloc[0]["runtime_nno"] == 21.0


def test_extract_new_failures_from_pdk_filters_to_oil_wells_across_schema_variants():
    source_a = pd.DataFrame([
        {"Скв.": "Vt_100", "Дата остановки": "2025-09-10", "Причина остановки": "Отказ", "Тип скважины": "НФ"},
        {"Скв.": "Vt_200", "Дата остановки": "2025-09-10", "Причина остановки": "Отказ", "Тип скважины": "ППД"},
    ])
    source_b = pd.DataFrame([
        {"Скв.": "Vt_300", "Дата остановки": "2025-09-11", "Причина остановки": "Отказ", "по назначению": "нефть"},
        {"Скв.": "Vt_400", "Дата остановки": "2025-09-11", "Причина остановки": "Отказ", "по назначению": "газ"},
    ])

    original_loader = extract_new_failures_from_pdk.__globals__["load_dataset_frames"]
    extract_new_failures_from_pdk.__globals__["load_dataset_frames"] = lambda path, dataset_type: [
        (Path("archive.xlsx"), source_a, "Архив", 4),
        (Path("new.xlsx"), source_b, "Основной", 1),
    ]
    try:
        extracted = extract_new_failures_from_pdk("dummy.xlsx", pd.Timestamp("2025-08-28"))
    finally:
        extract_new_failures_from_pdk.__globals__["load_dataset_frames"] = original_loader

    assert extracted["Скв."].tolist() == ["Vt_100", "Vt_300"]


def test_extract_new_failures_from_pdk_prefers_well_prefix_code_over_long_source_field_name():
    source_df = pd.DataFrame([
        {
            "Месторождение": "RemoteField",
            "Скв.": "Vt_100",
            "Дата остановки": "2025-09-10",
            "Причина остановки": "Отказ",
            "Отказавший узел": "Насос",
            "Тип скважины": "НФ",
        },
    ])

    original_loader = extract_new_failures_from_pdk.__globals__["load_dataset_frames"]
    extract_new_failures_from_pdk.__globals__["load_dataset_frames"] = lambda path, dataset_type: [
        (Path("pdk.xlsx"), source_df, "Sheet1", 0),
    ]
    try:
        extracted = extract_new_failures_from_pdk("dummy.xlsx", None)
    finally:
        extract_new_failures_from_pdk.__globals__["load_dataset_frames"] = original_loader

    assert extracted.iloc[0]["field"] == "Vt"
    assert extracted.iloc[0]["Месторождение"] == "Vt"


def test_normalize_contractor_folds_novye_tehnologii_variants():
    from analysis.ingest.svod.enrich import normalize_contractor

    assert normalize_contractor("Новые-технологии ООО «ИНК»") == "Новые технологии"
    assert normalize_contractor("Новые-технологии АО «ИНК-Запад»") == "Новые технологии"
    assert normalize_contractor("Новые технологии") == "Новые технологии"
    # Unrelated contractors are left untouched.
    assert normalize_contractor("Борец") == "Борец"
    assert normalize_contractor("ИНК") == "ИНК"


def test_clean_pad_value_blanks_values_without_a_number():
    assert _clean_pad_value("19") == "19"
    assert _clean_pad_value("Куст 8") == "Куст 8"
    assert _clean_pad_value("1А") == "1А"
    # A stray non-numeric pad (e.g. Mc_902 -> "р") does not fit the pattern.
    assert _clean_pad_value("р") is None
    assert _clean_pad_value(None) is None
    assert _clean_pad_value("") is None
    assert _clean_pad_value("nan") is None


def test_apply_autofilter_and_autofit_formats_dates_and_right_aligns():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Скв.", "Дата монтажа"])
    ws.append(["Vt_1", datetime(2025, 9, 1, 12, 30)])

    apply_autofilter_and_autofit(ws)

    assert ws.auto_filter.ref is not None
    assert ws.freeze_panes == "A2"
    date_cell = ws.cell(row=2, column=2)
    # Coerced to date-only with a date number format.
    assert isinstance(date_cell.value, date) and not isinstance(date_cell.value, datetime)
    assert date_cell.number_format == "DD.MM.YYYY"
    # Data cells are right-aligned; the header row keeps its default alignment.
    assert date_cell.alignment.horizontal == "right"
    assert ws.cell(row=1, column=1).alignment.horizontal in (None, "general")


def test_repair_pdk_well_ids_uses_field_prefix_for_numeric_only_wells():
    source_df = pd.DataFrame(
        [
            {"Месторождение": "ЯНГКМ", "Куст": "14", "Скв.": "Ya_160"},
            {"Месторождение": "ЯНГКМ", "Куст": "22", "Скв.": "Ya_680"},
            {"Месторождение": "ЯНГКМ", "Куст": "14", "Скв.": "160"},
        ]
    )

    repaired = repair_pdk_well_ids(source_df)

    assert repaired["Скв."].tolist() == ["Ya_160", "Ya_680", "Ya_160"]


def test_repair_pdk_well_ids_prefixes_unprefixed_ids_of_any_shape():
    source_df = pd.DataFrame(
        [
            {"Месторождение": "ЯНГКМ", "Куст": "14", "Скв.": "Ya_160"},
            {"Месторождение": "ЯНГКМ", "Куст": "14", "Скв.": "5"},
            {"Месторождение": "ЯНГКМ", "Куст": "14", "Скв.": "2ш"},
            {"Месторождение": "ЯНГКМ", "Куст": "14", "Скв.": "8-4"},
        ]
    )

    repaired = repair_pdk_well_ids(source_df)

    # Bare numbers, branch markers and hyphenated ids all gain the field prefix.
    assert repaired["Скв."].tolist() == ["Ya_160", "Ya_5", "Ya_2ш", "Ya_8-4"]


def test_repair_pdk_well_ids_prefers_cluster_specific_prefix_when_field_is_mixed():
    source_df = pd.DataFrame(
        [
            {"Месторождение": "MixedField", "Куст": "1", "Скв.": "Ya_101"},
            {"Месторождение": "MixedField", "Куст": "1", "Скв.": "Ya_102"},
            {"Месторождение": "MixedField", "Куст": "2", "Скв.": "Yb_201"},
            {"Месторождение": "MixedField", "Куст": "2", "Скв.": "201"},
        ]
    )

    repaired = repair_pdk_well_ids(source_df)

    assert repaired["Скв."].tolist() == ["Ya_101", "Ya_102", "Yb_201", "Yb_201"]


def test_extract_new_failures_from_pdk_repairs_numeric_only_wells_before_normalization():
    source_df = pd.DataFrame(
        [
            {
                "Месторождение": "ЯНГКМ",
                "Куст": "14",
                "Скв.": "Ya_160",
                "Дата остановки": "2025-09-09",
                "Причина остановки": "Отказ",
                "Отказавший узел": "Насос",
                "Тип скважины": "НФ",
            },
            {
                "Месторождение": "ЯНГКМ",
                "Куст": "14",
                "Скв.": "160",
                "Дата остановки": "2025-09-10",
                "Причина остановки": "Отказ",
                "Отказавший узел": "Насос",
                "Тип скважины": "НФ",
            },
        ]
    )

    original_loader = extract_new_failures_from_pdk.__globals__["load_dataset_frames"]
    extract_new_failures_from_pdk.__globals__["load_dataset_frames"] = lambda path, dataset_type: [
        (Path("pdk.xlsx"), source_df, "Sheet1", 0),
    ]
    try:
        extracted = extract_new_failures_from_pdk("dummy.xlsx", None)
    finally:
        extract_new_failures_from_pdk.__globals__["load_dataset_frames"] = original_loader

    assert extracted["Скв."].tolist() == ["Ya_160", "Ya_160"]
    assert extracted["well"].tolist() == ["Ya_160", "Ya_160"]
    assert extracted["normalized_well"].tolist() == ["ya_160", "ya_160"]


def test_flatten_lab_column_label_uses_lowest_meaningful_level():
    assert flatten_lab_column_label(("6К", "Состав воды", "Cl⁻, мг/л")) == "Cl⁻, мг/л"
    assert flatten_lab_column_label(("6К", "pH", "Unnamed: 55_level_2")) == "pH"
    assert flatten_lab_column_label(("Unnamed: 7_level_0", "Массовая концентрация хлористых солей в нефти, мг/дм³", "Unnamed: 7_level_2")) == "Массовая концентрация хлористых солей в нефти, мг/дм³"


def test_select_lab_rows_for_window_uses_recent_window_or_fallback():
    df = pd.DataFrame(
        {
            "Дата время отбора": [
                pd.Timestamp("2025-09-01"),
                pd.Timestamp("2025-09-05"),
                pd.Timestamp("2025-09-20"),
            ],
            "value": [1, 2, 3],
        }
    )

    selected = select_lab_rows_for_window(df, pd.Timestamp("2025-09-21"), av_window=10)
    assert selected["value"].tolist() == [3]

    selected = select_lab_rows_for_window(df, pd.Timestamp("2025-09-21"), av_window=30)
    assert selected["value"].tolist() == [1, 2, 3]

    selected = select_lab_rows_for_window(df, pd.Timestamp("2025-09-03"), av_window=2)
    assert selected["value"].tolist() == [1]

    selected = select_lab_rows_for_window(df, pd.Timestamp("2025-09-21"), av_window=-1)
    assert selected["value"].tolist() == [1, 2, 3]


def test_build_techregime_columns_from_header_rows_forward_fills_merged_groups():
    header_frame = pd.DataFrame(
        [
            [None, "Свойства флюида - PVT", None, "Текущий режим работы скважины", None, None],
            ["М/р", "Плотн. Воды", "Плотн. нефти", "Частота", "Дебит жидк.", "Газовый фактор"],
            [None, "г/см3", "г/см3", None, "м3/сут.", "м3/ тн"],
        ]
    )

    columns = build_techregime_columns_from_header_rows(header_frame)

    assert columns == [
        ("", "м/р"),
        ("свойства флюида - pvt", "плотн. воды"),
        ("свойства флюида - pvt", "плотн. нефти"),
        ("текущий режим работы скважины", "частота"),
        ("текущий режим работы скважины", "дебит жидк."),
        ("текущий режим работы скважины", "газовый фактор"),
    ]


def test_enrich_from_techregime_uses_latest_file_then_falls_back_to_earlier_values():
    records = pd.DataFrame(
        [
            {"well": "Vt_100", "failure_date": pd.Timestamp("2025-09-12")},
        ]
    )

    latest_df = pd.DataFrame(
        {
            ("", "id скважины"): ["Vt_100"],
            ("текущий режим работы скважины", "дебит жидк."): [120],
            ("текущий режим работы скважины", "дебит нефти"): [None],
            ("текущий режим работы скважины", "частота"): [55],
            ("свойства флюида - pvt", "плотн. нефти"): [0.83],
        }
    )
    older_df = pd.DataFrame(
        {
            ("", "id скважины"): ["Vt_100"],
            ("текущий режим работы скважины", "дебит нефти"): [15],
            ("текущий режим работы скважины", "дебит газа"): [2400],
            ("текущий режим работы скважины", "газовый фактор"): [800],
            ("текущий режим работы скважины", "рпл."): [250],
        }
    )

    original_list = enrich_from_techregime.__globals__["list_techregime_candidate_files"]
    original_load = enrich_from_techregime.__globals__["load_techregime_dataframe"]
    enrich_from_techregime.__globals__["list_techregime_candidate_files"] = lambda path, failure_date: [Path("latest.xlsx"), Path("older.xlsx")]
    enrich_from_techregime.__globals__["load_techregime_dataframe"] = lambda file_path: latest_df if "latest" in str(file_path) else older_df
    try:
        enriched = enrich_from_techregime(records, "dummy")
    finally:
        enrich_from_techregime.__globals__["list_techregime_candidate_files"] = original_list
        enrich_from_techregime.__globals__["load_techregime_dataframe"] = original_load

    row = enriched.iloc[0]
    assert row["Дебит жидк."] == 120
    assert row["liquid_rate"] == 120
    assert row["Дебит нефти"] == 15
    assert row["oil_rate"] == 15
    assert row["Дебит газа"] == 2400
    assert row["gas_rate"] == 2400
    assert row["Газовый фактор"] == 800
    assert row["gas_factor"] == 800
    assert row["ГЖФ"] == 20
    assert row["gas_liquid_ratio"] == 20
    assert row["Частота"] == 55
    assert row["frequency"] == 55
    assert row["Плотн. нефти"] == 0.83
    assert row["Рпл."] == 250


def test_enrich_from_techregime_sets_field_from_source_mr_not_well_prefix():
    records = pd.DataFrame(
        [
            {"well": "Vt_100", "failure_date": pd.Timestamp("2025-09-12"), "field": "Vt", "Месторождение": "Vt"},
        ]
    )

    tech_df = pd.DataFrame(
        {
            ("", "id скважины"): ["Vt_100"],
            ("", "м/р"): ["RemoteField"],
            ("текущий режим работы скважины", "дебит жидк."): [120],
        }
    )

    original_list = techregime_enrich_from_techregime.__globals__["list_techregime_candidate_files"]
    original_load = techregime_enrich_from_techregime.__globals__["load_techregime_dataframe"]
    techregime_enrich_from_techregime.__globals__["list_techregime_candidate_files"] = lambda path, failure_date: [Path("latest.xlsx")]
    techregime_enrich_from_techregime.__globals__["load_techregime_dataframe"] = lambda file_path: tech_df
    try:
        enriched = techregime_enrich_from_techregime(records, "dummy")
    finally:
        techregime_enrich_from_techregime.__globals__["list_techregime_candidate_files"] = original_list
        techregime_enrich_from_techregime.__globals__["load_techregime_dataframe"] = original_load

    row = enriched.iloc[0]
    assert row["field"] == "RemoteField"
    assert row["Месторождение"] == "RemoteField"


def test_enrich_from_techregime_uses_wellbore_column_not_well_type_column():
    records = pd.DataFrame(
        [
            {"well": "Vt_100", "failure_date": pd.Timestamp("2025-09-12")},
        ]
    )

    tech_df = pd.DataFrame(
        {
            ("", "id скважины"): ["Vt_100"],
            ("", "м/р"): ["RemoteField"],
            ("", "тип скв"): ["Нагнетательная"],
            ("", "тип ствола скв"): ["Наклонно-направленная"],
        }
    )

    original_list = techregime_enrich_from_techregime.__globals__["list_techregime_candidate_files"]
    original_load = techregime_enrich_from_techregime.__globals__["load_techregime_dataframe"]
    techregime_enrich_from_techregime.__globals__["list_techregime_candidate_files"] = lambda path, failure_date: [Path("latest.xlsx")]
    techregime_enrich_from_techregime.__globals__["load_techregime_dataframe"] = lambda file_path: tech_df
    try:
        enriched = techregime_enrich_from_techregime(records, "dummy")
    finally:
        techregime_enrich_from_techregime.__globals__["list_techregime_candidate_files"] = original_list
        techregime_enrich_from_techregime.__globals__["load_techregime_dataframe"] = original_load

    row = enriched.iloc[0]
    assert row["Тип ствола скв"] == "Наклонно-направленная"
    assert row["wellbore_type"] == "Наклонно-направленная"


def test_enrich_from_techregime_flat_sources_average_numeric_values_with_interval():
    records = pd.DataFrame(
        [
            {"well": "Vt_100", "failure_date": pd.Timestamp("2025-09-12")},
        ]
    )

    flat_df = pd.DataFrame(
        {
            "Дата": ["2025-09-10", "2025-09-12", "2025-08-20"],
            "ID скважины": ["Vt_100", "Vt_100", "Vt_100"],
            "М/р": ["FlatField", "FlatField", "FlatField"],
            "Текущий режим работы скважины | Дебит жидк.": [100, 130, 999],
            "Текущий режим работы скважины | Частота": [50, 52, 10],
        }
    )

    original_prepare = techregime_enrich_from_techregime.__globals__["_prepare_flat_techregime_frame"]
    techregime_enrich_from_techregime.__globals__["_prepare_flat_techregime_frame"] = lambda source: flat_df
    try:
        enriched = techregime_enrich_from_techregime(
            records,
            "dummy",
            prefer_sqlite=False,
            interval=7,
        )
    finally:
        techregime_enrich_from_techregime.__globals__["_prepare_flat_techregime_frame"] = original_prepare

    row = enriched.iloc[0]
    assert row["Дебит жидк."] == 115.0
    assert row["liquid_rate"] == 115.0
    assert row["Частота"] == 51.0
    assert row["field"] == "FlatField"
    assert row["Месторождение"] == "FlatField"


def test_enrich_from_techregime_prefers_runtime_window_over_fixed_interval():
    records = pd.DataFrame(
        [
            {"well": "Vt_100", "failure_date": pd.Timestamp("2025-09-12"), "runtime_nno": 30},
        ]
    )

    flat_df = pd.DataFrame(
        {
            "Дата": ["2025-08-01", "2025-08-20", "2025-09-10", "2025-09-12"],
            "ID скважины": ["Vt_100", "Vt_100", "Vt_100", "Vt_100"],
            "М/р": ["FlatField", "FlatField", "FlatField", "FlatField"],
            "Статус": ["В работе", "В работе", "В работе", "В работе"],
            "Текущий режим работы скважины | Дебит жидк.": [1000, 200, 100, 120],
            "Текущий режим работы скважины | Дебит газа": [10000, 2000, 1000, 1200],
            "Текущий режим работы скважины | Частота": [99, 60, 50, 52],
        }
    )

    original_prepare = techregime_enrich_from_techregime.__globals__["_prepare_flat_techregime_frame"]
    techregime_enrich_from_techregime.__globals__["_prepare_flat_techregime_frame"] = lambda source: flat_df
    try:
        enriched = techregime_enrich_from_techregime(
            records,
            "dummy",
            prefer_sqlite=False,
            interval=7,
        )
    finally:
        techregime_enrich_from_techregime.__globals__["_prepare_flat_techregime_frame"] = original_prepare

    row = enriched.iloc[0]
    assert row["Дебит жидк."] == 140.0
    assert row["liquid_rate"] == 140.0
    assert row["Дебит газа"] == 1400.0
    assert row["gas_rate"] == 1400.0
    assert row["ГЖФ"] == 10.0
    assert row["gas_liquid_ratio"] == 10.0
    assert row["Частота"] == 54.0
    assert row["frequency"] == 54.0


def test_enrich_from_techregime_ignores_ambiguous_multi_value_text_results():
    records = pd.DataFrame(
        [
            {"well": "Vt_100", "failure_date": pd.Timestamp("2025-09-12"), "runtime_nno": 30},
        ]
    )

    flat_df = pd.DataFrame(
        {
            "Дата": ["2025-09-10", "2025-09-12"],
            "ID скважины": ["Vt_100", "Vt_100"],
            "М/р": ["FlatField", "FlatField"],
            "Тип ствола скв": ["Наклонно-направленная", "Горизонтальная"],
            "Статус": ["В работе", "В работе"],
            "Текущий режим работы скважины | Дебит жидк.": [100, 120],
        }
    )

    original_prepare = techregime_enrich_from_techregime.__globals__["_prepare_flat_techregime_frame"]
    techregime_enrich_from_techregime.__globals__["_prepare_flat_techregime_frame"] = lambda source: flat_df
    try:
        enriched = techregime_enrich_from_techregime(
            records,
            "dummy",
            prefer_sqlite=False,
            interval=7,
        )
    finally:
        techregime_enrich_from_techregime.__globals__["_prepare_flat_techregime_frame"] = original_prepare

    row = enriched.iloc[0]
    assert pd.isna(row["Тип ствола скв"])
    assert pd.isna(row["wellbore_type"])
    assert row["Дебит жидк."] == 110.0


def test_enrich_from_opz_ignores_future_rows_and_keeps_gko_separate_from_gk():
    records = pd.DataFrame(
        [
            {"well": "Ic_329", "failure_date": pd.Timestamp("2022-01-01")},
        ]
    )
    opz_df = pd.DataFrame(
        {
            "Скважина": ["Ic_329", "Ic_329", "Ic_329"],
            "Дата ОПЗ": ["2021-12-20", "2021-12-25", "2022-02-01"],
            "Вид ОПЗ": ["СКО", "ГКО УЭЦН", "NaOH+СКО (2-й этап)"],
        }
    )

    original_load = current_opz_enrich.__globals__["load_dataset_frames"]
    current_opz_enrich.__globals__["load_dataset_frames"] = lambda source, dataset_type: [
        (Path("opz.xlsx"), opz_df, "Sheet1", 0)
    ]
    try:
        enriched = current_opz_enrich(records, "dummy")
    finally:
        current_opz_enrich.__globals__["load_dataset_frames"] = original_load

    row = enriched.iloc[0]
    assert row["СКО ОПЗ"] == 1
    assert row["ГКО УЭЦН"] == 1
    assert row["ГК ОПЗ"] == 0
    assert row["2-х этапка"] == 0


def test_enrich_from_opz_uses_runtime_lookback_when_available():
    records = pd.DataFrame(
        [
            {"well": "Ic_329", "failure_date": pd.Timestamp("2022-01-01"), "runtime_nno": 30},
        ]
    )
    opz_df = pd.DataFrame(
        {
            "Скважина": ["Ic_329", "Ic_329"],
            "Дата ОПЗ": ["2021-12-20", "2021-11-10"],
            "Вид ОПЗ": ["СКО", "ОПЗ"],
        }
    )

    original_load = current_opz_enrich.__globals__["load_dataset_frames"]
    current_opz_enrich.__globals__["load_dataset_frames"] = lambda source, dataset_type: [
        (Path("opz.xlsx"), opz_df, "Sheet1", 0)
    ]
    try:
        enriched = current_opz_enrich(records, "dummy")
    finally:
        current_opz_enrich.__globals__["load_dataset_frames"] = original_load

    row = enriched.iloc[0]
    assert row["СКО ОПЗ"] == 1
    assert row["ОПЗ"] == 0


def test_enrich_from_techregime_zero_safe_columns_use_status_and_ignore_zeros():
    records = pd.DataFrame(
        [
            {"well": "Vt_100", "failure_date": pd.Timestamp("2025-09-12")},
        ]
    )

    flat_df = pd.DataFrame(
        {
            "Дата": ["2025-09-10", "2025-09-11", "2025-09-12", "2025-09-12"],
            "ID скважины": ["Vt_100", "Vt_100", "Vt_100", "Vt_100"],
            "М/р": ["FlatField", "FlatField", "FlatField", "FlatField"],
            "Статус": ["В работе", "В работе", "В работе", "Остановлена"],
            "Текущий режим работы скважины | Дебит жидк.": [0, 100, 0, 500],
            "Текущий режим работы скважины | Частота": [0, 50, 52, 99],
        }
    )

    original_prepare = techregime_enrich_from_techregime.__globals__["_prepare_flat_techregime_frame"]
    techregime_enrich_from_techregime.__globals__["_prepare_flat_techregime_frame"] = lambda source: flat_df
    try:
        enriched = techregime_enrich_from_techregime(
            records,
            "dummy",
            prefer_sqlite=False,
            interval=7,
        )
    finally:
        techregime_enrich_from_techregime.__globals__["_prepare_flat_techregime_frame"] = original_prepare

    row = enriched.iloc[0]
    assert row["Дебит жидк."] == 100.0
    assert row["liquid_rate"] == 100.0
    assert row["Частота"] == 51.0
    assert row["frequency"] == 51.0


def test_enrich_from_techregime_zero_safe_columns_ignore_negative_values():
    records = pd.DataFrame(
        [
            {"well": "Vt_100", "failure_date": pd.Timestamp("2025-09-12")},
        ]
    )

    flat_df = pd.DataFrame(
        {
            "Дата": ["2025-09-10", "2025-09-11", "2025-09-12"],
            "ID скважины": ["Vt_100", "Vt_100", "Vt_100"],
            "М/р": ["FlatField", "FlatField", "FlatField"],
            "Статус": ["В работе", "В работе", "В работе"],
            "Текущий режим работы скважины | Рзаб": [-20, 0, 120],
            "Текущий режим работы скважины | Рпл.": [200, 200, 200],
        }
    )

    original_prepare = techregime_enrich_from_techregime.__globals__["_prepare_flat_techregime_frame"]
    techregime_enrich_from_techregime.__globals__["_prepare_flat_techregime_frame"] = lambda source: flat_df
    try:
        enriched = techregime_enrich_from_techregime(
            records,
            "dummy",
            prefer_sqlite=False,
            interval=7,
        )
    finally:
        techregime_enrich_from_techregime.__globals__["_prepare_flat_techregime_frame"] = original_prepare

    row = enriched.iloc[0]
    assert row["Рзаб"] == 120.0
    assert row["bottomhole_pressure"] == 120.0


def test_enrich_from_techregime_clears_bottomhole_pressure_when_it_exceeds_reservoir_pressure_for_producers():
    records = pd.DataFrame(
        [
            {"well": "Vt_100", "failure_date": pd.Timestamp("2025-09-12")},
        ]
    )

    flat_df = pd.DataFrame(
        {
            "Дата": ["2025-09-12"],
            "ID скважины": ["Vt_100"],
            "М/р": ["FlatField"],
            "Тип скв": ["Добывающая"],
            "Текущий режим работы скважины | Рзаб": [180],
            "Текущий режим работы скважины | Рпл.": [150],
        }
    )

    original_prepare = techregime_enrich_from_techregime.__globals__["_prepare_flat_techregime_frame"]
    techregime_enrich_from_techregime.__globals__["_prepare_flat_techregime_frame"] = lambda source: flat_df
    try:
        enriched = techregime_enrich_from_techregime(
            records,
            "dummy",
            prefer_sqlite=False,
            interval=7,
        )
    finally:
        techregime_enrich_from_techregime.__globals__["_prepare_flat_techregime_frame"] = original_prepare

    row = enriched.iloc[0]
    assert pd.isna(row["Рзаб"])
    assert pd.isna(row["bottomhole_pressure"])
    assert row["Рпл."] == 150.0


def test_enrich_from_techregime_zero_safe_columns_leave_blank_when_runtime_window_has_no_positive_values():
    records = pd.DataFrame(
        [
            {"well": "Vt_100", "failure_date": pd.Timestamp("2025-09-12"), "runtime_nno": 7},
        ]
    )

    flat_df = pd.DataFrame(
        {
            "Дата": ["2025-09-10", "2025-08-20", "2025-08-25", "2025-07-01"],
            "ID скважины": ["Vt_100", "Vt_100", "Vt_100", "Vt_100"],
            "М/р": ["FlatField", "FlatField", "FlatField", "FlatField"],
            "Статус": ["В работе", "В работе", "В работе", "В работе"],
            "Текущий режим работы скважины | Дебит жидк.": [0, 100, 140, 999],
            "Текущий режим работы скважины | Частота": [0, 48, 52, 10],
        }
    )

    original_prepare = techregime_enrich_from_techregime.__globals__["_prepare_flat_techregime_frame"]
    techregime_enrich_from_techregime.__globals__["_prepare_flat_techregime_frame"] = lambda source: flat_df
    try:
        enriched = techregime_enrich_from_techregime(
            records,
            "dummy",
            prefer_sqlite=False,
            interval=7,
        )
    finally:
        techregime_enrich_from_techregime.__globals__["_prepare_flat_techregime_frame"] = original_prepare

    row = enriched.iloc[0]
    assert pd.isna(row["Дебит жидк."])
    assert pd.isna(row["liquid_rate"])
    assert pd.isna(row["Частота"])
    assert pd.isna(row["frequency"])


def test_enrich_from_techregime_does_not_use_history_older_than_one_year():
    records = pd.DataFrame(
        [
            {"well": "Vt_100", "failure_date": pd.Timestamp("2025-09-12")},
        ]
    )

    flat_df = pd.DataFrame(
        {
            "Дата": ["2024-08-01", "2024-08-15"],
            "ID скважины": ["Vt_100", "Vt_100"],
            "М/р": ["FlatField", "FlatField"],
            "Статус": ["В работе", "В работе"],
            "Текущий режим работы скважины | Дебит жидк.": [100, 140],
            "Текущий режим работы скважины | Частота": [48, 52],
        }
    )

    original_prepare = techregime_enrich_from_techregime.__globals__["_prepare_flat_techregime_frame"]
    techregime_enrich_from_techregime.__globals__["_prepare_flat_techregime_frame"] = lambda source: flat_df
    try:
        enriched = techregime_enrich_from_techregime(
            records,
            "dummy",
            prefer_sqlite=False,
            interval=7,
        )
    finally:
        techregime_enrich_from_techregime.__globals__["_prepare_flat_techregime_frame"] = original_prepare

    row = enriched.iloc[0]
    assert pd.isna(row["Дебит жидк."])
    assert pd.isna(row["Частота"])


def test_enrich_from_techregime_can_skip_sqlite_even_if_one_is_available():
    records = pd.DataFrame(
        [
            {"well": "Vt_100", "failure_date": pd.Timestamp("2025-09-12")},
        ]
    )

    flat_df = pd.DataFrame(
        {
            "Дата": ["2025-09-12"],
            "ID скважины": ["Vt_100"],
            "М/р": ["FlatField"],
            "Текущий режим работы скважины | Дебит жидк.": [120],
        }
    )

    original_resolve_sqlite = techregime_enrich_from_techregime.__globals__["resolve_techregime_sqlite_path"]
    original_prepare = techregime_enrich_from_techregime.__globals__["_prepare_flat_techregime_frame"]
    techregime_enrich_from_techregime.__globals__["resolve_techregime_sqlite_path"] = lambda source: Path("techregime.sqlite")
    techregime_enrich_from_techregime.__globals__["_prepare_flat_techregime_frame"] = lambda source: flat_df
    try:
        enriched = techregime_enrich_from_techregime(
            records,
            "dummy",
            prefer_sqlite=False,
        )
    finally:
        techregime_enrich_from_techregime.__globals__["resolve_techregime_sqlite_path"] = original_resolve_sqlite
        techregime_enrich_from_techregime.__globals__["_prepare_flat_techregime_frame"] = original_prepare

    row = enriched.iloc[0]
    assert row["Дебит жидк."] == 120
    assert row["field"] == "FlatField"


def test_workflow_extract_pdk_can_filter_to_failures_only():
    source_df = pd.DataFrame(
        [
            {"well": "Vt_100", "failure_date": pd.Timestamp("2025-09-10"), "failure_flag": 1},
            {"well": "Vt_101", "failure_date": pd.Timestamp("2025-09-11"), "failure_flag": 0},
            {"well": "Vt_102", "failure_date": pd.Timestamp("2025-09-12"), "failure_flag": -1},
        ]
    )

    workflow = FailureUpdateWorkflow(
        target_path="target.xlsx",
        pdk_path="pdk.xlsx",
        event_scope="failures_only",
    )

    original_extract = FailureUpdateWorkflow._extract_pdk.__globals__["extract_new_failures_from_pdk"]
    FailureUpdateWorkflow._extract_pdk.__globals__["extract_new_failures_from_pdk"] = lambda path, cutoff, audit=None: source_df.copy()
    try:
        workflow._extract_pdk()
    finally:
        FailureUpdateWorkflow._extract_pdk.__globals__["extract_new_failures_from_pdk"] = original_extract

    assert workflow.new_failures["well"].tolist() == ["Vt_100"]
    assert workflow.pdk_df["well"].tolist() == ["Vt_100"]


def test_workflow_extract_pdk_can_exclude_ambiguous_events():
    source_df = pd.DataFrame(
        [
            {"well": "Vt_100", "failure_date": pd.Timestamp("2025-09-10"), "failure_flag": 1},
            {"well": "Vt_101", "failure_date": pd.Timestamp("2025-09-11"), "failure_flag": 0},
            {"well": "Vt_102", "failure_date": pd.Timestamp("2025-09-12"), "failure_flag": -1},
        ]
    )

    workflow = FailureUpdateWorkflow(
        target_path="target.xlsx",
        pdk_path="pdk.xlsx",
        event_scope="all",
        exclude_ambiguous_events=True,
    )

    original_extract = FailureUpdateWorkflow._extract_pdk.__globals__["extract_new_failures_from_pdk"]
    FailureUpdateWorkflow._extract_pdk.__globals__["extract_new_failures_from_pdk"] = lambda path, cutoff, audit=None: source_df.copy()
    try:
        workflow._extract_pdk()
    finally:
        FailureUpdateWorkflow._extract_pdk.__globals__["extract_new_failures_from_pdk"] = original_extract

    assert workflow.new_failures["well"].tolist() == ["Vt_100", "Vt_101"]
    assert workflow.pdk_df["well"].tolist() == ["Vt_100", "Vt_101"]


def test_detect_duplicates_uses_failure_date_not_launch_date_column():
    workflow = FailureUpdateWorkflow("target.xlsx", "pdk.xlsx", dry_run=True)
    workflow.target_df = pd.DataFrame(
        [
            {
                "Скв.": "Ic_329",
                "Дата запуска": pd.Timestamp("2018-07-15"),
                "Дата остановки": pd.Timestamp("2019-01-11"),
            }
        ]
    )
    workflow.new_failures = pd.DataFrame(
        [
            {
                "well": "Ic_329",
                "failure_date": pd.Timestamp("2019-01-11"),
            }
        ]
    )

    duplicates = workflow._detect_duplicates()

    assert duplicates[0].status == DuplicateStatus.EXACT_EXISTING


def test_detect_duplicates_prevents_duplicate_pairs_within_new_failures_batch():
    workflow = FailureUpdateWorkflow("target.xlsx", "pdk.xlsx", dry_run=True)
    workflow.target_df = pd.DataFrame(columns=["Скв.", "Дата остановки"])
    workflow.new_failures = pd.DataFrame(
        [
            {"well": "Ic_329", "failure_date": pd.Timestamp("2019-01-11")},
            {"well": "Ic_329", "failure_date": pd.Timestamp("2019-01-11")},
        ]
    )

    duplicates = workflow._detect_duplicates()

    assert duplicates[0].status == DuplicateStatus.NEW
    assert duplicates[1].status == DuplicateStatus.EXACT_EXISTING


def test_detect_duplicates_falls_back_to_installation_date_for_running_wells():
    workflow = FailureUpdateWorkflow("target.xlsx", "pdk.xlsx", dry_run=True)
    workflow.target_df = pd.DataFrame(
        [
            {"Скв.": "Vt_100", "Дата остановки": None, "Дата монтажа": pd.Timestamp("2025-09-01")},
        ]
    )
    workflow.new_failures = pd.DataFrame(
        [
            {"well": "Vt_100", "failure_date": None, "installation_date": pd.Timestamp("2025-09-01")},
        ]
    )

    duplicates = workflow._detect_duplicates()

    assert duplicates[0].status == DuplicateStatus.EXACT_EXISTING


def test_enrich_from_techregime_prefers_sqlite_found_inside_folder(tmp_path: Path):
    sqlite_path = tmp_path / "techregime.sqlite"
    sqlite_path.write_text("", encoding="utf-8")
    records = pd.DataFrame(
        [
            {"well": "Vt_100", "failure_date": pd.Timestamp("2025-09-12")},
        ]
    )

    original_history = techregime_enrich_from_techregime.__globals__["_load_sqlite_well_history"]
    techregime_enrich_from_techregime.__globals__["_load_sqlite_well_history"] = lambda sqlite_path, well_value=None, failure_date=None, interval=None: pd.DataFrame(
        [
            {
                "М/р": "SqlField",
                "ID скважины": "Vt_100",
                "Дата": "2025-09-10",
                "Текущий режим работы скважины | Дебит жидк.": 150,
                "_meta_report_date": "2025-09-10",
            }
        ]
    )
    try:
        enriched = techregime_enrich_from_techregime(records, [tmp_path])
    finally:
        techregime_enrich_from_techregime.__globals__["_load_sqlite_well_history"] = original_history

    row = enriched.iloc[0]
    assert row["field"] == "SqlField"
    assert row["Месторождение"] == "SqlField"
    assert row["Дебит жидк."] == 150
    assert row["liquid_rate"] == 150
    assert row["techregime_match_confidence"] == 0.8


def test_enrich_from_techregime_uses_new_style_flat_excel_folder(tmp_path: Path):
    first_file = tmp_path / "export_1.xlsx"
    second_file = tmp_path / "export_2.xlsx"

    common_header = [
        ["М/р", "ID скважины", "Дата", "Текущий режим работы скважины", "Текущий режим работы скважины"],
        ["", "", "", "Дебит жидк.", "Частота"],
    ]
    for workbook_path, rows in [
        (
            first_file,
            common_header
            + [
                ["FlatField", "Vt_100", "10.09.2025", 110, None],
                ["FlatField", "Vt_100", "11.09.2025", None, 51],
            ],
        ),
        (
            second_file,
            common_header
            + [
                ["FlatField", "Vt_100", "12.09.2025", 120, 52],
                ["FlatField", "Vt_200", "12.09.2025", 130, 53],
            ],
        ),
    ]:
        wb = Workbook()
        ws = wb.active
        for row in rows:
            ws.append(row)
        wb.save(workbook_path)

    records = pd.DataFrame(
        [
            {"well": "Vt_100", "failure_date": pd.Timestamp("2025-09-12")},
        ]
    )

    enriched = techregime_enrich_from_techregime(records, [tmp_path])
    row = enriched.iloc[0]
    assert row["field"] == "FlatField"
    assert row["Месторождение"] == "FlatField"
    assert row["Дебит жидк."] == 115
    assert row["liquid_rate"] == 115
    assert row["Частота"] == 51.5
    assert row["frequency"] == 51.5
    assert row["techregime_match_confidence"] == 0.78


def test_discover_available_techregime_columns_maps_builtin_targets_and_keeps_extra_columns():
    flat_df = pd.DataFrame(
        [
            {
                "Дата": pd.Timestamp("2025-09-11"),
                "ID скважины": "Vt_100",
                "М/р": "FlatField",
                "Текущий режим работы скважины | Дебит жидк.": 110,
                "Текущий режим работы скважины | Дебит газа": 230,
                "Текущий режим работы скважины | Давление буфера": 17.5,
            }
        ]
    )

    original_prepare = discover_available_techregime_columns.__globals__["_prepare_flat_techregime_frame"]
    original_resolve = discover_available_techregime_columns.__globals__["resolve_techregime_sqlite_path"]
    discover_available_techregime_columns.__globals__["_prepare_flat_techregime_frame"] = lambda source: flat_df
    discover_available_techregime_columns.__globals__["resolve_techregime_sqlite_path"] = lambda source: None
    try:
        columns = discover_available_techregime_columns(["dummy"])
    finally:
        discover_available_techregime_columns.__globals__["_prepare_flat_techregime_frame"] = original_prepare
        discover_available_techregime_columns.__globals__["resolve_techregime_sqlite_path"] = original_resolve

    assert "Месторождение" in columns
    assert "Дебит жидк." in columns
    assert "Дебит газа" in columns
    assert "ГЖФ" in columns
    assert "Давление буфера" in columns
    assert "Дата" not in columns
    assert "ID скважины" not in columns


def test_enrich_from_techregime_populates_selected_extra_flat_column():
    flat_df = pd.DataFrame(
        [
            {
                "Дата": pd.Timestamp("2025-09-10"),
                "ID скважины": "Vt_100",
                "Текущий режим работы скважины | Давление буфера": 16.0,
            },
            {
                "Дата": pd.Timestamp("2025-09-12"),
                "ID скважины": "Vt_100",
                "Текущий режим работы скважины | Давление буфера": 18.0,
            },
        ]
    )
    records = pd.DataFrame(
        [
            {
                "well": "Vt_100",
                "failure_date": pd.Timestamp("2025-09-12"),
            }
        ]
    )

    original_prepare = techregime_enrich_from_techregime.__globals__["_prepare_flat_techregime_frame"]
    techregime_enrich_from_techregime.__globals__["_prepare_flat_techregime_frame"] = lambda source: flat_df
    try:
        enriched = techregime_enrich_from_techregime(
            records,
            techregime_path=["dummy"],
            prefer_sqlite=False,
            interval=7,
            selected_columns=["Давление буфера"],
        )
    finally:
        techregime_enrich_from_techregime.__globals__["_prepare_flat_techregime_frame"] = original_prepare

    assert enriched.iloc[0]["Давление буфера"] == 17.0


def test_refresh_existing_target_techregime_fields_updates_target_dataframe():
    workflow = FailureUpdateWorkflow(
        "target.xlsx",
        "pdk.xlsx",
        techregime_path="dummy.sqlite",
        refresh_techregime_existing=True,
        dry_run=True,
    )
    workflow.target_df = pd.DataFrame(
        [
            {
                "A. Месторождение": "OldField",
                "C. Скв.": "Vt_100",
                "D. Дата остановки": pd.Timestamp("2025-09-12"),
                "E. Дебит жидк.": None,
                "F. ГЖФ": None,
            }
        ]
    )
    workflow.audit = workflow.audit.__class__()

    original_enrich = workflow._refresh_existing_target_techregime_fields.__globals__["enrich_from_techregime"]
    workflow._refresh_existing_target_techregime_fields.__globals__["enrich_from_techregime"] = lambda records, techregime_path, well_col="well", failure_date_col="failure_date", **kwargs: records.assign(
        field="RemoteField",
        **{
            "Месторождение": "RemoteField",
            "Дебит жидк.": 123.0,
            "ГЖФ": 9.5,
        },
    )
    try:
        workflow._refresh_existing_target_techregime_fields()
    finally:
        workflow._refresh_existing_target_techregime_fields.__globals__["enrich_from_techregime"] = original_enrich

    assert workflow.target_df.iloc[0]["A. Месторождение"] == "RemoteField"
    assert workflow.target_df.iloc[0]["E. Дебит жидк."] == 123.0
    assert workflow.target_df.iloc[0]["F. ГЖФ"] == 9.5


def test_refresh_existing_target_techregime_fields_adds_selected_extra_column_and_fills_it():
    workflow = FailureUpdateWorkflow(
        "target.xlsx",
        "pdk.xlsx",
        techregime_path="dummy.sqlite",
        refresh_techregime_existing=True,
        dry_run=True,
        techregime_selected_columns=["Месторождение", "Давление буфера"],
    )
    workflow.target_df = pd.DataFrame(
        [
            {
                "A. Месторождение": "OldField",
                "C. Скв.": "Vt_100",
                "D. Дата остановки": pd.Timestamp("2025-09-12"),
            }
        ]
    )
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "A. Месторождение"
    ws["B1"] = "C. Скв."
    ws["C1"] = "D. Дата остановки"
    ws["A2"] = "OldField"
    ws["B2"] = "Vt_100"
    ws["C2"] = pd.Timestamp("2025-09-12")
    workflow.target_wb = wb
    workflow.target_ws = ws
    workflow.audit = workflow.audit.__class__()

    original_enrich = workflow._refresh_existing_target_techregime_fields.__globals__["enrich_from_techregime"]
    workflow._refresh_existing_target_techregime_fields.__globals__["enrich_from_techregime"] = (
        lambda records, techregime_path, well_col="well", failure_date_col="failure_date", **kwargs: records.assign(
            field="RemoteField",
            **{
                "Месторождение": "RemoteField",
                "Давление буфера": 17.25,
            },
        )
    )
    try:
        added_columns = workflow._ensure_selected_techregime_target_columns()
        workflow._refresh_existing_target_techregime_fields()
    finally:
        workflow._refresh_existing_target_techregime_fields.__globals__["enrich_from_techregime"] = original_enrich

    assert added_columns == ["Давление буфера"]
    assert "Давление буфера" in workflow.target_df.columns
    assert workflow.target_df.iloc[0]["Давление буфера"] == 17.25
    assert ws.cell(row=1, column=4).value == "Давление буфера"
    assert ws.cell(row=2, column=4).value == 17.25


def test_discover_available_techregime_columns_strips_top_level_group_for_extra_columns():
    flat_df = pd.DataFrame(
        [
            {
                "Дата": pd.Timestamp("2025-09-11"),
                "ID скважины": "Vt_100",
                "Свойства флюида - PVT | Плотн. воды на приеме ЭЦН": 1005.0,
            }
        ]
    )

    original_prepare = discover_available_techregime_columns.__globals__["_prepare_flat_techregime_frame"]
    original_resolve = discover_available_techregime_columns.__globals__["resolve_techregime_sqlite_path"]
    discover_available_techregime_columns.__globals__["_prepare_flat_techregime_frame"] = lambda source: flat_df
    discover_available_techregime_columns.__globals__["resolve_techregime_sqlite_path"] = lambda source: None
    try:
        columns = discover_available_techregime_columns(["dummy"])
    finally:
        discover_available_techregime_columns.__globals__["_prepare_flat_techregime_frame"] = original_prepare
        discover_available_techregime_columns.__globals__["resolve_techregime_sqlite_path"] = original_resolve

    assert "Плотн. воды на приеме ЭЦН" in columns
    assert "Свойства флюида - PVT | Плотн. воды на приеме ЭЦН" not in columns


def test_map_columns_handles_prefixed_target_headers_and_acid_type():
    workflow = FailureUpdateWorkflow("target.xlsx", "pdk.xlsx", dry_run=True)
    workflow.target_df = pd.DataFrame(columns=[
        "A. Месторождение",
        "C. Скв.",
        "I. Наработка (сут)",
        "J. Причина остановки",
        "BW. Кислый/Некислый",
    ])

    mapped_row = workflow._map_columns(
        {
            "field": "Vt",
            "well": "Vt_123",
            "runtime_nno": 17.0,
            "failure_reason": "Разгерметизация",
            "acid_type": "Некислый",
        }
    )

    assert mapped_row == ["Vt", "Vt_123", 17.0, "Разгерметизация", "Некислый"]


def test_map_columns_handles_failure_flag_target_header():
    workflow = FailureUpdateWorkflow("target.xlsx", "pdk.xlsx", dry_run=True)
    workflow.target_df = pd.DataFrame(columns=[
        "C. Скв.",
        "J. Причина остановки",
        "K. Флаг отказа",
    ])

    mapped_row = workflow._map_columns(
        {
            "well": "Vt_123",
            "failure_reason": "Разгерметизация",
            "failure_flag": 1,
        }
    )

    assert mapped_row == ["Vt_123", "Разгерметизация", 1]


def test_map_columns_handles_failure_marker_target_header_separately_from_failure_flag():
    workflow = FailureUpdateWorkflow("target.xlsx", "pdk.xlsx", dry_run=True)
    workflow.target_df = pd.DataFrame(columns=[
        "K. Флаг отказа",
        "L. Признак отказа",
    ])

    mapped_row = workflow._map_columns(
        {
            "failure_flag": 1,
            "failure_marker": "Аварийная",
        }
    )

    assert mapped_row == [1, "Аварийная"]


def test_map_columns_handles_failure_flag_target_header_in_english():
    workflow = FailureUpdateWorkflow("target.xlsx", "pdk.xlsx", dry_run=True)
    workflow.target_df = pd.DataFrame(columns=[
        "Failure Flag",
    ])

    mapped_row = workflow._map_columns(
        {
            "failure_flag": -1,
        }
    )

    assert mapped_row == [-1]


def test_load_target_can_add_and_fill_failure_flag_column_for_existing_rows():
    workflow = FailureUpdateWorkflow(
        "target.xlsx",
        "pdk.xlsx",
        dry_run=True,
        include_event_failure_flag=True,
    )
    workflow.target_df = pd.DataFrame(
        [
            {
                "Причина остановки": "Отказ",
                "Отказавший узел": "Насос",
            },
            {
                "Причина остановки": "ГТМ",
                "Отказавший узел": "нет",
            },
        ]
    )
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Причина остановки"
    ws["B1"] = "Отказавший узел"
    ws["A2"] = "Отказ"
    ws["B2"] = "Насос"
    ws["A3"] = "ГТМ"
    ws["B3"] = "нет"
    workflow.target_wb = wb
    workflow.target_ws = ws

    added = workflow._ensure_event_failure_flag_column()
    workflow._populate_existing_event_failure_flag_column()

    assert added is True
    assert workflow.target_df.iloc[0][FailureUpdateWorkflow.EVENT_FAILURE_FLAG_COLUMN] == 1
    assert workflow.target_df.iloc[1][FailureUpdateWorkflow.EVENT_FAILURE_FLAG_COLUMN] == 0
    assert ws.cell(row=1, column=3).value == FailureUpdateWorkflow.EVENT_FAILURE_FLAG_COLUMN
    assert ws.cell(row=2, column=3).value == 1
    assert ws.cell(row=3, column=3).value == 0


def test_map_columns_handles_exact_russian_target_headers_from_pdk():
    workflow = FailureUpdateWorkflow("target.xlsx", "pdk.xlsx", dry_run=True)
    workflow.target_df = pd.DataFrame(columns=[
        "Принадлежность",
        "Тип УЭЦН",
        "Нспуска",
        "Дата монтажа",
        "Причина отказа УЭЦН",
    ])

    mapped_row = workflow._map_columns(
        {
            "Принадлежность": "Борец",
            "Тип УЭЦН": "5А-160-2400",
            "Нспуска": 2400,
            "Дата монтажа": pd.Timestamp("2025-09-01"),
            "Причина отказа УЭЦН": "Солеотложения",
        }
    )

    assert mapped_row == [
        "Борец",
        "5А-160-2400",
        2400,
        pd.Timestamp("2025-09-01"),
        "Солеотложения",
    ]


def test_map_columns_handles_additional_safe_pdk_to_target_mappings():
    workflow = FailureUpdateWorkflow("target.xlsx", "pdk.xlsx", dry_run=True)
    workflow.target_df = pd.DataFrame(columns=[
        "Габарит УЭЦН",
        "Ном. Произв. м₃/сут",
        "Глубина спуска УЭЦН, по НКТ",
    ])

    mapped_row = workflow._map_columns(
        {
            "Габарит": "5А",
            "Производительность": 250,
            "descent_depth": 2450,
        }
    )

    assert mapped_row == ["5А", 250, 2450]


def test_append_row_to_worksheet_copies_alignment_from_template_row():
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Месторождение"
    ws["B1"] = "Скв."
    ws["A2"] = "Vt"
    ws["B2"] = "Vt_001"
    alignment = copy(ws["A2"].alignment)
    alignment.horizontal = "center"
    ws["A2"].alignment = alignment
    ws["B2"].alignment = copy(alignment)

    inserted_row = append_row_to_worksheet(
        ws,
        {"Месторождение": "Au", "Скв.": "Au_001"},
        header_row=1,
        copy_format_from_row=2,
    )

    assert inserted_row == 3
    assert ws["A3"].value == "Au"
    assert ws["B3"].value == "Au_001"
    assert ws["A3"].alignment.horizontal == "center"
    assert ws["B3"].alignment.horizontal == "center"


def test_highlight_runtime_cell_if_needed_marks_only_values_below_90_red():
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "I. Наработка (сут)"
    ws["A2"] = 45
    ws["A3"] = 90

    highlight_runtime_cell_if_needed(ws, 2, "I. Наработка (сут)", 45, threshold_days=90)
    highlight_runtime_cell_if_needed(ws, 3, "I. Наработка (сут)", 90, threshold_days=90)

    assert ws["A2"].fill.fill_type == "solid"
    assert ws["A2"].fill.start_color.rgb == "FFFF0000"
    assert ws["A3"].fill.fill_type is None


def test_enrich_lab_chemistry_replaces_missing_with_dash():
    records = pd.DataFrame([{"well": "Vt_100", "pH": None}])
    enriched = enrich_lab_chemistry(records)
    assert enriched.iloc[0]["pH"] == "-"
    assert enriched.iloc[0]["Cl⁻, мг/л"] == "-"


def test_collapse_duplicate_lab_columns_coalesces_same_named_columns():
    lab_df = pd.DataFrame(
        [
            ["Vt_100", None, 7.1, None, 100.0],
            [None, "Vt_200", None, 6.8, 200.0],
        ],
        columns=["Скважина", "Скважина", "pH", "pH", "Cl⁻, мг/л"],
    )

    collapsed = collapse_duplicate_lab_columns(lab_df)

    assert list(collapsed.columns) == ["Скважина", "pH", "Cl⁻, мг/л"]
    assert collapsed.iloc[0]["Скважина"] == "Vt_100"
    assert collapsed.iloc[1]["Скважина"] == "Vt_200"
    assert collapsed.iloc[0]["pH"] == 7.1
    assert collapsed.iloc[1]["pH"] == 6.8


def test_build_and_fill_esp_specs_from_target_mapping():
    target_df = pd.DataFrame(
        [
            {"Тип УЭЦН": "\tMT5A-100DP", "Габарит УЭЦН": "5A", "Ном. Произв. м₃/сут": 93, "Ном.напор (50Гц)": 2021},
            {"Тип УЭЦН": "MT5A-100DP", "Габарит УЭЦН": "5А", "Ном. Произв. м₃/сут": 93, "Ном.напор (50Гц)": 2021},
            {"Тип УЭЦН": "DN5850", "Габарит УЭЦН": "DN-387", "Ном. Произв. м₃/сут": 776, "Ном.напор (50Гц)": 1605},
        ]
    )
    specs = build_esp_specs_from_target(target_df)
    assert specs["MT5A-100DP"]["productivity"] == "93"
    assert specs["MT5A-100DP"]["nominal_head"] == "2021"

    records = pd.DataFrame(
        [
            {"Тип УЭЦН": "MT5A-100DP", "Габарит УЭЦН": None, "Ном. Произв. м₃/сут": None, "Ном.напор (50Гц)": None},
            {"Тип УЭЦН": "DN5850", "Габарит УЭЦН": "-", "Ном. Произв. м₃/сут": "-", "Ном.напор (50Гц)": "-"},
        ]
    )
    filled = fill_specs_from_esp_type(records, specs)
    assert filled.iloc[0]["Габарит УЭЦН"] == "5А"
    assert filled.iloc[0]["Ном. Произв. м₃/сут"] == "93"
    assert filled.iloc[0]["Ном.напор (50Гц)"] == "2021"
    assert filled.iloc[1]["Габарит УЭЦН"] == "DN-387"
    assert filled.iloc[1]["Ном. Произв. м₃/сут"] == "776"
    assert filled.iloc[1]["Ном.напор (50Гц)"] == "1605"


def test_enrich_from_big_extracts_equipment_blocks_and_computes_rate_delta():
    big_path = Path("test_big_fixture.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Скважинное оборудование"

    row1 = [
        "Скважина",
        "Дата отказа",
        "Дата монтажа",
        "Дата запуска",
        "Дата демонтажа",
        "Насос (50Гц)",
        "Насос (50Гц)",
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        "ПЭД",
        None,
        None,
        None,
        "ФА/НКТ",
        None,
        None,
        None,
    ]
    row2 = [
        None,
        None,
        None,
        None,
        None,
        "Модель ГНО",
        "Собственник оборудования",
        "Габарит УЭЦН",
        "Длина УЭЦН/метр",
        "Работа в кривизне",
        "Ном. Произв. м3/сут",
        "Ном.напор (50Гц)",
        "Номинальная частота, Гц",
        "Кол.ступеней",
        "Коррозионная стойкость",
        "Мощность, кВт",
        "Ном. напряжение/В",
        "Ном. ток/ A",
        "Ток x.x",
        "Глубина спуска УЭЦН, по НКТ",
        "ВГ, м",
        "Диаметр НКТ",
        "Марка НКТ",
    ]

    for col_idx, value in enumerate(row1, start=1):
        ws.cell(row=1, column=col_idx, value=value)
    for col_idx, value in enumerate(row2, start=1):
        ws.cell(row=2, column=col_idx, value=value)

    data_rows = [
        ["Vt_100", "2025-09-10", "2025-09-01", "2025-09-02", "2025-09-11", "MT5A-100DP", "Борец", "5А", 12.5, "нет", 90, 2100, 50, 300, "KR", 45, 1200, 30, 8, 2500, 150, 73, "J55"],
        ["Vt_100", "2025-09-15", "2025-09-05", "2025-09-06", "2025-09-16", "MT5A-100DP", "Борец", "5А", 12.8, "да", 100, 2200, 55, 320, "KR2", 50, 1250, 31, 9, 2550, 155, 73, "N80"],
    ]
    for row_idx, values in enumerate(data_rows, start=3):
        for col_idx, value in enumerate(values, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    wb.save(big_path)

    records = pd.DataFrame(
        [
            {
                "well": "Vt_100",
                "failure_date": pd.Timestamp("2025-09-12"),
                "Дебит жидк.": 120,
            }
        ]
    )

    try:
        enriched = enrich_from_big(records, str(big_path))
        enriched = enrich_big_derived_fields(enriched)
        row = enriched.iloc[0]
        assert row["contractor"] == "Борец"
        assert row["esp_type"] == "MT5A-100DP"
        assert row["Габарит УЭЦН"] == "5А"
        assert row["Длина УЭЦН/метр"] == 12.5
        assert row["Работа в кривизне"] == "нет"
        assert row["Ном. Произв. м₃/сут"] == 90
        assert row["Дельта Дебита Ж и номинала, м3/сут"] == 30
        assert row["Ном.напор (50Гц)"] == 2100
        assert row["Номинальная частота, Гц"] == 50
        assert row["Кол.ступеней"] == 300
        assert row["Коррозионная стойкость"] == "KR"
        assert row["Мощность, кВт"] == 45
        assert row["Ном. напряжение/В"] == 1200
        assert row["Ном. ток/ A"] == 30
        assert row["Ток x.x"] == 8
        assert row["Глубина спуска УЭЦН, по НКТ"] == 2500
        assert row["ВГ, м"] == 150
        assert row["Диаметр НКТ"] == 73
        assert row["Марка НКТ"] == "J55"
    finally:
        if big_path.exists():
            try:
                big_path.unlink()
            except PermissionError:
                pass


def test_extract_running_wells_from_artificial_lift_returns_active_runs_only():
    big_path = Path("test_big_running_fixture.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Скважинное оборудование"

    row1 = [
        "Скважина",
        "Цель спуска",
        "Дата отказа",
        "Дата монтажа",
        "Дата запуска",
        "Дата демонтажа",
        "ННО",
        "Дизайн УЭЦН",
        "Насос (50Гц)",
    ]
    row2 = [
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        "Собственник оборудования",
    ]
    for col_idx, value in enumerate(row1, start=1):
        ws.cell(row=1, column=col_idx, value=value)
    for col_idx, value in enumerate(row2, start=1):
        ws.cell(row=2, column=col_idx, value=value)

    data_rows = [
        # active mechanized-production run -> kept
        ["Vt_100", "Мех. добыча", None, "2025-09-01", "2025-09-02", None, 20, "MT5A-100DP", "Борец"],
        # completed run (has failure/dismantling dates) -> dropped
        ["Vt_200", "Мех. добыча", "2025-09-10", "2025-09-01", "2025-09-02", "2025-09-11", 9, "MT5A-200DP", "Борец"],
        # active but non-oil (injection) run -> dropped by purpose filter
        ["Vt_300", "Нагнетательная", None, "2025-09-01", "2025-09-02", None, 15, "MT5A-300DP", "Борец"],
    ]
    for row_idx, values in enumerate(data_rows, start=3):
        for col_idx, value in enumerate(values, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    wb.save(big_path)

    try:
        running = extract_running_wells_from_artificial_lift(str(big_path))

        assert len(running) == 1
        row = running.iloc[0]
        assert row["well"] == "Vt_100"
        assert row["field"] == "Vt"
        assert row["failure_flag"] == 0
        assert row["runtime_nno"] == 20
        assert row["techregime_query_date"] == pd.Timestamp("2025-09-21")
        assert pd.isna(row["failure_date"])
        # Injection/non-oil active runs must be excluded, mirroring the PDK НФ filter.
        assert "Vt_300" not in set(running["well"])
    finally:
        if big_path.exists():
            try:
                big_path.unlink()
            except PermissionError:
                pass


def test_enrich_from_telemetry_supports_folder_lookup():
    tel_df = pd.DataFrame(
        [
            {"Скважина": "Vt_100", "Дата": "2025-09-10", "Частота вращения двиг., Гц": 50, "Загрузка, %": 70, "Р на приеме насоса, атм": 25},
            {"Скважина": "Vt_100", "Дата": "2025-09-11", "Частота вращения двиг., Гц": 51, "Загрузка, %": 71, "Р на приеме насоса, атм": 26},
        ]
    )

    records = pd.DataFrame(
        [{"field": "Vt", "well": "Vt_100", "failure_date": pd.Timestamp("2025-09-12")}]
    )
    original_discover = enrich_from_telemetry.__globals__["discover_source_files"]
    original_load = enrich_from_telemetry.__globals__["load_excel_file"]
    enrich_from_telemetry.__globals__["discover_source_files"] = lambda source, dataset_type: [Path("test_telemetry_Vt.xlsx")]
    enrich_from_telemetry.__globals__["load_excel_file"] = lambda path: (tel_df.copy(), "Sheet1", 0)
    try:
        enriched = enrich_from_telemetry(records, ".")
        row = enriched.iloc[0]
        assert row["frequency"] == 51
        assert row["motor_load"] == 71
        assert row["intake_pressure"] == 26
    finally:
        enrich_from_telemetry.__globals__["discover_source_files"] = original_discover
        enrich_from_telemetry.__globals__["load_excel_file"] = original_load


def test_enrich_from_telemetry_uses_runtime_window_average():
    tel_df = pd.DataFrame(
        [
            {"Скважина": "Vt_100", "Дата": "2025-08-01", "Частота вращения двиг., Гц": 99, "Загрузка, %": 99},
            {"Скважина": "Vt_100", "Дата": "2025-09-10", "Частота вращения двиг., Гц": 50, "Загрузка, %": 70},
            {"Скважина": "Vt_100", "Дата": "2025-09-11", "Частота вращения двиг., Гц": 52, "Загрузка, %": 74},
            {"Скважина": "Vt_100", "Дата": "2025-09-12", "Частота вращения двиг., Гц": 54, "Загрузка, %": 76},
        ]
    )

    records = pd.DataFrame(
        [{"field": "Vt", "well": "Vt_100", "failure_date": pd.Timestamp("2025-09-12"), "runtime_nno": 3}]
    )
    original_discover = enrich_from_telemetry.__globals__["discover_source_files"]
    original_load = enrich_from_telemetry.__globals__["load_excel_file"]
    enrich_from_telemetry.__globals__["discover_source_files"] = lambda source, dataset_type: [Path("test_telemetry_Vt_runtime.xlsx")]
    enrich_from_telemetry.__globals__["load_excel_file"] = lambda path: (tel_df.copy(), "Sheet1", 0)
    try:
        enriched = enrich_from_telemetry(records, ".")
        row = enriched.iloc[0]
        assert row["frequency"] == 52.0
        assert row["Частота"] == 52.0
        assert abs(row["motor_load"] - 73.33333333333333) < 1e-9
        assert abs(row["Загр, Двиг,"] - 73.33333333333333) < 1e-9
    finally:
        enrich_from_telemetry.__globals__["discover_source_files"] = original_discover
        enrich_from_telemetry.__globals__["load_excel_file"] = original_load


def test_enrich_from_telemetry_only_backfills_missing_techregime_overlap_fields():
    tel_df = pd.DataFrame(
        [
            {"Скважина": "Vt_100", "Дата": "2025-09-10", "Частота вращения двиг., Гц": 50, "Загрузка, %": 70},
            {"Скважина": "Vt_100", "Дата": "2025-09-11", "Частота вращения двиг., Гц": 52, "Загрузка, %": 74},
        ]
    )

    records = pd.DataFrame(
        [
            {
                "field": "Vt",
                "well": "Vt_100",
                "failure_date": pd.Timestamp("2025-09-12"),
                "runtime_nno": 5,
                "frequency": 60.0,
                "Частота": 60.0,
                "motor_load": None,
                "Загр, Двиг,": None,
            }
        ]
    )
    original_discover = enrich_from_telemetry.__globals__["discover_source_files"]
    original_load = enrich_from_telemetry.__globals__["load_excel_file"]
    enrich_from_telemetry.__globals__["discover_source_files"] = lambda source, dataset_type: [Path("test_telemetry_Vt_backfill.xlsx")]
    enrich_from_telemetry.__globals__["load_excel_file"] = lambda path: (tel_df.copy(), "Sheet1", 0)
    try:
        enriched = enrich_from_telemetry(records, ".")
        row = enriched.iloc[0]
        assert row["frequency"] == 60.0
        assert row["Частота"] == 60.0
        assert row["motor_load"] == 72.0
        assert row["Загр, Двиг,"] == 72.0
    finally:
        enrich_from_telemetry.__globals__["discover_source_files"] = original_discover
        enrich_from_telemetry.__globals__["load_excel_file"] = original_load


def test_enrich_from_lab_uses_label_mapping_and_averages_recent_rows():
    lab_path = Path("test_lab_fixture.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Нефть"

    header_rows = [
        [None] * 9,
        [
            "Скважина",
            "Дата время отбора",
            "Массовая концентрация хлористых солей в нефти, мг/дм³",
            "Массовая доля механических примесей в нефти, мг/дм³",
            "6К",
            "6К",
            "6К",
            "6К",
            "6К",
        ],
        [
            None,
            None,
            None,
            None,
            "Состав воды",
            "Состав воды",
            "Общая минерализация, г/л",
            "pH",
            "Механические примеси (КВЧ), мг/дм³",
        ],
        [
            None,
            None,
            None,
            None,
            "Cl⁻, мг/л",
            "SO₄²⁻, мг/л",
            None,
            None,
            None,
        ],
    ]
    for row_idx, values in enumerate(header_rows, start=1):
        for col_idx, value in enumerate(values, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    data_rows = [
        ["Vt_100", pd.Timestamp("2025-09-10"), 10, 100, 1000, 10, 50, 6.0, 7],
        ["Vt_100", pd.Timestamp("2025-09-11"), 30, 300, 3000, 30, 70, 8.0, 9],
        ["Vt_100", pd.Timestamp("2025-08-01"), 100, 900, 9000, 90, 90, 9.0, 11],
    ]
    for row_idx, values in enumerate(data_rows, start=5):
        for col_idx, value in enumerate(values, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    wb.save(lab_path)

    records = pd.DataFrame(
        [
            {
                "well": "Vt_100",
                "failure_date": pd.Timestamp("2025-09-12"),
            }
        ]
    )

    try:
        enriched = enrich_from_lab(records, str(lab_path), av_window=3)
        row = enriched.iloc[0]
        assert row["Массовая концентрация хлористых солей в нефти, мг/дм³"] == 20.0
        assert row["Массовая доля механических примесей в нефти, мг/дм³"] == 200.0
        assert row["Cl⁻, мг/л"] == 2000.0
        assert row["SO₄²⁻, мг/л"] == 20.0
        assert row["Общая минерализация, г/л"] == 60.0
        assert row["pH"] == 7.0
        assert row["Механические примеси (КВЧ), мг/дм³"] == 8.0
    finally:
        if lab_path.exists():
            try:
                lab_path.unlink()
            except PermissionError:
                pass


def test_collapse_never_lets_a_same_day_incident_erase_a_multi_year_run():
    """⚠⚠ Регрессия, найденная при переносе: предпочтение отказу — тай-брейк, не фильтр.

    ПДК записывает под одной датой монтажа и происшествие в день спуска, и
    плановый подъём годы спустя. Пока предпочтение подтверждённому отказу было
    ФИЛЬТРОМ, длинная строка выбрасывалась целиком, перенос максимальной
    наработки писал её ННО на короткую строку, а санация обрезала его до нуля по
    календарю. Реальный случай — Ya_622, монтаж 2019-07-30: пуск на 1359 суток
    становился отказом на 0 суток.

    Направление ошибки худшее из возможных: она фабрикует события в полосе
    0–5 суток, где решается вопрос о детской смертности.
    """
    source_df = pd.DataFrame(
        [
            {
                "Скв.": "Ya_622",
                "Дата монтажа": "2019-07-30",
                "Дата остановки": "2019-07-30",
                "Наработка (сут)": 0,
                "Причина остановки": "Прочие",
                "Отказавший узел": "Кабельная линия",
                "Причина отказа УЭЦН": "Механическое повреждение кабеля",
                "Тип скважины": "НФ",
            },
            {
                "Скв.": "Ya_622",
                "Дата монтажа": "2019-07-30",
                "Дата остановки": "2023-04-21",
                "Наработка (сут)": 1359,
                "Причина остановки": "ГТМ",
                "Отказавший узел": "Нет",
                "Тип скважины": "НФ",
            },
        ]
    )

    original_loader = extract_new_failures_from_pdk.__globals__["load_dataset_frames"]
    extract_new_failures_from_pdk.__globals__["load_dataset_frames"] = lambda path, dataset_type: [
        (Path("pdk.xlsx"), source_df, "Sheet1", 0),
    ]
    try:
        extracted = extract_new_failures_from_pdk("dummy.xlsx", None)
    finally:
        extract_new_failures_from_pdk.__globals__["load_dataset_frames"] = original_loader

    assert len(extracted) == 1
    row = extracted.iloc[0]
    assert row["Дата остановки"] == pd.Timestamp("2023-04-21"), "поздний подъём должен решать"
    assert row["Наработка (сут)"] == 1359.0, "наработка пуска не должна схлопываться в ноль"


def test_collapse_still_prefers_the_failure_when_the_dates_tie():
    """Предпочтение отказу сохраняется — но только среди строк с ОДНОЙ датой."""
    source_df = pd.DataFrame(
        [
            {
                "Скв.": "Vt_777",
                "Дата монтажа": "2024-01-01",
                "Дата остановки": "2024-03-01",
                "Наработка (сут)": 60,
                "Причина остановки": "ГТМ",
                "Отказавший узел": "нет",
                "Тип скважины": "НФ",
            },
            {
                "Скв.": "Vt_777",
                "Дата монтажа": "2024-01-01",
                "Дата остановки": "2024-03-01",
                "Наработка (сут)": 60,
                "Причина остановки": "Снижение изоляции",
                "Отказавший узел": "ПЭД",
                "Тип скважины": "НФ",
            },
        ]
    )

    original_loader = extract_new_failures_from_pdk.__globals__["load_dataset_frames"]
    extract_new_failures_from_pdk.__globals__["load_dataset_frames"] = lambda path, dataset_type: [
        (Path("pdk.xlsx"), source_df, "Sheet1", 0),
    ]
    try:
        extracted = extract_new_failures_from_pdk("dummy.xlsx", None)
    finally:
        extract_new_failures_from_pdk.__globals__["load_dataset_frames"] = original_loader

    assert len(extracted) == 1
    assert extracted.iloc[0]["failure_flag"] == 1
    assert extracted.iloc[0]["Отказавший узел"] == "ПЭД"
