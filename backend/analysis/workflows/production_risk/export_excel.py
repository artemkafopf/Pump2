"""Russian-language companion workbook for production-risk outputs."""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from analysis.workflows.production_risk import config as C

_HDR_FILL = PatternFill("solid", fgColor="1F4E79")
_HDR_FONT = Font(bold=True, color="FFFFFF")
_TITLE_FONT = Font(bold=True, size=13)
_NOTE_FONT = Font(italic=True, size=9, color="666666")

_COLS_RU = {
    "scenario": "Сценарий",
    "scenario_label": "Сценарий (RU)",
    "scenario_primary": "Основной сценарий",
    "month": "Месяц",
    "period_type": "Период",
    "report_field": "Поле отчета",
    "plan_field": "Месторождение",
    "field_label": "Поле/месторождение",
    "risk_model_field": "Поле риск-модели",
    "wid": "Скважина",
    "raw_id": "Well_ID",
    "model_field": "Поле модели",
    "stratum": "Страта модели",
    "state_label": "Состояние скважины",
    "scope_label": "Охват ESP-моделью",
    "confidence": "Доверие",
    "age_source": "Источник возраста",
    "age_mean": "Средний возраст, оп.сут",
    "cov_source": "Источник stress-ковариат",
    "current_status": "ТР: состояние по фонду",
    "current_in_operation": "ТР: в работе",
    "current_status_date": "ТР: дата статуса",
    "event_90d_flag": "ГТМ/КРС в 90 дней",
    "first_gtm_date": "Первая дата ГТМ/КРС",
    "planned_oil_t": "План нефти, т",
    "planned_liquid_m3": "План жидкости, м3",
    "planned_ql_m3d": "План Ql, м3/сут",
    "planned_qnom_m3d": "Qnom, м3/сут",
    "planned_kpod_raw": "Kpod raw, Ql/Qnom",
    "planned_op_days": "План работы, сут",
    "calendar_days": "Календарных дней",
    "expected_failures": "Ожид. отказов УЭЦН",
    "expected_lost_op_days": "Ожид. потеря раб. дней",
    "oil_loss_t": "Потери нефти, т",
    "liquid_loss_m3": "Потери жидкости, м3",
    "p_fail_90d": "P(отказ 90д)",
    "exp_value_at_risk_90d_t": "Добыча под риском 90д, т",
    "theta": "Stress-множитель θ",
    "theta_ql": "Ql-множитель θ",
    "theta_kpod": "Kpod-множитель θ",
    "theta_uncertainty": "Uncertainty-множитель θ",
    "downtime_days": "Ремонт, дней",
    "flag_changeout": "Флаг замены",
    "covered_planned_oil_t": "Охваченный план нефти, т",
    "covered_planned_liquid_m3": "Охваченный план жидкости, м3",
    "total_planned_oil_t": "Общий план нефти, т",
    "total_planned_liquid_m3": "Общий план жидкости, м3",
    "uncovered_planned_oil_t": "Неохваченный план нефти, т",
    "uncovered_planned_liquid_m3": "Неохваченный план жидкости, м3",
    "risk_adjusted_total_oil_t": "Нефть с учетом риска, т",
    "risk_adjusted_total_liquid_m3": "Жидкость с учетом риска, м3",
    "coverage_pct_oil": "Охват модели, %",
    "risk_pct_total_oil": "Доля риска по нефти, %",
    "planned_gtm_jobs": "План. ГТМ/КРС, шт",
    "reactive_esp_failures": "Ожид. аварийные КРС, шт",
    "total_workover_demand": "Итого спрос КРС, шт",
    "contractor_group": "Группа подрядчика",
    "sour_class": "Кислый/некислый",
    "contractor_source": "Источник подрядчика",
    "sour_source": "Источник H2S-класса",
    "runtime_anomaly_count": "Аномалии наработки",
    "registry_oil_positive_months": "Реестр: нефть, мес",
    "registry_liquid_positive_months": "Реестр: жидкость, мес",
    "summary_positive_months": "Сводные данные: мес",
    "mean_planned_ql_m3d": "Средний Ql, м3/сут",
    "mean_planned_oil_rate_td": "Средний Qo, т/сут",
    "mean_planned_qnom_m3d": "Средний Qnom, м3/сут",
    "mean_planned_kpod_raw": "Средний Kpod raw",
    "wells_with_kpod": "Скважин с Kpod",
    f"p_fail_90d_{C.PRIMARY_SCENARIO_ID}": "P(отказ 90д), база/П50",
    f"exp_value_at_risk_90d_t_{C.PRIMARY_SCENARIO_ID}": "Добыча под риском 90д, база/П50, т",
    f"flag_changeout_{C.PRIMARY_SCENARIO_ID}": "Флаг замены, база/П50",
    f"theta_{C.PRIMARY_SCENARIO_ID}": "θ, база/П50",
    f"p_fail_90d_{C.STRESS_SCENARIO_ID}": "P(отказ 90д), stress/П75",
    f"exp_value_at_risk_90d_t_{C.STRESS_SCENARIO_ID}": "Добыча под риском 90д, stress/П75, т",
    f"flag_changeout_{C.STRESS_SCENARIO_ID}": "Флаг замены, stress/П75",
    f"theta_{C.STRESS_SCENARIO_ID}": "θ, stress/П75",
}


def _localize(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    out.columns = [_COLS_RU.get(col, col) for col in out.columns]
    return out


def _write_table(ws, df: pd.DataFrame, title: str, note: str = "") -> None:
    ws["A1"] = title
    ws["A1"].font = _TITLE_FONT
    if note:
        ws["A2"] = note
        ws["A2"].font = _NOTE_FONT
    df = _localize(df)
    start_row = 4
    for col_idx, col_name in enumerate(df.columns, start=1):
        cell = ws.cell(row=start_row, column=col_idx, value=col_name)
        cell.fill = _HDR_FILL
        cell.font = _HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row_idx, (_, row) in enumerate(df.iterrows(), start=start_row + 1):
        for col_idx, value in enumerate(row, start=1):
            if isinstance(value, float):
                value = round(value, 3)
            ws.cell(row=row_idx, column=col_idx, value=value)
    for col_idx, col_name in enumerate(df.columns, start=1):
        width = max(12, min(36, len(str(col_name)) + 4))
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.freeze_panes = ws["A5"]


def _summary_sheet(ws, monthly: pd.DataFrame, audit: pd.DataFrame) -> None:
    ws["A1"] = "Интеграция прогноза отказов УЭЦН в план добычи 2026-2027"
    ws["A1"].font = _TITLE_FONT
    ws["A3"] = "Ключевые итоги"
    ws["A3"].font = Font(bold=True)

    primary = monthly[monthly["scenario"] == C.PRIMARY_SCENARIO_ID].copy()
    if primary.empty:
        primary = monthly.copy()
    annual = (
        primary.assign(year=primary["month"].str[:4])
        .groupby("year", as_index=False)
        .agg(
            total_planned_oil_t=("total_planned_oil_t", "sum"),
            uncovered_planned_oil_t=("uncovered_planned_oil_t", "sum"),
            oil_loss_t=("oil_loss_t", "sum"),
            expected_failures=("expected_failures", "sum"),
        )
    )
    header_row = 5
    headers = ["Год", "План нефти, т", "Неохвачено ESP-моделью, т", "Ожид. потери нефти, т", "Ожид. отказы УЭЦН"]
    for idx, value in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=idx, value=value)
        cell.fill = _HDR_FILL
        cell.font = _HDR_FONT
    for row_idx, (_, row) in enumerate(annual.iterrows(), start=header_row + 1):
        ws.cell(row=row_idx, column=1, value=row["year"])
        ws.cell(row=row_idx, column=2, value=round(float(row["total_planned_oil_t"]), 1))
        ws.cell(row=row_idx, column=3, value=round(float(row["uncovered_planned_oil_t"]), 1))
        ws.cell(row=row_idx, column=4, value=round(float(row["oil_loss_t"]), 1))
        ws.cell(row=row_idx, column=5, value=round(float(row["expected_failures"]), 2))

    included = int((audit["scope_label"] != "excluded_non_esp").sum()) if not audit.empty else 0
    excluded = int((audit["scope_label"] == "excluded_non_esp").sum()) if not audit.empty else 0
    ws["A10"] = f"Скважин в расчете: {included}"
    ws["A11"] = f"Скважин вне консервативного ESP-охвата: {excluded}"
    ws["A13"] = "Замечание"
    ws["A13"].font = Font(bold=True)
    ws["A14"] = (
        "Основной сценарий использует только валидированную стратифицированную baseline-модель. "
        "Прогноз добычи использует renewal-симуляцию с downtime, потерями нефти и нагрузкой ремонтов. "
        "Stress-сценарий помечен как чувствительность и не является OOS-validated прогнозом."
    )
    ws["A14"].alignment = Alignment(wrap_text=True)
    ws.column_dimensions["A"].width = 32
    for col in ("B", "C", "D", "E"):
        ws.column_dimensions[col].width = 22


def _method_sheet(ws) -> None:
    ws["A1"] = "Методика"
    ws["A1"].font = _TITLE_FONT
    lines = [
        "1. База плана: только master-файл ПП; ДФ_04 используется как контекст ГТМ/КРС и ЭЦН-признак.",
        "2. Горизонт прогноза: июль 2026 — декабрь 2027.",
        "3. Валидационный график интенсивности: условный месячный риск для парка, активного на начало месяца; одинаковая логика до и после forecast_start.",
        "4. Прогноз/прогнозис: renewal-симуляция отказов с downtime, возвратом скважины после ремонта, потерями нефти и нагрузкой ремонтов.",
        "5. Калибровка: фиксированные прозрачные коэффициенты, заданные в коде; при запуске не пересчитываются по факту.",
        "6. Основной риск: валидированная модель stratum + operating age из esp_models.csv.",
        "7. Stress-сценарий: direct hazard overlay из esp_cox_coeffs.csv, Ql/Kpod и new-launch uncertainty hazard, только как sensitivity.",
        "8. Часы модели: operating days (ttf_mix), а не календарные дни.",
        "9. Потери добычи считаются через ожидаемые потерянные operating days и плановые месячные объемы.",
        "10. Для скважин без подтвержденного ЭЦН-следа основной расчет консервативно не применяется.",
    ]
    for idx, line in enumerate(lines, start=3):
        ws.cell(row=idx, column=1, value=line)
    ws.column_dimensions["A"].width = 110


def _kpod_sheet(ws, kpod_by_field: pd.DataFrame) -> None:
    ws["A1"] = "Средний Kpod по месторождениям"
    ws["A1"].font = _TITLE_FONT
    ws["A2"] = "Raw Kpod = Ql / Qnominal. История и прогноз; это не freq-normalized Kpod."
    ws["A2"].font = _NOTE_FONT
    if kpod_by_field is None or kpod_by_field.empty:
        ws["A4"] = "Нет данных Kpod."
        return

    def _write_pivot_section(value_col: str, title: str, y_title: str, header_row: int) -> int:
        field_col = "field_label" if "field_label" in kpod_by_field.columns else "plan_field"
        pivot = (
            kpod_by_field.pivot_table(
                index="month",
                columns=field_col,
                values=value_col,
                aggfunc="first",
            )
            .sort_index()
        )
        fields = [str(c) for c in pivot.columns]
        ws.cell(row=header_row - 1, column=1, value=title).font = Font(bold=True)
        ws.cell(row=header_row, column=1, value="Месяц").fill = _HDR_FILL
        ws.cell(row=header_row, column=1).font = _HDR_FONT
        for col_idx, field in enumerate(fields, start=2):
            cell = ws.cell(row=header_row, column=col_idx, value=field)
            cell.fill = _HDR_FILL
            cell.font = _HDR_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row_idx, month in enumerate(pivot.index, start=header_row + 1):
            ws.cell(row=row_idx, column=1, value=str(month))
            for col_idx, field in enumerate(fields, start=2):
                value = pivot.at[month, field]
                ws.cell(row=row_idx, column=col_idx, value=None if pd.isna(value) else round(float(value), 3))

        data = Reference(ws, min_col=2, max_col=1 + len(fields), min_row=header_row, max_row=header_row + len(pivot))
        cats = Reference(ws, min_col=1, min_row=header_row + 1, max_row=header_row + len(pivot))
        chart = LineChart()
        chart.title = title
        chart.y_axis.title = y_title
        chart.x_axis.title = "Месяц"
        chart.height = 12
        chart.width = 28
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        ws.add_chart(chart, f"A{header_row + len(pivot) + 3}")
        return header_row + len(pivot) + 20

    header_row = 4
    next_row = _write_pivot_section("mean_planned_kpod_raw", "Средний Kpod по месторождениям", "Kpod raw", header_row)
    if "mean_planned_ql_m3d" in kpod_by_field.columns:
        next_row = _write_pivot_section("mean_planned_ql_m3d", "Средний дебит жидкости по скважинам", "Ql, м3/сут", next_row)
    if "mean_planned_oil_rate_td" in kpod_by_field.columns:
        _write_pivot_section("mean_planned_oil_rate_td", "Средний дебит нефти по скважинам", "Qo, т/сут", next_row)

    field_col = "field_label" if "field_label" in kpod_by_field.columns else "plan_field"
    fields = [str(c) for c in kpod_by_field[field_col].dropna().unique()]
    ws.cell(row=header_row, column=1, value="Месяц").fill = _HDR_FILL
    ws.freeze_panes = ws["B5"]
    ws.column_dimensions["A"].width = 12
    for col_idx in range(2, 2 + len(fields)):
        ws.column_dimensions[get_column_letter(col_idx)].width = 18


def _save_workbook_unlocked(wb: Workbook, path: Path) -> Path:
    try:
        wb.save(path)
        return path
    except PermissionError:
        stamped = path.with_name(f"{path.stem}_{datetime.now().strftime('%Y%m%d_%H%M%S')}{path.suffix}")
        wb.save(stamped)
        return stamped


def write(
    out_path: Path,
    production_tables: dict[str, pd.DataFrame],
    workover: pd.DataFrame,
    changeout: pd.DataFrame,
    audit: pd.DataFrame,
) -> Path:
    wb = Workbook()
    _summary_sheet(wb.active, production_tables["monthly"], audit)
    wb.active.title = "Итог"

    _write_table(
        wb.create_sheet("Помесячно"),
        production_tables["monthly"],
        "Помесячная оценка потерь добычи",
        "Total plan включает весь план, covered plan — только консервативно охваченные ESP-скважины.",
    )
    _write_table(
        wb.create_sheet("По_месторождениям"),
        production_tables["by_field"],
        "Помесячная оценка по месторождениям",
    )
    _kpod_sheet(wb.create_sheet("Kpod"), production_tables.get("kpod_by_field", pd.DataFrame()))
    primary_well = production_tables["by_well"].copy()
    if "scenario" in primary_well.columns:
        primary_well = primary_well[primary_well["scenario"] == C.PRIMARY_SCENARIO_ID]
    _write_table(
        wb.create_sheet("Рейтинг_скважин"),
        primary_well.head(500),
        "Рейтинг скважин по ожидаемым потерям",
        "Показан основной сценарий база/П50. Для stress-сравнения см. лист changeout и CSV.",
    )
    _write_table(
        wb.create_sheet("ГТМ_КРС"),
        workover,
        "План ГТМ/КРС и ожидаемая аварийная нагрузка",
    )
    _write_table(
        wb.create_sheet("Аудит"),
        audit,
        "Аудит покрытия и мэппинга",
    )
    _write_table(
        wb.create_sheet("Changeout"),
        changeout.head(500),
        "Приоритет упреждающей замены",
    )
    _method_sheet(wb.create_sheet("Методика"))

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    return _save_workbook_unlocked(wb, out_path)
