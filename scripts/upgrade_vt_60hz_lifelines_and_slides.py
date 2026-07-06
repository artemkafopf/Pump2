from __future__ import annotations

import json
import math
import sys
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from analysis.paths import results_dir
from lifelines import AalenJohansenFitter, CoxPHFitter, KaplanMeierFitter
from lifelines.statistics import logrank_test
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from scipy.stats import fisher_exact, mannwhitneyu


REPO_ROOT = Path(__file__).resolve().parents[1]
_SLUG = "vt_60hz_upgrade_slides"
_PROMPT_DIR = REPO_ROOT / "analysis_outputs" / "vt_60hz_prompt_analysis_2026_06_22"
BASE_OUTPUT_DIR = results_dir(_SLUG)
FIGURES_DIR = BASE_OUTPUT_DIR / "figures"
TABLES_DIR = BASE_OUTPUT_DIR / "tables"
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")


COLORS = {
    "baseline": "#2ca02c",
    "true60": "#d62728",
    "high55": "#ff7f0e",
    "vt": "#d62728",
    "global": "#1f77b4",
}


def _numeric(series: pd.Series | np.ndarray | list[object]) -> pd.Series:
    return pd.to_numeric(pd.Series(series), errors="coerce").replace([np.inf, -np.inf], np.nan)


def ensure_dirs() -> None:
    FIGURES_DIR.mkdir(parents=True, exist_ok=True)
    TABLES_DIR.mkdir(parents=True, exist_ok=True)


def load_dataset() -> pd.DataFrame:
    path = _PROMPT_DIR / "tables" / "analysis_dataset.csv"
    df = pd.read_csv(path)
    for column in [
        "duration_best_days",
        "freq_w_mean",
        "h2s_effective_mg_l",
        "avg_glf",
        "mount_year",
        "trf_per_day",
        "calcium_load_per_day",
        "chloride_load_per_day",
        "sulfate_load_per_day",
        "gypsum_proxy_per_day",
    ]:
        if column in df.columns:
            df[column] = pd.to_numeric(df[column], errors="coerce")
    df["true60_flag"] = df["freq_w_mean"] > 58.0
    df["baseline_50_55_flag"] = (df["freq_w_mean"] > 50.0) & (df["freq_w_mean"] <= 55.0)
    df["high55_flag"] = df["freq_w_mean"] > 55.0
    df["high_h2s_flag"] = df["h2s_effective_mg_l"] >= 3.0
    df["infant_90d"] = ((df["event"] == 1) & (df["duration_best_days"] < 90.0)).astype(int)
    return df


def describe_mode(subset: pd.DataFrame, population: str, mode: str) -> dict[str, object]:
    failures = int(subset["event"].sum())
    runs = int(len(subset))
    return {
        "population": population,
        "mode": mode,
        "runs": runs,
        "failures": failures,
        "failure_rate": float(failures / max(runs, 1)),
        "median_duration_days": float(_numeric(subset["duration_best_days"]).median()) if runs else float("nan"),
        "mean_duration_days": float(_numeric(subset["duration_best_days"]).mean()) if runs else float("nan"),
        "infant_rate_runs_90d": float(subset["infant_90d"].mean()) if runs else float("nan"),
        "infant_share_failures_90d": float(subset["infant_90d"].sum() / max(failures, 1)),
        "mean_h2s_mg_l": float(_numeric(subset["h2s_effective_mg_l"]).mean()) if runs else float("nan"),
        "mean_glf": float(_numeric(subset["avg_glf"]).mean()) if runs else float("nan"),
    }


def build_true60_summary(df: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for population, pop_df in [("Global", df), ("Vt", df.loc[df["field"] == "Vt"].copy())]:
        rows.append(describe_mode(pop_df.loc[pop_df["baseline_50_55_flag"]].copy(), population, "50-55 Hz baseline"))
        rows.append(describe_mode(pop_df.loc[pop_df["high55_flag"]].copy(), population, ">55 Hz"))
        rows.append(describe_mode(pop_df.loc[pop_df["true60_flag"]].copy(), population, ">58 Hz true60 proxy"))
    frame = pd.DataFrame(rows)
    frame.to_csv(TABLES_DIR / "true60_focus_summary.csv", index=False, encoding="utf-8-sig")
    return frame


def build_field_true60_summary(df: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for field_name, field_df in df.groupby("field", dropna=False):
        base = field_df.loc[field_df["baseline_50_55_flag"]].copy()
        true60 = field_df.loc[field_df["true60_flag"]].copy()
        if len(base) == 0 and len(true60) == 0:
            continue
        pval = float("nan")
        if int(base["event"].sum()) >= 6 and int(true60["event"].sum()) >= 6:
            result = logrank_test(base["duration_best_days"], true60["duration_best_days"], base["event"], true60["event"])
            pval = float(result.p_value)
        rows.append(
            {
                "field": field_name,
                "baseline_runs": int(len(base)),
                "baseline_failures": int(base["event"].sum()),
                "baseline_median_days": float(_numeric(base["duration_best_days"]).median()) if len(base) else float("nan"),
                "true60_runs": int(len(true60)),
                "true60_failures": int(true60["event"].sum()),
                "true60_median_days": float(_numeric(true60["duration_best_days"]).median()) if len(true60) else float("nan"),
                "logrank_p_true60_vs_baseline": pval,
            }
        )
    frame = pd.DataFrame(rows).sort_values(["true60_failures", "baseline_failures"], ascending=[False, False]).reset_index(drop=True)
    frame.to_csv(TABLES_DIR / "true60_focus_by_field.csv", index=False, encoding="utf-8-sig")
    return frame


def plot_true60_km(base: pd.DataFrame, true60: pd.DataFrame, title: str, output_path: Path) -> dict[str, float]:
    fig, ax = plt.subplots(figsize=(8.5, 5.8))
    km_base = KaplanMeierFitter()
    km_true = KaplanMeierFitter()
    km_base.fit(base["duration_best_days"], base["event"], label=f"50-55 Hz baseline (n={len(base)}, fail={int(base['event'].sum())})")
    km_true.fit(true60["duration_best_days"], true60["event"], label=f">58 Hz true60 proxy (n={len(true60)}, fail={int(true60['event'].sum())})")
    km_base.plot_survival_function(ax=ax, ci_show=True, color=COLORS["baseline"], linewidth=2)
    km_true.plot_survival_function(ax=ax, ci_show=True, color=COLORS["true60"], linewidth=2)
    result = logrank_test(base["duration_best_days"], true60["duration_best_days"], base["event"], true60["event"])
    ax.set_title(title)
    ax.set_xlabel("Duration (days)")
    ax.set_ylabel("Survival probability")
    ax.grid(True, alpha=0.25)
    ax.text(
        0.98,
        0.05,
        f"log-rank p = {result.p_value:.3f}\nMedian: {km_base.median_survival_time_:.0f} vs {km_true.median_survival_time_:.0f} d",
        transform=ax.transAxes,
        ha="right",
        va="bottom",
        fontsize=9,
        bbox={"boxstyle": "round,pad=0.3", "fc": "white", "ec": "#cccccc", "alpha": 0.9},
    )
    fig.tight_layout()
    fig.savefig(output_path, dpi=170)
    plt.close(fig)
    return {
        "logrank_p": float(result.p_value),
        "baseline_median": float(km_base.median_survival_time_),
        "true60_median": float(km_true.median_survival_time_),
    }


def plot_true60_infant_rates(summary: pd.DataFrame, output_path: Path) -> None:
    fig, axes = plt.subplots(1, 2, figsize=(10.5, 4.8))
    populations = ["Global", "Vt"]
    for ax, metric, title in zip(
        axes,
        ["infant_rate_runs_90d", "infant_share_failures_90d"],
        ["Infant failures / runs (<90d)", "Infant failures / failures (<90d)"],
        strict=False,
    ):
        width = 0.35
        x = np.arange(len(populations))
        base_vals = [
            float(summary.loc[(summary["population"] == pop) & (summary["mode"] == "50-55 Hz baseline"), metric].iloc[0])
            for pop in populations
        ]
        true_vals = [
            float(summary.loc[(summary["population"] == pop) & (summary["mode"] == ">58 Hz true60 proxy"), metric].iloc[0])
            for pop in populations
        ]
        ax.bar(x - width / 2, base_vals, width, label="50-55 Hz baseline", color=COLORS["baseline"])
        ax.bar(x + width / 2, true_vals, width, label=">58 Hz true60 proxy", color=COLORS["true60"])
        ax.set_xticks(x, populations)
        ax.set_title(title)
        ax.set_ylim(0, max(true_vals + base_vals) * 1.25)
        ax.grid(True, axis="y", alpha=0.25)
        for idx, value in enumerate(base_vals):
            ax.text(idx - width / 2, value + 0.01, f"{value:.2f}", ha="center", va="bottom", fontsize=8)
        for idx, value in enumerate(true_vals):
            ax.text(idx + width / 2, value + 0.01, f"{value:.2f}", ha="center", va="bottom", fontsize=8)
    axes[0].legend(fontsize=8, loc="upper left")
    fig.suptitle("60 Hz focus: infant-failure burden")
    fig.tight_layout()
    fig.savefig(output_path, dpi=170)
    plt.close(fig)


def fit_lifelines_cox_models(df: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    rows = []
    term_rows = []

    def add_model(name: str, frame: pd.DataFrame, include_field: bool) -> None:
        work = frame.copy()
        work["duration"] = _numeric(work["duration_best_days"])
        work["true60"] = work["true60_flag"].astype(int)
        work["high_h2s"] = work["high_h2s_flag"].astype(int)
        work["avg_glf"] = _numeric(work["avg_glf"])
        work["mount_year"] = _numeric(work["mount_year"])
        if include_field:
            field_counts = work["field"].astype("string").value_counts()
            keep_fields = field_counts.loc[field_counts >= 20].index.astype(str).tolist()
            work = work.loc[work["field"].astype("string").isin(keep_fields)].copy()
        keep = ["duration", "event", "true60", "high_h2s", "avg_glf", "mount_year", "Скв."]
        if include_field:
            keep.append("field")
        work = work[keep].dropna().copy()
        if include_field:
            formula = "true60 + high_h2s + avg_glf + mount_year + C(field)"
        else:
            formula = "true60 + high_h2s + avg_glf + mount_year"
        cph = CoxPHFitter()
        cph.fit(work, duration_col="duration", event_col="event", cluster_col="Скв.", robust=True, formula=formula)
        rows.append(
            {
                "model": name,
                "rows": int(len(work)),
                "events": int(work["event"].sum()),
                "log_likelihood": float(cph.log_likelihood_),
                "concordance": float(cph.concordance_index_),
            }
        )
        summary = cph.summary.reset_index().rename(columns={"covariate": "term"})
        summary["model"] = name
        term_rows.append(summary)

    global_compare = df.loc[df["baseline_50_55_flag"] | df["true60_flag"]].copy()
    vt_compare = df.loc[(df["field"] == "Vt") & (df["baseline_50_55_flag"] | df["true60_flag"])].copy()
    high55_compare = df.loc[df["baseline_50_55_flag"] | df["high55_flag"]].copy()
    high55_compare["true60_flag"] = high55_compare["high55_flag"]

    add_model("global_true60_vs_baseline_adjusted", global_compare, include_field=True)
    add_model("vt_true60_vs_baseline_adjusted", vt_compare, include_field=False)
    add_model("global_gt55_vs_baseline_adjusted", high55_compare, include_field=True)

    summary_df = pd.DataFrame(rows)
    terms_df = pd.concat(term_rows, ignore_index=True)
    summary_df.to_csv(TABLES_DIR / "lifelines_cox_summary.csv", index=False, encoding="utf-8-sig")
    terms_df.to_csv(TABLES_DIR / "lifelines_cox_terms.csv", index=False, encoding="utf-8-sig")
    return summary_df, terms_df


def plot_adjusted_hrs(terms_df: pd.DataFrame, output_path: Path) -> pd.DataFrame:
    focus = terms_df.loc[terms_df["term"] == "true60", ["model", "exp(coef)", "exp(coef) lower 95%", "exp(coef) upper 95%", "p"]].copy()
    rename_map = {
        "global_true60_vs_baseline_adjusted": "Global true60 vs 50-55",
        "vt_true60_vs_baseline_adjusted": "Vt true60 vs 50-55",
        "global_gt55_vs_baseline_adjusted": "Global >55 vs 50-55",
    }
    focus["label"] = focus["model"].map(rename_map)
    focus = focus.sort_values("label").reset_index(drop=True)
    y = np.arange(len(focus))
    fig, ax = plt.subplots(figsize=(8, 4.8))
    hr = focus["exp(coef)"].to_numpy(dtype=float)
    lo = focus["exp(coef) lower 95%"].to_numpy(dtype=float)
    hi = focus["exp(coef) upper 95%"].to_numpy(dtype=float)
    ax.errorbar(hr, y, xerr=[hr - lo, hi - hr], fmt="o", color=COLORS["true60"], ecolor="#555555", capsize=4)
    ax.axvline(1.0, color="black", linestyle="--", linewidth=1)
    ax.set_yticks(y, focus["label"].tolist())
    ax.set_xlabel("Adjusted hazard ratio")
    ax.set_title("60 Hz focus: adjusted hazard ratios")
    ax.grid(True, axis="x", alpha=0.25)
    for idx, row in focus.iterrows():
        ax.text(float(row["exp(coef) upper 95%"]) * 1.02, idx, f"p={float(row['p']):.3f}", va="center", fontsize=8)
    fig.tight_layout()
    fig.savefig(output_path, dpi=170)
    plt.close(fig)
    focus.to_csv(TABLES_DIR / "true60_adjusted_hr_focus.csv", index=False, encoding="utf-8-sig")
    return focus


def build_ajf_tables(df: pd.DataFrame) -> pd.DataFrame:
    category_order = ["КЛ (R-0)", "ПЭД (R-0)", "Слом вала", "Засорение РО"]
    compare = df.loc[df["baseline_50_55_flag"] | df["true60_flag"]].copy()
    compare = compare.loc[compare["event"].isin([0, 1])].copy()
    category_map = {category: idx + 1 for idx, category in enumerate(category_order)}
    compare["event_code"] = compare["failure_category_raw"].map(category_map).fillna(0).astype(int)
    rows = []
    for population, pop_df in [("Global", compare), ("Vt", compare.loc[compare["field"] == "Vt"].copy())]:
        for mode, mask in [("50-55 Hz baseline", pop_df["baseline_50_55_flag"]), (">58 Hz true60 proxy", pop_df["true60_flag"])]:
            subset = pop_df.loc[mask].copy()
            if subset.empty:
                continue
            for category, event_code in category_map.items():
                if int((subset["event_code"] == event_code).sum()) < 3:
                    continue
                ajf = AalenJohansenFitter()
                ajf.fit(subset["duration_best_days"], subset["event_code"], event_of_interest=event_code)
                cif = ajf.cumulative_density_.reset_index()
                cif.columns = ["time", "cif"]
                for horizon in [90, 180, 365]:
                    eligible = cif.loc[cif["time"] <= horizon]
                    value = float(eligible["cif"].iloc[-1]) if not eligible.empty else 0.0
                    rows.append(
                        {
                            "population": population,
                            "mode": mode,
                            "category": category,
                            "horizon_days": horizon,
                            "cif": value,
                        }
                    )
    frame = pd.DataFrame(rows)
    frame.to_csv(TABLES_DIR / "lifelines_ajf_cif_summary.csv", index=False, encoding="utf-8-sig")
    return frame


def plot_vt_chemistry_focus(df: pd.DataFrame, output_path: Path) -> pd.DataFrame:
    vt = df.loc[(df["field"] == "Vt") & (df["event"] == 1)].copy()
    rows = []
    for metric in ["trf_per_day", "calcium_load_per_day", "chloride_load_per_day", "sulfate_load_per_day", "gypsum_proxy_per_day"]:
        x = _numeric(vt[metric])
        y = _numeric(vt["duration_best_days"])
        mask = x.notna() & y.notna()
        if int(mask.sum()) < 20:
            continue
        rho = float(pd.Series(x.loc[mask]).corr(pd.Series(y.loc[mask]), method="spearman"))
        rows.append({"metric": metric, "rho": rho, "rows": int(mask.sum())})
    frame = pd.DataFrame(rows).sort_values("rho").reset_index(drop=True)
    fig, ax = plt.subplots(figsize=(8, 4.8))
    ax.barh(frame["metric"], frame["rho"], color=[COLORS["true60"] if "trf" in m else COLORS["global"] for m in frame["metric"]])
    ax.axvline(0.0, color="black", linestyle="--", linewidth=1)
    ax.set_title("Vt: raw timing signal of duty vs chemistry proxies")
    ax.set_xlabel("Spearman rho with TTF")
    ax.grid(True, axis="x", alpha=0.25)
    for idx, row in frame.iterrows():
        ax.text(float(row["rho"]) + (0.01 if row["rho"] >= 0 else -0.01), idx, f"{row['rho']:.2f}", va="center", ha="left" if row["rho"] >= 0 else "right", fontsize=8)
    fig.tight_layout()
    fig.savefig(output_path, dpi=170)
    plt.close(fig)
    frame.to_csv(TABLES_DIR / "vt_chemistry_focus.csv", index=False, encoding="utf-8-sig")
    return frame


def compare_vt_true60_environment(df: pd.DataFrame) -> pd.DataFrame:
    vt = df.loc[df["field"] == "Vt"].copy()
    vt["group"] = np.where(vt["true60_flag"], ">58 Hz", "Rest of Vt")
    vt["acid_flag"] = vt["Кислый/Некислый"].astype("string").str.contains("Кислый", case=False, na=False).astype(int)

    rows: list[dict[str, object]] = []
    continuous_metrics = [
        ("h2s_effective_mg_l", "H2S effective, mg/L"),
        ("avg_glf", "GLF"),
        ("Kpod", "Kpod"),
        ("trf_per_day", "TRF per day"),
        ("tlf_per_day", "TLF per day"),
        ("calcium_load_per_day", "Calcium load per day"),
        ("chloride_load_per_day", "Chloride load per day"),
        ("sulfate_load_per_day", "Sulfate load per day"),
        ("gypsum_proxy_per_day", "Gypsum proxy per day"),
        ("duration_best_days", "Best duration, days"),
    ]
    binary_metrics = [
        ("high_h2s_flag", "High H2S flag (>=3 mg/L)"),
        ("Кислый/Некислый", "Acidic well label"),
        ("corrosion_fill", "Corrosion protection indicated"),
    ]

    true60 = vt.loc[vt["group"] == ">58 Hz"].copy()
    rest = vt.loc[vt["group"] == "Rest of Vt"].copy()

    for column, label in continuous_metrics:
        x = _numeric(true60[column])
        y = _numeric(rest[column])
        x = x.dropna()
        y = y.dropna()
        p_value = float("nan")
        if len(x) >= 3 and len(y) >= 3:
            p_value = float(mannwhitneyu(x, y, alternative="two-sided").pvalue)
        rows.append(
            {
                "metric": label,
                "metric_type": "continuous",
                "true60_n": int(len(x)),
                "rest_n": int(len(y)),
                "true60_mean": float(x.mean()) if len(x) else float("nan"),
                "rest_mean": float(y.mean()) if len(y) else float("nan"),
                "true60_median": float(x.median()) if len(x) else float("nan"),
                "rest_median": float(y.median()) if len(y) else float("nan"),
                "difference_true60_minus_rest_median": float(x.median() - y.median()) if len(x) and len(y) else float("nan"),
                "p_value": p_value,
            }
        )

    for column, label in binary_metrics:
        use_column = "acid_flag" if column == "Кислый/Некислый" else column
        x = _numeric(true60[use_column]).fillna(0).astype(int)
        y = _numeric(rest[use_column]).fillna(0).astype(int)
        true60_pos = int(x.sum())
        true60_neg = int(len(x) - true60_pos)
        rest_pos = int(y.sum())
        rest_neg = int(len(y) - rest_pos)
        p_value = float("nan")
        odds_ratio = float("nan")
        if len(x) > 0 and len(y) > 0:
            odds_ratio, p_value = fisher_exact([[true60_pos, true60_neg], [rest_pos, rest_neg]], alternative="two-sided")
        rows.append(
            {
                "metric": label,
                "metric_type": "binary",
                "true60_n": int(len(x)),
                "rest_n": int(len(y)),
                "true60_mean": float(x.mean()) if len(x) else float("nan"),
                "rest_mean": float(y.mean()) if len(y) else float("nan"),
                "true60_median": float(x.mean()) if len(x) else float("nan"),
                "rest_median": float(y.mean()) if len(y) else float("nan"),
                "difference_true60_minus_rest_median": float(x.mean() - y.mean()) if len(x) and len(y) else float("nan"),
                "odds_ratio": float(odds_ratio),
                "p_value": float(p_value),
            }
        )

    frame = pd.DataFrame(rows)
    frame.to_csv(TABLES_DIR / "vt_true60_environment_vs_rest.csv", index=False, encoding="utf-8-sig")
    return frame


def plot_vt_true60_environment(env_df: pd.DataFrame, output_path: Path) -> pd.DataFrame:
    focus_labels = [
        "H2S effective, mg/L",
        "GLF",
        "Kpod",
        "Calcium load per day",
        "Chloride load per day",
        "Sulfate load per day",
        "Gypsum proxy per day",
    ]
    plot_df = env_df.loc[env_df["metric"].isin(focus_labels)].copy()
    plot_df["direction"] = np.where(plot_df["difference_true60_minus_rest_median"] > 0, "Higher at >58 Hz", "Lower at >58 Hz")

    fig, ax = plt.subplots(figsize=(9.4, 5.2))
    y = np.arange(len(plot_df))
    colors = [COLORS["true60"] if value > 0 else COLORS["baseline"] for value in plot_df["difference_true60_minus_rest_median"]]
    ax.barh(y, plot_df["difference_true60_minus_rest_median"], color=colors)
    ax.axvline(0.0, color="black", linestyle="--", linewidth=1)
    ax.set_yticks(y, plot_df["metric"].tolist())
    ax.set_xlabel("Median difference: >58 Hz minus rest of Vt")
    ax.set_title("Vt harsh-environment check for >58 Hz wells")
    ax.grid(True, axis="x", alpha=0.25)
    for idx, row in plot_df.reset_index(drop=True).iterrows():
        value = float(row["difference_true60_minus_rest_median"])
        p_text = f"p={float(row['p_value']):.3f}" if pd.notna(row["p_value"]) else "p=n/a"
        ax.text(value + (0.02 * abs(value) if value != 0 else 0.02), idx, p_text, va="center", ha="left" if value >= 0 else "right", fontsize=8)
    fig.tight_layout()
    fig.savefig(output_path, dpi=170)
    plt.close(fig)
    plot_df.to_csv(TABLES_DIR / "vt_true60_environment_plot_data.csv", index=False, encoding="utf-8-sig")
    return plot_df


def build_vt_true60_failure_mix(df: pd.DataFrame) -> pd.DataFrame:
    vt_failures = df.loc[(df["field"] == "Vt") & (df["event"] == 1)].copy()
    vt_failures["group"] = np.where(vt_failures["true60_flag"], ">58 Hz", "Rest of Vt")
    counts = (
        vt_failures.groupby(["group", "failure_category_raw"], dropna=False)
        .size()
        .rename("count")
        .reset_index()
        .rename(columns={"failure_category_raw": "category"})
    )
    totals = counts.groupby("group")["count"].transform("sum")
    counts["share"] = counts["count"] / totals
    counts = counts.sort_values(["group", "share", "count"], ascending=[True, False, False]).reset_index(drop=True)
    counts.to_csv(TABLES_DIR / "vt_true60_failure_mix.csv", index=False, encoding="utf-8-sig")
    return counts


def plot_vt_true60_failure_mix(mix_df: pd.DataFrame, output_path: Path) -> None:
    top_categories = (
        mix_df.groupby("category")["count"]
        .sum()
        .sort_values(ascending=False)
        .head(6)
        .index.tolist()
    )
    plot_df = mix_df.loc[mix_df["category"].isin(top_categories)].copy()
    pivot = plot_df.pivot(index="category", columns="group", values="share").fillna(0.0)
    pivot = pivot.reindex(top_categories).sort_values(">58 Hz", ascending=True)
    fig, ax = plt.subplots(figsize=(8.8, 5.4))
    y = np.arange(len(pivot))
    ax.barh(y - 0.18, pivot.get("Rest of Vt", pd.Series(0.0, index=pivot.index)), height=0.35, color=COLORS["baseline"], label="Rest of Vt")
    ax.barh(y + 0.18, pivot.get(">58 Hz", pd.Series(0.0, index=pivot.index)), height=0.35, color=COLORS["true60"], label=">58 Hz")
    ax.set_yticks(y, pivot.index.tolist())
    ax.set_xlabel("Failure share within Vt failures")
    ax.set_title("Vt failure mix: >58 Hz vs rest of Vt")
    ax.grid(True, axis="x", alpha=0.25)
    ax.legend(loc="lower right", fontsize=8)
    fig.tight_layout()
    fig.savefig(output_path, dpi=170)
    plt.close(fig)


def build_true60_focus_note(summary_df: pd.DataFrame, field_df: pd.DataFrame, hr_df: pd.DataFrame, env_df: pd.DataFrame) -> Path:
    global_base = summary_df.loc[(summary_df["population"] == "Global") & (summary_df["mode"] == "50-55 Hz baseline")].iloc[0]
    global_true = summary_df.loc[(summary_df["population"] == "Global") & (summary_df["mode"] == ">58 Hz true60 proxy")].iloc[0]
    vt_base = summary_df.loc[(summary_df["population"] == "Vt") & (summary_df["mode"] == "50-55 Hz baseline")].iloc[0]
    vt_true = summary_df.loc[(summary_df["population"] == "Vt") & (summary_df["mode"] == ">58 Hz true60 proxy")].iloc[0]
    h2s_row = env_df.loc[env_df["metric"] == "H2S effective, mg/L"].iloc[0]
    glf_row = env_df.loc[env_df["metric"] == "GLF"].iloc[0]
    acid_row = env_df.loc[env_df["metric"] == "Acidic well label"].iloc[0]
    za_row = field_df.loc[field_df["field"] == "Za"]
    za_text = ""
    if not za_row.empty:
        row = za_row.iloc[0]
        za_text = f"- `Za` is the clearest raw field-level warning sign: baseline median `{row['baseline_median_days']:.1f}` d vs true60 `{row['true60_median_days']:.1f}` d, log-rank p `{row['logrank_p_true60_vs_baseline']:.3f}`."
    note_lines = [
        "# 60 Hz Focus Note",
        "",
        "## Raw Comparison",
        "",
        f"- Global `50-55 Hz` baseline: `{int(global_base['runs'])}` runs / `{int(global_base['failures'])}` failures / median `{float(global_base['median_duration_days']):.1f}` d.",
        f"- Global `>58 Hz true60 proxy`: `{int(global_true['runs'])}` runs / `{int(global_true['failures'])}` failures / median `{float(global_true['median_duration_days']):.1f}` d.",
        f"- Vt `50-55 Hz` baseline: `{int(vt_base['runs'])}` runs / `{int(vt_base['failures'])}` failures / median `{float(vt_base['median_duration_days']):.1f}` d.",
        f"- Vt `>58 Hz true60 proxy`: `{int(vt_true['runs'])}` runs / `{int(vt_true['failures'])}` failures / median `{float(vt_true['median_duration_days']):.1f}` d.",
        za_text,
        "",
        "## Practical Interpretation",
        "",
        "- Raw `>58 Hz` wells look somewhat worse than the `50-55 Hz` baseline on median life and infant burden, especially in `Vt`.",
        "- After adjustment, the global `true60` hazard ratio is modest and not statistically strong; the `Vt`-only adjusted HR is higher but still not conclusive because the subgroup is small.",
        "- `H2S` is more consistently associated with shorter life than `true60` itself in the adjusted models.",
        f"- The new harsh-environment check argues against a simple confounding story inside `Vt`: median effective `H2S` is `{float(h2s_row['true60_median']):.3f}` at `>58 Hz` versus `{float(h2s_row['rest_median']):.3f}` in the rest of `Vt` (Mann-Whitney p `{float(h2s_row['p_value']):.3f}`), and median `GLF` is `{float(glf_row['true60_median']):.1f}` versus `{float(glf_row['rest_median']):.1f}` (p `{float(glf_row['p_value']):.3f}`).",
        f"- The acidic-label share also does not prove that `>58 Hz` wells were always harsher: `{float(acid_row['true60_mean']):.2f}` at `>58 Hz` versus `{float(acid_row['rest_mean']):.2f}` in the rest of `Vt`.",
        "- The best current reading is: `60 Hz` is a plausible stress amplifier in some contexts, but not a proven standalone root cause across the portfolio.",
        "",
        "## Recommendation",
        "",
        "- Do not use a blanket `60 Hz is always bad` rule.",
        "- Use a guarded rule: prioritize caution for `60 Hz` in fields/wells that are already chemically severe or already show fast-failing category mix, but do not assume every `>58 Hz` Vt well was simply the harshest environment case.",
        "- For `Vt`, treat `60 Hz` as a monitoring and review regime, not yet as a universally prohibited mode based on this evidence alone.",
        "",
    ]
    path = BASE_OUTPUT_DIR / "vt_60hz_focus_note.md"
    path.write_text("\n".join(line for line in note_lines if line != ""), encoding="utf-8")
    return path


def add_title_slide(prs: Presentation, title: str, subtitle: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    background.fill.solid()
    background.fill.fore_color.rgb = RGBColor(245, 247, 250)
    background.line.fill.background()
    tx = slide.shapes.add_textbox(Inches(0.7), Inches(0.8), Inches(11.2), Inches(1.4))
    p = tx.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RGBColor(25, 35, 52)
    sub = slide.shapes.add_textbox(Inches(0.75), Inches(2.0), Inches(10.8), Inches(1.2))
    p2 = sub.text_frame.paragraphs[0]
    p2.text = subtitle
    p2.font.size = Pt(16)
    p2.font.color.rgb = RGBColor(80, 91, 107)


def add_bullets_slide(prs: Presentation, title: str, bullets: list[str], image_path: Path | None = None) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.0), Inches(0.6))
    p = title_box.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = RGBColor(25, 35, 52)

    left_w = Inches(5.2) if image_path else Inches(11.5)
    body = slide.shapes.add_textbox(Inches(0.6), Inches(1.0), left_w, Inches(5.9))
    tf = body.text_frame
    tf.word_wrap = True
    for idx, bullet in enumerate(bullets):
        para = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        para.text = bullet
        para.level = 0
        para.font.size = Pt(17)
        para.font.color.rgb = RGBColor(47, 56, 70)
        para.space_after = Pt(12)
    if image_path is not None and image_path.exists():
        slide.shapes.add_picture(str(image_path), Inches(6.0), Inches(1.0), width=Inches(6.3))


def generate_slides(summary_df: pd.DataFrame, field_df: pd.DataFrame, hr_df: pd.DataFrame, env_df: pd.DataFrame) -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    global_base = summary_df.loc[(summary_df["population"] == "Global") & (summary_df["mode"] == "50-55 Hz baseline")].iloc[0]
    global_true = summary_df.loc[(summary_df["population"] == "Global") & (summary_df["mode"] == ">58 Hz true60 proxy")].iloc[0]
    vt_base = summary_df.loc[(summary_df["population"] == "Vt") & (summary_df["mode"] == "50-55 Hz baseline")].iloc[0]
    vt_true = summary_df.loc[(summary_df["population"] == "Vt") & (summary_df["mode"] == ">58 Hz true60 proxy")].iloc[0]
    global_hr = hr_df.loc[hr_df["label"] == "Global true60 vs 50-55"].iloc[0]
    vt_hr = hr_df.loc[hr_df["label"] == "Vt true60 vs 50-55"].iloc[0]
    h2s_row = env_df.loc[env_df["metric"] == "H2S effective, mg/L"].iloc[0]
    glf_row = env_df.loc[env_df["metric"] == "GLF"].iloc[0]
    acid_row = env_df.loc[env_df["metric"] == "Acidic well label"].iloc[0]

    add_title_slide(
        prs,
        "60 Hz Operating Mode Review",
        "Question: does the true 60 Hz regime shorten ESP life, and is Vt especially sensitive?",
    )

    add_bullets_slide(
        prs,
        "Data Sufficiency",
        [
            f"Global `>58 Hz` true60 proxy population: {int(global_true['runs'])} runs, {int(global_true['failures'])} failures.",
            f"Vt `>58 Hz` true60 proxy population: {int(vt_true['runs'])} runs, {int(vt_true['failures'])} failures.",
            "That is enough for descriptive survival comparison, but still limited for firm Vt-only adjustment.",
            "The strongest fields for raw true60 comparison are Ya, Vt, Ic, Az, and Za.",
        ],
        image_path=FIGURES_DIR / "true60_infant_rates.png",
    )

    add_bullets_slide(
        prs,
        "Global 60 Hz Read",
        [
            f"Global median life: `50-55 Hz` = {float(global_base['median_duration_days']):.0f} d, `>58 Hz` = {float(global_true['median_duration_days']):.0f} d.",
            f"Global infant burden is worse in raw terms at true60: infant/runs `{float(global_true['infant_rate_runs_90d']):.2f}` vs `{float(global_base['infant_rate_runs_90d']):.2f}`.",
            f"Adjusted Cox HR for true60 vs 50-55 = `{float(global_hr['exp(coef)']):.2f}` with p `{float(global_hr['p']):.3f}`.",
            "Interpretation: raw true60 looks somewhat worse, but the adjusted global penalty is not strong enough to call 60 Hz a standalone portfolio-wide root cause.",
        ],
        image_path=FIGURES_DIR / "true60_km_global.png",
    )

    add_bullets_slide(
        prs,
        "Vt 60 Hz Read",
        [
            f"Vt median life: `50-55 Hz` = {float(vt_base['median_duration_days']):.0f} d, `>58 Hz` = {float(vt_true['median_duration_days']):.0f} d.",
            f"Vt infant burden rises in raw terms at true60: infant/runs `{float(vt_true['infant_rate_runs_90d']):.2f}` vs `{float(vt_base['infant_rate_runs_90d']):.2f}`.",
            f"Adjusted Cox HR inside Vt = `{float(vt_hr['exp(coef)']):.2f}` with p `{float(vt_hr['p']):.3f}`.",
            "Interpretation: Vt looks directionally more sensitive to true60, but the subgroup is still small, so this is a warning signal rather than proof.",
        ],
        image_path=FIGURES_DIR / "true60_km_vt.png",
    )

    add_bullets_slide(
        prs,
        "Were >58 Hz Vt Wells Simply Harsher?",
        [
            f"No clear evidence of that on measured chemistry: median effective H2S is `{float(h2s_row['true60_median']):.3f}` at `>58 Hz` versus `{float(h2s_row['rest_median']):.3f}` in the rest of `Vt`.",
            f"GLF also does not look systematically harsher: median `{float(glf_row['true60_median']):.1f}` at `>58 Hz` versus `{float(glf_row['rest_median']):.1f}` in the rest of `Vt`.",
            f"Acidic-label share is `{float(acid_row['true60_mean']):.2f}` at `>58 Hz` versus `{float(acid_row['rest_mean']):.2f}` in the rest of `Vt`.",
            "So the worse raw Vt true60 outcome cannot be dismissed as 'those were always the harshest wells' based on the measured harshness proxies we have.",
        ],
        image_path=FIGURES_DIR / "vt_true60_environment_check.png",
    )

    za_row = field_df.loc[field_df["field"] == "Za"]
    za_bullet = "No single field dominates the story."
    if not za_row.empty:
        row = za_row.iloc[0]
        za_bullet = f"Za is the clearest raw field-specific caution: baseline median {float(row['baseline_median_days']):.0f} d vs true60 {float(row['true60_median_days']):.0f} d, p={float(row['logrank_p_true60_vs_baseline']):.3f}."

    add_bullets_slide(
        prs,
        "What Explains More Than 60 Hz",
        [
            "H2S is a stronger and more consistent driver than true60 in the adjusted models.",
            "Vt is chemically harsher than global, with much higher mean H2S and acidic-share.",
            "Separate ion-load proxies carry more timing signal than normalized TRF, and they should stay separate instead of being collapsed into one precipitate-only proxy.",
            za_bullet,
        ],
        image_path=FIGURES_DIR / "vt_chemistry_focus.png",
    )

    add_bullets_slide(
        prs,
        "Adjusted Hazard Ratios",
        [
            "Global true60 vs 50-55: modest HR above 1, but not statistically strong after field/H2S/GLF adjustment.",
            "Vt true60 vs 50-55: higher HR than global, but still not conclusive because of sample size.",
            "Sensitivity check using all `>55 Hz` instead of `>58 Hz` also does not create a clean universal penalty.",
        ],
        image_path=FIGURES_DIR / "true60_adjusted_hrs.png",
    )

    add_bullets_slide(
        prs,
        "Vt Failure Pattern at >58 Hz",
        [
            "Within Vt, the `>58 Hz` failures lean more toward `Засорение РО` and `ПЭД (R-0)` than the rest of the Vt failures.",
            "That pattern is directionally consistent with 60 Hz acting as a stress amplifier on top of specific failure mechanisms, not just shifting all categories equally.",
            "The sample is still small, so this is useful context, not proof.",
        ],
        image_path=FIGURES_DIR / "vt_true60_failure_mix.png",
    )

    add_bullets_slide(
        prs,
        "Operational Conclusion",
        [
            "Do not use a blanket rule that 60 Hz is universally harmful across the portfolio.",
            "Use a guarded rule: 60 Hz deserves review in chemically severe wells and in fields already biased toward fast-failing categories, but measured harshness alone does not explain the Vt `>58 Hz` signal.",
            "For Vt, treat true60 as a heightened-monitoring regime, especially when failure mix and chemistry both look unfavorable.",
            "Best next step: collect more Vt true60 history and pair it with package-backed competing-risk analysis before making hard operating limits.",
        ],
        image_path=FIGURES_DIR / "vt_vs_global_failure_mix.png",
    )

    path = BASE_OUTPUT_DIR / "vt_60hz_focus_slides.pptx"
    try:
        prs.save(path)
        return path
    except PermissionError:
        alt_path = BASE_OUTPUT_DIR / "vt_60hz_focus_slides_v2.pptx"
        prs.save(alt_path)
        return alt_path


def main() -> None:
    ensure_dirs()
    df = load_dataset()

    summary_df = build_true60_summary(df)
    field_df = build_field_true60_summary(df)

    global_base = df.loc[df["baseline_50_55_flag"]].copy()
    global_true = df.loc[df["true60_flag"]].copy()
    vt_base = df.loc[(df["field"] == "Vt") & (df["baseline_50_55_flag"])].copy()
    vt_true = df.loc[(df["field"] == "Vt") & (df["true60_flag"])].copy()

    plot_true60_km(global_base, global_true, "Global: true60 (>58 Hz) vs 50-55 Hz baseline", FIGURES_DIR / "true60_km_global.png")
    if len(vt_base) > 0 and len(vt_true) > 0:
        plot_true60_km(vt_base, vt_true, "Vt: true60 (>58 Hz) vs 50-55 Hz baseline", FIGURES_DIR / "true60_km_vt.png")

    plot_true60_infant_rates(summary_df, FIGURES_DIR / "true60_infant_rates.png")
    _, terms_df = fit_lifelines_cox_models(df)
    hr_df = plot_adjusted_hrs(terms_df, FIGURES_DIR / "true60_adjusted_hrs.png")
    build_ajf_tables(df)
    plot_vt_chemistry_focus(df, FIGURES_DIR / "vt_chemistry_focus.png")
    env_df = compare_vt_true60_environment(df)
    plot_vt_true60_environment(env_df, FIGURES_DIR / "vt_true60_environment_check.png")
    vt_true60_mix = build_vt_true60_failure_mix(df)
    plot_vt_true60_failure_mix(vt_true60_mix, FIGURES_DIR / "vt_true60_failure_mix.png")
    note_path = build_true60_focus_note(summary_df, field_df, hr_df, env_df)
    slides_path = generate_slides(summary_df, field_df, hr_df, env_df)

    summary = {
        "true60_global_runs": int(summary_df.loc[(summary_df["population"] == "Global") & (summary_df["mode"] == ">58 Hz true60 proxy"), "runs"].iloc[0]),
        "true60_global_failures": int(summary_df.loc[(summary_df["population"] == "Global") & (summary_df["mode"] == ">58 Hz true60 proxy"), "failures"].iloc[0]),
        "true60_vt_runs": int(summary_df.loc[(summary_df["population"] == "Vt") & (summary_df["mode"] == ">58 Hz true60 proxy"), "runs"].iloc[0]),
        "true60_vt_failures": int(summary_df.loc[(summary_df["population"] == "Vt") & (summary_df["mode"] == ">58 Hz true60 proxy"), "failures"].iloc[0]),
        "vt_true60_environment_table": str(TABLES_DIR / "vt_true60_environment_vs_rest.csv"),
        "note_path": str(note_path),
        "slides_path": str(slides_path),
    }
    (BASE_OUTPUT_DIR / "vt_60hz_focus_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"[true60] Wrote focus outputs to {BASE_OUTPUT_DIR}", flush=True)


if __name__ == "__main__":
    main()
