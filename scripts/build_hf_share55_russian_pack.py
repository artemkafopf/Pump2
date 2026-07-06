from __future__ import annotations

import json
import sys
from pathlib import Path

import pandas as pd
from analysis.paths import results_dir
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.util import Inches, Pt


REPO_ROOT = Path(__file__).resolve().parents[1]
DOCS_DIR = REPO_ROOT / "docs"
_PROMPT_DIR = REPO_ROOT / "analysis_outputs" / "vt_60hz_prompt_analysis_2026_06_22"
FIGURES_DIR = _PROMPT_DIR / "figures"
TABLES_DIR = _PROMPT_DIR / "tables"
_SLUG = "vt_hf_share55_ru_pack"
BASE_OUTPUT_DIR = results_dir(_SLUG)
FIGURES_DIR = BASE_OUTPUT_DIR / "figures"

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")


def read_csv(name: str) -> pd.DataFrame:
    return pd.read_csv(TABLES_DIR / name)


def load_inputs() -> dict[str, pd.DataFrame]:
    return {
        "population": read_csv("hf_days_share55_population_comparison.csv"),
        "axis": read_csv("hf_days_share55_axis_summary.csv"),
        "hr": read_csv("hf_days_share55_hr_summary.csv"),
        "threshold_tests": read_csv("hf_days_share55_threshold_tests.csv"),
        "mix": read_csv("hf_days_share55_vt_failure_mix.csv"),
        "env": read_csv("hf_days_share55_vt_environment.csv"),
    }


def pick(df: pd.DataFrame, **filters: object) -> pd.Series:
    mask = pd.Series(True, index=df.index)
    for key, value in filters.items():
        mask &= df[key] == value
    return df.loc[mask].iloc[0]


def build_global_mix() -> pd.DataFrame:
    path = TABLES_DIR / "analysis_dataset.csv"
    df = pd.read_csv(path)
    for col in ["duration_best_days", "n_freq_above_55hz", "event"]:
        df[col] = pd.to_numeric(df[col], errors="coerce")
    df["hf_flag"] = (df["n_freq_above_55hz"] / df["duration_best_days"]) > 0.5
    fails = df.loc[df["event"] == 1].copy()
    fails["group"] = fails["hf_flag"].map({True: "HF", False: "Rest"})
    counts = (
        fails.groupby(["group", "failure_category_raw"], dropna=False)
        .size()
        .rename("count")
        .reset_index()
        .rename(columns={"failure_category_raw": "category"})
    )
    counts["share"] = counts["count"] / counts.groupby("group")["count"].transform("sum")
    return counts


def build_markdown(inputs: dict[str, pd.DataFrame], global_mix: pd.DataFrame) -> Path:
    population = inputs["population"]
    axis = inputs["axis"]
    hr = inputs["hr"]
    threshold_tests = inputs["threshold_tests"]
    mix_vt = inputs["mix"]
    env = inputs["env"]

    pop_global_old = pick(population, population="Global", definition="mean >58 Hz proxy")
    pop_global_new = pick(population, population="Global", definition="days >55 / duration_best_days >0.5")
    pop_vt_old = pick(population, population="Vt", definition="mean >58 Hz proxy")
    pop_vt_new = pick(population, population="Vt", definition="days >55 / duration_best_days >0.5")

    glob_best_base = pick(axis, axis_name="best_available_days", population="Global", mode="50-55 Hz baseline")
    glob_best_hf = pick(axis, axis_name="best_available_days", population="Global", mode="HF by >55d share >0.5")
    glob_true_base = pick(axis, axis_name="ttf_true_days", population="Global", mode="50-55 Hz baseline")
    glob_true_hf = pick(axis, axis_name="ttf_true_days", population="Global", mode="HF by >55d share >0.5")
    vt_best_base = pick(axis, axis_name="best_available_days", population="Vt", mode="50-55 Hz baseline")
    vt_best_hf = pick(axis, axis_name="best_available_days", population="Vt", mode="HF by >55d share >0.5")
    vt_true_base = pick(axis, axis_name="ttf_true_days", population="Vt", mode="50-55 Hz baseline")
    vt_true_hf = pick(axis, axis_name="ttf_true_days", population="Vt", mode="HF by >55d share >0.5")

    glob_hr_true = pick(hr, axis_name="ttf_true_days", population="Global")
    vt_hr_true = pick(hr, axis_name="ttf_true_days", population="Vt")
    glob_hr_trf = pick(hr, axis_name="trf_50hz_equiv_days", population="Global")
    vt_hr_trf = pick(hr, axis_name="trf_50hz_equiv_days", population="Vt")

    glob_thr90 = pick(threshold_tests, threshold_days=90, population="Global")
    vt_thr90 = pick(threshold_tests, threshold_days=90, population="Vt")

    h2s_vt = env.loc[env["metric"] == "H2S effective, mg/L"].iloc[0]
    glf_vt = env.loc[env["metric"] == "GLF"].iloc[0]

    mix_vt_hf = mix_vt.loc[mix_vt["group"] == "HF share>0.5"].copy()
    mix_vt_rest = mix_vt.loc[mix_vt["group"] == "Rest of Vt"].copy()
    mix_vt_join = mix_vt_hf.merge(mix_vt_rest[["category", "share"]], on="category", how="left", suffixes=("_hf", "_rest"))
    mix_vt_join["share_diff"] = mix_vt_join["share_hf"] - mix_vt_join["share_rest"].fillna(0.0)
    mix_vt_join = mix_vt_join.sort_values("share_diff", ascending=False).reset_index(drop=True)

    mix_glob_hf = global_mix.loc[global_mix["group"] == "HF"].copy()
    mix_glob_rest = global_mix.loc[global_mix["group"] == "Rest"].copy()
    mix_glob_join = mix_glob_hf.merge(mix_glob_rest[["category", "share"]], on="category", how="left", suffixes=("_hf", "_rest"))
    mix_glob_join["share_diff"] = mix_glob_join["share_hf"] - mix_glob_join["share_rest"].fillna(0.0)
    mix_glob_join = mix_glob_join.sort_values("share_diff", ascending=False).reset_index(drop=True)

    vt_top1 = mix_vt_join.iloc[0]
    vt_top2 = mix_vt_join.iloc[1]
    glob_top1 = mix_glob_join.iloc[0]
    glob_top2 = mix_glob_join.iloc[1]

    lines = [
        "# Высокочастотные скважины: новое определение `дней >55 Гц / days_total > 0.5`",
        "",
        "## Короткий ответ",
        "",
        "**При новом определении для `Vt` не получается подтверждения, что высокочастотные скважины работают хуже, чем базовая группа `50–55 Гц`.**",
        "",
        "Это важное отличие от прежнего прокси `mean >58 Hz`. Новое определение заметно расширяет группу и фактически отвечает уже не на вопрос про режим, близкий к `60 Гц`, а на вопрос про **устойчивую работу выше `55 Гц` в течение большей части срока**.",
        "",
        "## Как изменился состав выборки",
        "",
        f"- `Global`, старый прокси `mean >58 Hz`: `{int(pop_global_old['runs'])}` прогонов / `{int(pop_global_old['failures'])}` отказов.",
        f"- `Global`, новое определение: `{int(pop_global_new['runs'])}` / `{int(pop_global_new['failures'])}`.",
        f"- `Vt`, старый прокси `mean >58 Hz`: `{int(pop_vt_old['runs'])}` / `{int(pop_vt_old['failures'])}`.",
        f"- `Vt`, новое определение: `{int(pop_vt_new['runs'])}` / `{int(pop_vt_new['failures'])}`.",
        "",
        "То есть новое определение **существенно** меняет популяцию.",
        "",
        "## Что получилось по `Global`",
        "",
        f"- По `best-available duration`: медиана `{float(glob_best_base['median_duration']):.1f}` дн. в базе против `{float(glob_best_hf['median_duration']):.1f}` дн. в новой HF-группе.",
        f"- По `TTF_true`: медиана `{float(glob_true_base['median_duration']):.1f}` дн. против `{float(glob_true_hf['median_duration']):.1f}` дн.",
        f"- Скорректированный HR по `TTF_true`: `{float(glob_hr_true['hr_hf']):.2f}`.",
        f"- Разница по доле ранних отказов на горизонте `90` дней: `{float(glob_thr90['share_difference_hf_minus_baseline']):.2f}`.",
        "",
        "**Вывод по Global:** новая HF-группа в среднем не показывает явного сокращения календарной наработки, но глобально остается слабый сигнал немного большей ранней аварийности.",
        "",
        "## Что получилось по `Vt`",
        "",
        f"- По `best-available duration`: медиана `{float(vt_best_base['median_duration']):.1f}` дн. в базе против `{float(vt_best_hf['median_duration']):.1f}` дн. в новой HF-группе.",
        f"- По `TTF_true`: медиана `{float(vt_true_base['median_duration']):.1f}` дн. против `{float(vt_true_hf['median_duration']):.1f}` дн.",
        f"- Скорректированный HR по `TTF_true`: `{float(vt_hr_true['hr_hf']):.2f}`.",
        f"- Разница по доле ранних отказов на `90` днях: `{float(vt_thr90['share_difference_hf_minus_baseline']):.2f}`.",
        "",
        "**Вывод по Vt:** при новом определении группа HF **не выглядит хуже** базовой группы `50–55 Гц` ни по медиане, ни по скорректированному риску.",
        "",
        "## В чем главная разница между `Global` и `Vt`",
        "",
        "- В `Global` новая HF-группа выглядит немного менее благоприятной по ранним отказам, но без сильного эффекта по общей наработке.",
        "- В `Vt` этого сигнала почти нет: новая HF-группа по факту оказывается близкой к базе или даже выглядит лучше по медиане.",
        "- Значит, новое определение не воспроизводит тот риск, который мы видели на старом узком прокси, ориентированном на режим, близкий к `60 Гц`.",
        "",
        "## Риски и проблемное оборудование",
        "",
        "### Global",
        "",
        f"- Наиболее переизбыточная категория в новой HF-группе: `{glob_top1['category']}` (`{float(glob_top1['share_hf']):.1%}` против `{float(glob_top1['share_rest']):.1%}` в остальной популяции).",
        f"- Следующая по отклонению: `{glob_top2['category']}` (`{float(glob_top2['share_hf']):.1%}` против `{float(glob_top2['share_rest']):.1%}`).",
        "",
        "Практически это значит, что на глобальном уровне новое определение HF не выделяет один явно «провальный» узел — профиль отказов довольно близок к остальному фонду.",
        "",
        "### Vt",
        "",
        f"- Наиболее переизбыточная категория: `{vt_top1['category']}` (`{float(vt_top1['share_hf']):.1%}` против `{float(vt_top1['share_rest']):.1%}`).",
        f"- Следующая по отклонению: `{vt_top2['category']}` (`{float(vt_top2['share_hf']):.1%}` против `{float(vt_top2['share_rest']):.1%}`).",
        "",
        "Для `Vt` это смещает фокус риска не на двигатель как главный первый риск, а скорее на механические / забивочные сценарии и вал.",
        "",
        "## Это более жесткая среда?",
        "",
        f"- В `Vt` новая HF-группа не выглядит химически тяжелее: медианный `H2S` `{float(h2s_vt['hf_median']):.3f}` против `{float(h2s_vt['rest_median']):.3f}`, медианный `GLF` `{float(glf_vt['hf_median']):.1f}` против `{float(glf_vt['rest_median']):.1f}`.",
        "",
        "То есть новый результат нельзя объяснить тем, что в HF-группу попали просто самые тяжелые по среде скважины.",
        "",
        "## Интерпретация по TRF",
        "",
        f"- По оси `TRF` скорректированный HR составляет `{float(glob_hr_trf['hr_hf']):.2f}` для `Global` и `{float(vt_hr_trf['hr_hf']):.2f}` для `Vt`.",
        "",
        "Это не подтверждает гипотезу о том, что при новом определении HF насосы обязательно «съедают» меньший накопленный частотный ресурс до отказа.",
        "",
        "## Финальный вывод",
        "",
        "- Если вопрос звучит как **«опасно ли для `Vt`, если скважина большую часть времени работает выше `55 Гц`?»**, то по новому определению ответ скорее **нет, явного ухудшения не видно**.",
        "- Если вопрос звучит как **«опасен ли именно режим, близкий к `60 Гц`?»**, тогда новое определение слишком широкое и размывает тот сигнал, который был на старом прокси.",
        "",
        "## Ключевые рисунки",
        "",
        "### Global: KM по `TTF_true`",
        "",
        "![Global HF share55 TTF_true](../analysis_outputs/vt_60hz_prompt_analysis_2026_06_22/figures/hf_share55_ttf_true_days_global_km.png)",
        "",
        "### Vt: KM по `TTF_true`",
        "",
        "![Vt HF share55 TTF_true](../analysis_outputs/vt_60hz_prompt_analysis_2026_06_22/figures/hf_share55_ttf_true_days_vt_km.png)",
        "",
        "### Ранние отказы: 30 / 60 / 90 дней",
        "",
        "![HF share55 thresholds](../analysis_outputs/vt_60hz_prompt_analysis_2026_06_22/figures/hf_share55_thresholds.png)",
        "",
        "### Vt: профиль отказов новой HF-группы",
        "",
        "![HF share55 Vt mix](../analysis_outputs/vt_60hz_prompt_analysis_2026_06_22/figures/hf_share55_vt_failure_mix.png)",
        "",
    ]
    path = DOCS_DIR / "vt_highfreq_share55_global_vt_ru.md"
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def add_title_slide(prs: Presentation, title: str, subtitle: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(246, 247, 249)
    bg.line.fill.background()
    box = slide.shapes.add_textbox(Inches(0.7), Inches(0.7), Inches(11.6), Inches(1.2))
    p = box.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(27)
    p.font.bold = True
    p.font.color.rgb = RGBColor(26, 36, 52)
    sub = slide.shapes.add_textbox(Inches(0.75), Inches(1.75), Inches(11.0), Inches(1.0))
    p2 = sub.text_frame.paragraphs[0]
    p2.text = subtitle
    p2.font.size = Pt(16)
    p2.font.color.rgb = RGBColor(74, 86, 104)


def add_bullets_slide(prs: Presentation, title: str, bullets: list[str], image_path: Path | None = None) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title_box = slide.shapes.add_textbox(Inches(0.45), Inches(0.3), Inches(12.0), Inches(0.6))
    p = title_box.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = RGBColor(26, 36, 52)
    body_width = Inches(5.4) if image_path else Inches(11.8)
    body = slide.shapes.add_textbox(Inches(0.55), Inches(1.0), body_width, Inches(5.9))
    tf = body.text_frame
    tf.word_wrap = True
    for idx, bullet in enumerate(bullets):
        para = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        para.text = bullet
        para.font.size = Pt(17)
        para.font.color.rgb = RGBColor(46, 56, 68)
        para.space_after = Pt(10)
    if image_path is not None and image_path.exists():
        slide.shapes.add_picture(str(image_path), Inches(6.0), Inches(1.0), width=Inches(6.4))


def build_slides(inputs: dict[str, pd.DataFrame], global_mix: pd.DataFrame) -> Path:
    population = inputs["population"]
    axis = inputs["axis"]
    hr = inputs["hr"]
    threshold_tests = inputs["threshold_tests"]
    mix_vt = inputs["mix"]

    pop_global_new = pick(population, population="Global", definition="days >55 / duration_best_days >0.5")
    pop_vt_new = pick(population, population="Vt", definition="days >55 / duration_best_days >0.5")
    glob_true_base = pick(axis, axis_name="ttf_true_days", population="Global", mode="50-55 Hz baseline")
    glob_true_hf = pick(axis, axis_name="ttf_true_days", population="Global", mode="HF by >55d share >0.5")
    vt_true_base = pick(axis, axis_name="ttf_true_days", population="Vt", mode="50-55 Hz baseline")
    vt_true_hf = pick(axis, axis_name="ttf_true_days", population="Vt", mode="HF by >55d share >0.5")
    glob_hr_true = pick(hr, axis_name="ttf_true_days", population="Global")
    vt_hr_true = pick(hr, axis_name="ttf_true_days", population="Vt")
    glob_thr90 = pick(threshold_tests, threshold_days=90, population="Global")
    vt_thr90 = pick(threshold_tests, threshold_days=90, population="Vt")

    mix_vt_hf = mix_vt.loc[mix_vt["group"] == "HF share>0.5"].copy()
    mix_vt_rest = mix_vt.loc[mix_vt["group"] == "Rest of Vt"].copy()
    mix_vt_join = mix_vt_hf.merge(mix_vt_rest[["category", "share"]], on="category", how="left", suffixes=("_hf", "_rest"))
    mix_vt_join["share_diff"] = mix_vt_join["share_hf"] - mix_vt_join["share_rest"].fillna(0.0)
    mix_vt_join = mix_vt_join.sort_values("share_diff", ascending=False).reset_index(drop=True)
    vt_top1 = mix_vt_join.iloc[0]
    vt_top2 = mix_vt_join.iloc[1]

    mix_glob_hf = global_mix.loc[global_mix["group"] == "HF"].copy()
    mix_glob_rest = global_mix.loc[global_mix["group"] == "Rest"].copy()
    mix_glob_join = mix_glob_hf.merge(mix_glob_rest[["category", "share"]], on="category", how="left", suffixes=("_hf", "_rest"))
    mix_glob_join["share_diff"] = mix_glob_join["share_hf"] - mix_glob_join["share_rest"].fillna(0.0)
    mix_glob_join = mix_glob_join.sort_values("share_diff", ascending=False).reset_index(drop=True)
    glob_top1 = mix_glob_join.iloc[0]
    glob_top2 = mix_glob_join.iloc[1]

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    add_title_slide(
        prs,
        "Новая HF-дефиниция: Global и Vt",
        "Высокочастотная скважина = более 50% дней работы выше 55 Гц",
    )
    add_bullets_slide(
        prs,
        "Как изменилась популяция",
        [
            f"Global: новая HF-группа = {int(pop_global_new['runs'])} прогонов и {int(pop_global_new['failures'])} отказов.",
            f"Vt: новая HF-группа = {int(pop_vt_new['runs'])} прогонов и {int(pop_vt_new['failures'])} отказов.",
            "Это существенно шире, чем старый прокси, поэтому результат меняется.",
        ],
    )
    add_bullets_slide(
        prs,
        "Global: итог по TTF",
        [
            f"TTF_true медиана: {float(glob_true_base['median_duration']):.0f} дн. в базе против {float(glob_true_hf['median_duration']):.0f} дн. в новой HF-группе.",
            f"Скорректированный HR по TTF_true: {float(glob_hr_true['hr_hf']):.2f}.",
            f"Разница по ранним отказам на 90 днях: {float(glob_thr90['share_difference_hf_minus_baseline']):.2f}.",
            "Вывод: глобально есть лишь слабый неблагоприятный сигнал, но без явного падения TTF.",
        ],
        image_path=FIGURES_DIR / "hf_share55_ttf_true_days_global_km.png",
    )
    add_bullets_slide(
        prs,
        "Vt: итог по TTF",
        [
            f"TTF_true медиана: {float(vt_true_base['median_duration']):.0f} дн. в базе против {float(vt_true_hf['median_duration']):.0f} дн. в новой HF-группе.",
            f"Скорректированный HR по TTF_true: {float(vt_hr_true['hr_hf']):.2f}.",
            f"Разница по ранним отказам на 90 днях: {float(vt_thr90['share_difference_hf_minus_baseline']):.2f}.",
            "Вывод: в Vt новая HF-группа не показывает ухудшения относительно базы 50–55 Гц.",
        ],
        image_path=FIGURES_DIR / "hf_share55_ttf_true_days_vt_km.png",
    )
    add_bullets_slide(
        prs,
        "Профиль рисков: Global",
        [
            f"Наиболее переизбыточная категория: {glob_top1['category']} ({float(glob_top1['share_hf']):.0%} против {float(glob_top1['share_rest']):.0%}).",
            f"Следующая категория: {glob_top2['category']} ({float(glob_top2['share_hf']):.0%} против {float(glob_top2['share_rest']):.0%}).",
            "Но в целом профиль отказов у новой HF-группы в Global близок к остальному фонду.",
        ],
    )
    add_bullets_slide(
        prs,
        "Профиль рисков: Vt",
        [
            f"Наиболее переизбыточная категория: {vt_top1['category']} ({float(vt_top1['share_hf']):.0%} против {float(vt_top1['share_rest']):.0%}).",
            f"Следующая категория: {vt_top2['category']} ({float(vt_top2['share_hf']):.0%} против {float(vt_top2['share_rest']):.0%}).",
            "Для Vt акцент смещается в сторону механических / забивочных сценариев, а не к явному общему сокращению TTF.",
        ],
        image_path=FIGURES_DIR / "hf_share55_vt_failure_mix.png",
    )
    add_bullets_slide(
        prs,
        "Главный вывод",
        [
            "Если вопрос — опасно ли, что скважина большую часть времени работает выше 55 Гц, то по новой дефиниции для Vt явного ухудшения не видно.",
            "Если вопрос — опасен ли именно режим, близкий к 60 Гц, новая дефиниция слишком широкая и размывает прежний риск-сигнал.",
            "Поэтому для ответа про sustained >55 Hz эта дефиниция подходит, а для ответа про near-60 Hz — нет.",
        ],
        image_path=FIGURES_DIR / "hf_share55_thresholds.png",
    )

    path = BASE_OUTPUT_DIR / "vt_highfreq_share55_ru_slides.pptx"
    prs.save(path)
    return path


def main() -> None:
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    inputs = load_inputs()
    global_mix = build_global_mix()
    md_path = build_markdown(inputs, global_mix)
    slides_path = build_slides(inputs, global_mix)
    summary = {
        "markdown_path": str(md_path),
        "slides_path": str(slides_path),
    }
    (BASE_OUTPUT_DIR / "vt_highfreq_share55_ru_pack.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(f"[hf-share55-ru] Wrote markdown to {md_path}", flush=True)
    print(f"[hf-share55-ru] Wrote slides to {slides_path}", flush=True)


if __name__ == "__main__":
    main()
