"""Data loading and crosswalks for the production-risk workflow."""
from __future__ import annotations

import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from functools import lru_cache
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

REPO_ROOT = Path(__file__).resolve().parents[4]
for _p in (str(REPO_ROOT), str(REPO_ROOT / "backend")):
    if _p not in sys.path:
        sys.path.insert(0, _p)

from analysis.paths import (
    resolve_gtm_schedule_path,
    resolve_pp_master_path,
    resolve_prediction_workbook_path,
    resolve_techregime_workbook_path,
)
from analysis.workflows.production_risk import config as C

_CODE_RE = re.compile(r"^([A-ZА-Я]+_\d+)")


def norm_well(x) -> str | None:
    if x is None:
        return None
    s = str(x).strip().upper().replace(" ", "")
    if not s:
        return None
    m = _CODE_RE.match(s)
    return m.group(1) if m else s


def contractor_group(v) -> str:
    s = str(v or "")
    if "Борец" in s:
        return "brt"
    if "Шлюмберже" in s:
        return "slb"
    if not s.strip():
        return "Pooled"
    return "oth"


def sour_group(v) -> str:
    s = str(v or "")
    return "sour" if ("исл" in s and "екисл" not in s) else "nonsour"


def map_model_field_from_well(code: str) -> str | None:
    prefix = code.split("_", 1)[0].upper() if "_" in code else code.upper()
    if prefix in C.EXPLICIT_GLOBAL_FALLBACK:
        return None
    return C.FIELD_PREFIX_MAP.get(prefix)


def model_prefix_from_well(code: str) -> str:
    return code.split("_", 1)[0].upper() if "_" in code else code.upper()


def is_explicit_global_fallback(code: str) -> bool:
    return model_prefix_from_well(code) in C.EXPLICIT_GLOBAL_FALLBACK


def _as_float(x) -> float | None:
    if isinstance(x, (int, float, np.integer, np.floating)) and np.isfinite(float(x)):
        return float(x)
    return None


def _as_positive_float(x) -> float | None:
    value = _as_float(x)
    return value if value is not None and value > 0.0 else None


@lru_cache(maxsize=8)
def _qnominal_tables(path_key: str = "") -> tuple[dict[str, float], dict[str, float]]:
    """Installed Qnom by well plus field-typical medians from WellsArtificialLiftBig."""
    try:
        from analysis.data.equipment_big import load_equipment_big
        df = load_equipment_big(Path(path_key) if path_key else None)
    except Exception:
        return {}, {}
    if df.empty or "well_key" not in df.columns or "q_nom_m3d" not in df.columns:
        return {}, {}

    focus = df.copy()
    if "is_esp" in focus.columns:
        focus = focus[focus["is_esp"].fillna(False)].copy()
    focus["well_norm"] = focus["well_key"].map(norm_well)
    focus["q_nom_m3d"] = pd.to_numeric(focus["q_nom_m3d"], errors="coerce")
    focus = focus[focus["well_norm"].notna() & (focus["q_nom_m3d"] > 0)].copy()
    if focus.empty:
        return {}, {}

    by_well: dict[str, float] = {}
    if "install_date" in focus.columns:
        focus = focus.sort_values(["well_norm", "install_date"], na_position="last")
    for well, grp in focus.groupby("well_norm"):
        vals = grp["q_nom_m3d"].dropna()
        if not vals.empty:
            by_well[str(well)] = float(vals.iloc[-1])

    focus["model_field"] = focus["well_norm"].map(map_model_field_from_well)
    by_field = {
        str(field): float(grp["q_nom_m3d"].median())
        for field, grp in focus.dropna(subset=["model_field"]).groupby("model_field")
        if grp["q_nom_m3d"].notna().any()
    }
    return by_well, by_field


def field_typical_qnominal_map(equipment_big_path: Path | None = None) -> dict[str, float]:
    """Median installed Qnom by model field, used for planned/new wells."""
    _, by_field = _qnominal_tables(str(equipment_big_path) if equipment_big_path else "")
    return dict(by_field)


def resolve_qnominal(
    code: str,
    model_field: str | None = None,
    equipment_big_path: Path | None = None,
) -> float | None:
    """Resolve raw Kpod denominator Qnominal for a well.

    Existing wells use installed ``q_nom_m3d`` from Big.  New/planned wells use the
    field-typical median so planned high-rate wells can still fire the overload side
    instead of self-neutralizing via a back-calculated nominal.
    """
    by_well, by_field = _qnominal_tables(str(equipment_big_path) if equipment_big_path else "")
    norm = norm_well(code) or str(code)
    exact = _as_positive_float(by_well.get(norm))
    if exact is not None:
        return exact
    field = str(model_field or map_model_field_from_well(norm) or "")
    typical = _as_positive_float(by_field.get(field))
    return typical


def _as_datetime(x) -> datetime | None:
    return x if isinstance(x, datetime) else None


def _month_id(x) -> str | None:
    if isinstance(x, datetime):
        return x.strftime("%Y-%m")
    if isinstance(x, date):
        return datetime(x.year, x.month, x.day).strftime("%Y-%m")
    if x is None:
        return None
    s = str(x).strip()
    if not s:
        return None
    try:
        return pd.to_datetime(s, dayfirst=True).strftime("%Y-%m")
    except Exception:
        return None


def _header_map(ws, row_idx: int) -> dict[str, int]:
    headers = next(ws.iter_rows(min_row=row_idx, max_row=row_idx, values_only=True))
    out: dict[str, int] = {}
    for idx, value in enumerate(headers):
        if value is None:
            continue
        out[str(value).strip()] = idx
    return out


def _find_col(header: dict[str, int], names: list[str]) -> int:
    for name in names:
        if name in header:
            return header[name]
    raise KeyError(f"missing required columns: {names}")


def _month_columns(ws, header_row: int = 3) -> tuple[list[str], list[int]]:
    header = next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True))
    months: list[str] = []
    cols: list[int] = []
    for idx, value in enumerate(header):
        month = _month_id(value)
        if month is not None:
            months.append(month)
            cols.append(idx)
    return months, cols


def _frame_from_records(records: list[tuple], months: list[str]) -> pd.DataFrame:
    if not records:
        return pd.DataFrame(columns=months, dtype=float)
    df = pd.DataFrame.from_records(records, columns=["wid", *months]).groupby("wid", as_index=True).sum()
    return df.astype(float)


@dataclass
class PlanData:
    master_path: Path
    months: list[str]
    fwd_months: list[str]
    oil_volume: pd.DataFrame
    oil_volume_m3: pd.DataFrame
    liquid_volume: pd.DataFrame
    gas_volume: pd.DataFrame
    produced_water_volume: pd.DataFrame
    produced_water_rate: pd.DataFrame
    op_days: pd.DataFrame
    op_days_raw: pd.DataFrame
    registry_oil_rate: pd.DataFrame
    registry_oil_m3_rate: pd.DataFrame
    registry_liquid_rate: pd.DataFrame
    registry_water_rate: pd.DataFrame
    cal_days: pd.Series
    producer_meta: pd.DataFrame
    producers: list[str]
    anomalies: pd.DataFrame


@dataclass
class EspRun:
    well_code: str
    field_raw: str | None
    ctr_raw: str | None
    sour_raw: str | None
    run_seq: int | None
    mount: datetime | None
    stop: datetime | None
    demo: datetime | None
    age_op: float | None
    failure_flag: int | None
    glf: float | None
    load_mean: float | None
    curvature: float | None


@dataclass
class EspState:
    well_code: str
    field: str | None
    ctr: str
    sour: str
    run_seq: int | None
    mount: datetime | None
    stop: datetime | None
    demo: datetime | None
    age_op: float | None
    failure_flag: int | None
    is_active: bool
    covariates: dict[str, float]
    cov_source: str


@dataclass
class EspSource:
    workbook_path: Path
    source_cutoff: datetime | None
    runs_by_well: dict[str, list[EspRun]]
    states_by_well: dict[str, EspState]


@dataclass
class DowntimeStats:
    n: int
    p25: int
    p50: int
    p75: int
    mean: int


@dataclass
class CurrentTechregimeStatus:
    well_code: str
    status: str
    in_operation: bool
    lift: str | None
    age_op: float | None
    report_date: datetime | None


def _read_plan_sheet(
    ws,
    months: list[str],
    month_cols: list[int],
    *,
    rate_mode: bool,
) -> tuple[dict[str, pd.DataFrame], dict[str, dict]]:
    oil_records: list[tuple] = []
    oil_m3_records: list[tuple] = []
    liq_records: list[tuple] = []
    water_records: list[tuple] = []
    gas_records: list[tuple] = []
    op_records: list[tuple] = []
    meta: dict[str, dict] = {}
    for row in ws.iter_rows(min_row=4, values_only=True):
        wid = row[11]
        if wid is None:
            continue
        code = norm_well(wid)
        if code is None:
            continue
        meta.setdefault(
            code,
            {
                "raw_id": str(wid).strip(),
                "plan_field": str(row[0]).strip() if row[0] is not None else "",
                "license_area": str(row[1]).strip() if row[1] is not None else "",
                "pad": str(row[12]).strip() if row[12] is not None else "",
            },
        )
        indicator = str(row[8] or "").strip()
        unit = str(row[9] or "").strip().lower()
        values = [float(v) if isinstance(v, (int, float)) else 0.0 for i, v in enumerate(row) if i in month_cols]
        is_rate = "/сут" in unit
        if indicator == "Добыча нефти" and is_rate == rate_mode:
            oil_records.append((code, *values))
        elif indicator == "Добыча нефти м3" and is_rate == rate_mode:
            oil_m3_records.append((code, *values))
        elif indicator == "Добыча жидкости" and is_rate == rate_mode:
            liq_records.append((code, *values))
        elif _is_produced_water_indicator(indicator) and is_rate == rate_mode:
            water_records.append((code, *values))
        elif _is_associated_gas_indicator(indicator) and is_rate == rate_mode:
            # Stored in тыс. м3; scaled to м3 so it shares units with ``qgas``.
            gas_records.append((code, *[v * 1000.0 for v in values]))
        elif indicator == "Отработанное время" and not rate_mode:
            op_records.append((code, *values))
    return (
        {
            "oil": _frame_from_records(oil_records, months),
            "oil_m3": _frame_from_records(oil_m3_records, months),
            "liq": _frame_from_records(liq_records, months),
            "water": _frame_from_records(water_records, months),
            "gas": _frame_from_records(gas_records, months),
            "op": _frame_from_records(op_records, months),
        },
        meta,
    )


def _is_associated_gas_indicator(indicator: str) -> bool:
    """True only for associated (попутный) gas produced through the well.

    Three lookalikes must not match.  ``ППД_газ`` / ``ППД_газ (Сайклинг)`` are injection.
    ``Добыча_ПГ`` is natural gas from gas wells, a different stream.  ``Добыча C2..C5+``
    are separated fractions already counted inside ПНГ.
    """
    text = str(indicator or "").strip().casefold()
    if not text or "ппд" in text:
        return False
    return text == "добыча пнг"


def _is_produced_water_indicator(indicator: str) -> bool:
    """True for produced-water rows, false for pressure-maintenance water.

    The plan contains ``Добыча воды (ППД)`` / ``ППД_вода`` rows.  Those are water
    injection / pressure-maintenance volumes, not water produced through the ESP,
    so they must not feed produced-water audit columns.
    """
    text = str(indicator or "").strip().casefold()
    if not text:
        return False
    if "ппд" in text:
        return False
    return "попутной воды" in text or text == "добыча воды"


def _align_frame(df: pd.DataFrame, index: pd.Index, months: list[str]) -> pd.DataFrame:
    return df.reindex(index=index, columns=months, fill_value=0.0).astype(float)


def load_plan(
    forecast_start: date = C.FORECAST_START,
    horizon_end: date = C.HORIZON_END,
    master_path: Path | None = None,
) -> PlanData:
    master_path = Path(master_path or resolve_pp_master_path())
    wb = openpyxl.load_workbook(master_path, read_only=True, data_only=True)
    ws_summary = wb["Сводные данные"]
    ws_registry = wb["Реестр"]
    months, month_cols = _month_columns(ws_summary, header_row=3)
    summary_frames, meta = _read_plan_sheet(ws_summary, months, month_cols, rate_mode=False)
    registry_frames, _ = _read_plan_sheet(ws_registry, months, month_cols, rate_mode=True)
    wb.close()

    all_index = pd.Index(
        sorted(
            set(summary_frames["oil"].index)
            | set(summary_frames["liq"].index)
            | set(summary_frames["op"].index)
            | set(registry_frames["oil"].index)
            | set(registry_frames["liq"].index)
        ),
        name="wid",
    )
    oil_volume = _align_frame(summary_frames["oil"], all_index, months)
    oil_volume_m3 = _align_frame(summary_frames["oil_m3"], all_index, months)
    liquid_volume = _align_frame(summary_frames["liq"], all_index, months)
    gas_volume = _align_frame(summary_frames["gas"], all_index, months)
    direct_water_volume = _align_frame(summary_frames["water"], all_index, months)
    op_days_raw = _align_frame(summary_frames["op"], all_index, months)
    registry_oil_rate = _align_frame(registry_frames["oil"], all_index, months)
    registry_oil_m3_rate = _align_frame(registry_frames["oil_m3"], all_index, months)
    registry_liquid_rate = _align_frame(registry_frames["liq"], all_index, months)
    registry_water_rate = _align_frame(registry_frames["water"], all_index, months)

    cal_days = pd.Series(
        {month: float(pd.Period(month, freq="M").days_in_month) for month in months},
        dtype=float,
    )
    cal_matrix = pd.DataFrame(
        np.repeat(cal_days.to_numpy()[None, :], len(all_index), axis=0),
        index=all_index,
        columns=months,
    )
    op_days = op_days_raw.clip(lower=0.0)
    op_days = pd.DataFrame(
        np.minimum(op_days.to_numpy(dtype=float), cal_matrix.to_numpy(dtype=float)),
        index=all_index,
        columns=months,
    )
    derived_water_volume = (liquid_volume - oil_volume_m3).where(oil_volume_m3 > 0.0, 0.0).clip(lower=0.0)
    produced_water_volume = direct_water_volume.where(direct_water_volume > 0.0, derived_water_volume)
    derived_water_rate = np.divide(
        produced_water_volume.to_numpy(dtype=float),
        op_days.to_numpy(dtype=float),
        out=np.zeros_like(produced_water_volume.to_numpy(dtype=float), dtype=float),
        where=op_days.to_numpy(dtype=float) > 0.0,
    )
    produced_water_rate = pd.DataFrame(derived_water_rate, index=all_index, columns=months)
    produced_water_rate = registry_water_rate.where(registry_water_rate > 0.0, produced_water_rate)

    anomaly_rows: list[dict] = []
    for wid in all_index:
        count = 0
        for month in months:
            raw = float(op_days_raw.at[wid, month]) if month in op_days_raw.columns else 0.0
            clipped = float(op_days.at[wid, month])
            oil = float(oil_volume.at[wid, month])
            liq = float(liquid_volume.at[wid, month])
            cal = float(cal_days[month])
            reason = None
            if raw < 0:
                reason = "negative_runtime"
            elif raw > cal:
                reason = "runtime_gt_calendar"
            elif (oil > 0 or liq > 0) and clipped <= 0:
                reason = "positive_volume_zero_runtime"
            if reason is not None:
                count += 1
                anomaly_rows.append(
                    {
                        "wid": wid,
                        "month": month,
                        "reason": reason,
                        "runtime_raw": raw,
                        "runtime_model": clipped,
                        "calendar_days": cal,
                        "oil_volume_t": oil,
                        "liquid_volume_m3": liq,
                    }
                )
        meta.setdefault(wid, {"raw_id": wid, "plan_field": "", "license_area": "", "pad": ""})
        meta[wid]["runtime_anomaly_count"] = count

    fwd_months = [m for m in months if forecast_start.strftime("%Y-%m") <= m <= horizon_end.strftime("%Y-%m")]
    producers: list[str] = []
    meta_rows: list[dict] = []
    for wid in all_index:
        oil_total = float(oil_volume.loc[wid, fwd_months].sum()) if fwd_months else 0.0
        liq_total = float(liquid_volume.loc[wid, fwd_months].sum()) if fwd_months else 0.0
        summary_pos = (oil_volume.loc[wid, fwd_months] > 0) | (liquid_volume.loc[wid, fwd_months] > 0)
        reg_oil_pos = registry_oil_rate.loc[wid, fwd_months] > 0 if fwd_months else pd.Series(dtype=bool)
        reg_liq_pos = registry_liquid_rate.loc[wid, fwd_months] > 0 if fwd_months else pd.Series(dtype=bool)
        first_active = None
        if fwd_months:
            pos_months = [m for m in fwd_months if bool(summary_pos[m])]
            if pos_months:
                first_active = pos_months[0]
        if oil_total > 0 or liq_total > 0:
            producers.append(wid)
        row = dict(meta.get(wid, {}))
        row.update(
            {
                "wid": wid,
                "planned_oil_t": oil_total,
                "planned_liquid_m3": liq_total,
                "planned_produced_water_m3": float(produced_water_volume.loc[wid, fwd_months].sum()) if fwd_months else 0.0,
                "produced_water_positive_months": int((produced_water_rate.loc[wid, fwd_months] > 0).sum()) if fwd_months else 0,
                "first_active_month": first_active,
                "summary_positive_months": int(summary_pos.sum()) if fwd_months else 0,
                "registry_oil_positive_months": int(reg_oil_pos.sum()) if fwd_months else 0,
                "registry_liquid_positive_months": int(reg_liq_pos.sum()) if fwd_months else 0,
            }
        )
        meta_rows.append(row)
    producer_meta = pd.DataFrame(meta_rows).set_index("wid").sort_index()

    return PlanData(
        master_path=master_path,
        months=months,
        fwd_months=fwd_months,
        oil_volume=oil_volume,
        oil_volume_m3=oil_volume_m3,
        liquid_volume=liquid_volume,
        gas_volume=gas_volume,
        produced_water_volume=produced_water_volume,
        produced_water_rate=produced_water_rate,
        op_days=op_days,
        op_days_raw=op_days_raw,
        registry_oil_rate=registry_oil_rate,
        registry_oil_m3_rate=registry_oil_m3_rate,
        registry_liquid_rate=registry_liquid_rate,
        registry_water_rate=registry_water_rate,
        cal_days=cal_days,
        producer_meta=producer_meta,
        producers=sorted(producers),
        anomalies=pd.DataFrame(anomaly_rows),
    )


@lru_cache(maxsize=4)
def _well_deposit_pairs(master_path_key: str = "") -> tuple[tuple[str, str, str], ...]:
    master_path = Path(master_path_key) if master_path_key else resolve_pp_master_path()
    wb = openpyxl.load_workbook(master_path, read_only=True, data_only=True)
    ws = wb["Сводные данные"]
    header = _header_map(ws, 3)
    c_dep = _find_col(header, ["Месторождение"])
    c_un = _find_col(header, ["УН"])
    c_well = _find_col(header, ["Well_ID*", "Well_ID"])
    seen: dict[str, tuple[str, str]] = {}
    for row in ws.iter_rows(min_row=4, values_only=True):
        code = norm_well(row[c_well]) if c_well < len(row) else None
        if not code or code in seen:
            continue
        seen[code] = (str(row[c_dep] or "").strip(), str(row[c_un] or "").strip())
    wb.close()
    return tuple((code, dep, un) for code, (dep, un) in seen.items())


def well_deposit_map(master_path: Path | None = None) -> dict[str, str]:
    """Well code -> «Месторождение», the deposit the well actually produces from.

    The well-code prefix is the УН (operating unit), not the deposit: Верхнетирский
    УН wells sit on both Большетирское and Ичёдинское.  Свод's own «Месторождение»
    column carries one value per prefix, so the ПП plan is the only source that
    resolves the two apart.
    """
    return {code: dep for code, dep, _ in _well_deposit_pairs(str(master_path or ""))}


def well_un_map(master_path: Path | None = None) -> dict[str, str]:
    """Well code -> «УН». Matches the code prefix; kept explicit for hazard tests."""
    return {code: un for code, _, un in _well_deposit_pairs(str(master_path or ""))}


def load_gtm(path: Path | None = None) -> pd.DataFrame:
    path = Path(path or resolve_gtm_schedule_path())
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["ДФ04_cводный"]
    header_row = None
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=15, values_only=True), start=1):
        if row and str(row[0]).strip() == "Источник":
            header_row = i
            break
    recs: list[dict] = []
    if header_row is not None:
        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            well = row[6]
            if well is None:
                continue
            recs.append(
                {
                    "gtm_group": row[1],
                    "gtm_kind": row[2],
                    "plan_field": row[3],
                    "pad": row[5],
                    "well": str(well).strip(),
                    "well_code": norm_well(well),
                    "well_cat": row[8],  # col I «Категория скважины»; col H is «Конструкция»
                    "krs_date": row[13],
                    "startup_date": row[14],
                    "liquid_delta_m3": _as_float(row[15]),
                    "oil_delta_t": _as_float(row[17]),
                    "lift": row[28],
                    "h2s_ppm": _as_float(row[29]),
                    "priority": row[30],
                    "is_esp": "ЭЦН" in str(row[28] or ""),
                }
            )
    wb.close()
    df = pd.DataFrame(recs)
    if not df.empty:
        for col in ("krs_date", "startup_date"):
            df[col] = pd.to_datetime(df[col], errors="coerce")
    return df


def load_current_techregime_status(path: Path | None = None) -> dict[str, CurrentTechregimeStatus]:
    path = Path(path or resolve_techregime_workbook_path())
    if not path.exists():
        return {}
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb["ТР НЕФТЬ"] if "ТР НЕФТЬ" in wb.sheetnames else wb.worksheets[0]
        report_date = _as_datetime(ws.cell(row=3, column=4).value) or _as_datetime(ws.cell(row=4, column=4).value)
        header = _header_map(ws, 8)
        c_well = _find_col(header, ["ID скважины", "Скважина", "№ скважины"])
        c_status = _find_col(header, ["Состояние по фонду"])
        c_lift = header.get("Способ эксплуатации")
        c_age = header.get("Время наработки (ННО)")
        statuses: dict[str, CurrentTechregimeStatus] = {}
        for row in ws.iter_rows(min_row=9, values_only=True):
            code = norm_well(row[c_well])
            if code is None:
                continue
            status = str(row[c_status] or "").strip()
            if not status:
                continue
            lift = str(row[c_lift]).strip() if c_lift is not None and row[c_lift] is not None else None
            statuses[code] = CurrentTechregimeStatus(
                well_code=code,
                status=status,
                in_operation=status.casefold() == "в работе",
                lift=lift,
                age_op=_as_float(row[c_age]) if c_age is not None else None,
                report_date=report_date,
            )
        return statuses
    finally:
        wb.close()


def _fallback_covariates(
    run: EspRun, prev_run: EspRun | None, run_seq_ordinal: int | None = None
) -> dict[str, float]:
    cov: dict[str, float] = {}
    if run.glf is not None and run.glf > 0:
        cov["log_glf_mean_opdays"] = float(math.log(run.glf))
    if run.load_mean is not None:
        cov["load_mean"] = float(run.load_mean)
    cov["curvature_deg10m"] = float(run.curvature) if run.curvature is not None else 0.0
    # log_run_seq must use the within-well ordinal (1..N), not «Нспуска» (a global run-ID).
    if run_seq_ordinal is not None and run_seq_ordinal >= 0:
        cov["log_run_seq"] = float(math.log(1.0 + run_seq_ordinal))
    if run.mount is not None:
        cov["install_pre2020"] = 1.0 if run.mount.year < 2020 else 0.0
        cov["install_2023plus"] = 1.0 if run.mount.year >= 2023 else 0.0
    if prev_run is not None and run.mount is not None:
        prev_end = prev_run.stop or prev_run.demo
        if prev_end is not None:
            gap = (run.mount - prev_end).days
            if gap > 0:
                cov["log_days_since_prev_failure"] = float(math.log(gap))
    return cov


def _load_run_covariates(bundle_date: str) -> dict[tuple[str, int], dict[str, float]]:
    path = C.run_covariates_path(bundle_date)
    if not path.exists():
        return {}
    df = pd.read_csv(path, encoding="utf-8-sig")
    out: dict[tuple[str, int], dict[str, float]] = {}
    for _, row in df.iterrows():
        code = norm_well(row.get("well_key_norm"))
        seq = _as_float(row.get("run_seq"))
        if code is None or seq is None:
            continue
        cov = {}
        for key, value in row.items():
            if key in {"well_key_norm", "run_seq", "install_date", "stratum_key", "covariate_available", "clock"}:
                continue
            num = _as_float(value)
            if num is not None:
                cov[key] = num
        out[(code, int(seq))] = cov
    return out


def load_esp_source(
    bundle_date: str = C.BUNDLE_DATE,
    prediction_workbook_path: Path | None = None,
) -> EspSource:
    workbook_path = Path(prediction_workbook_path or resolve_prediction_workbook_path())
    wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    ws = wb["Свод"]
    header = _header_map(ws, 1)
    c_field = _find_col(header, ["Месторождение"])
    c_well = _find_col(header, ["Скв."])
    c_ctr = _find_col(header, ["Принадлежность"])
    c_seq = _find_col(header, ["Нспуска"])
    c_mount = _find_col(header, ["Дата монтажа"])
    c_stop = _find_col(header, ["Дата остановки"])
    c_age = _find_col(header, ["Наработка (сут)"])
    c_demo = _find_col(header, ["Дата демонтажа"])
    c_sour = _find_col(header, ["Кислый/Некислый"])
    # «Признак отказа» is NOT a fallback: in the reworked Свод it is a text
    # sub-classification (Преждевременный/Многосуточный/…), and matching it would
    # silently yield failure_flag=None on every row.
    c_ff = _find_col(header, ["Флаг отказа", "Failure Flag"])
    c_glf = _find_col(header, ["ГЖФ", "Газовый фактор"])
    c_load = _find_col(header, ["Загр, Двиг,"])
    c_curve = _find_col(header, ["Работа в кривизне"])

    run_cov_map = _load_run_covariates(bundle_date)
    runs_by_well: dict[str, list[EspRun]] = defaultdict(list)
    cutoff_dates: list[datetime] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        code = norm_well(row[c_well])
        if code is None:
            continue
        run = EspRun(
            well_code=code,
            field_raw=str(row[c_field]).strip() if row[c_field] is not None else None,
            ctr_raw=str(row[c_ctr]).strip() if row[c_ctr] is not None else None,
            sour_raw=str(row[c_sour]).strip() if row[c_sour] is not None else None,
            run_seq=int(row[c_seq]) if isinstance(row[c_seq], (int, float)) else None,
            mount=_as_datetime(row[c_mount]),
            stop=_as_datetime(row[c_stop]),
            demo=_as_datetime(row[c_demo]),
            age_op=_as_float(row[c_age]),
            failure_flag=int(row[c_ff]) if isinstance(row[c_ff], (int, float)) else None,
            glf=_as_float(row[c_glf]),
            load_mean=_as_float(row[c_load]),
            curvature=_as_float(row[c_curve]),
        )
        runs_by_well[code].append(run)
        # Cutoff = latest event the workbook knows about.  Prefer the stop date; the
        # demo date only stands in when the stop is missing — demo typos otherwise
        # poison the cutoff (seen: демонтаж 2026-11 against остановка 2026-04).
        if run.stop is not None:
            cutoff_dates.append(run.stop)
        elif run.demo is not None:
            cutoff_dates.append(run.demo)
    wb.close()

    source_cutoff = max(cutoff_dates) if cutoff_dates else None

    states_by_well: dict[str, EspState] = {}
    for code, runs in runs_by_well.items():
        runs.sort(key=lambda r: (r.mount or datetime.min, r.run_seq or -1))
        last = runs[-1]
        prev = runs[-2] if len(runs) > 1 else None
        # esp_run_covariates.csv is keyed by within-well ordinal (mount order), NOT «Нспуска».
        last_ord = len([r for r in runs if r.mount is not None]) if last.mount is not None else None
        cov = dict(run_cov_map.get((code, last_ord), {})) if last_ord is not None else {}
        cov_source = "bundle" if cov else "none"
        fallback_cov = _fallback_covariates(last, prev, last_ord)
        missing_keys = {k: v for k, v in fallback_cov.items() if k not in cov}
        if missing_keys:
            cov.update(missing_keys)
            cov_source = "bundle+fallback" if cov_source == "bundle" else "fallback"
        # The reworked Свод marks the running pump by an EMPTY stop date; the flag no
        # longer identifies it (0 is shared by open rows and closed ГТМ/ППР pulls).
        is_active = last.stop is None and last.demo is None
        age_op = last.age_op
        if is_active and age_op is None and last.mount is not None:
            # Open rows carry no «Наработка (сут)» — the running pump's age is the
            # calendar age at the techregime snapshot the Свод was built from
            # (same convention as esp_population.load_svod_runs).
            snapshot = datetime.combine(C.SVOD_OPEN_ASOF, datetime.min.time())
            age_op = float(max((snapshot - last.mount).days, 0))
        states_by_well[code] = EspState(
            well_code=code,
            field=last.field_raw,
            ctr=contractor_group(last.ctr_raw),
            sour=sour_group(last.sour_raw),
            run_seq=last.run_seq,
            mount=last.mount,
            stop=last.stop,
            demo=last.demo,
            age_op=age_op,
            failure_flag=last.failure_flag,
            is_active=is_active,
            covariates=cov,
            cov_source=cov_source,
        )
    return EspSource(
        workbook_path=workbook_path,
        source_cutoff=source_cutoff,
        runs_by_well=dict(runs_by_well),
        states_by_well=states_by_well,
    )


def derive_downtime_quantiles(
    source: EspSource,
    max_gap_days: int = C.MAX_REPAIR_GAP_DAYS,
) -> tuple[DowntimeStats, dict[str, DowntimeStats]]:
    global_gaps: list[float] = []
    field_gaps: dict[str, list[float]] = defaultdict(list)
    for runs in source.runs_by_well.values():
        rr = [r for r in runs if r.mount is not None]
        rr.sort(key=lambda r: (r.mount or datetime.min, r.run_seq or -1))
        for i in range(len(rr) - 1):
            end = rr[i].stop or rr[i].demo
            nxt = rr[i + 1].mount
            if end is None or nxt is None:
                continue
            gap = (nxt - end).days
            if 0 <= gap <= max_gap_days:
                global_gaps.append(float(gap))
                field = str(rr[i].field_raw or "").strip()
                if field:
                    field_gaps[field].append(float(gap))

    def _stats(values: list[float], fallback: dict[str, int] | None = None) -> DowntimeStats:
        if not values:
            fb = fallback or C.GLOBAL_DOWNTIME_FALLBACK
            return DowntimeStats(
                n=0,
                p25=int(fb["p25"]),
                p50=int(fb["p50"]),
                p75=int(fb["p75"]),
                mean=int(round((float(fb["p25"]) + float(fb["p50"]) + float(fb["p75"])) / 3.0)),
            )
        arr = np.asarray(values, dtype=float)
        return DowntimeStats(
            n=int(arr.size),
            p25=int(round(np.percentile(arr, 25))),
            p50=int(round(np.percentile(arr, 50))),
            p75=int(round(np.percentile(arr, 75))),
            mean=int(round(float(np.mean(arr)))),
        )

    global_stats_raw = _stats(global_gaps, fallback=C.GLOBAL_DOWNTIME_FALLBACK)
    global_stats = DowntimeStats(
        n=global_stats_raw.n,
        p25=int(C.GLOBAL_DOWNTIME_FALLBACK["p25"]),
        p50=int(C.GLOBAL_DOWNTIME_FALLBACK["p50"]),
        p75=int(C.GLOBAL_DOWNTIME_FALLBACK["p75"]),
        mean=global_stats_raw.mean,
    )
    field_stats = {field: _stats(values) for field, values in field_gaps.items()}
    return global_stats, field_stats
